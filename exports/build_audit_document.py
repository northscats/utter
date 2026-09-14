"""Create the audit's Word source for native DOC conversion with LibreOffice."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).parent
source = (ROOT / "audit-results.md").read_text()
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.top_margin = section.bottom_margin = Cm(1.7)
section.left_margin = section.right_margin = Cm(1.7)
normal = doc.styles["Normal"]
normal.font.name = "DejaVu Sans"
normal.font.size = Pt(10)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.12
for name, size in [("Title", 26), ("Heading 1", 17), ("Heading 2", 13), ("Heading 3", 11)]:
    style = doc.styles[name]
    style.font.name = "DejaVu Sans"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string("174A6E")
    style.paragraph_format.keep_with_next = True

def inline(paragraph, text):
    for n, part in enumerate(text.split("**")):
        run = paragraph.add_run(part)
        run.bold = bool(n % 2)

header = section.header.paragraphs[0]
header.text = "brain4sales  |  Итоги аудита  |  13.09.2026"
header.style = doc.styles["Caption"]
footer = section.footer.paragraphs[0]
footer.alignment = 2
footer.add_run("Страница ")
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)
doc.core_properties.title = "Итоги аудита"
doc.core_properties.subject = "Источники данных, финансовые показатели и риски предприятий"
doc.core_properties.author = "brain4sales"

lines = source.splitlines()
i = 0
tables = 0
while i < len(lines):
    text = lines[i].strip()
    if not text:
        i += 1
        continue
    if text.startswith("|"):
        rows = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            cells = [x.strip() for x in lines[i].strip().strip("|").split("|")]
            if not all(re.fullmatch(r"[:\s-]+", c) for c in cells):
                rows.append(cells)
            i += 1
        table = doc.add_table(rows=0, cols=len(rows[0]))
        table.style = "Table Grid"
        table.autofit = False
        widths = [3.4, 1.8, 5.4, 7.0] if tables in (2, 3) else (
            [3.4, 6.1, 2.2, 5.9] if tables == 0 else (
            [5.2, 1.8, 10.6] if len(rows[0]) == 3 else [3.0, 4.8, 4.8, 5.0]))
        for col, width in zip(table.columns, widths):
            col.width = Cm(width)
        for index, row in enumerate(rows):
            cells = table.add_row().cells
            for c, width in zip(cells, widths):
                c.width = Cm(width)
            for c, value in zip(cells, row):
                p = c.paragraphs[0]
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.line_spacing = 1.05
                inline(p, value)
                for run in p.runs:
                    run.font.size = Pt(8.5)
                    if index == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
                if index == 0:
                    shading = OxmlElement("w:shd")
                    shading.set(qn("w:fill"), "174A6E")
                    c._tc.get_or_add_tcPr().append(shading)
            props = table.rows[-1]._tr.get_or_add_trPr()
            props.append(OxmlElement("w:cantSplit"))
            if index == 0:
                props.append(OxmlElement("w:tblHeader"))
        doc.add_paragraph()
        tables += 1
        continue
    if text.startswith("#"):
        level = len(text) - len(text.lstrip("#"))
        value = text[level:].strip()
        if i == 0:
            doc.add_paragraph(value, "Title")
        else:
            doc.add_heading(value, level=min(level, 3))
    elif text.startswith("- "):
        inline(doc.add_paragraph(style="List Bullet"), text[2:])
    else:
        # Keep the explicit ordered item number from the source.
        inline(doc.add_paragraph(), text)
    i += 1

assert tables == 7, tables
out = ROOT / "Итоги аудита.docx"
doc.save(out)
print(f"Created {out}: {tables} tables, {len(doc.paragraphs)} paragraphs")