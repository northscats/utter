"""Shared formatting for narrative audit spreadsheets; no application dependencies."""
from pathlib import Path
from datetime import datetime
from zoneinfo import ZoneInfo
import textwrap
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


def parse_tsv(text):
    rows = [line.split("\t") for line in text.strip().splitlines()]
    for i, row in enumerate(rows, 1):
        assert len(row) == len(rows[0]), (i, len(row), len(rows[0]))
    return rows


def build_audit(path, title, notes, sections):
    wb = Workbook()
    cover = wb.active
    cover.title = "Об отчёте"
    wb.properties.title = title
    wb.properties.creator = "Replit Agent"
    counts = {}
    for index, (name, rows, widths) in enumerate(sections, 1):
        ws = wb.create_sheet(name)
        assert all(len(row) == len(widths) for row in rows)
        count = len(widths)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=count)
        ws.cell(1, 1, name).font = Font(name="Calibri", size=19, bold=True, color="FFFFFF")
        ws.cell(1, 1).fill = PatternFill("solid", fgColor="153454")
        ws.row_dimensions[1].height = 34
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=count)
        ws.cell(2, 1, "Аудит кода и методики • Фактические финансовые суммы предприятия не сверялись")
        ws.cell(2, 1).font = Font(name="Calibri", size=11, italic=True, color="536577")
        ws.row_dimensions[2].height = 25
        ws.cell(3, 1, "К оглавлению").hyperlink = "#'Об отчёте'!A1"
        ws.cell(3, 1).style = "Hyperlink"
        for ri, row in enumerate(rows, 5):
            for ci, value in enumerate(row, 1):
                cell = ws.cell(ri, ci, value)
                cell.font = Font(name="Calibri", size=11, color="203449", bold=ci <= 2)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if ri == 5:
                    cell.fill = PatternFill("solid", fgColor="153454")
                    cell.font = Font(name="Calibri", size=11, color="FFFFFF", bold=True)
                elif ri % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="F0F5FA")
                if value == "Критический":
                    cell.fill = PatternFill("solid", fgColor="FCE4E4")
                    cell.font = Font(name="Calibri", size=11, color="A62626", bold=True)
            lines = max(max(1, len(textwrap.wrap(str(v), width=max(7, int(w * .83)))))
                        for v, w in zip(row, widths))
            ws.row_dimensions[ri].height = max(44 if ri == 5 else 70, min(390, lines * 15 + 16))
        for ci, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(ci)].width = width
        table = Table(displayName=f"AuditTable{index}", ref=f"A5:{get_column_letter(count)}{ws.max_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)
        ws.freeze_panes = "C6"
        ws.sheet_view.showGridLines = False
        ws.sheet_view.zoomScale = 75
        ws.print_title_rows = "1:5"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_area = f"A1:{get_column_letter(count)}{ws.max_row}"
        ws.oddFooter.center.text = "Страница &P из &N"
        counts[name] = len(rows) - 1
    cover.sheet_view.showGridLines = False
    cover.sheet_view.zoomScale = 85
    for col, width in {"A": 4, "B": 38, "C": 78, "D": 15}.items():
        cover.column_dimensions[col].width = width
    cover.merge_cells("B2:D2")
    cover["B2"] = title
    cover["B2"].font = Font(name="Calibri", size=22, color="FFFFFF", bold=True)
    cover["B2"].fill = PatternFill("solid", fgColor="153454")
    cover.row_dimensions[2].height = 45
    cover.merge_cells("B3:D3")
    cover["B3"] = "Критический анализ • " + datetime.now(ZoneInfo("Europe/Moscow")).strftime("%d.%m.%Y %H:%M МСК")
    cover["B3"].font = Font(name="Calibri", size=14, color="536577")
    cover.row_dimensions[3].height = 28
    for ri, (label, value) in enumerate(notes, 5):
        cover.cell(ri, 2, label)
        cover.merge_cells(start_row=ri, start_column=3, end_row=ri, end_column=4)
        cover.cell(ri, 3, value)
        for ci in (2, 3):
            c = cover.cell(ri, ci)
            c.font = Font(name="Calibri", size=12, bold=ci == 2, color="153454")
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.fill = PatternFill("solid", fgColor="F0F5FA")
        cover.row_dimensions[ri].height = max(52, len(textwrap.wrap(value, width=87)) * 17 + 14)
    nav = len(notes) + 7
    cover.cell(nav, 2, "СОДЕРЖАНИЕ").font = Font(name="Calibri", size=14, bold=True, color="153454")
    for ri, (name, count) in enumerate(counts.items(), nav + 1):
        cover.cell(ri, 2, name).hyperlink = f"#'{name}'!A1"
        cover.cell(ri, 2).style = "Hyperlink"
        cover.cell(ri, 3, f"{count} строк")
        cover.row_dimensions[ri].height = 27
    cover.freeze_panes = "C5"
    cover.print_area = f"B2:D{cover.max_row}"
    cover.page_setup.orientation = "landscape"
    cover.page_setup.paperSize = cover.PAPERSIZE_A3
    cover.page_setup.fitToWidth = 1
    cover.page_setup.fitToHeight = 1
    cover.sheet_properties.pageSetUpPr.fitToPage = True
    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    actual = load_workbook(out)
    assert actual.sheetnames == ["Об отчёте"] + list(counts)
    for name, expected_rows, _ in sections:
        ws = actual[name]
        assert ws.max_row == len(expected_rows) + 4
        assert len(ws.tables) == 1 and ws.freeze_panes == "C6"
        for ri, values in enumerate(expected_rows, 5):
            for ci, value in enumerate(values, 1):
                assert ws.cell(ri, ci).value == value
                assert ws.cell(ri, ci).data_type != "f"
    print(f"Verified {out}: {len(actual.sheetnames)} sheets; {counts}; {out.stat().st_size} bytes")