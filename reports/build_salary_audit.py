"""Build and verify the payroll settlement audit; no application imports."""
import json
from hashlib import sha256
from pathlib import Path
from zipfile import ZipFile
from openpyxl import load_workbook
from audit_workbook import build_audit
from salary_audit_content import METRICS
from salary_audit_ui import TABLES, CHARTS
from salary_audit_assessment import (
    CONCLUSIONS, QUALITY, REQUIRED, GLOSSARY, ALTERNATES,
)

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "reports/payroll_settlements_metrics_audit.xlsx"
probes = json.loads((ROOT / "reports/salary_audit_probes.json").read_text())
for key, item in probes["source_files"].items():
    path = item.get("path", key)
    assert sha256((ROOT / path).read_bytes()).hexdigest() == item["sha256"], path
assert probes["case_count"] == probes["assertions_verified"] == len(probes["cases"])
TESTS = [["ID", "Проверка", "Искусственный вход", "Фактический результат",
          "Оценка результата", "Границы и источник"]]
for i, case in enumerate(probes["cases"], 1):
    assert case["assertion_verified"]
    assert case["actual"] == case["expected_actual"]
    TESTS.append([
        f"P{i:02d}", case["title_ru"], case["input_ru"],
        json.dumps(case["actual"], ensure_ascii=False),
        case["assessment_ru"],
        case["scope_ru"] + f" R9/R10: {case['test_id']}. "
        "actual=expected; null означает Н/Д, не0. Искусственный пример, не реальные работники.",
    ])

specs = [
    ("R1", "osv-dashboard/key_indicators.py",
     "110–263;312–397;510–615;920–931;1044–1063;1104–1216;1280–1285",
     "Источник, известность, остатки, обороты, аналитики, отношение69/70, ряды и дельта"),
    ("R2", "osv-dashboard/static/key_indicators.js",
     "415–469;790–822;2589–2802;3531–3645",
     "Все KPI, четыре таблицы, подсказки, заглушки и фактическая SVG-линия"),
    ("R3", "osv-dashboard/key_indicators_service.py", "33–224",
     "Принятые файлы, слияние отчётов и суточных данных, ошибки и метка обновления"),
    ("R4", "osv-dashboard/osvparser.py", "45–54;207–313",
     "Чтение чисел, аналитики, сальдо и суточные обороты"),
    ("R5", "osv-dashboard/annual_key_indicators.py", "18–79;181–266",
     "Отдельные исторические показатели, без зарплатной карточки"),
    ("R6", "osv-dashboard/static/annual_key_indicators.js", "Годовое представление",
     "Не источник для текущей salary-модалки"),
    ("R7", "osv-dashboard/builder.py", "100–170;292–306",
     "Legacy зарплатные обороты и таблица налогов68/69"),
    ("R8", "osv-dashboard/templates/dashboard_template.html", "996–1002;1393–1405;1423–1433",
     "Legacy-график начислено / выплачено и удержано"),
    ("R9", "reports/salary_audit_probes.py", "Полный файл",
     "Изолированные синтетические проверки реального расчёта"),
    ("R10", "reports/salary_audit_probes.json", "source_files,cases",
     "Точные результаты actual/expected, русские оценки и хеши"),
]
SOURCES = [["ID", "Файл", "Строки / разделы", "Что подтверждает", "SHA-256", "Ограничение"]]
for ref, path, lines, purpose in specs:
    SOURCES.append([
        ref, path, lines, purpose, sha256((ROOT / path).read_bytes()).hexdigest(),
        "Локальная версия. Не доказывает совпадение с опубликованной версией, "
        "правильность реальных зарплат и соблюдение сроков выплаты.",
    ])

notes = [
    ("Предмет", "Критический аудит «Расчёты по оплате труда» действующего "
     "Flask-экрана «Основные показатели → ОСВ»: каждая метрика, "
     "каждый столбец таблиц, график и подсказки."),
    ("Состав", f"{len(METRICS)-1} паспортов показателей/состояний, "
     "все20 столбцов четырёх таблиц,12 недоступных управленческих показателей; "
     f"{len(TESTS)-1} контрольных примеров. Отдельно рассмотрены legacy и годовой режим."),
    ("Главный вывод", "Экран полезен для контроля бухгалтерских расчётов70/69, "
     "но не является подтверждённым отчётом о ФОТ, выплатах, просрочке, "
     "средней зарплате и тарифах взносов."),
    ("Главные риски", "Обе стороны сальдо показаны после сворачивания. "
     "Обороты не классифицированы по экономическому содержанию. "
     "График — толькоОбК70, а его подсказка наследует формулу остатка. "
     "Некоторые табличные подсказки показывают не значение выбранной ячейки."),
    ("Положительные защиты", "Точные даты сальдо; Н/Д не подменяется "
     "числовым нулём на уровне отсутствующей границы. Неизвестный/нулевой "
     "знаменатель блокирует отношение оборотов. Демонстрационные "
     "зарплаты, численность и тарифы не используются."),
    ("Необходимые данные", "Регистры начислений и составаФОТ, ведомости, "
     "сроки и распределённые платежи, удержания, кадровые данные/FTE, "
     "классификация персонала, расчёт и уплата взносов, бюджет и сопоставимыйАПГ."),
    ("Граница достоверности", "Первичка, реальные работники/зарплаты, "
     "платежи, закрытие, production-БД и опубликованная версия не проверялись. "
     "Проценты точности/полноты не назначены. Не является финансовым или юридическим заверением."),
    ("Контрольные примеры", "Изолированные входы подтверждают фактическое "
     "поведение кода. Успешныйassert не означает методологическую "
     "безупречность. Ошибка на искусственном входе не объявляется "
     "реальным финансовым инцидентом."),
    ("Единицы и формулы", "Рубли в приложении отображаются вмлн₽; "
     "формулы в книге текстовые. Остаток на дату не суммируется по месяцам "
     "и не равен потоку. ФОТ/средняя зарплата требуют согласованных состава и периода."),
    ("Что изменено", "Созданы только материалы аудита. Приложение, "
     "расчёты и данные не менялись; рекомендации не внедрены."),
    ("Навигация", "Начните с «Заключение». Далее паспорта, все столбцы "
     "и графики, качество источников и необходимые данные. R1–R10 — "
     "исходники с хешами; P — контрольные примеры. Вкладки содержат "
     "фильтры, закреплённые заголовки и настройки печати."),
]

sections = [
    ("Заключение", CONCLUSIONS, [8,16,50,76,74,38]),
    ("Метрики", METRICS, [8,43,78,66,68,66,54,68,66,38]),
    ("Таблицы", TABLES, [8,45,80,78,54,74,38]),
    ("Графики", CHARTS, [8,45,80,78,54,74,38]),
    ("Качество источников", QUALITY, [8,40,78,80,78,38]),
    ("Необходимые данные", REQUIRED, [8,40,90,76,30]),
    ("Другие экраны", ALTERNATES, [8,42,80,80,74,38]),
    ("Контрольные примеры", TESTS, [8,48,80,110,90,80]),
    ("Источники анализа", SOURCES, [8,67,48,72,70,74]),
    ("Термины и оценки", GLOSSARY, [38,85,105]),
]
build_audit(OUT, "Расчёты по оплате труда: критический аудит", notes, sections)
with ZipFile(OUT) as z:
    assert z.testzip() is None
    assert not any(n.startswith("xl/externalLinks/") for n in z.namelist())
    assert "xl/vbaProject.bin" not in z.namelist()
wb = load_workbook(OUT)
assert len(wb.sheetnames) == 11 and not wb._external_links
assert all(c.data_type != "f" for ws in wb for row in ws for c in row)
print("Verified cell values, source hashes, assertions, ZIP; no formulas, macros or external links.")