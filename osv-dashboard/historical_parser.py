# -*- coding: utf-8 -*-
"""Standalone parser for annual financial-analysis XLSX workbooks.

The parser deliberately does not depend on the dashboard or its database.  Its
output consists only of JSON-compatible dictionaries, numbers and strings.
"""
from __future__ import annotations

from datetime import date, datetime
import os
import re
import unicodedata
import zipfile
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


class HistoricalParserError(ValueError):
    """The workbook could not be read as a historical analytics workbook."""


_SPACE_RE = re.compile(r"[\s\u00a0]+")
_PREFIX_RE = re.compile(r"^\s*\d+\s*[.)\-–—:]\s*")
_YEAR_RE = re.compile(r"(?<!\d)((?:19|20)\d{2})(?!\d)")
_MISSING_RE = re.compile(r"^[\-–—−]+$")
_NUMBER_RE = re.compile(r"^[+−-]?(?:\d[\d\s\u00a0]*)(?:[.,]\d+)?$")
_UPPER_RE = re.compile(r"^\s*<\s*([+−-]?\d[\d\s\u00a0]*(?:[.,]\d+)?)\s*(%)?\s*$")

MAX_XLSX_MEMBERS = 2000
MAX_XLSX_UNCOMPRESSED = 100 * 1024 * 1024
MAX_XLSX_MEMBER_SIZE = 25 * 1024 * 1024
MAX_XLSX_COMPRESSION_RATIO = 500
MAX_XLSX_WORKSHEETS = 100
MAX_XLSX_CELLS = 500_000


def normalize_name(value):
    """Normalize a Russian heading for matching, without transliteration."""
    text = unicodedata.normalize("NFKC", str(value or "")).replace("ё", "е").lower()
    text = _PREFIX_RE.sub("", text)
    text = text.replace("–", "-").replace("—", "-").replace("−", "-")
    text = re.sub(r"[«»\"']", "", text)
    text = re.sub(r"\s*[-:]\s*", " ", text)
    text = re.sub(r"[^\w%]+", " ", text, flags=re.UNICODE)
    return _SPACE_RE.sub(" ", text).strip()


def canonical_metric_key(value):
    """Return a deterministic, readable key for a metric in any workbook."""
    key = normalize_name(value)
    # Common editorial abbreviations and explanatory parentheticals must not
    # create a second key for the same ratio on another thematic sheet.
    key = re.sub(r"\bкоэф\b", "коэффициент", key)
    key = re.sub(r"\b(общей|промежуточной)\b", "", key)
    key = re.sub(r"\bв\s+т\s+ч\b", "", key)
    key = _SPACE_RE.sub("_", key).strip("_")
    return key or "unnamed_metric"


SECTION_ALIASES = {
    "assets": ("активы",),
    "liabilities": ("пассивы",),
    "net_assets": ("чистые активы",),
    "financial_stability": ("фин устойчивость", "финансовая устойчивость"),
    "working_capital": ("сос", "собственные оборотные средства"),
    "liquidity": ("ликвидность",),
    "asset_liability": ("актив пассив",),
    "financial_results": ("фин результаты", "финансовые результаты"),
    "profitability": ("рентабельность",),
    "turnover": ("оборачиваемость",),
    "rating": ("рейтинг",),
    "insolvent_structure": ("неудовлетв структура", "неудовлетворительная структура"),
    "creditworthiness": ("кредитоспособность",),
    "bankruptcy_forecast": ("прогноз банкротства",),
    "dynamics": ("динамика",),
}
_SECTION_LOOKUP = {
    normalize_name(alias): key
    for key, aliases in SECTION_ALIASES.items()
    for alias in aliases
}

# The vocabulary is intentionally explicit.  A new row is imported, but it is
# visible in unknown_metrics until an operator confirms an alias.
_KNOWN_LABELS = {
    "assets": (
        "ВНЕОБОРОТНЫЕ АКТИВЫ", "Внеоборотные активы, всего",
        "в т.ч. основные средства", "нематериальные активы",
        "отложенные налоговые активы", "ОБОРОТНЫЕ АКТИВЫ",
        "Оборотные активы, всего", "в т.ч. запасы",
        "НДС по приобретенным ценностям", "дебиторская задолженность",
        "денежные средства и краткосрочные фин. вложения", "ИТОГО",
        "ВАЛЮТА БАЛАНСА",
    ),
    "liabilities": (
        "Собственный капитал", "в т.ч. уставный капитал",
        "добавочный капитал (без дооценки)", "нераспределенная прибыль",
        "Долгосрочные обязательства, всего",
        "в т.ч. заемные средства (долгосрочные)",
        "отложенные налоговые обязательства",
        "Краткосрочные обязательства, всего",
        "в т.ч. заемные средства (краткосрочные)",
        "кредиторская задолженность", "ВАЛЮТА БАЛАНСА",
    ),
    "net_assets": (
        "Чистые активы", "Уставный капитал",
        "Превышение чистых активов над УК (разница)",
        "Превышение ЧА над УК (раз, разов)",
    ),
    "financial_stability": (
        "Коэффициент автономии", "Коэффициент финансового левериджа",
        "Коэффициент обеспеченности СОС", "Индекс постоянного актива",
        "Коэффициент покрытия инвестиций",
        "Коэффициент маневренности собственного капитала",
        "Коэффициент мобильности имущества",
        "Коэффициент мобильности оборотных средств",
        "Коэффициент обеспеченности запасов",
        "Коэффициент краткосрочной задолженности",
    ),
    "working_capital": (
        "СОС1 (без учета долгосрочных и краткосрочных пассивов)",
        "СОС2 (с учетом долгосрочных пассивов, NWC)",
        "СОС3 (с учетом долгосрочных и краткосрочных займов)",
        "Излишек/недостаток СОС1 (к запасам)",
        "Излишек/недостаток СОС2 (к запасам)",
        "Излишек/недостаток СОС3 (к запасам)",
    ),
    "liquidity": (
        "Коэффициент текущей (общей) ликвидности",
        "Коэффициент быстрой (промежуточной) ликвидности",
        "Коэффициент абсолютной ликвидности",
    ),
    "asset_liability": (
        "А1. Высоколиквидные активы", "А2. Быстрореализуемые (дебиторка)",
        "А3. Медленно реализуемые", "А4. Труднореализуемые (внеоборотные)",
        "П1. Наиболее срочные обязательства", "П2. Среднесрочные обязательства",
        "П3. Долгосрочные обязательства",
        "П4. Постоянные пассивы (собств. капитал)",
        "Излишек/недостаток (А1 – П1)", "Излишек/недостаток (А2 – П2)",
        "Излишек/недостаток (А3 – П3)", "Излишек/недостаток (А4 – П4)",
    ),
    "financial_results": (
        "Выручка", "Расходы по обычным видам деятельности",
        "Себестоимость продаж", "Валовая прибыль (убыток)",
        "Управленческие расходы", "Прибыль (убыток) от продаж",
        "Прочие доходы и расходы (кроме %)", "Проценты к получению",
        "EBIT (прибыль до % и налогов)", "Проценты к уплате",
        "Налог на прибыль и прочее", "Прибыль до налогообложения",
        "Чистая прибыль (убыток)",
        "Изменение нераспределенной прибыли (стр. 1370)",
    ),
    "profitability": (
        "Рентабельность продаж (прибыль от продаж / выручка)",
        "Рентабельность продаж по EBIT",
        "Рентабельность продаж по чистой прибыли",
        "Прибыль от продаж на рубль вложенного капитала",
        "Коэффициент покрытия процентов (ICR)",
        "Рентабельность собственного капитала (ROE)",
        "Рентабельность активов (ROA)",
        "Прибыль на задействованный капитал (ROCE)",
        "Рентабельность производственных фондов", "Фондоотдача (коэф.)",
    ),
    "turnover": (
        "Оборачиваемость оборотных средств", "Оборачиваемость запасов",
        "Оборачиваемость дебиторской задолженности",
        "Оборачиваемость кредиторской задолженности",
        "Оборачиваемость активов", "Оборачиваемость собственного капитала",
    ),
    "rating": (
        "Балл финансового положения", "Качественная оценка фин. положения",
        "Балл результатов деятельности", "Качественная оценка результатов",
        "Итоговая рейтинговая оценка (балл)", "ИТОГОВЫЙ РЕЙТИНГ",
        "Оценка фин. положения (I блок)", "Оценка результатов (II блок)",
    ),
    "insolvent_structure": (
        "Коэффициент текущей ликвидности",
        "Коэффициент обеспеченности собственными средствами",
        "Коэффициент восстановления платежеспособности",
    ),
    "creditworthiness": (
        "Коэф. абсолютной ликвидности", "Коэф. быстрой ликвидности",
        "Коэф. текущей ликвидности", "Коэф. наличия собственных средств",
        "Рентабельность продукции", "Рентабельность деятельности",
        "ИТОГО сумма баллов", "Класс заемщика",
    ),
    "bankruptcy_forecast": (
        "Z-счет Альтмана (4-факторная модель)",
        "T1 — оборотный капитал / активы",
        "T2 — нераспределенная прибыль / активы", "T3 — EBIT / активы",
        "T4 — собственный капитал / заемный", "Z-счет Альтмана",
        "Вероятность банкротства (Альтман)", "Z-счет Таффлера",
        "X1 — прибыль до налогообложения / краткосрочные обязательства",
        "X2 — оборотные активы / обязательства",
        "X3 — краткосрочные обязательства / активы",
        "X4 — выручка / активы", "Вероятность банкротства (Таффлер)",
        "Модель Сайфуллина–Кадыкова (R)",
        "К1 — обеспеченность собственными средствами",
        "К2 — текущая ликвидность", "К3 — оборачиваемость активов",
        "К4 — коммерческая маржа",
        "К5 — рентабельность собственного капитала", "Итого (R)",
        "Вероятность банкротства (Сайфуллин–Кадыков)",
    ),
    "dynamics": (
        "Изменение выручки, тыс. руб.", "Изменение выручки, %",
        "Изменение расходов по обычным видам, тыс. руб.",
        "Изменение расходов по обычным видам, %",
        "Изменение прибыли от продаж, тыс. руб.",
        "Изменение EBIT, тыс. руб.",
        "Изменение процентов к уплате, тыс. руб.",
        "Изменение чистой прибыли, тыс. руб.",
    ),
}
_KNOWN_KEYS = {
    section: {canonical_metric_key(label) for label in labels}
    for section, labels in _KNOWN_LABELS.items()
}
_MONEY_SECTIONS = {
    "assets", "liabilities", "net_assets", "working_capital",
    "asset_liability", "financial_results",
}
_EBIT_LABEL = "EBIT (прибыль до % и налогов)"


def _raw(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")
    return str(value)


def _validate_xlsx_archive(path):
    """Reject XLSX ZIP bombs and pathologically large workbooks before loading."""
    try:
        with zipfile.ZipFile(path, "r") as archive:
            members = archive.infolist()
            if len(members) > MAX_XLSX_MEMBERS:
                raise HistoricalParserError(
                    f"XLSX contains too many archive members ({len(members)})")
            total_size = sum(item.file_size for item in members)
            if total_size > MAX_XLSX_UNCOMPRESSED:
                raise HistoricalParserError(
                    "XLSX decompressed content exceeds the 100 MB limit")
            worksheets = [
                item for item in members
                if item.filename.startswith("xl/worksheets/")
                and item.filename.endswith(".xml")
            ]
            if len(worksheets) > MAX_XLSX_WORKSHEETS:
                raise HistoricalParserError(
                    f"XLSX contains too many worksheets ({len(worksheets)})")
            cells = 0
            for item in members:
                if item.file_size > MAX_XLSX_MEMBER_SIZE:
                    raise HistoricalParserError(
                        f"XLSX member is too large: {item.filename}")
                if (item.file_size > 1024
                        and item.file_size / max(item.compress_size, 1)
                        > MAX_XLSX_COMPRESSION_RATIO):
                    raise HistoricalParserError(
                        f"XLSX compression ratio is unsafe: {item.filename}")
            for item in worksheets:
                tail = b""
                with archive.open(item) as worksheet:
                    while True:
                        chunk = worksheet.read(1024 * 1024)
                        if not chunk:
                            break
                        sample = tail + chunk
                        cells += sample.count(b"<c ")
                        tail = sample[-3:]
                        if cells > MAX_XLSX_CELLS:
                            raise HistoricalParserError(
                                "XLSX contains too many populated cells")
    except HistoricalParserError:
        raise
    except (BadZipFile, OSError) as exc:
        raise HistoricalParserError(
            f"Cannot read XLSX workbook {os.path.basename(str(path))}: {exc}"
        ) from exc


def _year(value):
    if isinstance(value, (date, datetime)):
        return value.year
    if isinstance(value, int) and 1900 <= value <= 2100:
        return value
    if isinstance(value, float) and value.is_integer() and 1900 <= value <= 2100:
        return int(value)
    found = _YEAR_RE.findall(_raw(value))
    return int(found[-1]) if found else None


def _aliases(confirmed_aliases):
    source = confirmed_aliases or {}
    if not isinstance(source, dict):
        raise HistoricalParserError("confirmed_aliases must be a mapping")
    sections = source.get("sections", {})
    metrics = source.get("metrics", {})
    if not sections and not metrics:
        sections = source
    section_map = {normalize_name(k): str(v) for k, v in sections.items()}
    metric_map = {}
    for key, value in metrics.items():
        if isinstance(value, dict):
            metric_map[str(key)] = {
                normalize_name(alias): str(target) for alias, target in value.items()
            }
        else:
            metric_map.setdefault("*", {})[normalize_name(key)] = str(value)
    return section_map, metric_map


def _section_key(title, section_aliases):
    norm = normalize_name(title)
    if norm in section_aliases:
        return section_aliases[norm], True
    if norm in _SECTION_LOOKUP:
        return _SECTION_LOOKUP[norm], False
    return "unknown_" + (canonical_metric_key(title) or "section"), False


def _header(ws):
    """Find the best row containing annual columns."""
    best = None
    for row in ws.iter_rows():
        years = [(cell.column, _year(cell.value), _raw(cell.value))
                 for cell in row if _year(cell.value) is not None]
        if years and (best is None or len(years) > len(best[1])):
            best = (row[0].row, years)
    if best is None:
        return None, []
    row_number, annual = best
    # A year often spans a value column followed by "% к валюте".  There are
    # no merged cells in many exports, so inherit the preceding year for
    # populated header cells until the next explicit year.  Columns before the
    # first year (Norm, Weight, Category) remain support metadata.
    explicit = {column: (year, period) for column, year, period in annual}
    expanded = []
    inherited_year = None
    for column in range(min(explicit), ws.max_column + 1):
        if column in explicit:
            inherited_year, period = explicit[column]
            expanded.append((column, inherited_year, period))
            continue
        header = _raw(ws.cell(row_number, column).value).strip()
        if inherited_year is not None and header:
            expanded.append((column, inherited_year, header))
    return row_number, expanded


def _unit(label, header, section_key, raw_value):
    combined = (normalize_name(label) + " " + normalize_name(header)).strip()
    raw_text = _raw(raw_value).strip()
    # The confirmed EBIT label itself contains a percent sign as shorthand for
    # "проценты".  It is a money row, not a percentage series.  Explicit
    # percent evidence in the value or annual-column header still wins; this
    # narrow exception must not relax unit validation for any other label.
    if "%" in raw_text or "%" in normalize_name(header):
        return "percent", 1.0
    if (
        section_key == "financial_results"
        and normalize_name(label) == normalize_name(_EBIT_LABEL)
    ):
        # Ignore only the abbreviation in the label, not explicit money units
        # in the header (millions/rubles must retain their scaling).
        combined = normalize_name(header)
    if "%" in combined:
        return "percent", 1.0
    if re.search(r"\bмлн(?:\.| )*\s*руб", combined):
        return "thousand_rubles", 1000.0
    if re.search(r"\bтыс(?:\.| )*\s*руб", combined):
        return "thousand_rubles", 1.0
    if re.search(r"\bруб", combined):
        return "thousand_rubles", 0.001
    if section_key in _MONEY_SECTIONS:
        return "thousand_rubles", 1.0
    return "ratio", 1.0


def _value(value, unit, factor, year, period, coordinate):
    raw = _raw(value)
    stripped = raw.strip()
    result = {
        "year": year, "period": period, "raw": raw, "value": None,
        "unit": unit, "is_missing": False, "is_upper_bound": False,
        "cell": coordinate,
    }
    if not stripped or _MISSING_RE.fullmatch(stripped):
        result["is_missing"] = True
        return result
    upper = _UPPER_RE.fullmatch(stripped)
    numeric = upper.group(1) if upper else stripped.rstrip("%").strip()
    if upper:
        result["is_upper_bound"] = True
        if upper.group(2):
            result["unit"] = "percent"
    normalized = numeric.replace("\u00a0", "").replace(" ", "")
    normalized = normalized.replace("−", "-").replace(",", ".")
    if _NUMBER_RE.fullmatch(numeric):
        try:
            result["value"] = float(normalized) * factor
            return result
        except ValueError:
            pass
    result["value"] = stripped
    result["unit"] = "text"
    return result


def _metric_alias(section_key, label, metric_aliases):
    normalized = normalize_name(label)
    mapping = metric_aliases.get(section_key, {})
    target = mapping.get(normalized)
    if target is None:
        target = metric_aliases.get("*", {}).get(normalized)
    return target


def _parse_sheet(ws, key, title, metric_aliases):
    header_row, year_columns = _header(ws)
    section = {
        "key": key, "title": title, "raw_title": title,
        "sheet_name": ws.title, "is_known": not key.startswith("unknown_"),
        "rows": [],
    }
    if header_row is None:
        return section
    for row_number in range(header_row + 1, ws.max_row + 1):
        label_value = ws.cell(row_number, 1).value
        label = _raw(label_value).strip()
        if not label:
            continue
        vals = []
        for column, year, period in year_columns:
            cell = ws.cell(row_number, column)
            unit, factor = _unit(label, period, key, cell.value)
            vals.append(_value(
                cell.value, unit, factor, year, period,
                f"{get_column_letter(column)}{row_number}",
            ))
        alias = _metric_alias(key, label, metric_aliases)
        metric_key = alias or canonical_metric_key(label)
        known = metric_key in _KNOWN_KEYS.get(key, set()) or alias is not None
        has_content = any(not value["is_missing"] for value in vals)
        kind = "metric" if has_content else "group"
        units = sorted({value["unit"] for value in vals if not value["is_missing"]})
        section["rows"].append({
            "key": metric_key, "label": label, "raw_label": _raw(label_value),
            "kind": kind, "is_known": known, "alias_used": alias,
            "unit": units[0] if len(units) == 1 else ("mixed" if units else None),
            "values": vals,
        })
    return section


def _find_metric(section, *needles):
    if not section:
        return None
    normalized_needles = [normalize_name(n) for n in needles]
    for row in section["rows"]:
        name = normalize_name(row["label"])
        if any(needle in name for needle in normalized_needles):
            return row
    return None


def _numeric_years(row):
    if not row:
        return {}
    return {
        value["year"]: value["value"] for value in row["values"]
        if isinstance(value["value"], (int, float)) and value["year"] is not None
    }


def _has_numeric_for_year(row, year):
    return any(
        value["year"] == year and isinstance(value["value"], (int, float))
        for value in (row or {}).get("values", [])
    )


def _validate(sections, warnings):
    by_key = {section["key"]: section for section in sections}
    years = sorted({
        value["year"]
        for section in sections
        for row in section["rows"]
        for value in row["values"]
        if value["year"] is not None
    })
    checks = []
    assets = by_key.get("assets")
    liabilities = by_key.get("liabilities")
    details = []
    av = _numeric_years(_find_metric(assets, "валюта баланса"))
    pv = _numeric_years(_find_metric(liabilities, "валюта баланса"))
    balance_ok = bool(years)
    for year in years:
        if year not in av or year not in pv:
            balance_ok = False
            details.append({
                "year": year, "assets": av.get(year),
                "liabilities": pv.get(year), "reason": "missing_numeric_value",
            })
        elif abs(av[year] - pv[year]) > 0.01:
            balance_ok = False
            details.append({
                "year": year, "assets": av[year], "liabilities": pv[year],
                "reason": "balance_mismatch",
            })
    checks.append({"group": "balance", "ok": balance_ok, "details": details})
    financial = by_key.get("financial_results")
    revenue = _find_metric(financial, "выручка")
    net_profit = _find_metric(financial, "чистая прибыль")
    financial_missing = [
        year for year in years
        if not (_has_numeric_for_year(revenue, year)
                and _has_numeric_for_year(net_profit, year))
    ]
    financial_ok = bool(years) and not financial_missing
    checks.append({
        "group": "financial_results", "ok": financial_ok,
        "details": [{"year": year, "reason": "missing_numeric_value"}
                    for year in financial_missing],
    })
    liquidity = by_key.get("liquidity")
    liquidity_rows = [
        _find_metric(liquidity, "текущей"),
        _find_metric(liquidity, "быстрой"),
        _find_metric(liquidity, "абсолютной"),
    ]
    liquidity_missing = [
        year for year in years
        if not all(_has_numeric_for_year(row, year) for row in liquidity_rows)
    ]
    liquidity_ok = bool(years) and not liquidity_missing
    checks.append({
        "group": "liquidity", "ok": liquidity_ok,
        "details": [{"year": year, "reason": "missing_numeric_value"}
                    for year in liquidity_missing],
    })
    for check in checks:
        if not check["ok"]:
            warnings.append("Required group validation failed: " + check["group"])
    return {"ok": all(check["ok"] for check in checks), "checks": checks}


def _reconcile(summary_ws, sections, section_aliases, metric_aliases, warnings):
    """Compare summary blocks with thematic sheets, without emitting them."""
    by_key = {section["key"]: section for section in sections}
    compared = conflicts = 0
    conflict_rows = []
    current_key = None
    header_row = None
    year_columns = []
    for row_number in range(1, summary_ws.max_row + 1):
        first = _raw(summary_ws.cell(row_number, 1).value).strip()
        possible_key, _ = _section_key(first, section_aliases)
        if first and possible_key in by_key:
            current_key = possible_key
            header_row = None
            year_columns = []
            continue
        years = [
            (cell.column, _year(cell.value), _raw(cell.value))
            for cell in summary_ws[row_number] if _year(cell.value) is not None
        ]
        if current_key and years:
            header_row, year_columns = row_number, years
            continue
        if not current_key or not year_columns or row_number <= (header_row or 0) or not first:
            continue
        alias = _metric_alias(current_key, first, metric_aliases)
        metric_key = alias or canonical_metric_key(first)
        target = next(
            (row for row in by_key[current_key]["rows"] if row["key"] == metric_key),
            None,
        )
        if target is None:
            continue
        thematic = {(v["year"], v["unit"]): v for v in target["values"]}
        for column, year, period in year_columns:
            cell = summary_ws.cell(row_number, column)
            unit, factor = _unit(first, period, current_key, cell.value)
            source = _value(cell.value, unit, factor, year, period, cell.coordinate)
            other = thematic.get((year, source["unit"]))
            if other is None or (source["is_missing"] and other["is_missing"]):
                continue
            compared += 1
            if source["value"] != other["value"] or source["is_upper_bound"] != other["is_upper_bound"]:
                conflicts += 1
                item = {
                    "section": current_key, "metric": metric_key, "year": year,
                    "summary_raw": source["raw"], "thematic_raw": other["raw"],
                }
                conflict_rows.append(item)
                warnings.append(
                    "Summary conflict: {section}/{metric}/{year}".format(**item)
                )
    return {
        "sheet_name": summary_ws.title, "emitted_values": 0,
        "compared_values": compared, "conflict_count": conflicts,
        "conflicts": conflict_rows,
    }


def parse_historical_xlsx(path, confirmed_aliases=None):
    """Parse an annual historical analytics workbook.

    Args:
        path: filesystem path to an ``.xlsx`` file.
        confirmed_aliases: optional mapping with ``sections`` and ``metrics``.
            Metric aliases may be global or grouped by canonical section key.

    Returns:
        A JSON-compatible dictionary containing ``sections``, flattened
        ``values``, ``preview``, validation, reconciliation and warnings.
    """
    if not str(path).lower().endswith(".xlsx"):
        raise HistoricalParserError("Historical analytics parser requires an .xlsx file")
    section_aliases, metric_aliases = _aliases(confirmed_aliases)
    _validate_xlsx_archive(path)
    try:
        workbook = load_workbook(path, read_only=False, data_only=True)
    except (BadZipFile, OSError, KeyError, ValueError) as exc:
        raise HistoricalParserError(
            f"Cannot read XLSX workbook {os.path.basename(str(path))}: {exc}"
        ) from exc
    except Exception as exc:
        raise HistoricalParserError(
            f"Corrupt or unsupported XLSX workbook {os.path.basename(str(path))}: {exc}"
        ) from exc

    warnings = []
    sections = []
    summary_ws = None
    aliases_applied = []
    unknown_sections = []
    unknown_metrics = []
    for ws in workbook.worksheets:
        title = _raw(ws.title)
        if normalize_name(title) == normalize_name("Сводная таблица"):
            summary_ws = ws
            continue
        key, section_alias_used = _section_key(title, section_aliases)
        section = _parse_sheet(ws, key, title, metric_aliases)
        section["alias_used"] = key if section_alias_used else None
        sections.append(section)
        if section_alias_used:
            aliases_applied.append({
                "type": "section", "raw": title, "canonical": key,
            })
        if not section["is_known"]:
            unknown_sections.append({"raw": title, "normalized": normalize_name(title)})
            warnings.append("Unknown section: " + title)
        for row in section["rows"]:
            if row["alias_used"]:
                aliases_applied.append({
                    "type": "metric", "section": key, "raw": row["raw_label"],
                    "canonical": row["key"],
                })
            if not row["is_known"]:
                item = {
                    "section": key, "raw": row["raw_label"],
                    "suggested_key": canonical_metric_key(row["raw_label"]),
                }
                unknown_metrics.append(item)
                warnings.append(f"Unknown metric in {key}: {row['raw_label']}")

    validation = _validate(sections, warnings)
    if summary_ws is None:
        reconciliation = {
            "sheet_name": None, "emitted_values": 0, "compared_values": 0,
            "conflict_count": 0, "conflicts": [],
        }
        warnings.append("Reconciliation sheet 'Сводная таблица' was not found")
    else:
        reconciliation = _reconcile(
            summary_ws, sections, section_aliases, metric_aliases, warnings
        )

    values = []
    for section in sections:
        for row in section["rows"]:
            for value in row["values"]:
                values.append({
                    "section_key": section["key"], "metric_key": row["key"],
                    "metric_label": row["raw_label"], **value,
                })
    years = sorted({value["year"] for value in values if value["year"] is not None})
    preview = {
        "section_count": len(sections),
        "metric_count": sum(
            1 for section in sections for row in section["rows"]
            if row["kind"] == "metric"
        ),
        "value_count": len(values),
        "non_missing_value_count": sum(not value["is_missing"] for value in values),
        "years": years,
        "unknown_section_count": len(unknown_sections),
        "unknown_metric_count": len(unknown_metrics),
        "warning_count": len(warnings),
    }
    return {
        "format": "annual_historical_analytics_xlsx",
        "source_file": os.path.basename(str(path)),
        "sections": sections, "values": values, "preview": preview,
        "unknown_sections": unknown_sections, "unknown_metrics": unknown_metrics,
        "aliases_applied": aliases_applied, "validation": validation,
        "reconciliation": reconciliation, "warnings": warnings,
    }


# Short, convenient public alias.
parse = parse_historical_xlsx


__all__ = [
    "HistoricalParserError", "SECTION_ALIASES", "canonical_metric_key",
    "normalize_name", "parse", "parse_historical_xlsx",
]