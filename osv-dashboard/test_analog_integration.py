# -*- coding: utf-8 -*-
"""Route-level regressions for the analogue module."""
import os
import pathlib
import tempfile
import unittest
import json
import hashlib
import re
import shutil
import zipfile
from types import SimpleNamespace
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from datetime import date
from decimal import Decimal
from unittest.mock import patch
from xml.etree import ElementTree as ET
from openpyxl import load_workbook

_tmp = tempfile.mkdtemp(prefix='analog_route_test_')
os.environ['DATABASE_URL'] = f'sqlite:///{_tmp}/analog.db'
os.environ['INITIAL_ADMIN_PASSWORD'] = 'analog-admin-passphrase-2026'
os.environ.setdefault('ENCRYPTION_KEY',
                      'MDAwMDAwMDAwMDAwMDAwMDAwMDAwMDAwMDAwMDAwMDA=')
os.environ.setdefault('REPLIT_DOMAINS', 'preview.example')

from app import app, init_db  # noqa: E402
from database import (AnalogClimateReference, AnalogCostIndexReference,
                       AnalogMonthlyIndex,
                        AnalogCostRow, AnalogDocument, AnalogImport,
                        AnalogDiagnostic, AnalogSource,
                       AnalogSSRChapter, AnalogSSRItem,
                       AnalogNetworkAttribute,
                      AnalogPassport, AnalogPassportVersion, User, db)  # noqa: E402
from analog_migration import import_json_v4  # noqa: E402
from analog_parser import _tei_rows  # noqa: E402
from analog_monthly_indices import indexed_cost, price_period, seed_monthly_indices  # noqa: E402
from analog_cost_export import build_cost_workbook, estimate_values  # noqa: E402
from analog_routes import (_backfill_analog_total_costs,
                           _backfill_inferred_object_types, _format_ru_number,
                           _result_row, _table_321_total)  # noqa: E402
import analog_routes  # noqa: E402


ROOT = pathlib.Path(__file__).parent.parent / 'attached_assets'


class AnalogRoutesTest(unittest.TestCase):
    def test_catalog_omits_manual_creation_and_system_import(self):
        response = self.client.get('/analog-search/simple?view=base')
        self.assertEqual(response.status_code, 200)
        html = response.get_data(as_text=True)
        for text in ('Ручное создание и системный импорт',
                     'Создать пустой паспорт', 'Импорт JSON v4',
                     'create-object-type', 'create-network-fields',
                     'Импортировать JSON'):
            self.assertNotIn(text, html)
        for text in ('Добавить объект', 'Все объекты', 'catalog-tbody',
                     'Поиск в каталоге по адресу', 'analog-import-form'):
            self.assertIn(text, html)

    def test_similarity_column_only_after_search_filter(self):
        path = '/analog-search/simple'
        for params in ({}, {'view': 'search', 'page': '2'},
                       {'view': 'search', 'capacity_unit': 'mw',
                        'tolerance_percent': '20'}):
            page = self.client.get(path, query_string=params)
            self.assertEqual(page.status_code, 200)
            html = page.get_data(as_text=True)
            self.assertNotIn('<th>Аналогичность</th>', html)
            self.assertNotIn('class="similarity"', html)
        for params in ({'q': 'несуществующий объект'},
                       {'object_type': 'котельная'},
                       {'work_type': 'строительство'},
                       {'climate_zone': 'IIВ'},
                       {'climate_region': 'II'},
                       {'capacity_value': '1', 'capacity_unit': 'mw'},
                       {'title': 'несуществующий объект'},
                       {'address': 'несуществующий адрес'},
                       {'q': 'несуществующий объект', 'page': '2'}):
            page = self.client.get(path, query_string=params)
            self.assertEqual(page.status_code, 200)
            self.assertIn('<th>Аналогичность</th>',
                          page.get_data(as_text=True))
        catalog = self.client.get(path, query_string={
            'view': 'base', 'q': 'котельная'}).get_data(as_text=True)
        self.assertNotIn('<th>Аналогичность</th>', catalog)

    def test_cost_workbook_per_object_and_missing_breakdown(self):
        chapter = SimpleNamespace(number=8, building=Decimal('27381.19163'),
                                  mounting=Decimal('4161.972214'),
                                  equipment=Decimal('10135.59086'))
        row = lambda ch, title, total: SimpleNamespace(
            chapter_number=ch, name=title, total=Decimal(total))
        items = [
            row(9, 'Пусконаладочные работы', '988.88576'),
            row(12, 'Разработка проектной документации', '3751.07866'),
            row(12, 'Инженерные изыскания', '1347.15036'),
            row(12, 'Экспертиза проекта', '401.03'),
            row(12, 'Авторский надзор', '25'),
        ]
        values = estimate_values([chapter], items)
        self.assertEqual(values[11][0], Decimal('32532049.604000'))
        self.assertEqual(values[15][0], Decimal('10135590.86000'))
        self.assertEqual(values[16][0], Decimal('5499259.02000'))
        self.assertEqual(values[17][0], Decimal('3751078.66000'))
        self.assertEqual(values[18][0], Decimal('1347150.36000'))
        self.assertEqual(values[19][0], Decimal('401030.00'))
        self.assertIsNone(estimate_values([], [])[11][0])
        self.assertIsNone(estimate_values([chapter], [])[11][0])
        self.assertIsNone(estimate_values(
            [chapter], [row(8, 'Утилизация отходов', '1')])[12][0])
        self.assertIsNone(estimate_values([], [
            row(12, 'Проектно-изыскательские работы', '100'),
            row(12, 'Разработка проектной документации', '70')])[16][0])
        empty_book = load_workbook(build_cost_workbook(
            SimpleNamespace(title='Другой аналог'), None, None, [], []),
            data_only=True)
        self.assertEqual(empty_book.worksheets[0]['D9'].value, 'Другой аналог')
        self.assertIsNone(empty_book.worksheets[0]['D11'].value)
        self.assertIsNone(empty_book.worksheets[0]['D19'].value)
        rounded_book = load_workbook(build_cost_workbook(
            SimpleNamespace(title='БМК'), None,
            Decimal('1.80567497850387'), [], []), data_only=True)
        self.assertEqual(rounded_book.worksheets[0]['D10'].value, 1.81)
        self.assertEqual(rounded_book.worksheets[0]['D10'].number_format, '0.00')

        with app.app_context():
            passport = AnalogPassport(
                title='Объект для выгрузки Excel', capacity_kw=Decimal('1163'),
                created_by=self.admin_id)
            db.session.add(passport)
            db.session.flush()
            version = AnalogPassportVersion(
                passport_id=passport.id, version_no=1,
                conclusion_number='export-test-2026',
                approval_date=date(2026, 5, 19),
                examination_result='positive', source_hash='export-source-test',
                ssr_hash='export-ssr-test', imported_by=self.admin_id)
            db.session.add(version)
            db.session.flush()
            passport.current_version_id = version.id
            db.session.add(AnalogSSRChapter(
                version_id=version.id, number=8, building=chapter.building,
                mounting=chapter.mounting, equipment=chapter.equipment))
            for item in items:
                db.session.add(AnalogSSRItem(
                    version_id=version.id, chapter_number=item.chapter_number,
                    name=item.name, total=item.total))
            db.session.commit()
            passport_id = passport.id
        try:
            for view in ('search', 'base'):
                listing = self.client.get(
                    '/analog-search/simple', query_string={
                        'view': view, 'q': 'Объект для выгрузки Excel'})
                self.assertEqual(listing.status_code, 200)
                self.assertIn(
                    f'/analog-search/passports/{passport_id}/estimate.xlsx',
                    listing.get_data(as_text=True))
                self.assertIn('M12 10v7m-3-3 3 3 3-3',
                              listing.get_data(as_text=True))
            response = self.client.get(
                f'/analog-search/passports/{passport_id}/estimate.xlsx')
            self.assertEqual(response.status_code, 200)
            self.assertEqual(
                response.headers['Content-Type'],
                'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
            self.assertIn(f'analog-estimate-{passport_id}.xlsx',
                          response.headers['Content-Disposition'])
            workbook = load_workbook(BytesIO(response.data), data_only=True)
            self.assertEqual(workbook.sheetnames[:3],
                             ['НМЦК СМР', 'индекс факт.инфл.',
                              'расчет индекса прогнозной инф'])
            sheet = workbook.worksheets[0]
            self.assertEqual(sheet['D9'].value, 'Объект для выгрузки Excel')
            self.assertEqual(sheet['D10'].value, 1)
            self.assertEqual(sheet['D10'].number_format, '0.00')
            self.assertEqual(sheet['D11'].value, 32532049.604)
            self.assertEqual(sheet['D16'].value, 5499259.02)
            self.assertEqual(sheet['D17'].value, 3751078.66)
            self.assertEqual(sheet['E18'].value, 1347150.36)
            self.assertEqual(sheet['D19'].value, 401030)
            self.assertIn('не подтверждено', sheet['E15'].comment.text)
            self.assertIn('главе 9',
                          workbook['Источники сумм']['D2'].value)
            self.assertIn('export-test-2026', sheet['A6'].value)
            with app.app_context():
                db.session.get(AnalogPassport, passport_id).archived = True
                db.session.commit()
            self.assertEqual(self.client.get(
                f'/analog-search/passports/{passport_id}/estimate.xlsx').status_code,
                404)
        finally:
            with app.app_context():
                db.session.delete(db.session.get(AnalogPassport, passport_id))
                db.session.commit()

    def test_monthly_index_register_and_conversion(self):
        with app.app_context():
            seed_monthly_indices()
            self.assertEqual(
                AnalogMonthlyIndex.query.filter_by(year=2026, month=7).one().percent,
                Decimal('101.670'))
            self.assertIsNone(
                AnalogMonthlyIndex.query.filter_by(year=2026, month=8).first())
            months = {(row.year, row.month): row.percent
                      for row in AnalogMonthlyIndex.query.all()}
            self.assertEqual(max(months), (2026, 7))
            version = SimpleNamespace(
                total_cost=Decimal('1000'),
                price_year=2026,
                payload_json=json.dumps({'ssr': {
                    'price_year': 2026, 'price_month': 5,
                    'approved_date': '2026-06-25'}}))
            cost, factor, start = indexed_cost(version, months, (2026, 7))
            expected = (Decimal('100.13') / 100 *
                        Decimal('100.88') / 100 *
                        Decimal('101.67') / 100)
            self.assertEqual(start, (2026, 5))
            self.assertEqual(factor, expected)
            self.assertEqual(cost, Decimal('1000') * expected)
            self.assertIsNone(indexed_cost(version, {
                (2026, 5): Decimal('100.13'),
                (2026, 7): Decimal('101.67')}, (2026, 7)))
            self.assertIsNone(indexed_cost(version, months, (2026, 4)))
            quarter_only = SimpleNamespace(
                price_year=2025, total_cost=100,
                payload_json=json.dumps({'ssr': {
                    'price_year': 2025, 'price_quarter': 2,
                    'approved_date': '2025-05-01'}}))
            self.assertEqual(price_period(quarter_only), (2025, 5))
            self.assertIsNone(price_period(SimpleNamespace(
                price_year=2025, payload_json=json.dumps({'ssr': {
                    'price_year': 2025, 'price_quarter': 2,
                    'approved_date': '2026-05-01'}}))))

        response = self.client.get('/analog-search/monthly-indices')
        self.assertEqual(response.status_code, 200)
        html = response.get_data(as_text=True)
        self.assertIn('Последний загруженный индекс:', html)
        self.assertIn('07.2026', html)
        self.assertIn('101.67', html)
        with app.app_context():
            passport = AnalogPassport(
                title='Помесячный тестовый объект', object_type='котельная',
                work_type='строительство', created_by=self.admin_id)
            db.session.add(passport)
            db.session.flush()
            version = AnalogPassportVersion(
                passport_id=passport.id, version_no=1, conclusion_number='monthly-test',
                examination_result='positive', source_hash='monthly-source',
                ssr_hash='monthly-ssr', total_cost=1000,
                price_year=2026, price_quarter=2, imported_by=self.admin_id,
                payload_json=json.dumps({'ssr': {
                    'price_year': 2026, 'price_quarter': 2,
                    'approved_date': '2026-05-01'}}))
            db.session.add(version)
            db.session.flush()
            passport.current_version_id = version.id
            db.session.commit()
            passport_id, version_id = passport.id, version.id
        listing = self.client.get('/analog-search/simple')
        self.assertEqual(listing.status_code, 200)
        page = listing.get_data(as_text=True)
        self.assertIn('Индексы', page)
        self.assertIn('Реестр индексов ↗', page)
        self.assertIn('Стоимость (текущие цены), тыс. руб', page)
        self.assertIn('Помесячный тестовый объект', page)
        self.assertIn('1 000,0', page)
        self.assertIn('Уровень цен: 05.2026; коэффициент:', page)
        catalog = self.client.get('/analog-search/simple?view=base')
        self.assertIn('Стоимость (текущие цены), тыс. руб',
                      catalog.get_data(as_text=True))
        with app.app_context():
            db.session.delete(db.session.get(AnalogPassportVersion, version_id))
            db.session.delete(db.session.get(AnalogPassport, passport_id))
            db.session.commit()
        invalid = self.client.post('/analog-search/monthly-indices', data={
            'year': '2026', 'month': '8', 'percent': 'nan'})
        self.assertEqual(invalid.status_code, 400)
        saved = self.client.post('/analog-search/monthly-indices', data={
            'year': '2026', 'month': '8', 'percent': '100,88'})
        self.assertEqual(saved.status_code, 302)
        try:
            with app.app_context():
                seed_monthly_indices()
                row = AnalogMonthlyIndex.query.filter_by(year=2026, month=8).one()
                self.assertEqual(row.percent, Decimal('100.880'))
            self.assertIn('08.2026', self.client.get(
                '/analog-search/monthly-indices').get_data(as_text=True))
        finally:
            with app.app_context():
                AnalogMonthlyIndex.query.filter_by(year=2026, month=8).delete()
                db.session.commit()

    def test_add_monthly_indices_year_atomically_and_restrict_editing(self):
        path = '/analog-search/monthly-indices'
        listing = self.client.get(path)
        self.assertEqual(listing.status_code, 200)
        html = listing.get_data(as_text=True)
        self.assertIn('action="/analog-search/monthly-indices/year"', html)
        self.assertIn('name="percent_1"', html)
        self.assertNotIn('id="year-2098"', html)
        self.assertIn('name="year" min="2017" max="2100"', html)

        year_path = path + '/year'
        empty = self.client.post(year_path, data={'year': '2098'},
                                 follow_redirects=True)
        self.assertEqual(empty.status_code, 200)
        self.assertIn('Укажите индекс хотя бы за один месяц',
                      empty.get_data(as_text=True))
        self.assertIn('class="alert error"', empty.get_data(as_text=True))
        for data in ({'year': '2101', 'percent_1': '101'},
                     {'year': '2098', 'percent_1': '101',
                      'percent_2': 'не число'},
                     {'year': '2098', 'percent_1': '101',
                      'percent_2': 'NaN'},
                     {'year': '2098', 'percent_1': '101.1234'}):
            self.assertEqual(self.client.post(year_path, data=data).status_code,
                             400)
        with app.app_context():
            self.assertEqual(AnalogMonthlyIndex.query.filter_by(year=2098).count(), 0)

        try:
            saved = self.client.post(year_path, data={
                'year': '2098', 'percent_1': '101,125',
                'percent_7': '99.5', 'percent_12': '0'})
            self.assertEqual(saved.status_code, 302)
            self.assertTrue(saved.headers['Location'].endswith('#year-2098'))
            with app.app_context():
                values = {(row.month, row.percent) for row in
                          AnalogMonthlyIndex.query.filter_by(year=2098).all()}
                self.assertEqual(values, {
                    (1, Decimal('101.125')), (7, Decimal('99.500')),
                    (12, Decimal('0.000'))})
                _, latest = analog_routes._monthly_index_context()
                self.assertEqual(latest, (2098, 7))
            updated = self.client.get(path).get_data(as_text=True)
            self.assertIn('id="year-2098"', updated)
            self.assertIn('101.125', updated)
            self.assertEqual(self.client.post(year_path, data={
                'year': '2098', 'percent_1': '77'}).status_code, 409)
            with app.app_context():
                self.assertEqual(
                    AnalogMonthlyIndex.query.filter_by(
                        year=2098, month=1).one().percent, Decimal('101.125'))

            with app.app_context():
                viewer = User(username='indices-viewer', password_hash='unused',
                              role='viewer')
                db.session.add(viewer)
                db.session.commit()
                viewer_id = viewer.id
                viewer_version = int(viewer.password_version or 1)
            try:
                self.login_values(viewer_id, viewer_version)
                self.assertNotIn('action="/analog-search/monthly-indices/year"',
                                 self.client.get(path).get_data(as_text=True))
                self.assertEqual(self.client.post(year_path, data={
                    'year': '2099', 'percent_1': '100'}).status_code, 403)
            finally:
                self.login_values(self.admin_id, self.admin_version)
                with app.app_context():
                    db.session.delete(db.session.get(User, viewer_id))
                    db.session.commit()
        finally:
            with app.app_context():
                AnalogMonthlyIndex.query.filter_by(year=2098).delete()
                db.session.commit()

    def test_clear_index_and_zero_do_not_reappear_or_zero_out_cost(self):
        path = '/analog-search/monthly-indices'
        with app.app_context():
            seed_monthly_indices()
            row = AnalogMonthlyIndex.query.filter_by(year=2026, month=7).one()
            original_percent, original_source = row.percent, row.source
        try:
            cleared = self.client.post(path, data={
                'year': '2026', 'month': '7', 'percent': ''},
                follow_redirects=True)
            self.assertEqual(cleared.status_code, 200)
            self.assertIn('Индекс за 07.2026 очищен',
                          cleared.get_data(as_text=True))
            with app.app_context():
                seed_monthly_indices()
                indices, latest = analog_routes._monthly_index_context()
                self.assertIsNone(indices[(2026, 7)])
                self.assertEqual(latest, (2026, 6))
                self.assertEqual(
                    AnalogMonthlyIndex.query.filter_by(
                        year=2026, month=7).one().percent, Decimal('0.000'))
            self.assertIn('value="" placeholder="—"', self.client.get(
                path).get_data(as_text=True))

            zero = self.client.post(path, data={
                'year': '2026', 'month': '7', 'percent': '0'},
                follow_redirects=True)
            self.assertEqual(zero.status_code, 200)
            html = zero.get_data(as_text=True)
            self.assertIn('Индекс за 07.2026 сохранён', html)
            self.assertIn('value="0" placeholder="—"', html)
            with app.app_context():
                indices, latest = analog_routes._monthly_index_context()
                self.assertEqual(indices[(2026, 7)], Decimal('0.000'))
                self.assertEqual(latest, (2026, 6))
                version = SimpleNamespace(
                    total_cost=1000, price_year=2026,
                    payload_json=json.dumps({'ssr': {
                        'price_year': 2026, 'price_month': 7}}))
                self.assertIsNone(indexed_cost(version, indices, (2026, 7)))
        finally:
            with app.app_context():
                row = AnalogMonthlyIndex.query.filter_by(
                    year=2026, month=7).one()
                row.percent, row.source = original_percent, original_source
                db.session.commit()

    @classmethod
    def setUpClass(cls):
        os.environ['INITIAL_ADMIN_PASSWORD'] = 'analog-admin-passphrase-2026'
        cls.addClassCleanup(shutil.rmtree, _tmp, ignore_errors=True)
        cls.storage_patch = patch.object(
            analog_routes, 'ANALOG_STORAGE', os.path.join(_tmp, 'raw_sources'))
        cls.storage_patch.start()
        cls.addClassCleanup(cls.storage_patch.stop)
        app.config.update(TESTING=True, WTF_CSRF_ENABLED=False)
        init_db()
        cls.conclusion = (ROOT / '50-1-1-2-018685-2026_26-04-29_11-32-17_signature_1789461292785.xml').read_bytes()
        cls.ssr = (ROOT / 'ССР_Кривцово_1789461305316.gge').read_bytes()
        cls.conclusion_101 = (ROOT / '56-1-1-2-044346-2024_24-08-06_09-12-42_signature_1789537746147.xml').read_bytes()
        cls.ssr_101 = (ROOT / 'Сводный_сметный_расчет_1789537733540.gge').read_bytes()

    def setUp(self):
        # Keep fixture identities distinct between tests sharing this temporary
        # database, while preserving the real XML/GGE payloads within each test.
        marker = int(hashlib.sha256(self._testMethodName.encode()).hexdigest()[:8], 16)
        self.conclusion = type(self).conclusion.replace(
            b'018685-2026', f'{marker % 1000000:06d}-2026'.encode())
        self.conclusion_101 = type(self).conclusion_101.replace(
            b'044346-2024', f'{(marker + 1) % 1000000:06d}-2024'.encode())
        self.client = app.test_client()
        with app.app_context():
            self.admin = User.query.filter_by(role='admin').first()
            if not self.admin:
                self.admin = User(username='analog-admin', password_hash='unused',
                                  role='admin')
                db.session.add(self.admin)
            self.admin.password_change_required = False
            db.session.commit()
            self.admin_id = self.admin.id
            self.admin_version = int(self.admin.password_version or 1)
        self.login_values(self.admin_id, self.admin_version)

    def login_values(self, user_id, version=1):
        with self.client.session_transaction() as session:
            session['_user_id'] = str(user_id)
            session['_fresh'] = True
            session['password_version'] = version

    def files(self, conclusion=None, ssr=None):
        return {'conclusion': (BytesIO(conclusion or self.conclusion), 'conclusion.xml'),
                'ssr': (BytesIO(ssr or self.ssr), 'summary.gge')}

    def test_passport_omits_documentation_for_current_and_previous_versions(self):
        with app.app_context():
            passport = AnalogPassport(
                title='Паспорт без технического перечня', object_type='котельная',
                created_by=self.admin_id)
            db.session.add(passport)
            db.session.flush()
            version_ids = []
            for number in (1, 2):
                version = AnalogPassportVersion(
                    passport_id=passport.id, version_no=number,
                    conclusion_number=f'ui-docs-{number}',
                    examination_result='positive',
                    source_hash=f'ui-docs-source-{number}',
                    ssr_hash=f'ui-docs-ssr-{number}', imported_by=self.admin_id,
                    payload_json=json.dumps({'ssr': {'name': f'ССР версии {number}'}}))
                db.session.add(version)
                db.session.flush()
                version_ids.append(version.id)
                db.session.add(AnalogSource(
                    version_id=version.id,
                    source_hash=f'ui-docs-source-{number}',
                    ssr_hash=f'ui-docs-ssr-{number}',
                    source_filename='hidden-source.xml',
                    ssr_filename='hidden-summary.gge'))
                db.session.add(AnalogDiagnostic(
                    version_id=version.id, code='ui_test',
                    message='Сохранённое сообщение диагностики'))
                db.session.add(AnalogDocument(
                    passport_id=passport.id, version_id=version.id,
                    ordinal='3.1.1.987', name='hidden-document-name',
                    checksum='hidden-document-checksum',
                    signature_name='hidden-signature.sig',
                    signature_checksum='hidden-signature-checksum'))
            passport.current_version_id = version_ids[-1]
            db.session.commit()
            passport_id = passport.id

        base = f'/analog-search/passports/{passport_id}'
        for suffix, number in (('', 2), (f'?version={version_ids[0]}', 1),
                               (f'?version={version_ids[1]}', 2)):
            with self.subTest(version=number, url=suffix):
                response = self.client.get(base + suffix)
                self.assertEqual(response.status_code, 200)
                html = response.get_data(as_text=True)
                for text in ('Документация 3.1.1', 'Контрольная сумма',
                             '3.1.1.987', 'hidden-document-name',
                             'hidden-document-checksum', 'hidden-signature.sig',
                             'hidden-signature-checksum',
                             'Диагностика и происхождение',
                             'Нет сообщений диагностики.',
                             'Сырые байты сохранены по защищённым путям.',
                             'hidden-source.xml', 'hidden-summary.gge',
                             'ui-docs-source-', 'ui-docs-ssr-',
                             'Сохранённое сообщение диагностики'):
                    self.assertNotIn(text, html)
                for text in ('Основные сведения',
                             'Организации (Все)', 'Геология и прочий климат',
                             'Технико-экономические показатели', 'Таблица 3.2.1',
                             'Сводный сметный расчёт (GGE)',
                             f'ССР версии {number}',
                             'data-target="brief-tab"', 'data-target="detail-tab"',
                             'id="brief-tab" class="tab-content"',
                             'id="detail-tab" class="tab-content" hidden',
                             "document.getElementById(btn.dataset.target).hidden = false"):
                    self.assertIn(text, html)
                self.assertNotIn('Версии и происхождение', html)
                self.assertNotIn('cost-adjust-form', html)
                self.assertNotIn('Переимпортировать', html)
        with app.app_context():
            self.assertEqual(AnalogDocument.query.filter_by(
                passport_id=passport_id).count(), 2)
            for number, version_id in enumerate(version_ids, 1):
                source = AnalogSource.query.filter_by(version_id=version_id).one()
                self.assertEqual(source.source_filename, 'hidden-source.xml')
                self.assertEqual(source.ssr_filename, 'hidden-summary.gge')
                self.assertEqual(source.source_hash, f'ui-docs-source-{number}')
                self.assertEqual(source.ssr_hash, f'ui-docs-ssr-{number}')
                diagnostic = AnalogDiagnostic.query.filter_by(
                    version_id=version_id).one()
                self.assertEqual(diagnostic.code, 'ui_test')
                self.assertEqual(diagnostic.message,
                                 'Сохранённое сообщение диагностики')

    def test_summary_estimate_101_preview_and_confirm(self):
        response = self.client.post(
            '/analog-search/import/preview',
            data=self.files(self.conclusion_101, self.ssr_101),
            content_type='multipart/form-data')
        self.assertEqual(response.status_code, 200)
        preview = response.get_json()
        self.assertTrue(preview['can_confirm'], preview['diagnostics'])
        pair = preview['preview']
        self.assertEqual(pair['ssr']['schema_version'], '1.01')
        self.assertEqual(pair['ssr']['price_quarter'], 4)
        self.assertEqual(pair['ssr']['summary']['total'], '31372.62')
        self.assertEqual(len(pair['ssr']['chapters']), 14)
        self.assertEqual(len(pair['ssr']['items']), 17)
        html = self.client.get(
            f"/analog-search/imports/{preview['import_id']}/preview")
        self.assertEqual(html.status_code, 200)
        page = html.get_data(as_text=True)
        self.assertIn('31372.62', page)
        self.assertIn('2023 / квартал 4', page)
        self.assertIn('<option selected>БМК</option>', page)

        confirmed = self.client.post(
            f"/analog-search/import/{preview['import_id']}/confirm", json={})
        self.assertEqual(confirmed.status_code, 200)
        passport_id = confirmed.get_json()['passport_id']
        with app.app_context():
            version = AnalogPassportVersion.query.filter_by(
                passport_id=passport_id).order_by(
                    AnalogPassportVersion.version_no.desc()).first()
            passport = version.passport
            self.assertEqual(passport.object_type, 'БМК')
            self.assertEqual(version.schema_version, '1.01')
            self.assertEqual((version.price_year, version.price_quarter), (2023, 4))
            self.assertEqual(version.total_cost, Decimal('31372.62'))
            self.assertEqual(passport.capacity_kw, Decimal('1800.000000'))
            payload = json.loads(version.payload_json)
            self.assertEqual(len(payload['ssr']['chapters']), 14)
            self.assertEqual(len(payload['ssr']['items']), 17)
        passport_page = self.client.get(
            f'/analog-search/passports/{passport_id}').get_data(as_text=True)
        self.assertIn('1.8 мВт', passport_page)
        self.assertIn('Сводный сметный расчёт (GGE)', passport_page)
        self.assertIn('Конструктивные решения.gsfx', passport_page)
        self.assertIn('Оплата труда рабочих', passport_page)
        self.assertNotIn('WorkersSalary', passport_page)
        self.assertIn('111.13', passport_page)
        self.assertIn('ГРАНД-Смета', passport_page)
        self.assertIn('Стоимость по ССР', passport_page)
        self.assertIn('Строительные работы', passport_page)
        self.assertIn('Монтажные работы', passport_page)
        self.assertIn('Оборудование', passport_page)
        self.assertIn('Прочие затраты', passport_page)
        self.assertNotIn('Текущая стоимость', passport_page)
        self.assertNotIn('Предложить по адресу', passport_page)

    def test_explicit_import_type_overrides_bmk_suggestion(self):
        conclusion = self.conclusion_101.replace(
            b'56-1-1-2-044346-2024', b'56-1-1-2-044346-2099')
        response = self.client.post(
            '/analog-search/import/preview',
            data=self.files(conclusion, self.ssr_101 + b' '),
            content_type='multipart/form-data')
        preview = response.get_json()
        confirmed = self.client.post(
            f"/analog-search/import/{preview['import_id']}/confirm",
            json={'object_type': 'ЦТП'})
        self.assertEqual(confirmed.status_code, 200)
        with app.app_context():
            passport = db.session.get(
                AnalogPassport, confirmed.get_json()['passport_id'])
            self.assertEqual(passport.object_type, 'ЦТП')

    def test_legacy_default_types_are_reclassified_from_saved_text(self):
        with app.app_context():
            rows = [
                AnalogPassport(
                    title='Строительство БМК на 1,2 МВт',
                    object_type='котельная', created_by=self.admin_id),
                AnalogPassport(
                    title=('Блочно-модульная котельная 2,1МВт, '
                           'взамен действующей г.о. Электросталь'),
                    extracted_json=json.dumps({
                        'documents': ['Реконструкция ЦТП № 4']
                    }, ensure_ascii=False),
                    object_type='котельная', created_by=self.admin_id),
                AnalogPassport(
                    title='Реконструкция объекта',
                    construction_object='Центральный тепловой пункт № 3',
                    object_type='котельная', created_by=self.admin_id),
                AnalogPassport(
                    title='Обычная котельная',
                    object_type='котельная', created_by=self.admin_id),
            ]
            db.session.add_all(rows)
            db.session.flush()
            ids = [row.id for row in rows]
            self.assertEqual(_backfill_inferred_object_types(), 3)
            db.session.commit()
            self.assertEqual(
                [db.session.get(AnalogPassport, row_id).object_type
                 for row_id in ids],
                ['БМК', 'БМК', 'ЦТП', 'котельная'])

    def test_capacity_edit_is_immediately_reflected_in_catalog(self):
        title = 'catalog-capacity-after-passport-edit'
        with app.app_context():
            passport = AnalogPassport(
                title=title,
                object_type='котельная',
                work_type='строительство',
                capacity_kw=Decimal('1.8'),
                created_by=self.admin_id,
            )
            db.session.add(passport)
            db.session.commit()
            passport_id = passport.id

        edited = self.client.post(
            f'/analog-search/passports/{passport_id}/edit',
            data={'capacity_kw': '1800'},
        )
        self.assertEqual(edited.status_code, 302)

        catalog = self.client.get(
            '/analog-search/simple',
            query_string={'view': 'base', 'title': title, 'format': 'json'},
        )
        self.assertEqual(catalog.status_code, 200)
        payload = catalog.get_json()
        self.assertEqual(payload['total'], 1)
        self.assertEqual(payload['items'][0]['capacity_kw'], '1800.000000')
        self.assertEqual(payload['items'][0]['capacity_mw_display'], '1,8')

        with app.app_context():
            db.session.delete(db.session.get(AnalogPassport, passport_id))
            db.session.commit()

    def test_html_confirm_redirects_to_analog_search(self):
        response = self.client.post(
            '/analog-search/import/preview',
            data=self.files(self.conclusion_101, self.ssr_101 + b'\n '),
            content_type='multipart/form-data')
        preview = response.get_json()
        confirmed = self.client.post(
            f"/analog-search/import/{preview['import_id']}/confirm",
            data={}, follow_redirects=False)
        self.assertEqual(confirmed.status_code, 302)
        self.assertEqual(
            confirmed.headers['Location'],
            '/analog-search/simple?view=base')

        duplicate = self.client.post(
            '/analog-search/import/preview',
            data=self.files(self.conclusion_101, self.ssr_101 + b'\n '),
            content_type='multipart/form-data',
            headers={'Accept': 'text/html'},
            follow_redirects=False)
        self.assertEqual(duplicate.status_code, 302)
        self.assertEqual(duplicate.headers['Location'], '/analog-search/simple')
        page = self.client.get('/analog-search/simple').get_data(as_text=True)
        self.assertIn('Объект уже находится в общей таблице аналогов', page)

    def test_preview_confirm_duplicate_and_raw_paths(self):
        conclusion = self.conclusion + b'\n'
        response = self.client.post('/analog-search/import/preview',
                                    data=self.files(conclusion),
                                    content_type='multipart/form-data')
        self.assertEqual(response.status_code, 200)
        preview = response.get_json()
        self.assertEqual(preview['status'], 'preview')
        confirmed = self.client.post(
            f"/analog-search/import/{preview['import_id']}/confirm",
            json={})
        self.assertEqual(confirmed.status_code, 200)
        passport_id = confirmed.get_json()['passport_id']
        with app.app_context():
            row = db.session.get(AnalogImport, preview['import_id'])
            raw_path = os.path.realpath(os.path.join(
                os.path.dirname(analog_routes.__file__), row.raw_conclusion_path))
            self.assertTrue(raw_path.startswith(
                os.path.realpath(os.path.join(_tmp, 'raw_sources')) + os.sep))
            self.assertTrue(AnalogPassport.query.get(passport_id).current_version_id)
            documents = AnalogDocument.query.filter_by(
                passport_id=passport_id).all()
            self.assertTrue(any(item.parent_id is not None for item in documents))
            self.assertTrue(any(item.node_kind == 'section' for item in documents))
            passport = db.session.get(AnalogPassport, passport_id)
            self.assertEqual(
                passport.address,
                'Солнечногорский район, с.п. Кривцовское, д. Кривцово')
            self.assertEqual(passport.capacity_kw, Decimal('7594.390000'))
            self.assertEqual(passport.teps[-1].value_text, '6,53')
            self.assertEqual(passport.teps[-1].unit, 'Гкал/ч')
        duplicate = self.client.post('/analog-search/import/preview',
                                     data=self.files(conclusion),
                                     content_type='multipart/form-data')
        self.assertEqual(duplicate.get_json()['status'], 'duplicate')

    def test_global_abs_boundary_and_pagination(self):
        with app.app_context():
            before = AnalogPassport.query.filter_by(archived=False).count()
            for number, capacity in enumerate((90, 110, 89, 111), 1):
                db.session.add(AnalogPassport(
                    title=f'boundary-{number}', capacity_kw=capacity,
                    object_type='котельная', created_by=self.admin.id))
            db.session.commit()
        response = self.client.get('/analog-search/api/search?capacity_kw=100&per_page=2')
        payload = response.get_json()
        self.assertEqual(payload['good'], 2)
        self.assertEqual(payload['total'], before + 4)
        self.assertEqual(len(payload['items']), 2)
        self.assertTrue(all(item['good'] for item in payload['items']))

    def test_ui_search_accepts_gcal_and_excludes_over_ten_percent(self):
        with app.app_context():
            db.session.add_all([
                AnalogPassport(
                    title='gcal-inside', capacity_kw=Decimal('1163'),
                    object_type='котельная', created_by=self.admin_id),
                AnalogPassport(
                    title='gcal-boundary', capacity_kw=Decimal('1279.3'),
                    object_type='котельная', created_by=self.admin_id),
                AnalogPassport(
                    title='gcal-outside', capacity_kw=Decimal('1280'),
                    object_type='котельная', created_by=self.admin_id),
            ])
            db.session.commit()
        response = self.client.get(
            '/analog-search/simple',
            query_string={'view': 'search', 'capacity_value': '1',
                          'capacity_unit': 'gcal_h',
                          'q': 'gcal-'})
        self.assertEqual(response.status_code, 200)
        html = response.get_data(as_text=True)
        self.assertIn('gcal-inside', html)
        self.assertIn('gcal-boundary', html)
        self.assertNotIn('gcal-outside', html)

    def test_configurable_capacity_tolerance(self):
        with app.app_context():
            db.session.add_all([
                AnalogPassport(title='tolerance-exact', capacity_kw=1000,
                               created_by=self.admin_id),
                AnalogPassport(title='tolerance-wide', capacity_kw=1200,
                               created_by=self.admin_id),
            ])
            db.session.commit()
        for endpoint in ('/analog-search/simple', '/analog-search/api/search'):
            params = dict(capacity_value='1', capacity_unit='mw',
                          q='tolerance-', format='json')
            for value, expected in [('10', 1), ('20', 2), ('0', 1), ('19,5', 1)]:
                response = self.client.get(endpoint, query_string={
                    **params, 'tolerance_percent': value})
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.get_json()['total'], expected)
                self.assertTrue(all(row['good'] for row in response.get_json()['items']))
            for value in ('-1', '101', 'NaN', 'Infinity', 'abc'):
                response = self.client.get(endpoint, query_string={
                    **params, 'tolerance_percent': value})
                self.assertEqual(response.status_code, 400)
        html = self.client.get('/analog-search/simple').get_data(as_text=True)
        self.assertIn('name="tolerance_percent"', html)
        self.assertNotIn('Хорошие аналоги', html)

    def test_search_type_options_reflect_only_active_passports(self):
        with app.app_context():
            passports = [
                AnalogPassport(title='filter-type-live-1',
                               object_type='тип-доступен', created_by=self.admin_id),
                AnalogPassport(title='filter-type-live-2',
                               object_type='тип-доступен', created_by=self.admin_id),
                AnalogPassport(title='filter-type-archived',
                               object_type='тип-архивный', archived=True,
                               created_by=self.admin_id),
            ]
            db.session.add_all(passports)
            db.session.commit()
            ids = [passport.id for passport in passports]

        def options(query=None):
            response = self.client.get('/analog-search/simple', query_string=query)
            self.assertEqual(response.status_code, 200)
            html = response.get_data(as_text=True)
            select = re.search(
                r'<select name="object_type"[^>]*>(.*?)</select>',
                html, re.DOTALL)
            self.assertIsNotNone(select)
            return html, re.findall(r'<option value="([^"]*)"', select.group(1))

        try:
            html, values = options({'object_type': 'тип-доступен'})
            with app.app_context():
                expected = {value for (value,) in db.session.query(
                    AnalogPassport.object_type).filter(
                        AnalogPassport.archived.is_(False)).distinct().all()
                    if value and value.strip()}
            self.assertEqual(set(values), expected | {''})
            self.assertEqual(values.count('тип-доступен'), 1)
            self.assertNotIn('тип-архивный', values)
            self.assertIn('value="тип-доступен" selected', html)

            with app.app_context():
                for passport_id in ids[:2]:
                    db.session.get(AnalogPassport, passport_id).archived = True
                db.session.commit()
            _, values = options()
            self.assertNotIn('тип-доступен', values)
        finally:
            with app.app_context():
                for passport_id in ids:
                    passport = db.session.get(AnalogPassport, passport_id)
                    if passport:
                        db.session.delete(passport)
                db.session.commit()

    def test_fractional_capacity_accepts_dot_and_comma_and_filters_results(self):
        with app.app_context():
            db.session.add_all([
                AnalogPassport(
                    title='fraction-inside', capacity_kw=Decimal('1200'),
                    object_type='котельная', created_by=self.admin_id),
                AnalogPassport(
                    title='fraction-outside', capacity_kw=Decimal('1400'),
                    object_type='котельная', created_by=self.admin_id),
            ])
            db.session.commit()
        for value in ('1.2', '1,2'):
            response = self.client.get(
                '/analog-search/simple',
                query_string={'view': 'search', 'capacity_value': value,
                              'capacity_unit': 'mw', 'q': 'fraction-'})
            self.assertEqual(response.status_code, 200)
            html = response.get_data(as_text=True)
            self.assertIn(f'value="{value}"', html)
            self.assertIn('fraction-inside', html)
            self.assertNotIn('fraction-outside', html)

            api_response = self.client.get(
                '/analog-search/api/search',
                query_string={'capacity_value': value, 'capacity_unit': 'mw',
                              'q': 'fraction-'})
            self.assertEqual(api_response.status_code, 200)
            api = api_response.get_json()
            self.assertEqual(api['total'], 1)
            self.assertEqual(api['items'][0]['title'], 'fraction-inside')
            self.assertTrue(api['items'][0]['good'])
            self.assertNotIn(
                'fraction-outside',
                [item['title'] for item in api['items']])

    def test_capacity_search_rejects_invalid_and_nonpositive_values(self):
        for value in ('invalid', '0', '-1', 'Infinity', 'NaN'):
            for path in ('/analog-search/simple',
                         '/analog-search/api/search'):
                response = self.client.get(
                    path,
                    query_string={'view': 'search',
                                  'capacity_value': value,
                                  'capacity_unit': 'mw'})
                self.assertEqual(response.status_code, 400)
                self.assertIn(
                    'Мощность должна быть положительным числом',
                    response.get_data(as_text=True))

    def test_power_parser_normalizes_legacy_formats_to_kw(self):
        object_node = ET.fromstring('''
        <Object>
          <TEI><Name>Мощность котельной</Name><Measure>Гигакалория в час</Measure><Value>14</Value></TEI>
          <TEI><Name>Мощность котельной Гкал</Name><Value>6,53</Value></TEI>
          <TEI><Name>Мощность котельной</Name><Measure>МВт/Гкал/ч</Measure><Value>31/26,66</Value></TEI>
        </Object>
        ''')
        rows = _tei_rows(object_node)
        self.assertEqual(rows[0]['normalized_value'], Decimal('16282'))
        self.assertEqual(rows[1]['normalized_value'], Decimal('7594.39'))
        self.assertEqual(rows[2]['normalized_value'], Decimal('31000'))
        self.assertTrue(all(row['normalized_unit'] == 'кВт' for row in rows))

    def test_catalog_formats_total_cost_and_dual_capacity(self):
        with app.app_context():
            passport = AnalogPassport(
                title='Формат каталога', object_type='котельная',
                work_type='строительство', capacity_kw=31000,
                created_by=self.admin_id)
            db.session.add(passport)
            db.session.flush()
            version = AnalogPassportVersion(
                passport_id=passport.id, version_no=1,
                conclusion_number='catalog-format', examination_result='positive',
                source_hash='catalog-format-source',
                ssr_hash='catalog-format-ssr', total_cost=Decimal('1234.5'),
                imported_by=self.admin_id)
            db.session.add(version)
            db.session.flush()
            passport.current_version_id = version.id
            db.session.commit()
            row = _result_row(passport)
        self.assertEqual(_format_ru_number(Decimal('1234.5'), 1), '1 234,5')
        self.assertEqual(row['capacity_mw_display'], '31')
        self.assertEqual(row['capacity_gcal_h_display'], '26,66')
        self.assertEqual(row['total_cost_display'], '1 234,5')
        response = self.client.get(
            '/analog-search/simple',
            query_string={'view': 'base', 'q': 'Формат каталога'})
        html = response.get_data(as_text=True)
        self.assertIn('МВт/<strong>Гкал/ч</strong>', html)
        self.assertIn('31/<strong>26,66</strong>', html)
        self.assertIn('1 234,5', html)
        search = self.client.get(
            '/analog-search/simple',
            query_string={'view': 'search', 'q': 'Формат каталога'})
        search_html = search.get_data(as_text=True)
        for heading in ('Наименование', 'Тип', 'Вид работ', 'Аналогичность'):
            self.assertIn(f'<th>{heading}</th>', search_html)
        self.assertIn('<th>Стоимость, тыс. руб.</th>', search_html)
        self.assertIn('<th>Стоимость (текущие цены), тыс. руб</th>',
                      search_html)
        self.assertIn('МВт/<strong>Гкал/ч</strong>', search_html)
        self.assertIn('<td>котельная</td>', search_html)
        self.assertIn('31/<strong>26,66</strong>', search_html)
        self.assertIn('<td>1 234,5</td>', search_html)

    def test_network_total_cost_comes_from_table_321_total(self):
        with app.app_context():
            passport = AnalogPassport(
                title='Сеть со стоимостью 3.2.1',
                object_type='тепловая сеть', created_by=self.admin_id)
            db.session.add(passport)
            db.session.flush()
            version = AnalogPassportVersion(
                passport_id=passport.id, version_no=1,
                conclusion_number='network-table-321',
                examination_result='positive',
                source_hash='network-table-321-source',
                ssr_hash='network-table-321-ssr',
                imported_by=self.admin_id)
            db.session.add(version)
            db.session.flush()
            passport.current_version_id = version.id
            db.session.add(AnalogCostRow(
                passport_id=passport.id, version_id=version.id,
                ordinal='ctot', name='total',
                approved_cost=Decimal('40036.33')))
            db.session.flush()
            self.assertEqual(_table_321_total(passport.costs),
                             Decimal('40036.33'))
            self.assertEqual(_backfill_analog_total_costs(), 1)
            db.session.commit()
            self.assertEqual(version.total_cost, Decimal('40036.330000'))
            row = _result_row(passport)
            self.assertEqual(row['total_cost_display'], '40 036,33')

    def test_thousand_records_are_paginated_by_db(self):
        with app.app_context():
            before = AnalogPassport.query.filter_by(archived=False).count()
            db.session.bulk_save_objects([
                AnalogPassport(title=f'bulk-{index}', capacity_kw=100 + (index % 30),
                               object_type='котельная', created_by=self.admin_id)
                for index in range(1000)])
            db.session.commit()
        response = self.client.get(
            '/analog-search/api/search?capacity_kw=100&page=20&per_page=25')
        self.assertEqual(response.status_code, 200)
        payload = response.get_json()
        self.assertEqual(payload['total'], before + 1000)
        self.assertEqual(len(payload['items']), 25)

    def test_json_v4_preserves_legacy_and_network_fields(self):
        archive = next(ROOT.glob('*Объекты_аналоги*.zip'))
        with zipfile.ZipFile(archive) as stream:
            name = next(item for item in stream.namelist()
                        if item.endswith('.json') and not item.startswith('__'))
            payload = json.loads(stream.read(name))
        with app.app_context():
            report = import_json_v4(payload, user_id=self.admin_id)
            self.assertEqual(report['transformed'], len(payload['records']))
            self.assertGreater(report['fields_preserved'], 0)
            self.assertGreater(AnalogPassport.query.filter(
                AnalogPassport.legacy_json != '{}').count(), 0)
            self.assertGreater(AnalogPassportVersion.query.filter(
                AnalogPassportVersion.total_cost.is_not(None)).count(), 0)

    def test_reimport_target_and_same_identity_are_rejected(self):
        first = self.client.post('/analog-search/import/preview',
                                 data=self.files(self.conclusion + b'\t'),
                                 content_type='multipart/form-data').get_json()
        pending = self.client.post(
            '/analog-search/api/import/preview',
            data=self.files(self.conclusion + b' \n'),
            content_type='multipart/form-data').get_json()
        self.assertEqual(pending['status'], 'preview')
        confirmed = self.client.post(
            f"/analog-search/import/{first['import_id']}/confirm", json={}).get_json()
        target = confirmed['passport_id']
        with app.app_context():
            baseline = len(AnalogPassport.query.get(target).versions)
            baseline_imports = AnalogImport.query.count()
        for endpoint in ('/analog-search/import/preview',
                         '/analog-search/import',
                         '/analog-search/api/import/preview'):
            for key in ('passport_id', 'target_passport_id'):
                for value in ('', str(target)):
                    response = self.client.post(
                        f'{endpoint}?{key}={value}',
                        data=self.files(self.conclusion + b' \n'),
                        content_type='multipart/form-data')
                    self.assertEqual(response.status_code, 400)
                response = self.client.post(
                    endpoint, data={**self.files(self.conclusion + b' \n'),
                                    key: str(target)},
                    content_type='multipart/form-data')
                self.assertEqual(response.status_code, 400)
            response = self.client.post(
                endpoint, data=self.files(self.conclusion + b' \n'),
                content_type='multipart/form-data')
            self.assertEqual(response.status_code, 409)
        for endpoint in (f'/analog-search/import/{pending["import_id"]}/confirm',
                         f'/analog-search/imports/{pending["import_id"]}/confirm',
                         f'/analog-search/api/import/{pending["import_id"]}/confirm'):
            self.assertEqual(self.client.post(endpoint, json={}).status_code, 409)
            self.assertEqual(self.client.post(
                endpoint, json={'passport_id': target}).status_code, 400)
            self.assertEqual(self.client.post(
                endpoint, data={'target_passport_id': target}).status_code, 400)
        self.assertEqual(self.client.post(
            f'/analog-search/passports/{target}/reimport',
            data=self.files(), content_type='multipart/form-data').status_code, 409)
        duplicate = self.client.post(
            '/analog-search/import/preview',
            data=self.files(self.conclusion + b'\t'),
            content_type='multipart/form-data')
        self.assertEqual(duplicate.status_code, 200)
        self.assertEqual(duplicate.get_json()['status'], 'duplicate')
        self.assertEqual(duplicate.get_json()['passport_id'], target)
        with app.app_context():
            self.assertEqual(len(AnalogPassport.query.get(target).versions), baseline)
            self.assertEqual(AnalogImport.query.count(), baseline_imports)
            self.assertEqual(db.session.get(AnalogImport, pending['import_id']).status, 'preview')

    def test_old_targeted_preview_cannot_confirm(self):
        first = self.client.post('/analog-search/import/preview',
                                 data=self.files(self.conclusion_101, self.ssr_101),
                                 content_type='multipart/form-data').get_json()
        self.assertEqual(first['status'], 'preview')
        with app.app_context():
            target = AnalogPassport(title='legacy-target', object_type='ЦТП',
                                    created_by=self.admin_id)
            db.session.add(target)
            db.session.flush()
            row = db.session.get(AnalogImport, first['import_id'])
            row.target_passport_id = target.id
            db.session.commit()
            target_id = target.id
        response = self.client.post(
            f"/analog-search/import/{first['import_id']}/confirm", json={})
        self.assertEqual(response.status_code, 409)
        self.assertEqual(self.client.get(
            f"/analog-search/imports/{first['import_id']}/preview").status_code, 409)
        with app.app_context():
            self.assertEqual(AnalogPassportVersion.query.filter_by(
                passport_id=target_id).count(), 0)

    def test_concurrent_pending_confirms_create_only_one_passport(self):
        first = self.client.post('/analog-search/import/preview',
                                 data=self.files(self.conclusion + b'\t'),
                                 content_type='multipart/form-data').get_json()
        second = self.client.post('/analog-search/import/preview',
                                  data=self.files(self.conclusion + b'\n '),
                                  content_type='multipart/form-data').get_json()
        self.assertEqual((first['status'], second['status']), ('preview', 'preview'))

        def confirm(import_id):
            client = app.test_client()
            with client.session_transaction() as session:
                session['_user_id'] = str(self.admin_id)
                session['_fresh'] = True
                session['password_version'] = self.admin_version
            return client.post(
                f'/analog-search/api/import/{import_id}/confirm', json={}).status_code

        with ThreadPoolExecutor(max_workers=2) as pool:
            statuses = list(pool.map(confirm, (first['import_id'], second['import_id'])))
        self.assertEqual(sorted(statuses), [200, 409])
        with app.app_context():
            self.assertEqual(AnalogPassportVersion.query.filter_by(
                conclusion_number=first['preview']['conclusion']['conclusion_number']
            ).count(), 1)

    def test_network_create_api_and_html_persist_typed_attrs(self):
        api = self.client.post('/analog-search/passports', json={
            'title': 'Сеть API integration', 'object_type': 'тепловая сеть',
            'network_length_m': '125.5', 'network_diameter_mm': '273'})
        self.assertEqual(api.status_code, 201)
        api_id = api.get_json()['passport_id']
        html = self.client.post('/analog-search/passports', data={
            'title': 'Сеть HTML integration', 'object_type': 'тепловая сеть',
            'network_length_m': '80', 'network_diameter_mm': '159'},
            headers={'Accept': 'text/html'})
        self.assertEqual(html.status_code, 302)
        with app.app_context():
            attrs = AnalogNetworkAttribute.query.filter_by(passport_id=api_id).all()
            self.assertEqual({item.key for item in attrs},
                             {'length_m', 'diameter_mm'})
            self.assertEqual({item.unit for item in attrs}, {'м', 'мм'})
            self.assertIsNotNone(AnalogPassport.query.filter_by(
                title='Сеть HTML integration').first())

    def test_html_preview_and_confirm_redirect(self):
        response = self.client.post('/analog-search/import/preview',
                                    data=self.files(self.conclusion + b'\n \n'),
                                    content_type='multipart/form-data',
                                    headers={'Accept': 'text/html'},
                                    follow_redirects=False)
        self.assertEqual(response.status_code, 303)
        preview_url = response.headers['Location']
        self.assertRegex(preview_url,
                         r'/analog-search/imports/\d+/preview$')
        response = self.client.get(preview_url)
        self.assertEqual(response.status_code, 200)
        self.assertIn('Загрузить объект в базу данных',
                      response.get_data(as_text=True))
        html = response.get_data(as_text=True)
        self.assertIn('analog-preview-page', html)
        self.assertIn('Солнечногорский район, с.п. Кривцовское, д. Кривцово', html)
        self.assertIn('Технико-экономические показатели', html)
        self.assertIn('Мощность котельной', html)
        self.assertIn('6,53 Гкал', html)
        self.assertIn('Климатический район, подрайон', html)
        self.assertIn('Геологические условия', html)
        self.assertIn('Сейсмическая активность (баллов)', html)

    def test_module_entry_opens_quick_search(self):
        page = self.client.get('/analog-search')
        self.assertEqual(page.status_code, 302)
        self.assertEqual(page.headers['Location'], '/analog-search/simple')
        page = self.client.get('/analog-search', follow_redirects=True)
        self.assertEqual(page.status_code, 200)
        html = page.get_data(as_text=True)
        self.assertIn('Поиск объектов-аналогов', html)
        self.assertIn('name="capacity_value"', html)
        self.assertNotIn('Поиск по нормативно-правовому акту', html)
        self.assertNotIn('Выберите способ', html)
        self.assertNotIn('/analog-search/regulatory', html)

    def test_regulatory_search_applies_135pr_mandatory_criteria(self):
        with app.app_context():
            passport = AnalogPassport(
                title='Нормативный аналог', object_type='тепловая сеть',
                climate_zone='IIВ', work_type='Строительство',
                created_by=self.admin_id)
            db.session.add(passport)
            db.session.flush()
            version = AnalogPassportVersion(
                passport_id=passport.id, version_no=1,
                conclusion_number=f'normative-{passport.id}',
                examination_result='positive',
                source_hash=f'normative:{passport.id}',
                ssr_hash=f'normative:{passport.id}',
                imported_by=self.admin_id)
            db.session.add(version)
            db.session.flush()
            passport.current_version_id = version.id
            db.session.add_all([
                AnalogNetworkAttribute(
                    passport_id=passport.id, version_id=version.id,
                    key='length_m', value='100', value_number=Decimal('100'),
                    unit='м'),
                AnalogNetworkAttribute(
                    passport_id=passport.id, version_id=version.id,
                    key='diameter_mm', value='200',
                    value_number=Decimal('200'), unit='мм'),
            ])
            db.session.commit()
            passport_id = passport.id

        response = self.client.post(
            '/analog-search/regulatory/search',
            json={'work_type': 'Строительство', 'length_m': '105',
                  'diameter_mm': '210', 'climate_region': 'II',
                  'climate_subzone': 'В'})
        self.assertEqual(response.status_code, 200)
        item = next(row for row in response.get_json()['items']
                    if row['id'] == passport_id)
        self.assertTrue(item['suitable'])
        self.assertEqual(item['algorithm_version'], '135pr-604pr-v1')

    def test_old_regulatory_pages_redirect_to_quick_search(self):
        for path in ('/analog-search/regulatory', '/analog-search/advanced'):
            with self.subTest(path=path):
                page = self.client.get(path)
                self.assertEqual(page.status_code, 302)
                self.assertEqual(page.headers['Location'], '/analog-search/simple')
                submitted = self.client.post(
                    path + '/search', data={},
                    headers={'Accept': 'text/html'})
                self.assertEqual(submitted.status_code, 303)
                self.assertEqual(submitted.headers['Location'],
                                 '/analog-search/simple')

    def test_climate_filter_supports_region_and_exact_subregion(self):
        with app.app_context():
            for zone in ('IIА', 'IIВ', 'IIIА'):
                db.session.add(AnalogPassport(
                    title=f'climate-filter-fixture-{zone}',
                    object_type='котельная', climate_zone=zone,
                    created_by=self.admin_id))
            db.session.commit()

        region = self.client.get(
            '/analog-search/api/search',
            query_string={'q': 'climate-filter-fixture',
                          'climate_region': 'II'}).get_json()
        self.assertEqual(region['total'], 2)
        self.assertEqual(
            {item['climate_zone'] for item in region['items']},
            {'IIА', 'IIВ'})

        exact = self.client.get(
            '/analog-search/api/search',
            query_string={'q': 'climate-filter-fixture',
                          'climate_region': 'III',
                          'climate_zone': 'IIВ'}).get_json()
        self.assertEqual(exact['total'], 1)
        self.assertEqual(exact['items'][0]['climate_zone'], 'IIВ')

    def test_simple_climate_input_suggestions_and_combined_filters(self):
        prefix = 'climate-ui-fixture-'
        with app.app_context():
            fixtures = [
                (prefix + 'match-1', 'IIВ', 'котельная', 'строительство', 1000, False),
                (prefix + 'match-2', 'IIВ', 'котельная', 'строительство', 1000, False),
                (prefix + 'latin-duplicate', 'IIB', 'ЦТП', 'реконструкция', 1000, False),
                (prefix + 'other-subzone', 'IIБ', 'котельная', 'строительство', 1000, False),
                (prefix + 'wrong-power', 'IIВ', 'котельная', 'строительство', 1500, False),
                (prefix + 'wrong-type', 'IIВ', 'ЦТП', 'строительство', 1000, False),
                (prefix + 'wrong-work', 'IIВ', 'котельная', 'реконструкция', 1000, False),
                (prefix + 'archived', 'архивный-район', 'котельная', 'строительство', 1000, True),
            ]
            db.session.add_all([
                AnalogPassport(title=title, climate_zone=zone, object_type=kind,
                               work_type=work, capacity_kw=power, archived=archived,
                               created_by=self.admin_id)
                for title, zone, kind, work, power, archived in fixtures])
            db.session.commit()

        path = '/analog-search/simple'
        html = self.client.get(path).get_data(as_text=True)
        self.assertRegex(
            html, r'<input type="text" name="climate_zone" list="climate-zone-options"'
                  r'\s+aria-label="Климатический район" placeholder="Климатический район"')
        datalist = re.search(
            r'<datalist id="climate-zone-options">(.*?)</datalist>', html, re.DOTALL)
        self.assertIsNotNone(datalist)
        options = re.findall(r'<option value="([^"]*)"', datalist.group(1))
        self.assertTrue({
            'IА', 'IБ', 'IВ', 'IГ', 'IД',
            'IIА', 'IIБ', 'IIВ', 'IIГ',
            'IIIА', 'IIIБ', 'IIIВ',
            'IVА', 'IVБ', 'IVВ', 'IVГ',
        }.issubset(options))
        self.assertEqual(options.count('IIВ'), 1)
        self.assertNotIn('IIB', options)
        self.assertNotIn('архивный-район', options)
        self.assertNotIn('', options)
        self.assertNotIn('температура января', datalist.group(1))

        params = {'view': 'search', 'q': prefix, 'object_type': 'котельная',
                  'work_type': 'строительство', 'capacity_value': '1',
                  'capacity_unit': 'mw', 'tolerance_percent': '0',
                  'climate_zone': '  iiB  '}
        response = self.client.get(path, query_string={**params, 'per_page': 1})
        self.assertEqual(response.status_code, 200)
        page = response.get_data(as_text=True)
        self.assertIn('value="  iiB  "', page)
        self.assertIn(prefix + 'match-1', page)
        self.assertNotIn(prefix + 'match-2', page)
        self.assertIn('climate_zone=++iiB++', page)  # pagination retains raw input
        second = self.client.get(path, query_string={
            **params, 'per_page': 1, 'page': 2}).get_data(as_text=True)
        self.assertIn('value="  iiB  "', second)
        self.assertIn(prefix + 'match-2', second)
        self.assertNotIn(prefix + 'match-1', second)
        self.assertEqual(self.client.get(path, query_string={
            **params, 'format': 'json'}).get_json()['total'], 2)

        for zone, count in (('IIВ', 6), ('IIB', 6), ('ii б', 1),
                            ('неизвестный', 0), ('', 7)):
            with self.subTest(zone=zone):
                result = self.client.get(path, query_string={
                    'q': prefix, 'climate_zone': zone, 'format': 'json'}).get_json()
                self.assertEqual(result['total'], count)
        for endpoint in (path, '/analog-search/api/search'):
            # The legacy region mode remains unchanged; it does not include
            # the Latin-spelled legacy record.
            for extra, count in (({'climate_region': 'II'}, 6),
                                 ({'climate_region': 'III',
                                   'climate_zone': ' IIB '}, 6),
                                 ({'climate_zone': 'IIБ'}, 1),
                                 ({'climate_zone': 'unknown'}, 0)):
                with self.subTest(endpoint=endpoint, extra=extra):
                    result = self.client.get(endpoint, query_string={
                        'q': prefix, 'format': 'json', **extra}).get_json()
                    self.assertEqual(result['total'], count)

    def test_admin_can_remove_passport_from_analogs_table(self):
        with app.app_context():
            passport = AnalogPassport(
                title='Аналог для удаления', object_type='котельная',
                work_type='строительство', capacity_kw=100,
                created_by=self.admin_id)
            db.session.add(passport)
            db.session.commit()
            passport_id = passport.id

        page = self.client.get('/analog-search/simple?view=base')
        html = page.get_data(as_text=True)
        self.assertIn('Аналог для удаления', html)
        self.assertIn(
            f'/analog-search/passports/{passport_id}/archive', html)
        self.assertIn('Удалить', html)

        response = self.client.post(
            f'/analog-search/passports/{passport_id}/archive',
            headers={'Accept': 'text/html'}, follow_redirects=False)
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.headers['Location'], '/analog-search/simple')
        self.assertNotIn(
            'Аналог для удаления',
            self.client.get('/analog-search/simple').get_data(as_text=True))
        with app.app_context():
            self.assertTrue(db.session.get(AnalogPassport, passport_id).archived)

    def test_regular_user_cannot_remove_passport(self):
        with app.app_context():
            passport = AnalogPassport(
                title='Защищённый аналог', object_type='котельная',
                work_type='строительство', capacity_kw=100,
                created_by=self.admin_id)
            user = User(username='analog-delete-user', password_hash='unused',
                        role='user')
            user.password_change_required = False
            db.session.add_all([passport, user])
            db.session.commit()
            passport_id = passport.id
            user_id, user_version = user.id, int(user.password_version or 1)
        self.login_values(user_id, user_version)
        self.assertEqual(
            self.client.post(
                f'/analog-search/passports/{passport_id}/archive').status_code,
            403)

    def test_preview_get_is_owned_and_confirmed_preview_redirects(self):
        response = self.client.post('/analog-search/import/preview',
                                    data=self.files(self.conclusion + b'\n \n \n'),
                                    content_type='multipart/form-data')
        import_id = response.get_json()['import_id']
        preview_url = f'/analog-search/imports/{import_id}/preview'
        self.assertEqual(self.client.get(preview_url).status_code, 200)

        with app.app_context():
            user = User(username='analog-preview-other', password_hash='unused',
                        role='user')
            user.password_change_required = False
            db.session.add(user)
            db.session.commit()
            user_id, user_version = user.id, int(user.password_version or 1)
        self.login_values(user_id, user_version)
        self.assertEqual(self.client.get(preview_url).status_code, 403)

        self.login_values(self.admin_id, self.admin_version)
        confirmed = self.client.post(
            f'/analog-search/import/{import_id}/confirm', json={})
        passport_id = confirmed.get_json()['passport_id']
        reopened = self.client.get(preview_url, follow_redirects=False)
        self.assertEqual(reopened.status_code, 302)
        self.assertEqual(reopened.headers['Location'], '/analog-search/simple')

    def test_reupload_restores_archived_analog_to_table(self):
        conclusion = self.conclusion + b'\n  \n'
        preview = self.client.post('/analog-search/import/preview',
                                   data=self.files(conclusion),
                                   content_type='multipart/form-data').get_json()
        confirmed = self.client.post(
            f"/analog-search/import/{preview['import_id']}/confirm", json={})
        passport_id = confirmed.get_json()['passport_id']
        with app.app_context():
            version_count = len(
                db.session.get(AnalogPassport, passport_id).versions)
        archived = self.client.post(
            f'/analog-search/passports/{passport_id}/archive',
            headers={'Accept': 'text/html'})
        self.assertEqual(archived.status_code, 302)
        with app.app_context():
            self.assertTrue(db.session.get(AnalogPassport, passport_id).archived)

        duplicate = self.client.post(
            '/analog-search/import/preview', data=self.files(conclusion),
            content_type='multipart/form-data',
            headers={'Accept': 'text/html'})
        self.assertEqual(duplicate.status_code, 302)
        self.assertEqual(duplicate.headers['Location'], '/analog-search/simple')
        page = self.client.get('/analog-search/simple').get_data(as_text=True)
        self.assertIn('снова добавлен в общую таблицу аналогов', page)
        with app.app_context():
            passport = db.session.get(AnalogPassport, passport_id)
            self.assertFalse(passport.archived)
            self.assertEqual(len(passport.versions), version_count)

    def test_html_xhr_parse_error_is_actionable_json(self):
        response = self.client.post(
            '/analog-search/import/preview',
            data=self.files(b'<broken'),
            content_type='multipart/form-data',
            headers={'Accept': 'text/html',
                     'X-Requested-With': 'XMLHttpRequest'})
        self.assertEqual(response.status_code, 422)
        payload = response.get_json()
        self.assertEqual(payload['status'], 'error')
        self.assertEqual(payload['diagnostics'][0]['code'], 'malformed_xml')

    def test_import_form_has_upload_and_processing_states(self):
        page = self.client.get('/analog-search/simple?view=base')
        html = page.get_data(as_text=True)
        self.assertIn('analog-import-progress', html)
        self.assertIn('xhr.upload.onprogress', html)
        self.assertIn('Сервер проверяет и разбирает данные', html)
        self.assertIn('button.disabled = true', html)
        self.assertIn('window.location.assign(finalUrl)', html)
        self.assertNotIn('document.write(xhr.responseText)', html)

    def test_confirm_is_owned_by_creator_or_admin(self):
        response = self.client.post('/analog-search/import/preview',
                                    data=self.files(self.conclusion + b'\n\n'),
                                    content_type='multipart/form-data')
        import_id = response.get_json()['import_id']
        with app.app_context():
            user = User(username='analog-regular', password_hash='unused', role='user')
            user.password_change_required = False
            db.session.add(user)
            db.session.commit()
            user_id, user_version = user.id, int(user.password_version or 1)
        self.login_values(user_id, user_version)
        denied = self.client.post(f'/analog-search/import/{import_id}/confirm',
                                  json={})
        self.assertEqual(denied.status_code, 403)

    def test_climate_suggestion_requires_explicit_confirmation(self):
        with app.app_context():
            db.session.add_all([
                AnalogClimateReference(
                    address_pattern='Омская область', climate_zone='IВ',
                    climate_subzone='', source='СП 131.13330',
                    reference_version='2025-general', reference_date=date(2025, 1, 1)),
                AnalogClimateReference(
                    address_pattern='город Омск', climate_zone='IВ',
                    climate_subzone='IВ-2', source='СП 131.13330, приложение А',
                    reference_version='2025-omsk', reference_date=date(2025, 2, 1)),
            ])
            passport = AnalogPassport(
                title='Климатический тест', address='Омская область, город Омск',
                object_type='котельная', work_type='строительство', capacity_kw=10,
                created_by=self.admin_id)
            db.session.add(passport)
            db.session.commit()
            passport_id = passport.id
        proposal = self.client.get(
            '/analog-search/api/climate-suggestion',
            query_string={'address': 'Омская область, город Омск'}).get_json()
        self.assertEqual(proposal['reference_version'], '2025-omsk')
        rejected = self.client.post(
            f'/analog-search/passports/{passport_id}/climate/confirm',
            json={'reference_version': '2025-omsk'})
        self.assertEqual(rejected.status_code, 400)
        confirmed = self.client.post(
            f'/analog-search/passports/{passport_id}/climate/confirm',
            json={'confirmed': True, 'reference_version': '2025-omsk'})
        self.assertEqual(confirmed.status_code, 200)
        with app.app_context():
            saved = db.session.get(AnalogPassport, passport_id)
            self.assertEqual(saved.climate_subzone, 'IВ-2')
            self.assertEqual(saved.climate_reference_version, '2025-omsk')
            self.assertEqual(saved.climate_reference_date, date(2025, 2, 1))

    def test_versioned_cost_index_is_selected_and_saved_with_passport_version(self):
        with app.app_context():
            db.session.add_all([
                AnalogCostIndexReference(
                    effective_date=date(2025, 1, 1), price_level='2020/Q1',
                    factor='1.25', source='Согласованный индекс Минстроя',
                    reference_version='indices-2025-q1', reference_date=date(2025, 1, 10)),
                AnalogCostIndexReference(
                    effective_date=date(2025, 7, 1), price_level='2020/Q1',
                    factor='1.40', source='Согласованный индекс Минстроя',
                    reference_version='indices-2025-q3', reference_date=date(2025, 7, 10)),
            ])
            passport = AnalogPassport(
                title='Индексный тест', object_type='котельная',
                work_type='строительство', capacity_kw=10, created_by=self.admin_id)
            db.session.add(passport)
            db.session.flush()
            version = AnalogPassportVersion(
                passport_id=passport.id, version_no=1, conclusion_number='cost-test',
                examination_result='positive', source_hash='cost-source',
                ssr_hash='cost-ssr', total_cost=100, imported_by=self.admin_id)
            db.session.add(version)
            db.session.flush()
            passport.current_version_id = version.id
            db.session.commit()
            passport_id, version_id = passport.id, version.id
        missing = self.client.get('/analog-search/api/cost-index-suggestion',
                                  query_string={'price_level': 'unknown',
                                                'date': '2025-08-01'}).get_json()
        self.assertEqual(missing['status'], 'missing')
        rejected = self.client.post(
            f'/analog-search/passports/{passport_id}/versions/{version_id}/adjust',
            json={'price_level': '2020/Q1', 'calculation_date': '2025-08-01',
                  'reference_version': 'indices-2025-q3'})
        self.assertEqual(rejected.status_code, 400)
        saved = self.client.post(
            f'/analog-search/passports/{passport_id}/versions/{version_id}/adjust',
            json={'price_level': '2020/Q1', 'calculation_date': '2025-08-01',
                  'reference_version': 'indices-2025-q3', 'confirmed': True})
        self.assertEqual(saved.status_code, 200)
        with app.app_context():
            version = db.session.get(AnalogPassportVersion, version_id)
            self.assertEqual(str(version.adjusted_cost), '140.000000')
            self.assertEqual(version.adjustment_reference_version, 'indices-2025-q3')
            self.assertEqual(version.adjustment_reference_date, date(2025, 7, 10))


if __name__ == '__main__':
    unittest.main()
