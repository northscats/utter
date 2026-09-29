"""Export one analogue's estimate into the owner's NMCK example workbook.

SSR chapter subtotals are cumulative, not additive. Its stored cost figures
are in thousands of rubles; the example workbook uses rubles. Unverifiable
categories stay empty rather than turning missing data into zero.
"""
import re
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill


TEMPLATE = (Path(__file__).resolve().parent.parent / 'attached_assets'
            / 'ЖуковкаАналог_(2)_1790673815845.xlsx')
THOUSAND = Decimal('1000')
MONEY_CELLS = {11: 'D11', 12: 'D12', 13: 'D13', 14: 'D14',
               15: 'E15', 16: 'D16', 17: 'D17', 18: 'E18', 19: 'D19'}
EXCLUDED = re.compile(r'зимн|утилиз|отход|непредвид|резерв средств', re.I)
COMMISSION = re.compile(r'пуско.?налад', re.I)
AUTHOR = re.compile(r'авторск\w*\s+надзор', re.I)
EXPERTISE = re.compile(r'экспертиз', re.I)
SURVEY = re.compile(r'изыскан|обследован', re.I)
COMBINED_DESIGN = re.compile(r'проектно.?изыскательск|\bпир\b', re.I)
DESIGN = re.compile(r'проектн|проектир|техническ\w*\s+документац|рабоч\w*\s+документац', re.I)
NON_PIR = re.compile(r'строительн\w*\s+контрол|публичн\w*\s+аудит', re.I)


def _money(value):
    amount = Decimal(str(value)) if value is not None else None
    return amount * THOUSAND if amount is not None and amount.is_finite() else None


def estimate_values(chapters, items):
    """Map workbook row to (rubles or None, source/explanation)."""
    result = {row: (None, 'Нет подтверждённой детализации в ССР')
              for row in MONEY_CELLS}
    chapter8 = next((ch for ch in chapters if ch.number == 8), None)
    if chapter8 is not None:
        excluded_in_chapter = [
            item for item in items
            if 1 <= item.chapter_number <= 8
            and EXCLUDED.search(item.name or '')
            and item.total and item.total > 0
        ]
        for row, field in ((12, 'building'), (13, 'mounting'), (15, 'equipment')):
            value = _money(getattr(chapter8, field))
            if excluded_in_chapter:
                # Do not guess which cost column contains an excluded item.
                result[row] = (None, 'В главах 1–8 есть исключаемые позиции: нужна ручная сверка')
            elif value is not None and value != 0:
                label = {'building': 'строительные работы',
                         'mounting': 'монтажные работы',
                         'equipment': 'оборудование'}[field]
                result[row] = (value, f'ССР, накопительный итог главы 8: {label}')

    commissioning = [
        item for item in items
        if 1 <= item.chapter_number <= 13
        and COMMISSION.search(item.name or '') and item.total is not None
    ]
    if commissioning:
        result[14] = (sum((_money(item.total) or Decimal('0')
                           for item in commissioning), Decimal('0')),
                      'ССР, отдельные позиции «Пусконаладочные работы», главы '
                      + ', '.join(str(item.chapter_number) for item in commissioning))
    if all(result[row][0] is not None for row in (12, 13, 14)):
        result[11] = (sum(result[row][0] for row in (12, 13, 14)),
                      'Строительные + монтажные + пусконаладочные; '
                      'как в предоставленном образце (ПНР могут находиться в главе 9)')

    # Chapter 12 usually holds design, surveys, expertise and author supervision.
    # A general chapter total is not PIR: it can include unrelated expenses.
    classified = {17: [], 18: [], 19: []}
    combined = []
    unknown = []
    for item in items:
        if item.chapter_number != 12 or item.total is None or item.total <= 0:
            continue
        name = item.name or ''
        if AUTHOR.search(name) or NON_PIR.search(name):
            continue
        if EXPERTISE.search(name):
            classified[19].append(item)
        elif COMBINED_DESIGN.search(name):
            combined.append(item)
        elif SURVEY.search(name):
            classified[18].append(item)
        elif DESIGN.search(name):
            classified[17].append(item)
        else:
            unknown.append(item)
    for row, matches in classified.items():
        if matches:
            result[row] = (sum(_money(item.total) for item in matches),
                           'ССР, глава 12: ' +
                           '; '.join(item.name or item.ordinal or 'позиция'
                                     for item in matches))
    if (combined or any(classified.values())) and not unknown:
        if combined and (classified[17] or classified[18]):
            result[16] = (None, 'Есть общий ПИР и отдельные позиции: '
                               'возможен двойной счёт, нужна сверка')
        else:
            total = sum((_money(item.total) for item in combined), Decimal('0'))
            total += sum((result[row][0] or 0 for row in (17, 18, 19)),
                         Decimal('0'))
            result[16] = (total, 'ССР, глава 12: ПИР и экспертиза без авторского надзора')
    elif unknown:
        result[16] = (None, 'В главе 12 есть неклассифицированные затраты: '
                            + '; '.join(item.name or 'без названия'
                                        for item in unknown))
    return result


def build_cost_workbook(passport, version, capacity_gcal_h, chapters, items):
    """Reuse the provided workbook's formatting; clear every example amount."""
    workbook = load_workbook(TEMPLATE)
    sheet = workbook.worksheets[0]
    title = (passport.title or version.expert_object_name
             if version else passport.title) or 'Объект-аналог'
    sheet['D9'] = title
    # A source title beginning with "=" must not become an Excel formula.
    sheet['D9'].data_type = 's'
    if version:
        source = (f'1.2. Объект-аналог: заключение экспертизы '
                  f'№{version.conclusion_number or "не указан"}'
                  + (f' от {version.approval_date:%d.%m.%Y}'
                     if version.approval_date else '')
                  + f'. {title}')
    else:
        source = f'1.2. Объект-аналог: {title}. Нет утверждённой версии ССР.'
    sheet['A6'] = source
    # Write the same two-decimal capacity the operator sees in the catalog,
    # not the long repeating result of the kW -> Гкал/ч conversion.
    sheet['D10'] = (float(Decimal(str(capacity_gcal_h)).quantize(
        Decimal('0.01'), rounding=ROUND_HALF_UP))
        if capacity_gcal_h is not None else None)
    sheet['D10'].number_format = '0.00'
    results = estimate_values(chapters, items) if version else {
        row: (None, 'Нет утверждённой версии ССР') for row in MONEY_CELLS}
    if results[15][0] is not None:
        results[15] = (results[15][0],
                       results[15][1] + '; в ССР не подтверждено, что '
                       'всё оборудование является технологическим')
    for row, cell in MONEY_CELLS.items():
        amount, provenance = results[row]
        sheet[cell] = float(amount) if amount is not None else None
        sheet[cell].comment = Comment(provenance, 'Источник суммы')
    if results[15][0] is not None:
        sheet['E15'].fill = PatternFill('solid', fgColor='FFF0BF')

    audit = workbook.create_sheet('Источники сумм')
    audit.append(['Строка образца', 'Наименование', 'Сумма, руб.',
                  'Источник / причина пропуска'])
    for row, cell in MONEY_CELLS.items():
        amount, provenance = results[row]
        audit.append([row, sheet[f'B{row}'].value,
                      float(amount) if amount is not None else None, provenance])
        audit.cell(audit.max_row, 3).number_format = '#,##0.00'
    audit.append([])
    audit.append(['Методика', 'Суммы ССР хранятся в тыс. руб.; '
                  'в этом файле они переведены в рубли (× 1000).'])
    audit.append(['Методика', 'Итоги глав ССР накопительные; главы 1–8 не суммируются.'])
    audit.append(['Методика', 'СМР включает отдельно указанную ПНР, '
                  'даже если она находится в главе 9, по образцу владельца.'])
    audit.append(['Методика', 'Пустые суммы не означают ноль. '
                  'Справочные листы индексов исходного файла не пересчитывались.'])
    audit.freeze_panes = 'A2'
    audit.column_dimensions['A'].width = 21
    audit.column_dimensions['B'].width = 69
    audit.column_dimensions['C'].width = 23
    audit.column_dimensions['D'].width = 105
    for cell in audit[1]:
        cell.fill = PatternFill('solid', fgColor='174A7E')
        cell.font = Font(bold=True, color='FFFFFF')
    for row in audit.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    audit.auto_filter.ref = 'A1:D10'
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output