# -*- coding: utf-8 -*-
"""ОСВ Дашборд — веб-приложение с PostgreSQL, авторизацией и мультипредприятием."""
import json
import gzip
import os
import shutil
import calendar
import hashlib
from io import BytesIO
from datetime import datetime, timedelta, timezone
import re
import time
import traceback
from urllib.parse import urlparse

import threading

_limits_history_ready = threading.Event()
_limits_parse_slots = threading.Semaphore(1)

import bcrypt
from flask import (Flask, Response, abort, flash, g, has_request_context, jsonify,
                   redirect, render_template, request, send_file, send_from_directory,
                   session, url_for)
from flask_login import (LoginManager, current_user, login_required,
                         login_user, logout_user)
from flask_wtf.csrf import CSRFProtect, CSRFError, generate_csrf
from functools import wraps
from itsdangerous import BadSignature, SignatureExpired, URLSafeTimedSerializer
from sqlalchemy import and_, func, or_
from sqlalchemy.exc import IntegrityError

from jobs import create_job, get_job, finish_job, fail_job

from werkzeug.middleware.dispatcher import DispatcherMiddleware
from werkzeug.middleware.proxy_fix import ProxyFix

from database import (AiProvider, AiSummary, OsvFileAiSummary, AuditLog, Company, CompanyContact,
                       CompanyDadataProfile, CompanyDadataSnapshot,
                       DadataDiagnosticEvent,
                       CompanyFinance, CompanyLicense, CompanyMembership,
                       CompanyOkved, DailyRow, EmailRule, HistoricalImport,
                       HistoricalMetricValue, HistoricalRawValue, HistoricalMapping,
                       HistoricalYearVersion, RatingProfile, RevenueAnalyticMapping,
                       CostAnalyticMapping,
                        DebtorAnalyticMapping, ConsumerCollectionAccrual,
                        ConsumerPaymentAllocation, ConsumerCollectionNorm,
                       MetricLayout,
                       Setting, Upload, User,
                       UserModuleAccess, OneCChannel, OneCAttempt,
                       OneCIdempotency, AnalogPassport, AnalogPassportVersion,
                       AnalogSource, AnalogOrganizationRole, AnalogTep,
                       AnalogDocument, AnalogCostRow, AnalogSSRChapter,
                       AnalogSSRItem, AnalogSSRSignature, AnalogDiagnostic,
                       AnalogNetworkAttribute, AnalogImport,
                       LimitChange, LimitComparison,
                       LimitComparisonResolution, LimitDataset, LimitDelivery,
                       LimitDeliveryAttempt, LimitMaxEvent, LimitBotDialog,
                       LimitBotOutbound, LimitModerationRequest, LimitObject,
                       LimitRecipient, LimitRowSnapshot, LimitSubscription,
                       LimitVersion, LimitVersionTombstone,
                       LimitDeliveryPart, db, utcnow)
from directory_service import (accept_dadata, apply_normalized_to_company,
                                clean_identifier, company_identity_key,
                                duplicate_companies,
                                save_snapshot)
from company_creation import resolve_existing_company, serialized_company_creation
from module_registry import (module_key_for_path, registered_modules,
                             valid_module_keys)
from osvparser import extract_org_name, parse_many, parse_zip
from source_manifest import accepted_parsed_files, build_manifest, dump_manifest
from builder import build_dataset
from dashboard_compat import (json_for_script,
                               upgrade_legacy_salary_daily_labels,
                               upgrade_saved_dashboard_template)
from metric_layouts import (
    RESERVED_DASHBOARD_IDS, LayoutValidationError, validate_layout_payload,
    validate_mode_metric,
)
from password_policy import validate_password

# ---------------------------------------------------------------------------
BASE = os.path.dirname(os.path.abspath(__file__))
UPLOADS = os.path.join(BASE, 'uploads')
OUTPUT = os.path.join(BASE, 'output')
DOCS_STORAGE = os.path.join(BASE, 'uploads', 'docs')
os.makedirs(UPLOADS, exist_ok=True)
os.makedirs(OUTPUT, exist_ok=True)
os.makedirs(DOCS_STORAGE, exist_ok=True)

ANALYTICAL_ACCOUNT_NAMES = {
    '01': 'Основные средства',
    '08': 'Вложения во внеоборотные активы',
    '10': 'Материалы',
    '20': 'Основное производство',
    '26': 'Общехозяйственные расходы',
    '44': 'Расходы на продажу',
    '50': 'Касса',
    '51': 'Расчётные счета',
    '52': 'Валютные счета',
    '55': 'Специальные счета в банках',
    '60': 'Расчёты с поставщиками и подрядчиками',
    '62': 'Расчёты с покупателями и заказчиками',
    '66': 'Краткосрочные кредиты и займы',
    '67': 'Долгосрочные кредиты и займы',
    '68': 'Расчёты по налогам и сборам',
    '69': 'Расчёты по социальному страхованию и обеспечению',
    '70': 'Расчёты с персоналом по оплате труда',
    '76': 'Расчёты с разными дебиторами и кредиторами',
    '90': 'Продажи',
    '91': 'Прочие доходы и расходы',
}

_DURABLE_MODELS = (
    User, Company, Upload, DailyRow, HistoricalImport, HistoricalYearVersion,
    HistoricalMetricValue, HistoricalRawValue, HistoricalMapping,
    RevenueAnalyticMapping, CostAnalyticMapping, DebtorAnalyticMapping,
    ConsumerCollectionAccrual, ConsumerPaymentAllocation,
    ConsumerCollectionNorm,
    AnalogPassport, AnalogPassportVersion, AnalogSource,
    AnalogOrganizationRole, AnalogTep, AnalogDocument, AnalogCostRow,
    AnalogSSRChapter, AnalogSSRItem, AnalogDiagnostic,
    AnalogSSRSignature, AnalogNetworkAttribute, AnalogImport,
    LimitDataset, LimitVersion, LimitObject, LimitRowSnapshot, LimitComparison,
    LimitVersionTombstone,
    LimitComparisonResolution, LimitChange, LimitRecipient, LimitSubscription,
    LimitModerationRequest, LimitMaxEvent, LimitBotDialog, LimitBotOutbound,
    LimitDelivery, LimitDeliveryAttempt, LimitDeliveryPart,
)
def _durable_backup(reason, allow_removals=False, required=False,
                    allowed_removed_tables=()):
    """A backup failure must be visible, but cannot undo a committed upload."""
    if has_request_context():
        g.durable_backup_done = True
    if app.config.get('TESTING') and not os.environ.get('DURABLE_BACKUP_DIR'):
        return
    from durable_data import backup
    try:
        result = backup(
            db, UPLOADS, OUTPUT, _DURABLE_MODELS, reason,
            allow_removals=allow_removals,
            allowed_removed_tables=allowed_removed_tables)
        if not result['promoted'] and result['shrink_rejected']:
            message = (
                f'{reason}: latest не обновлён из-за уменьшения таблиц: '
                f'{", ".join(result["shrink_rejected"])}')
            app.logger.error('durable backup rejected: %s', message)
            Setting.set('durable_backup_error', message)
            db.session.commit()
        elif result.get('superseded'):
            app.logger.info(
                'durable backup superseded by a newer snapshot after %s',
                reason)
        app.logger.info('durable backup completed after %s: %s', reason, result)
        return result
    except Exception as exc:
        app.logger.exception('durable backup failed after %s', reason)
        try:
            Setting.set('durable_backup_error', f'{reason}: {exc}')
            db.session.commit()
        except Exception:
            db.session.rollback()
            app.logger.exception('could not record durable backup error')
        if required:
            raise
        return None


_durable_safety_lock = threading.Lock()
_durable_safety_pending = []
_durable_safety_thread = None


def _run_durable_safety_backups():
    """Serialize and coalesce safety-net snapshots outside HTTP workers."""
    global _durable_safety_thread
    while True:
        with _durable_safety_lock:
            reasons = list(dict.fromkeys(_durable_safety_pending))
            _durable_safety_pending.clear()
            if not reasons:
                _durable_safety_thread = None
                return
        reason = '; '.join(reasons[:5])
        if len(reasons) > 5:
            reason += f'; ещё изменений: {len(reasons) - 5}'
        with app.app_context():
            _durable_backup(f'асинхронная страховка: {reason}')


def _schedule_durable_safety_backup(reason):
    """Queue one background snapshot without delaying the browser response."""
    global _durable_safety_thread
    if (app.config.get('TESTING')
            and not app.config.get('DURABLE_BACKUP_ASYNC_TEST')):
        _durable_backup(reason)
        return
    with _durable_safety_lock:
        _durable_safety_pending.append(reason)
        if _durable_safety_thread and _durable_safety_thread.is_alive():
            return
        _durable_safety_thread = threading.Thread(
            target=_run_durable_safety_backups,
            name='durable-safety-backup',
            daemon=True,
        )
        _durable_safety_thread.start()


def _durable_restore():
    """Merge durable data after schema creation and report its status."""
    if app.config.get('TESTING') and not os.environ.get('DURABLE_BACKUP_DIR'):
        return True
    from durable_data import restore
    try:
        result = restore(db, UPLOADS, OUTPUT, _DURABLE_MODELS)
        app.logger.info('durable restore completed: %s', result)
        return True
    except Exception as exc:
        db.session.rollback()
        app.logger.exception('durable restore failed')
        try:
            Setting.set('durable_restore_error', str(exc))
            db.session.commit()
        except Exception:
            db.session.rollback()
            app.logger.exception('could not record durable restore error')
        return False

# Тот же набор из 20 аналитических счетов используется для оценки
# комплектности загрузок и для блока покрытия данными.
REQUIRED_ACCOUNTS = list(ANALYTICAL_ACCOUNT_NAMES)

# Выходные файлы, которые заменяются при успешной обработке ОСВ и поэтому
# должны быть восстановлены, если фоновая обработка завершится ошибкой.
_COMPANY_OUTPUT_NAMES = ('dashboard_current.html', 'kpi_cache.json')

# Префикс пути, если приложение работает не в корне домена.
# Задаётся через переменную окружения: SCRIPT_NAME=/osv
# При развёртывании в корне оставьте переменную пустой или не задавайте её.
SCRIPT_NAME = os.environ.get('SCRIPT_NAME', '').rstrip('/')

app = Flask(__name__)
app.secret_key = os.environ.get('SESSION_SECRET', 'dev-secret-change-me')
# 200 МБ: покрывает ZIP-пакеты документов (до 100 МБ на файл внутри архива
# + служебные данные ZIP); лимиты распакованного содержимого — в doc_pipeline.py
app.config['MAX_CONTENT_LENGTH'] = 200 * 1024 * 1024
app.config['WTF_CSRF_TIME_LIMIT'] = None   # токен живёт до истечения сессии
app.config['HISTORICAL_IMPORT_ENABLED'] = (
    os.environ.get('HISTORICAL_IMPORT_ENABLED', '1').strip().lower()
    not in ('0', 'false', 'no', 'off'))
app.config['HISTORICAL_EXPERT_ENABLED'] = (
    os.environ.get('HISTORICAL_EXPERT_ENABLED', '1').strip().lower()
    not in ('0', 'false', 'no', 'off'))
app.config['HISTORICAL_EXECUTIVE_ENABLED'] = (
    os.environ.get('HISTORICAL_EXECUTIVE_ENABLED', '1').strip().lower()
    not in ('0', 'false', 'no', 'off'))
app.config['PORTFOLIO_HOME_ENABLED'] = (
    os.environ.get('PORTFOLIO_HOME_ENABLED', '1').strip().lower()
    not in ('0', 'false', 'no', 'off'))
app.config['COMPANY_COMPARISON_ENABLED'] = (
    os.environ.get('COMPANY_COMPARISON_ENABLED', '1').strip().lower()
    not in ('0', 'false', 'no', 'off'))

# Если задан подпуть — сообщаем Flask, где он «живёт»
if SCRIPT_NAME:
    app.config['APPLICATION_ROOT'] = SCRIPT_NAME
    app.config['SESSION_COOKIE_PATH'] = SCRIPT_NAME
    # Доверяем заголовкам X-Forwarded-* от nginx/Apache
    app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1, x_proto=1, x_host=1, x_prefix=1)

# --- PostgreSQL —  DATABASE_URL из переменной окружения ---
app.config['SQLALCHEMY_DATABASE_URI'] = os.environ['DATABASE_URL']
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False
app.config['SQLALCHEMY_ENGINE_OPTIONS'] = {'hide_parameters': True}

db.init_app(app)
csrf = CSRFProtect(app)


@app.after_request
def _durable_mutation_safety_net(response):
    """Capture successful user mutations not covered by an explicit backup."""
    if (request.method in ('POST', 'PUT', 'PATCH', 'DELETE')
            and response.status_code < 400
             and not getattr(g, 'durable_backup_done', False)
             and not getattr(g, 'skip_durable_backup', False)):
        _schedule_durable_safety_backup(
            f'http_{request.method.lower()}:{request.endpoint or "unknown"}')
    return response
_LOGIN_CSRF = URLSafeTimedSerializer(
    app.secret_key, salt='login-csrf-v1')

# --- Универсальный парсер документов (/docs/*) ---
app.config['DOCS_STORAGE'] = DOCS_STORAGE
from docs_routes import docs_bp
app.register_blueprint(docs_bp)
from analog_routes import analog_bp
app.register_blueprint(analog_bp)

# ---------------------------------------------------------------------------
# Шифрование значений в БД (Fernet / AES-128-CBC)
# ---------------------------------------------------------------------------
import sys as _sys
from cryptography.fernet import Fernet as _Fernet

_ENCRYPTION_KEY_RAW = os.environ.get('ENCRYPTION_KEY', '')
if not _ENCRYPTION_KEY_RAW:
    print(
        '\n[FATAL] Переменная окружения ENCRYPTION_KEY не задана.\n'
        'Сгенерируйте ключ:\n'
        '  python3 -c "import secrets,base64; '
        'print(base64.urlsafe_b64encode(secrets.token_bytes(32)).decode())"\n'
        'и сохраните как Replit Secret с именем ENCRYPTION_KEY.\n',
        file=_sys.stderr, flush=True,
    )
    _sys.exit(1)
try:
    _FERNET = _Fernet(_ENCRYPTION_KEY_RAW.encode())
except Exception as _fernet_err:
    print(f'[FATAL] Неверный формат ENCRYPTION_KEY: {_fernet_err}',
          file=_sys.stderr, flush=True)
    _sys.exit(1)

# ---------------------------------------------------------------------------
# IMAP background thread state (single-worker gunicorn keeps this in one process)
# ---------------------------------------------------------------------------
_imap_stop = threading.Event()           # set() to request graceful shutdown
_imap_thread: threading.Thread | None = None
_imap_thread_started = False
_imap_thread_lock = threading.Lock()
_limits_history_ready = threading.Event()
_limits_parse_slots = threading.Semaphore(1)
_limits_recovery_started = False
_limits_recovery_lock = threading.Lock()

login_manager = LoginManager(app)
login_manager.login_view = 'login'
login_manager.login_message = None


@login_manager.user_loader
def load_user(user_id):
    return db.session.get(User, int(user_id))


def _ensure_imap_thread() -> None:
    """Start the IMAP background thread once per process (idempotent)."""
    global _imap_thread, _imap_thread_started
    if _imap_thread_started:
        return
    with _imap_thread_lock:
        if _imap_thread_started:
            return
        _imap_thread_started = True
    _imap_thread = threading.Thread(target=_imap_loop, daemon=True, name='imap-loop')
    _imap_thread.start()

def _ensure_limits_recovery() -> None:
    """Start consumers only after restore and history reconciliation."""
    global _limits_recovery_started
    if app.config.get('TESTING') or not _limits_history_ready.is_set():
        return
    if _limits_recovery_started:
        return
    with _limits_recovery_lock:
        if _limits_recovery_started:
            return
        _limits_recovery_started = True

    def work():
        with app.app_context():
            for row in LimitVersion.query.filter(
                    LimitVersion.status.in_(('queued', 'parsing'))).all():
                row.status = 'queued'
            db.session.commit()
            for row in LimitVersion.query.filter_by(status='queued').all():
                _limit_parse_in_background(row.id)
            from limits_bot import start_worker
            start_worker(app, _limits_adapter, lambda: _limits_history_ready.is_set())
        while True:
            time.sleep(3600)
            with app.app_context():
                try:
                    from limits_service import cleanup_purged_files, cleanup_temporary_versions
                    cleanup_purged_files()
                    cleanup_temporary_versions(int(os.environ.get('LIMITS_TEMP_TTL_HOURS', '72')))
                    if Setting.get('limits_rotation_needs_backup') == '1':
                        _schedule_limits_rotation_backup()
                except Exception:
                    db.session.rollback()
                    app.logger.exception('limits periodic cleanup failed')

    threading.Thread(target=work, name='limits-recovery', daemon=True).start()
@app.before_request
def _maybe_start_imap_thread():
    """Lazy IMAP thread starter — fires on the first HTTP request."""
    if request.endpoint in ('healthz', 'healthz_watchdog_recovery'):
        return None
    _ensure_imap_thread()
    _ensure_limits_recovery()


@app.before_request
def _enforce_password_lifecycle():
    """Bind sessions to the password version and enforce forced changes."""
    if request.endpoint == 'static' or request.path.startswith('/static/'):
        return None
    if not current_user.is_authenticated:
        return None

    current_version = int(current_user.password_version or 1)
    session_version = session.get('password_version')
    if session_version is None:
        # Do not bless a session created before the version binding existed:
        # it may be a stolen/restored cookie from before a password change.
        logout_user()
        flash('Сессия устарела. Войдите снова.')
        return redirect(url_for('login'))
    try:
        version_matches = int(session_version) == current_version
    except (TypeError, ValueError):
        version_matches = False
    if not version_matches:
        logout_user()
        flash('Пароль был изменён. Войдите снова.')
        return redirect(url_for('login'))

    if request.endpoint in ('login', 'logout', 'healthz',
                            'healthz_watchdog_recovery'):
        return None
    if (current_user.password_change_required
            and request.endpoint != 'password_change_required'):
        return redirect(url_for('password_change_required'))


@app.route('/healthz', methods=['GET'])
def healthz():
    """Process-local liveness probe with no database or storage access."""
    return {'status': 'ok'}, 200
def admin_required(f):
    """Декоратор: доступ только для администраторов."""
    @wraps(f)
    @login_required
    def decorated(*args, **kwargs):
        if not current_user.is_admin:
            abort(403)
        return f(*args, **kwargs)
    return decorated


# Машинный POST освобождён от CSRF внутри init_one_c; административные
# операции намеренно остаются под общей Flask-WTF защитой.
from one_c import init_one_c
init_one_c(app, csrf, admin_required)


@app.before_request
def _enforce_module_access():
    """Защищает маршруты областями, объявленными в едином реестре."""
    if not current_user.is_authenticated:
        return None
    module_key = module_key_for_path(request.path)
    if module_key and not current_user.has_module_access(module_key):
        abort(403, description='Нет доступа к этому модулю')
    company_match = re.match(r'^/(?:company|dashboard)/(\d+)(?:/|$)', request.path)
    if company_match and not _company_is_visible(int(company_match.group(1))):
        abort(403, description='Нет доступа к предприятию')
    any_company_match = re.match(
        r'^/(?:company|dashboard|directories/companies|api/companies)/(\d+)(?:/|$)',
        request.path)
    if any_company_match and request.method not in ('GET', 'HEAD', 'OPTIONS'):
        company = db.session.get(Company, int(any_company_match.group(1)))
        is_restore = (
            request.path.endswith('/archive')
            and current_user.is_admin
            and request.form.get('archive') == '0')
        is_permanent_delete = (
            request.path.startswith('/directories/companies/')
            and request.path.endswith('/delete'))
        if (company and company.is_archived
                and not (is_restore or is_permanent_delete)):
            abort(409, description='Архивное предприятие доступно только для чтения')


def _company_is_visible(company_id):
    """Серверная область видимости предприятия для всех новых и legacy URL."""
    if current_user.is_admin:
        return True
    return CompanyMembership.query.filter_by(
        user_id=current_user.id, company_id=company_id).first() is not None


def _require_company_access(company_id):
    company = db.session.get(Company, company_id) or abort(404)
    if not _company_is_visible(company_id):
        abort(403, description='Нет доступа к предприятию')
    return company


# ---------------------------------------------------------------------------
# Jinja-фильтр для русского склонения
# ---------------------------------------------------------------------------
@app.template_filter('pluralize')
def pluralize_filter(n, one, few, many):
    n = abs(int(n))
    if 11 <= n % 100 <= 19:
        return many
    r = n % 10
    if r == 1:
        return one
    if 2 <= r <= 4:
        return few
    return many


# ---------------------------------------------------------------------------
# Вспомогательные функции
# ---------------------------------------------------------------------------
def render_dashboard(dataset, back_url='/'):
    tpl_path = os.path.join(BASE, 'templates', 'dashboard_template.html')
    with open(tpl_path, encoding='utf-8') as f:
        tpl = f.read()
    tpl = tpl.replace('/*__DATA__*/', json_for_script(dataset), 1)
    tpl = tpl.replace('/*__BACK_URL__*/', back_url, 1)
    return tpl


def _compute_completeness(upload):
    """Возвращает словарь с данными о комплектности счетов для upload."""
    if upload is None:
        return None
    return _compute_completeness_from_accounts(upload.accounts_list)


def _dataset_accounts(dataset):
    """Нормализованный список счетов, фактически вошедших в dataset."""
    return sorted({
        str(account).strip()
        for account in (dataset.get('accounts_loaded', []) or [])
        if account is not None and str(account).strip()
    })


def _compute_completeness_from_accounts(accounts):
    """Оценивает комплектность по совокупному набору номеров счетов."""
    loaded = set(str(a) for a in accounts)
    required = set(str(a) for a in REQUIRED_ACCOUNTS)
    found = len(loaded & required)
    total = len(required)
    missing = sorted(required - loaded, key=lambda x: int(x))
    if found == total:
        status = 'full'
    elif found == 0:
        status = 'none'
    else:
        status = 'partial'
    return dict(found=found, total=total, missing=missing, status=status)


def _compute_company_completeness(uploads, additional_accounts=None):
    """Комплектность предприятия по всем принятым загрузкам, а не последней."""
    loaded = set()
    for upload in uploads:
        loaded.update(str(a) for a in upload.accounts_list)
    loaded.update(str(a) for a in (additional_accounts or []))
    if not uploads and not loaded:
        return None
    return _compute_completeness_from_accounts(loaded)


def _compute_status_color(kpis):
    """Вычисляет цвет статус-индикатора по данным KPI: green / yellow / red / grey."""
    if kpis is None:
        return 'grey'
    score = kpis.get('score', 0)
    if score >= 70:
        return 'green'
    elif score >= 40:
        return 'yellow'
    return 'red'


def _save_kpi_cache(dataset, out_dir):
    """Вычисляет KPI из dataset и сохраняет kpi_cache.json рядом с дашбордом."""
    from builder import extract_kpis
    try:
        kpis = extract_kpis(dataset)
        kpis['accounts_loaded'] = _dataset_accounts(dataset)
        with open(os.path.join(out_dir, 'kpi_cache.json'), 'w', encoding='utf-8') as f:
            json.dump(kpis, f, ensure_ascii=False)
    except Exception as exc:
        app.logger.warning(f'kpi_cache save failed: {exc}')

def _dataset_metrics_json(dataset):
    from builder import extract_kpis
    return json.dumps(extract_kpis(dataset)['metrics'], ensure_ascii=False)

def _accepted_monthly_files(paths, parsed_data, audit):
    """Return ``(path, parsed)`` only for files accepted by ``parse_many``.

    ``parse_many`` deliberately keeps only the first file for a duplicate
    account and reports parser failures/plans in its audit.  The persisted
    source manifest must follow that acceptance decision exactly; in
    particular, a file that merely exists on disk must not become provenance
    evidence.
    """
    return accepted_parsed_files(paths, parsed_data, audit)


def _accepted_daily_files(paths):
    """Return manifest-ready evidence for each successfully parsed daily OSV."""
    from osvparser import extract_daily_date, parse_osv
    accepted = []
    for path in paths:
        try:
            parsed = dict(parse_osv(path))
        except Exception:
            continue
        observed = list(parsed.get('daily') or [])
        direct_date = extract_daily_date(path)
        if direct_date is not None:
            observed.append({'date': direct_date.isoformat()})
        if not observed:
            observed.extend({
                'date': item['date'].isoformat()
            } for item in _derive_cash_daily_rows(parsed))
        parsed['daily'] = observed
        accepted.append((path, parsed))
    return accepted


def _schedule_osv_file_summaries(company_id):
    """Best-effort post-commit orchestration shared by every ingestion path."""
    try:
        from osv_ai_summary import ensure_file_summaries
        ensure_file_summaries(company_id, UPLOADS)
    except Exception as exc:
        db.session.rollback()
        app.logger.warning(
            'OSV file AI scheduling failed for company %s: %s',
            company_id, exc)
_PASSWORD_HASH_PREFIX = 'bcrypt-sha256$1$'


def _hash_password(plain: str) -> str:
    """Hash the complete UTF-8 password, including passwords over 72 bytes."""
    digest = hashlib.sha256(plain.encode('utf-8')).digest()
    return _PASSWORD_HASH_PREFIX + bcrypt.hashpw(
        digest, bcrypt.gensalt()).decode('ascii')


def _check_password(plain: str, hashed: str) -> bool:
    try:
        if hashed.startswith(_PASSWORD_HASH_PREFIX):
            hashed = hashed[len(_PASSWORD_HASH_PREFIX):]
            candidate = hashlib.sha256(plain.encode('utf-8')).digest()
        else:
            # Existing accounts used raw bcrypt and remain verifiable.
            candidate = plain.encode('utf-8')
        return bcrypt.checkpw(candidate, hashed.encode())
    except (ValueError, TypeError):
        return False


def _set_user_password(user, password: str, *, require_change: bool = False):
    """Update password metadata atomically with the hash.

    ``password_version`` is intentionally bumped for every change so sessions
    established before an administrator reset cannot remain usable.
    """
    user.password_hash = _hash_password(password)
    user.password_changed_at = utcnow()
    user.password_change_required = bool(require_change)
    user.password_version = (user.password_version or 0) + 1


def _audit_in_transaction(action: str, details=None, user_id=None):
    """Add a security audit row to the caller's transaction.

    Password mutation callers commit only after this row has been added, so
    the mutation and its audit record cannot silently diverge.
    """
    uid = user_id
    if uid is None and current_user.is_authenticated:
        uid = current_user.id
    db.session.add(AuditLog(
        user_id=uid,
        action=action,
        details=json.dumps(details, ensure_ascii=False, default=str)
                if details is not None else None,
        ip=request.remote_addr if has_request_context() else None,
    ))


def _company_summary(company):
    """Агрегирует данные для карточки на главной странице."""
    uploads = company.uploads  # уже отсортированы по убыванию
    last = uploads[0] if uploads else None
    files_total = sum(u.files_count for u in uploads)
    out_dir = os.path.join(OUTPUT, str(company.id))
    dash_path = os.path.join(out_dir, 'dashboard_current.html')
    kpi_path  = os.path.join(out_dir, 'kpi_cache.json')
    kpis = None
    dashboard_accounts = []
    if os.path.isfile(kpi_path):
        try:
            with open(kpi_path, encoding='utf-8') as _kf:
                kpis = json.load(_kf)
            dashboard_accounts = (
                kpis.get('accounts_loaded', [])
                if isinstance(kpis, dict) else [])
        except Exception:
            pass
    if not dashboard_accounts and os.path.isfile(dash_path):
        # Совместимость с дашбордами, созданными до сохранения accounts_loaded
        # в KPI-кэше. В HTML dataset уже содержит фактический список счетов.
        try:
            with open(dash_path, encoding='utf-8') as _df:
                match = re.search(
                    r'"accounts_loaded"\s*:\s*(\[[^\]]*\])',
                    _df.read())
            if match:
                parsed_accounts = json.loads(match.group(1))
                if isinstance(parsed_accounts, list):
                    dashboard_accounts = parsed_accounts
        except (OSError, ValueError, TypeError):
            pass
    return dict(
        id=company.id,
        name=company.name,
        inn=company.inn or '',
        full_name=company.full_name or '',
        address=company.address or '',
        director_name=company.director_name or '',
        director_phone=company.director_phone or '',
        accountant_name=company.accountant_name or '',
        accountant_phone=company.accountant_phone or '',
        note=company.note or '',
        last_period=last.period_label if last else None,
        files_total=files_total,
        closure=last.closure_status if last else None,
        profit_sign=last.profit_sign if last else 'neutral',
        last_upload=(last.uploaded_at.strftime('%d.%m.%Y') if last else None),
        has_dashboard=os.path.isfile(dash_path),
        kpis=kpis,
        status_color=_compute_status_color(kpis),
        last_completeness=_compute_company_completeness(
            uploads, dashboard_accounts),
    )


def _is_daily_xls(path):
    """True если XLS — суточная ОСВ (дата в шапке), False — периодическая (будет базой)."""
    from osvparser import extract_daily_date
    return extract_daily_date(path) is not None


def _compute_discrepancy(company_id, new_data):
    """Сравнивает СН новой ОСВ с расчётным СК из последних суточных записей.
    Возвращает список расхождений по счетам, отсортированный по убыванию дельты."""
    results = []
    for acc, d in new_data.items():
        new_snd = d['top']['snd']
        new_snk = d['top']['snk']
        if new_snd == 0.0 and new_snk == 0.0:
            continue  # нет начального сальдо — сравнивать нечего
        last_dr = DailyRow.query.filter_by(company_id=company_id, acc=acc)\
                                .order_by(DailyRow.date.desc()).first()
        if last_dr is None:
            continue
        delta_d = new_snd - last_dr.skd
        delta_k = new_snk - last_dr.skk
        if abs(delta_d) > 1.0 or abs(delta_k) > 1.0:
            results.append(dict(
                acc=acc,
                last_date=last_dr.date.strftime('%d.%m.%Y'),
                calc_skd=last_dr.skd,
                calc_skk=last_dr.skk,
                new_snd=new_snd,
                new_snk=new_snk,
                delta_d=delta_d,
                delta_k=delta_k,
            ))
    return sorted(results, key=lambda x: abs(x['delta_d']) + abs(x['delta_k']), reverse=True)


def _org_name_from_filename(path):
    """Пробует извлечь название орг. из имени суточного файла.

    Например: «ОСВ по счету 20 за 1 июля 2026 г. МУП КГХ.xls» → «МУП КГХ»
    """
    name = os.path.splitext(os.path.basename(path))[0]
    m = re.search(r'\d{4}\s*г\.?\s+(.+)', name)
    if m:
        org = m.group(1).strip()
        if org:
            return org
    return None


def _save_files(paths_src, company_id, ts):
    """Сохраняет файлы в uploads/<company_id>/<ts>/."""
    batch_dir = os.path.join(UPLOADS, str(company_id), ts)
    os.makedirs(batch_dir, exist_ok=True)
    saved = []
    for f in paths_src:
        name = os.path.basename(f.filename)
        if not name.lower().endswith(('.xls', '.xlsx')):
            continue
        dest = os.path.join(batch_dir, name)
        f.save(dest)
        saved.append(dest)
    return saved

def _derive_cash_daily_rows(parsed, through_date=None):
    """Derive exact end-of-day cash balances from dated report turnovers.

    This is valid only when the report's dated rows reconcile to its complete
    opening, turnover and closing totals. It does not spread a period total:
    every movement stays on its source date, and no-movement days preserve the
    preceding balance.
    """
    from collections import defaultdict
    from datetime import date as _date
    from math import isfinite
    from key_indicators import _period_bounds

    if str(parsed.get('acc')) not in {'50', '51', '52', '55'}:
        return []
    start, end = _period_bounds(parsed.get('period', ''))
    top = parsed.get('top')
    if start is None or end is None or not isinstance(top, dict):
        return []
    numeric_keys = ('snd', 'snk', 'obd', 'obk', 'skd', 'skk')
    try:
        totals = {key: float(top[key]) for key in numeric_keys}
    except (KeyError, TypeError, ValueError):
        return []
    if not all(isfinite(value) for value in totals.values()):
        return []
    report_end = end
    end = min(end, through_date or _date.today())
    if end < start:
        return []

    movements = defaultdict(lambda: [0.0, 0.0])
    for item in parsed.get('daily') or []:
        try:
            day = _date.fromisoformat(item['date'])
        except (KeyError, TypeError, ValueError):
            return []
        if day < start or day > report_end:
            return []
        debit, credit = item.get('obd'), item.get('obk')
        try:
            debit, credit = float(debit), float(credit)
        except (TypeError, ValueError):
            return []
        if not isfinite(debit) or not isfinite(credit):
            return []
        movements[day][0] += debit
        movements[day][1] += credit

    report_debit = round(sum(value[0] for value in movements.values()), 2)
    report_credit = round(sum(value[1] for value in movements.values()), 2)
    if (abs(report_debit - totals['obd']) > 0.01 or
            abs(report_credit - totals['obk']) > 0.01):
        return []

    running = totals['snd'] - totals['snk']
    result = []
    cursor = start
    while cursor <= end:
        debit, credit = movements[cursor]
        opening = running
        running = round(running + debit - credit, 2)
        result.append({
            'date': cursor,
            'snd': round(max(opening, 0.0), 2),
            'snk': round(max(-opening, 0.0), 2),
            'obd': round(debit, 2),
            'obk': round(credit, 2),
            'skd': round(max(running, 0.0), 2),
            'skk': round(max(-running, 0.0), 2),
        })
        cursor += timedelta(days=1)

    # The report total describes its actual generated state even when its
    # nominal period endpoint lies in the future (for example "September").
    closing = totals['skd'] - totals['skk']
    if abs(running - closing) > 0.01:
        return []
    return result


_BUNDLED_MUP_KGH_ACCOUNT55 = os.path.join(
    BASE, 'data',
    'mup-kgh-account-55-2026-07-01_2026-09-15.xls')


def _import_bundled_mup_kgh_account55():
    """Idempotently install the accepted account-55 ОСВ for МУП КГХ.

    The source is bundled because database-only imports disappear on a clean
    deployment.  It is still registered through the normal immutable source
    manifest and converted to DailyRow observations by the regular importer.
    """
    from osvparser import parse_osv
    from source_manifest import sha256_file

    if not os.path.isfile(_BUNDLED_MUP_KGH_ACCOUNT55):
        raise RuntimeError('Не найден поставляемый файл ОСВ счёта 55 МУП КГХ')
    companies = [
        company for company in Company.query.all()
        if re.sub(r'\s+', ' ', company.name or '').strip().upper() == 'МУП КГХ'
    ]
    if not companies:
        return False
    if len(companies) != 1:
        raise RuntimeError('Нельзя однозначно выбрать предприятие МУП КГХ')
    company = companies[0]
    parsed = parse_osv(_BUNDLED_MUP_KGH_ACCOUNT55)
    observations = _derive_cash_daily_rows(parsed)
    if (
        str(parsed.get('acc')) != '55'
        or len(observations) != 77
        or observations[0]['date'].isoformat() != '2026-07-01'
        or observations[-1]['date'].isoformat() != '2026-09-15'
    ):
        raise RuntimeError('Поставляемая ОСВ счёта 55 не прошла контроль периода')

    digest = sha256_file(_BUNDLED_MUP_KGH_ACCOUNT55)
    batch = f'bundled_account55_{digest[:12]}'
    company_root = os.path.join(UPLOADS, str(company.id))
    batch_dir = os.path.join(company_root, batch)
    os.makedirs(batch_dir, exist_ok=True)
    destination = os.path.join(
        batch_dir, os.path.basename(_BUNDLED_MUP_KGH_ACCOUNT55))
    if not os.path.isfile(destination) or sha256_file(destination) != digest:
        shutil.copy2(_BUNDLED_MUP_KGH_ACCOUNT55, destination)

    manifests = [
        upload.source_manifest for upload in Upload.query.filter_by(
            company_id=company.id).all()
        if isinstance(upload.source_manifest, list)
    ]
    already_registered = any(
        item.get('sha256') == digest
        for manifest in manifests for item in manifest
        if isinstance(item, dict)
    )
    if not already_registered:
        uploader = User.query.filter_by(role='admin').order_by(User.id).first()
        if uploader is None:
            raise RuntimeError('Нет администратора для регистрации ОСВ счёта 55')
        manifest = build_manifest([(destination, parsed)], company_root)
        db.session.add(Upload(
            company_id=company.id,
            uploaded_by=uploader.id,
            period_label='01.07.2026–15.09.2026',
            accounts_loaded=json.dumps(['55']),
            files_count=1,
            source='manual',
            dashboard_file='',
            metrics_json='{}',
            source_manifest_json=dump_manifest(manifest),
        ))
    existing_rows = {
        row.date: row for row in DailyRow.query.filter_by(
            company_id=company.id, acc='55').all()
    }
    balance_keys = ('snd', 'snk', 'obd', 'obk', 'skd', 'skk')
    rows_match = len(existing_rows) == len(observations) and all(
        observation['date'] in existing_rows
        and all(
            abs(float(getattr(existing_rows[observation['date']], key) or 0)
                - float(observation[key])) <= 0.01
            for key in balance_keys
        )
        for observation in observations
    )
    if not rows_match and _upsert_daily_rows(company.id, [destination]) != 1:
        raise RuntimeError('Не удалось построить дневные остатки счёта 55')
    db.session.commit()
    if not already_registered:
        _schedule_osv_file_summaries(company.id)
    return True


def _upsert_daily_rows(company_id, paths):
    """Если XLS-файл является суточным ОСВ (дата извлекается из шапки),
    записывает/обновляет строку DailyRow для этого дня и счёта.
    Возвращает количество обработанных файлов."""
    from osvparser import extract_daily_date, parse_osv
    count = 0
    for path in paths:
        dt = extract_daily_date(path)
        try:
            parsed = parse_osv(path)
        except Exception:
            continue
        acc = parsed['acc']
        top = parsed['top']
        observations = ([{
            'date': dt,
            'source_type': 'direct_daily_osv',
            **{key: top[key] for key in ('snd', 'snk', 'obd', 'obk', 'skd', 'skk')},
        }] if dt is not None else [
            {**observation, 'source_type': 'reconciled_turnovers'}
            for observation in _derive_cash_daily_rows(parsed)
        ])
        if not observations:
            continue
        for observation in observations:
            row = DailyRow.query.filter_by(
                company_id=company_id, date=observation['date'], acc=acc).first()
            if row is None:
                row = DailyRow(
                    company_id=company_id, date=observation['date'], acc=acc)
                db.session.add(row)
            for key in ('snd', 'snk', 'obd', 'obk', 'skd', 'skk'):
                setattr(row, key, observation[key])
            row.source_type = observation['source_type']
            row.uploaded_at = utcnow()
        count += 1
    return count


def _process_files(paths, company_id, ts, back_url, company_name=''):
    """Парсит файлы, строит дашборд, возвращает (upload_record_kwargs, audit)."""
    # Суточные XLS хранятся рядом с месячными, но не являются входом
    # parse_many.  Передача полного списка здесь особенно опасна для первого
    # импорта предприятия: audit тогда ссылается на месячные файлы, которых
    # нет в списке манифеста (или, наоборот, на суточный файл).
    monthly_paths = [path for path in paths if not _is_daily_xls(path)]
    daily_paths = [path for path in paths if path not in monthly_paths]
    if not monthly_paths:
        return None, []

    data, audit = parse_many(monthly_paths)
    if not data:
        return None, audit

    dataset = build_dataset(data, company_name=company_name)
    html = render_dashboard(dataset, back_url=back_url)

    out_dir = os.path.join(OUTPUT, str(company_id))
    os.makedirs(out_dir, exist_ok=True)
    out_name = 'dashboard_current.html'
    with open(os.path.join(out_dir, out_name), 'w', encoding='utf-8') as f:
        f.write(html)
    _save_kpi_cache(dataset, out_dir)

    brief = dataset.get('brief', {})
    closure_status = _dataset_closure_status(dataset)

    # Финансовый результат
    lead = brief.get('lead', '')
    if 'убыток' in lead.lower() or 'убыточна' in lead.lower():
        profit_sign = 'loss'
    elif 'прибыль' in lead.lower() or 'прибыльна' in lead.lower():
        profit_sign = 'profit'
    else:
        profit_sign = 'neutral'

    want = ['20', '25', '26', '44', '50', '51', '60', '62', '66', '67', '68', '69', '70', '76', '90', '91']
    missing = [a for a in want if a not in data]
    if missing:
        audit.append(dict(file='—', acc=', '.join(missing),
                          status='счета не загружены — соответствующие блоки будут пустыми',
                          ok=False))

    record_kwargs = dict(
        dashboard_file=out_name,
        period_label=dataset.get('period_label', ''),
        accounts_loaded=json.dumps(sorted(data.keys())),
        closure_status=closure_status,
        profit_sign=profit_sign,
        files_count=len(paths),
        metrics_json=_dataset_metrics_json(dataset),
        source_manifest_json=_monthly_manifest_json(
            monthly_paths, data, audit, company_id,
            extra_accepted=_accepted_daily_files(daily_paths)),
    )
    return record_kwargs, audit


# ---------------------------------------------------------------------------
# IMAP sync helper and background loop
# ---------------------------------------------------------------------------
def _get_admin_id() -> int:
    """Return the ID of the first admin user (used for automated uploads)."""
    u = User.query.filter_by(role='admin').order_by(User.id).first()
    return u.id if u else 1


def _do_imap_sync(uploader_id: int) -> dict:
    """Core IMAP sync logic. Must be called inside an app context.
    Returns dict(ok, msg, processed, updated).
    """
    import imaplib, email, ssl, fnmatch, tempfile, shutil as _shutil
    from email.header import decode_header as _dh
    from email.utils import getaddresses as _getaddresses

    def _decode(val):
        if val is None:
            return ''
        parts = []
        for chunk, charset in _dh(val):
            if isinstance(chunk, bytes):
                parts.append(chunk.decode(charset or 'utf-8', errors='replace'))
            else:
                parts.append(chunk)
        return ''.join(parts)

    host     = Setting.get('imap_host', '').strip()
    port_str = Setting.get('imap_port', '993').strip()
    login    = Setting.get('imap_login', '').strip()
    password = decrypt_value(Setting.get('imap_password', '').strip())
    folder   = Setting.get('imap_folder', 'INBOX').strip() or 'INBOX'

    if not host or not login or not password:
        return dict(ok=False, msg='IMAP не настроен: заполните хост, логин и пароль и сохраните.',
                    processed=0, updated=0)

    try:
        port_int = int(port_str)
    except ValueError:
        return dict(ok=False, msg=f'Неверный порт: {port_str}', processed=0, updated=0)

    rules = EmailRule.query.order_by(EmailRule.id).all()

    def _match_rule(rule, sender, subject):
        if rule.sender_pattern:
            pattern = rule.sender_pattern.strip().lower()
            # Заголовок From обычно имеет вид «Имя <user@example.com>».
            # Правила администраторы задают как чистый адрес, поэтому
            # сопоставляем и полный заголовок, и извлечённые email-адреса.
            sender_candidates = [sender.strip().lower()]
            sender_candidates.extend(
                address.strip().lower()
                for _name, address in _getaddresses([sender])
                if address
            )
            if not any(fnmatch.fnmatch(candidate, pattern)
                       for candidate in sender_candidates):
                return False
        if rule.subject_pattern and not fnmatch.fnmatch(
                subject.strip().lower(), rule.subject_pattern.strip().lower()):
            return False
        return True

    stats = dict(emails_checked=0, emails_matched=0,
                 files_processed=0, files_skipped=0,
                 companies_updated=set(), errors=[])

    tmp_root = tempfile.mkdtemp(prefix='imap_sync_')
    try:
        ctx = ssl.create_default_context()
        with imaplib.IMAP4_SSL(host, port_int, ssl_context=ctx, timeout=20) as M:
            M.login(login, password)
            status, _ = M.select(folder)
            if status != 'OK':
                return dict(ok=False, msg=f'Папка «{folder}» не найдена',
                            processed=0, updated=0)

            _, uid_list = M.search(None, 'UNSEEN')
            uids = uid_list[0].split() if uid_list and uid_list[0] else []
            stats['emails_checked'] = len(uids)

            for uid in uids:
                _, data = M.fetch(uid, '(RFC822)')
                raw = data[0][1] if data and data[0] else None
                if not raw:
                    continue
                msg = email.message_from_bytes(raw)
                sender  = _decode(msg.get('From', ''))
                subject = _decode(msg.get('Subject', ''))

                matched_company_id = None
                for rule in rules:
                    if (rule.company and not rule.company.is_archived
                            and _match_rule(rule, sender, subject)):
                        matched_company_id = rule.company_id
                        break

                if matched_company_id is None:
                    # Не помечаем письмо прочитанным: после исправления или
                    # добавления правила маршрутизации его можно обработать
                    # повторно без пересылки.
                    continue

                stats['emails_matched'] += 1
                msg_tmp = os.path.join(tmp_root, uid.decode())
                os.makedirs(msg_tmp, exist_ok=True)

                attached_paths = []
                attached_names = set()
                for part in msg.walk():
                    fname_raw = part.get_filename()
                    if not fname_raw:
                        continue
                    fname = _decode(fname_raw)
                    ext = os.path.splitext(fname)[1].lower()
                    if ext not in ('.xls', '.xlsx', '.zip'):
                        continue
                    payload = part.get_payload(decode=True)
                    if not payload:
                        continue
                    safe_name = _unique_attachment_name(fname, attached_names)
                    dest = os.path.join(msg_tmp, safe_name)
                    with open(dest, 'wb') as fh:
                        fh.write(payload)
                    attached_paths.append(dest)

                if not attached_paths:
                    M.store(uid, '+FLAGS', '\\Seen')
                    continue

                ts = datetime.now().strftime('%Y%m%d_%H%M%S')
                company_id = matched_company_id
                for path in attached_paths:
                    ext = os.path.splitext(path)[1].lower()
                    attachment_output_snapshot = {}
                    attachment_batch_dirs = []
                    attachment_committed = False
                    try:
                        attachment_output_snapshot = _snapshot_company_outputs(
                            company_id)
                        if ext == '.zip':
                            ok, validation_error = _validate_upload(
                                path, os.path.basename(path))
                            if not ok:
                                raise ValueError(validation_error)

                            # ZIP может содержать как суточные, так и обычные
                            # периодические ОСВ, в том числе во вложенных папках.
                            # После проверки безопасности извлекаем только XLS/XLSX.
                            import zipfile as _zipfile
                            zip_dir = _reserve_upload_batch_dir(
                                company_id, f'zip_{ts}')
                            attachment_batch_dirs.append(zip_dir)
                            archive_name = _unique_attachment_name(
                                os.path.basename(path), set())
                            archive_path = os.path.join(zip_dir, archive_name)
                            _shutil.copy2(path, archive_path)
                            extracted_dir = os.path.join(
                                msg_tmp, f'archive_{len(stats["errors"])}')
                            os.makedirs(extracted_dir, exist_ok=True)
                            extracted_xls = []
                            extracted_names = {}
                            extracted_used_names = set()
                            with _zipfile.ZipFile(archive_path, 'r') as zf:
                                for index, info in enumerate(zf.infolist()):
                                    if info.is_dir():
                                        continue
                                    inner_name = info.filename.replace('\\', '/')
                                    if not inner_name.lower().endswith(('.xls', '.xlsx')):
                                        continue
                                    # The technical extraction index is kept
                                    # only in msg_tmp; persisted names retain
                                    # the archive basename and get a suffix
                                    # for duplicate basenames.
                                    base_name = _unique_attachment_name(
                                        os.path.basename(inner_name),
                                        extracted_used_names)
                                    target = os.path.join(
                                        extracted_dir, f'{index:04d}_{base_name}')
                                    with zf.open(info, 'r') as src, open(target, 'wb') as dst:
                                        _shutil.copyfileobj(src, dst)
                                    extracted_xls.append(target)
                                    extracted_names[target] = base_name

                            if not extracted_xls:
                                raise ValueError(
                                    'в архиве не найдено файлов .xls или .xlsx')

                            daily_xls = [p for p in extracted_xls if _is_daily_xls(p)]
                            monthly_xls = [p for p in extracted_xls
                                           if p not in daily_xls]
                            imported_count = 0
                            monthly_manifest_files = []
                            daily_manifest_files = []

                            if daily_xls:
                                daily_dir = _reserve_upload_batch_dir(
                                    company_id, f'daily_imap_{ts}')
                                attachment_batch_dirs.append(daily_dir)
                                saved_daily = []
                                daily_names = set()
                                for source_path in daily_xls:
                                    saved_name = _unique_attachment_name(
                                        extracted_names[source_path], daily_names)
                                    destination = os.path.join(
                                        daily_dir, saved_name)
                                    _shutil.copy2(source_path, destination)
                                    saved_daily.append(destination)
                                imported_count += _upsert_daily_rows(
                                    company_id, saved_daily)
                                daily_manifest_files.extend(
                                    _accepted_daily_files(saved_daily))

                            if monthly_xls:
                                save_dir = _reserve_upload_batch_dir(
                                    company_id, f'imap_{ts}')
                                attachment_batch_dirs.append(save_dir)
                                saved_monthly = []
                                saved_names = set()
                                for source_path in monthly_xls:
                                    saved_name = _unique_attachment_name(
                                        extracted_names[source_path],
                                        saved_names)
                                    destination = os.path.join(save_dir, saved_name)
                                    _shutil.copy2(source_path, destination)
                                    saved_monthly.append(destination)
                                parsed_monthly, parse_audit = parse_many(saved_monthly)
                                monthly_manifest_files.extend(
                                    _accepted_monthly_files(
                                        saved_monthly, parsed_monthly, parse_audit))
                                for audit_item in parse_audit:
                                    if not audit_item.get('ok'):
                                        stats['errors'].append(
                                            f'{os.path.basename(path)}: '
                                            f'{audit_item.get("file", "файл")}: '
                                            f'{audit_item.get("status", "ошибка разбора")}')
                                imported_count += len(parsed_monthly)

                            if imported_count == 0:
                                raise ValueError(
                                    'ни один Excel-файл в архиве не удалось разобрать как ОСВ')

                            _co = db.session.get(Company, company_id)
                            dataset, rebuild_audit = _rebuild_dashboard_with_daily(
                                company_id,
                                company_name=_co.name if _co else '')
                            if dataset is None:
                                detail = (rebuild_audit[-1].get('status')
                                          if rebuild_audit else 'нет данных')
                                raise ValueError(
                                    f'не удалось перестроить дашборд: {detail}')
                            html_content = render_dashboard(
                                dataset, back_url=f'/company/{company_id}')
                            out_dir = os.path.join(OUTPUT, str(company_id))
                            os.makedirs(out_dir, exist_ok=True)
                            with open(os.path.join(
                                    out_dir, 'dashboard_current.html'),
                                    'w', encoding='utf-8') as fh:
                                fh.write(html_content)
                            _save_kpi_cache(dataset, out_dir)
                            stats['companies_updated'].add(company_id)

                            brief = dataset.get('brief', {})
                            lead = brief.get('lead', '')
                            _up = Upload(
                                company_id=company_id,
                                uploaded_by=uploader_id,
                                period_label=dataset.get('period_label', ''),
                                accounts_loaded=json.dumps(
                                    _dataset_accounts(dataset)),
                                closure_status=_dataset_closure_status(dataset),
                                profit_sign=(
                                    'loss' if 'убыток' in lead.lower() else
                                    'profit' if 'прибыль' in lead.lower() else
                                    'neutral'),
                                files_count=len(extracted_xls),
                                dashboard_file='dashboard_current.html',
                                source='auto_imap',
                                metrics_json=_dataset_metrics_json(dataset),
                                source_manifest_json=dump_manifest(build_manifest(
                                    monthly_manifest_files + daily_manifest_files,
                                    os.path.join(UPLOADS, str(company_id)))),
                            )
                            db.session.add(_up)
                            db.session.commit()
                            attachment_committed = True
                            _schedule_osv_file_summaries(company_id)
                            _durable_backup('IMAP ZIP ОСВ')
                            stats['files_processed'] += len(extracted_xls)
                        else:  # XLS / XLSX
                            is_daily = _is_daily_xls(path)
                            save_dir = _reserve_upload_batch_dir(
                                company_id,
                                f'daily_imap_{ts}' if is_daily else f'imap_{ts}')
                            attachment_batch_dirs.append(save_dir)
                            dest_name = _unique_attachment_name(
                                os.path.basename(path), set())
                            dest = os.path.join(save_dir, dest_name)
                            _shutil.copy2(path, dest)
                            monthly_manifest_files = []

                            if is_daily:
                                imported = _upsert_daily_rows(company_id, [dest])
                                if imported == 0:
                                    raise ValueError(
                                        'не удалось разобрать суточную ОСВ')
                                from osvparser import parse_osv as _parse_osv
                                loaded_accounts = [_parse_osv(dest)['acc']]
                                monthly_manifest_files.extend(
                                    _accepted_daily_files([dest]))
                            else:
                                parsed_file, parse_audit = parse_many([dest])
                                if not parsed_file:
                                    detail = (parse_audit[-1].get('status')
                                              if parse_audit else
                                              'неизвестная ошибка разбора')
                                    raise ValueError(
                                        f'не удалось разобрать ОСВ: {detail}')
                                loaded_accounts = sorted(parsed_file.keys())
                                monthly_manifest_files.extend(
                                    _accepted_monthly_files(
                                        [dest], parsed_file, parse_audit))

                            _co = db.session.get(Company, company_id)
                            dataset, rebuild_audit = _rebuild_dashboard_with_daily(
                                company_id,
                                company_name=_co.name if _co else '')
                            if dataset is None:
                                detail = (rebuild_audit[-1].get('status')
                                          if rebuild_audit else 'нет данных')
                                raise ValueError(
                                    f'не удалось перестроить дашборд: {detail}')
                            html_content = render_dashboard(
                                dataset, back_url=f'/company/{company_id}')
                            out_dir = os.path.join(OUTPUT, str(company_id))
                            os.makedirs(out_dir, exist_ok=True)
                            with open(os.path.join(
                                    out_dir, 'dashboard_current.html'),
                                    'w', encoding='utf-8') as fh:
                                fh.write(html_content)
                            _save_kpi_cache(dataset, out_dir)
                            stats['companies_updated'].add(company_id)
                            brief = dataset.get('brief', {})
                            lead = brief.get('lead', '')
                            _up = Upload(
                                company_id=company_id,
                                uploaded_by=uploader_id,
                                period_label=dataset.get('period_label', ''),
                                accounts_loaded=json.dumps(loaded_accounts),
                                closure_status=_dataset_closure_status(dataset),
                                profit_sign=(
                                    'loss' if 'убыток' in lead.lower() else
                                    'profit' if 'прибыль' in lead.lower() else
                                    'neutral'),
                                files_count=1,
                                dashboard_file='dashboard_current.html',
                                source='auto_imap',
                                metrics_json=_dataset_metrics_json(dataset),
                                source_manifest_json=dump_manifest(build_manifest(
                                    monthly_manifest_files,
                                    os.path.join(UPLOADS, str(company_id)))),
                            )
                            db.session.add(_up)
                            db.session.commit()
                            attachment_committed = True
                            _schedule_osv_file_summaries(company_id)
                            _durable_backup('IMAP ОСВ')
                            stats['files_processed'] += 1
                    except Exception as exc:
                        db.session.rollback()
                        if not attachment_committed:
                            _remove_pending_batch_dirs(attachment_batch_dirs)
                            _restore_company_outputs(
                                attachment_output_snapshot)
                        stats['errors'].append(f"{os.path.basename(path)}: {exc}")
                        stats['files_skipped'] += 1

                M.store(uid, '+FLAGS', '\\Seen')

            M.logout()
    except imaplib.IMAP4.error as e:
        return dict(ok=False, msg=f'Ошибка IMAP-аутентификации: {e}',
                    processed=0, updated=0)
    except OSError as e:
        return dict(ok=False, msg=f'Сетевая ошибка: {e}', processed=0, updated=0)
    except Exception as e:
        return dict(ok=False, msg=f'Ошибка: {e}', processed=0, updated=0)
    finally:
        _shutil.rmtree(tmp_root, ignore_errors=True)

    upd = len(stats['companies_updated'])
    lines = []
    if stats['emails_checked'] == 0:
        lines.append('Непрочитанных писем нет')
    else:
        lines.append(f"Писем проверено: {stats['emails_checked']}, "
                     f"совпало с правилами: {stats['emails_matched']}")
    if stats['files_processed']:
        lines.append(f"Файлов обработано: {stats['files_processed']}, "
                     f"предприятий обновлено: {upd}")
    if stats['files_skipped']:
        lines.append(f"Пропущено с ошибкой: {stats['files_skipped']}")
    for err in stats['errors'][:3]:
        lines.append(f"  ⚠ {err}")
    if len(stats['errors']) > 3:
        lines.append(f"  … и ещё {len(stats['errors'])-3} ошибок")

    ok = not stats['errors'] or stats['files_processed'] > 0
    return dict(ok=ok, msg=' | '.join(lines),
                processed=stats['files_processed'], updated=upd)


def _imap_loop() -> None:
    """Daemon thread: poll IMAP inbox on a configurable interval."""
    while not _imap_stop.is_set():
        try:
            with app.app_context():
                interval_str = Setting.get('imap_interval', '30')
                try:
                    interval = max(5, int(interval_str)) * 60
                except ValueError:
                    interval = 30 * 60
                host = Setting.get('imap_host', '').strip()
                if host:
                    uid = _get_admin_id()
                    result = _do_imap_sync(uid)
                    if result['processed']:
                        print(f'[imap-loop] {result["msg"]}')
                else:
                    interval = 60  # check again in 1 min if not configured
        except Exception as exc:
            print(f'[imap-loop] unhandled error: {exc}')
            interval = 60
        _imap_stop.wait(interval)

def _snapshot_company_outputs(company_id):
    """Capture replaceable dashboard outputs before a pending acceptance."""
    out_dir = os.path.join(OUTPUT, str(company_id))
    snapshot = {}
    for name in _COMPANY_OUTPUT_NAMES:
        path = os.path.join(out_dir, name)
        if not os.path.isfile(path):
            snapshot[path] = None
            continue
        with open(path, 'rb') as stream:
            snapshot[path] = stream.read()
    return snapshot
def _bg_upload_all(job_id: str, company_id: int, company_name: str,
                   stage_dir: str, ts: str, uploader_id: int) -> None:
    """Background thread: classify + parse all staged files, rebuild dashboard."""
    import shutil as _sh
    with app.app_context():
        output_snapshot = {}
        batch_dirs = []
        committed = False
        try:
            output_snapshot = _snapshot_company_outputs(company_id)
            audit_out = []
            zip_dates = []
            monthly_manifest_files = []
            total_files = 0

            all_staged = [os.path.join(stage_dir, f)
                          for f in os.listdir(stage_dir)
                          if os.path.isfile(os.path.join(stage_dir, f))]
            xls_staged = [p for p in all_staged
                          if os.path.splitext(p)[1].lower() in ('.xls', '.xlsx')]
            zip_staged = [p for p in all_staged if p.lower().endswith('.zip')]

            # — XLS: classify daily vs monthly, move to permanent dirs
            if xls_staged:
                total_files += len(xls_staged)
                ts_dir = _reserve_upload_batch_dir(company_id, ts)
                batch_dirs.append(ts_dir)
                daily_dir = None
                full_xls, daily_xls = [], []
                for path in xls_staged:
                    dest = os.path.join(ts_dir, os.path.basename(path))
                    _sh.move(path, dest)
                    if _is_daily_xls(dest):
                        if daily_dir is None:
                            daily_dir = _reserve_upload_batch_dir(
                                company_id, f'daily_{ts}')
                            batch_dirs.append(daily_dir)
                        final = os.path.join(daily_dir, os.path.basename(dest))
                        os.rename(dest, final)
                        daily_xls.append(final)
                    else:
                        full_xls.append(dest)

                _upsert_daily_rows(company_id, full_xls + daily_xls)
                monthly_manifest_files.extend(
                    _accepted_daily_files(daily_xls))

                if full_xls:
                    data_check, parse_audit = parse_many(full_xls)
                    monthly_manifest_files.extend(
                        _accepted_monthly_files(
                            full_xls, data_check, parse_audit))
                    if data_check:
                        for disc in _compute_discrepancy(company_id, data_check):
                            audit_out.append(dict(
                                file=f'сч.{disc["acc"]}',
                                acc=disc['acc'],
                                status=(
                                    f'РАСХОЖДЕНИЕ: расчётное СК на {disc["last_date"]} — '
                                    f'Дт {disc["calc_skd"]:,.0f} / Кт {disc["calc_skk"]:,.0f} | '
                                    f'новая ОСВ СН — Дт {disc["new_snd"]:,.0f} / Кт {disc["new_snk"]:,.0f} | '
                                    f'дельта Дт {disc["delta_d"]:+,.0f} / Кт {disc["delta_k"]:+,.0f}'
                                ),
                                ok='disc',
                            ))

            # — ZIP: move and parse
            for zp in zip_staged:
                total_files += 1
                zip_dir = _reserve_upload_batch_dir(company_id, f'zip_{ts}')
                batch_dirs.append(zip_dir)
                dest = os.path.join(zip_dir, os.path.basename(zp))
                _sh.move(zp, dest)
                accepted_zip_dir = os.path.join(zip_dir, 'accepted')
                successes, errors = parse_zip(
                    dest, accepted_dir=accepted_zip_dir)
                audit_out += [{'file': e['file'], 'acc': '?',
                               'status': f'ОШИБКА: {e["error"]}', 'ok': False}
                              for e in errors]
                for item in successes:
                    dt, acc, top = item['date'], item['parsed']['acc'], item['parsed']['top']
                    row = DailyRow.query.filter_by(
                        company_id=company_id, date=dt, acc=acc).first()
                    if row is None:
                        row = DailyRow(company_id=company_id, date=dt, acc=acc)
                        db.session.add(row)
                    row.snd = top['snd']; row.snk = top['snk']
                    row.obd = top['obd']; row.obk = top['obk']
                    row.skd = top['skd']; row.skk = top['skk']
                    row.source_type = 'direct_daily_osv'
                    row.uploaded_at = utcnow()
                    zip_dates.append(dt)
                monthly_manifest_files.extend(
                    _accepted_daily_files([
                        item['path'] for item in successes if item.get('path')
                    ]))

            # Cleanup staging dir
            _sh.rmtree(stage_dir, ignore_errors=True)

            if not xls_staged and not zip_staged:
                raise ValueError(
                    'Среди загруженных нет файлов .xls, .xlsx или .zip.')

            # Rebuild dashboard
            co = db.session.get(Company, company_id)
            back_url = f'/company/{company_id}'
            dataset, rebuild_audit = _rebuild_dashboard_with_daily(
                company_id, company_name=co.name if co else company_name)
            audit_out = rebuild_audit + audit_out

            if dataset is None:
                msg = rebuild_audit[-1]['status'] if rebuild_audit else 'нет данных'
                raise RuntimeError(f'Дашборд не построен: {msg}')

            html_content = render_dashboard(dataset, back_url=back_url)
            out_dir = os.path.join(OUTPUT, str(company_id))
            os.makedirs(out_dir, exist_ok=True)
            with open(os.path.join(out_dir, 'dashboard_current.html'),
                      'w', encoding='utf-8') as fh:
                fh.write(html_content)

            period_lbl = dataset.get('period_label', '')
            if zip_dates:
                _MRU = {1:'янв',2:'фев',3:'мар',4:'апр',5:'май',6:'июн',
                        7:'июл',8:'авг',9:'сен',10:'окт',11:'ноя',12:'дек'}
                zd = sorted(set(zip_dates))
                d0, d1 = zd[0], zd[-1]
                zl = (f'{d0.day} {_MRU[d0.month]}' if d0 == d1
                      else f'{d0.strftime("%d.%m")}–{d1.strftime("%d.%m")}')
                period_lbl = f'{period_lbl} + {zl}'.strip(' +')

            brief = dataset.get('brief', {})
            lead = brief.get('lead', '')
            profit_sign = ('loss'   if 'убыток'  in lead.lower() else
                           'profit' if 'прибыль' in lead.lower() else 'neutral')
            closure_status = _dataset_closure_status(dataset)

            upload = Upload(
                company_id=company_id,
                uploaded_by=uploader_id,
                period_label=period_lbl,
                accounts_loaded=json.dumps(_dataset_accounts(dataset)),
                closure_status=closure_status,
                profit_sign=profit_sign,
                files_count=total_files,
                dashboard_file='dashboard_current.html',
                source='manual',
                metrics_json=_dataset_metrics_json(dataset),
                source_manifest_json=dump_manifest(build_manifest(
                    monthly_manifest_files,
                    os.path.join(UPLOADS, str(company_id)))),
            )
            db.session.add(upload)
            db.session.commit()
            committed = True
            _schedule_osv_file_summaries(company_id)
            _durable_backup('ОСВ background upload')

            discrepancy = [a for a in audit_out if a.get('ok') == 'disc']
            log_event('file_upload',
                      {'company_id': company_id, 'files': total_files, 'period': period_lbl},
                      user_id=uploader_id)
            finish_job(job_id, dict(
                audit=audit_out,
                discrepancy=discrepancy,
                period_label=period_lbl,
            ))
        except Exception as exc:
            db.session.rollback()
            if not committed:
                _remove_pending_batch_dirs(batch_dirs)
                _sh.rmtree(stage_dir, ignore_errors=True)
                _restore_company_outputs(output_snapshot)
            traceback.print_exc()
            fail_job(job_id, str(exc))


def _bg_company_new(job_id: str, company_id: int, final_dir: str,
                    final_paths: list, ts: str, back_url: str,
                    company_name: str, uploader_id: int) -> None:
    """Background thread: process new company files and create first Upload."""
    with app.app_context():
        output_snapshot = {}
        committed = False
        owns_new_batch = False
        try:
            output_snapshot = _snapshot_company_outputs(company_id)
            # A first-import worker owns this directory.  If a caller ever
            # retries against a company that already has an Upload, do not
            # remove that company's historical source directory on failure.
            owns_new_batch = (
                Upload.query.filter_by(company_id=company_id).first() is None)
            record_kwargs, audit = _process_files(
                final_paths, company_id, ts, back_url,
                company_name=company_name)
            if record_kwargs is None:
                raise ValueError(
                    'Ни один файл не удалось разобрать как ОСВ 1С')

            _upsert_daily_rows(company_id, final_paths)
            upload = Upload(company_id=company_id, uploaded_by=uploader_id,
                            source='manual', **record_kwargs)
            db.session.add(upload)
            db.session.commit()
            committed = True
            _schedule_osv_file_summaries(company_id)
            _durable_backup('первый ОСВ предприятия')

            log_event('company_created',
                      {'company_id': company_id, 'name': company_name},
                      user_id=uploader_id)
            finish_job(job_id, dict(
                audit=audit,
                discrepancy=[],
                period_label=record_kwargs.get('period_label', ''),
            ))
        except Exception as exc:
            db.session.rollback()
            if not committed:
                if owns_new_batch:
                    shutil.rmtree(final_dir, ignore_errors=True)
                _restore_company_outputs(output_snapshot)
            traceback.print_exc()
            fail_job(job_id, str(exc))


# ---------------------------------------------------------------------------
# Шифрование / дешифрование (хранилище чувствительных настроек в БД)
# ---------------------------------------------------------------------------
_XLS_MAGIC = b'\xd0\xcf\x11\xe0'  # CFBF — Excel 97-2003
_ZIP_MAGIC  = b'PK'                # ZIP / OOXML (xlsx, zip)


def encrypt_value(plain: str) -> str:
    """Шифрует строку. Возвращает «enc:<ciphertext>»."""
    if not plain:
        return plain
    return 'enc:' + _FERNET.encrypt(plain.encode()).decode()


def decrypt_value(cipher: str) -> str:
    """Расшифровывает значение, зашифрованное через encrypt_value().
    Открытый текст (без префикса enc:) возвращается как есть — миграция без перезаписи."""
    if not cipher or not cipher.startswith('enc:'):
        return cipher
    try:
        return _FERNET.decrypt(cipher[4:].encode()).decode()
    except Exception:
        return ''


def log_event(action: str, details=None, user_id: int | None = None,
              ip: str | None = None) -> bool:
    """Записывает событие безопасности в таблицу audit_log.
    Никогда не вызывает исключений — ошибки ведения журнала не должны ломать запросы."""
    from sqlalchemy.orm import Session as _AuditSession
    try:
        uid = user_id
        _ip = ip
        if uid is None:
            try:
                if current_user.is_authenticated:
                    uid = current_user.id
            except Exception:
                pass
        if _ip is None:
            try:
                _ip = request.remote_addr
            except Exception:
                pass
        entry = AuditLog(
            user_id=uid,
            action=action,
            details=json.dumps(details, ensure_ascii=False, default=str)
                    if details is not None else None,
            ip=_ip,
        )
        with _AuditSession(db.engine) as _sess:
            _sess.add(entry)
            _sess.commit()
        return True
    except Exception:
        return False


def _validate_upload(path: str, filename: str) -> tuple[bool, str]:
    """Проверяет magic bytes файла. Для ZIP проверяет zip-bomb и path traversal.
    Возвращает (ok, сообщение_об_ошибке)."""
    ext = os.path.splitext(filename)[1].lower()
    try:
        with open(path, 'rb') as fh:
            head = fh.read(8)
    except OSError as e:
        return False, f'Не удалось прочитать «{filename}»: {e}'

    if ext == '.xls':
        if not head.startswith(_XLS_MAGIC):
            return False, f'«{filename}» не является Excel 97-2003 (.xls) — неверная сигнатура файла'
    elif ext == '.xlsx':
        if not head.startswith(_ZIP_MAGIC):
            return False, f'«{filename}» не является Excel 2007+ (.xlsx) — неверная сигнатура файла'
    elif ext == '.zip':
        if not head.startswith(_ZIP_MAGIC):
            return False, f'«{filename}» не является ZIP-архивом — неверная сигнатура файла'
        ok, err = _check_zip_safety(path, filename)
        if not ok:
            return False, err

    return True, ''


def _check_zip_safety(path: str, filename: str) -> tuple[bool, str]:
    """Zip-bomb (>500 МБ распакованных данных) и path traversal."""
    import zipfile as _zf
    MAX_UNCOMPRESSED = 500 * 1024 * 1024
    try:
        with _zf.ZipFile(path, 'r') as zfile:
            total = sum(i.file_size for i in zfile.infolist())
            if total > MAX_UNCOMPRESSED:
                return False, (
                    f'«{filename}» отклонён: распакованный объём '
                    f'{total // 1024 // 1024} МБ превышает лимит 500 МБ (возможный zip-bomb)'
                )
            for info in zfile.infolist():
                name = info.filename
                parts = name.replace('\\', '/').split('/')
                if name.startswith('/') or '..' in parts:
                    return False, (
                        f'«{filename}» содержит опасный путь в архиве '
                        f'(path traversal): «{name}»'
                    )
    except _zf.BadZipFile as e:
        return False, f'«{filename}» — повреждённый ZIP-архив: {e}'
    return True, ''


# ---------------------------------------------------------------------------
# Обработчики ошибок
# ---------------------------------------------------------------------------
def _wants_json():
    """True если клиент ожидает JSON (AJAX-запрос)."""
    return (request.path.startswith('/settings/1c/') or
            request.path.startswith('/api/1c/') or
            request.headers.get('X-Requested-With') == 'XMLHttpRequest' or
            'application/json' in request.headers.get('Accept', ''))


@app.errorhandler(CSRFError)
def handle_csrf_error(e):
    if _wants_json():
        return jsonify(error='CSRF token invalid or expired'), 403
    return render_template('403.html'), 403


@app.errorhandler(403)
def forbidden(e):
    if _wants_json():
        return jsonify(error='Доступ запрещён'), 403
    return render_template('403.html'), 403


@app.errorhandler(500)
def internal_error(e):
    app.logger.error('500 Internal Server Error:\n%s', traceback.format_exc())
    if _wants_json():
        return jsonify(error='Внутренняя ошибка сервера'), 500
    return render_template('500.html'), 500


# ---------------------------------------------------------------------------
# Вспомогательный эндпоинт для JS/HTMX-клиентов: получить CSRF-токен
# ---------------------------------------------------------------------------
@app.route('/rating-api/csrf', methods=['GET'])
@login_required
def get_csrf_token():
    return jsonify(csrf_token=generate_csrf())


# ---------------------------------------------------------------------------
# Маршруты авторизации
# ---------------------------------------------------------------------------
def _configured_login_hosts():
    """Return explicitly configured browser origins, never proxy headers."""
    hosts = {request.host.lower()}
    for raw in (
            os.environ.get('REPLIT_DOMAINS', ''),
            os.environ.get('REPLIT_DEV_DOMAIN', ''),
    ):
        for item in raw.split(','):
            value = item.strip()
            if not value:
                continue
            parsed = urlparse(value if '://' in value else f'//{value}')
            if parsed.netloc:
                hosts.add(parsed.netloc.lower())
    return hosts


@app.route('/login', methods=['GET', 'POST'])
@csrf.exempt
def login():
    # Login uses a signed stateless token because the Replit preview proxy can
    # scope the GET and POST session cookies differently. A strict same-origin
    # check additionally prevents login-CSRF. Other POST routes still use the
    # session-bound Flask-WTF protection.
    if current_user.is_authenticated:
        return redirect(url_for('modules_dashboard'))
    error = None
    username = ''
    if request.method == 'POST':
        source = request.headers.get('Origin') or request.headers.get('Referer')
        source_host = urlparse(source).netloc.lower() if source else ''
        allowed_hosts = _configured_login_hosts()
        if not source_host or source_host not in allowed_hosts:
            abort(403, description='Недопустимый источник запроса входа')
        try:
            payload = _LOGIN_CSRF.loads(
                request.form.get('login_csrf_token', ''), max_age=15 * 60)
            if payload != 'login':
                raise BadSignature('invalid login token purpose')
        except (BadSignature, SignatureExpired):
            abort(403, description='Недействительный токен формы входа')
        username = request.form.get('username', '').strip()
        password = request.form.get('password', '')
        user = User.query.filter_by(username=username).first()
        if user and _check_password(password, user.password_hash):
            login_user(user)
            session['password_version'] = int(user.password_version or 1)
            log_event('login_ok', {'username': username})
            if user.password_change_required:
                return redirect(url_for('password_change_required'))
            return redirect(url_for('modules_dashboard'))
        log_event('login_fail', {'username': username})
        error = 'Неверный логин или пароль.'
    return render_template(
        'login.html', error=error, username=username,
        login_csrf_token=_LOGIN_CSRF.dumps('login'))


@app.route('/password/change-required', methods=['GET', 'POST'])
@login_required
def password_change_required():
    """The only application page available to an account awaiting a change."""
    user = current_user
    error = None
    if request.method == 'POST':
        current = request.form.get('current_password', '')
        new_password = request.form.get('new_password', '')
        repeat = request.form.get('new_password_repeat', '')
        if not _check_password(current, user.password_hash):
            error = 'Текущий пароль указан неверно.'
        elif new_password != repeat:
            error = 'Новые пароли не совпадают.'
        elif _check_password(new_password, user.password_hash):
            error = 'Новый пароль должен отличаться от текущего.'
        else:
            error = validate_password(new_password, username=user.username)
        if not error:
            _set_user_password(user, new_password, require_change=False)
            _audit_in_transaction('password_changed', {'reason': 'forced'})
            db.session.commit()
            session['password_version'] = int(user.password_version or 1)
            return redirect(url_for('modules_dashboard'))
    return render_template('password_change.html', error=error,
                           forced=True, username=user.username)


@app.route('/profile/password', methods=['GET', 'POST'])
@login_required
def profile_password():
    """Voluntary self-service password change for already active users."""
    if current_user.password_change_required:
        return redirect(url_for('password_change_required'))
    error = None
    if request.method == 'POST':
        current = request.form.get('current_password', '')
        new_password = request.form.get('new_password', '')
        repeat = request.form.get('new_password_repeat', '')
        if not _check_password(current, current_user.password_hash):
            error = 'Текущий пароль указан неверно.'
        elif new_password != repeat:
            error = 'Новые пароли не совпадают.'
        elif _check_password(new_password, current_user.password_hash):
            error = 'Новый пароль должен отличаться от текущего.'
        else:
            error = validate_password(new_password, username=current_user.username)
        if not error:
            _set_user_password(current_user, new_password)
            _audit_in_transaction(
                'password_changed', {'reason': 'self_service'})
            db.session.commit()
            session['password_version'] = int(current_user.password_version or 1)
            flash('Пароль изменён.')
            return redirect(url_for('modules_dashboard'))
    return render_template('password_change.html', error=error,
                           forced=False, username=current_user.username)


@app.route('/logout', methods=['POST'])
@login_required
def logout():
    log_event('logout')
    logout_user()
    return redirect(url_for('landing'))


# ---------------------------------------------------------------------------
# Посадочная страница (публичная)
# ---------------------------------------------------------------------------
@app.route('/')
def landing():
    if current_user.is_authenticated:
        return redirect(url_for('modules_dashboard'))
    v09_label = Setting.get('landing_v09_label', 'Версия 0.9 alfa')
    v09_url   = Setting.get('landing_v09_url', '')
    return render_template('landing.html', second_url=None,
                           v09_label=v09_label, v09_url=v09_url)


@app.route('/favicon.ico')
def favicon():
    """Serve a lightweight built-in icon without a missing static asset."""
    icon = (
        '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 64 64">'
        '<rect width="64" height="64" rx="14" fill="#0757a6"/>'
        '<path d="M35 8 19 35h11l-3 21 18-30H34z" fill="#fff"/>'
        '</svg>'
    )
    return Response(
        icon,
        mimetype='image/svg+xml',
        headers={'Cache-Control': 'public, max-age=86400'},
    )


# ---------------------------------------------------------------------------
# Каталог модулей
# ---------------------------------------------------------------------------
@app.route('/modules')
@login_required
def modules_dashboard():
    modules = [
        dict(module, url=url_for(module['endpoint']))
        for module in registered_modules()
        if current_user.has_module_access(module['key'])
    ]
    return render_template('modules.html', modules=modules)

@app.route('/investment-reconciliation', endpoint='investment_reconciliation')
@login_required
def module_start():
    module = next(
        item for item in registered_modules()
        if item['endpoint'] == request.endpoint)
    return render_template('module_start.html', module=module)

def _limit_dataset_or_404(dataset_id):
    return db.session.get(LimitDataset, dataset_id) or abort(404)
def _rating_profile_payload(row=None):
    from rating import DEFAULT_PROFILE, validate_profile
    config = validate_profile(row.config if row else DEFAULT_PROFILE)
    return dict(id=row.id if row else 'base', name=config['name'],
                is_default=bool(row and row.is_default), builtin=row is None,
                config=config)
@app.route('/home')
@login_required
def home():
    query = Company.query.filter(
        Company.status != 'archived', Company.is_rso.is_(True))
    if not current_user.is_admin:
        query = query.join(CompanyMembership).filter(
            CompanyMembership.user_id == current_user.id)
    companies = query.order_by(Company.created_at.desc()).all()
    if (app.config['PORTFOLIO_HOME_ENABLED']
            and request.args.get('legacy') != '1'):
        from portfolio import build_portfolio
        portfolio = build_portfolio(
            companies, selected_year=request.args.get('year', type=int))
        response = Response(render_template(
            'portfolio.html', portfolio=portfolio), mimetype='text/html')
        response.headers['Cache-Control'] = 'no-store, no-cache, must-revalidate'
        return response
    # Автоматически восстанавливаем kpi_cache.json для компаний, у которых
    # уже есть дашборд, но кэш КПЭ отсутствует (например, после деплоя).
    for c in companies:
        out_dir   = os.path.join(OUTPUT, str(c.id))
        dash_path = os.path.join(out_dir, 'dashboard_current.html')
        kpi_path  = os.path.join(out_dir, 'kpi_cache.json')
        if os.path.isfile(dash_path) and not os.path.isfile(kpi_path):
            try:
                dataset, _err = _rebuild_dashboard_with_daily(c.id, c.name)
                if dataset is not None:
                    _save_kpi_cache(dataset, out_dir)
            except Exception as exc:
                app.logger.warning(f'auto kpi_cache regen failed for company {c.id}: {exc}')
    cards = [_company_summary(c) for c in companies]
    response = Response(
        render_template('home.html', companies=cards,
                        rating_profiles=_rating_profiles_for_user(),
                        historical_expert_enabled=(
                            app.config['HISTORICAL_IMPORT_ENABLED']
                             and app.config['HISTORICAL_EXPERT_ENABLED']),
                         comparison_enabled=(
                             app.config['COMPANY_COMPARISON_ENABLED']
                             and app.config['HISTORICAL_IMPORT_ENABLED']
                             and app.config['HISTORICAL_EXPERT_ENABLED'])),
        mimetype='text/html')
    # В шаблоне находится клиент API рейтинга. Не позволяем открытой вкладке
    # после обновления сервера продолжать выполнять устаревший JavaScript.
    response.headers['Cache-Control'] = 'no-store, no-cache, must-revalidate'
    return response


# ---------------------------------------------------------------------------
# Создание нового предприятия + первая загрузка
# ---------------------------------------------------------------------------
@app.route('/company/new', methods=['GET', 'POST'])
@login_required
@serialized_company_creation
def company_new():
    if request.method == 'GET':
        return render_template('company_new.html')

    ts = time.strftime('%Y%m%d_%H%M%S')

    # Если на второй POST пришёл tmp_key — читаем уже сохранённые файлы
    tmp_key = request.form.get('tmp_key', '').strip()
    if tmp_key:
        tmp_dir = os.path.join(UPLOADS, 'tmp', tmp_key)
        tmp_paths = sorted([
            os.path.join(tmp_dir, f)
            for f in os.listdir(tmp_dir)
            if f.lower().endswith(('.xls', '.xlsx'))
        ]) if os.path.isdir(tmp_dir) else []
    else:
        # Первый POST — получаем файлы из запроса
        files = [f for f in request.files.getlist('files') if f.filename]
        if not files:
            return render_template('company_new.html',
                                   error='Выберите хотя бы один файл ОСВ.')

        tmp_dir = os.path.join(UPLOADS, 'tmp', ts)
        os.makedirs(tmp_dir, exist_ok=True)
        tmp_paths = []
        staged_names = set()
        for f in files:
            name = _unique_attachment_name(f.filename, staged_names)
            if not name.lower().endswith(('.xls', '.xlsx')):
                continue
            dest = os.path.join(tmp_dir, name)
            f.save(dest)
            ok, err = _validate_upload(dest, name)
            if not ok:
                import shutil as _sh_cn; _sh_cn.rmtree(tmp_dir, ignore_errors=True)
                return render_template('company_new.html', error=err)
            tmp_paths.append(dest)
        tmp_key = ts

    if not tmp_paths:
        return render_template('company_new.html',
                               error='Среди загруженных нет .xls/.xlsx файлов.')

    # Название: из формы → из шапки файла → из имени файла
    org_name = (request.form.get('org_name') or '').strip()
    if not org_name:
        org_name = extract_org_name(tmp_paths[0]) or ''
    if not org_name:
        org_name = _org_name_from_filename(tmp_paths[0]) or ''
    if not org_name:
        return render_template('company_new.html',
                               ask_name=True,
                               tmp_key=tmp_key,
                               error='Не удалось определить название организации из файла. '
                                     'Введите его вручную.')

    # Do not run the new-company worker against a restored company's data.
    try:
        existing = resolve_existing_company(org_name)
    except ValueError as exc:
        db.session.rollback()
        return render_template('company_new.html', error=str(exc),
                               ask_name=True, tmp_key=tmp_key)
    if existing:
        shutil.rmtree(tmp_dir, ignore_errors=True)
        flash('Загруженные файлы не импортированы. При необходимости добавьте '
              'их через загрузку ОСВ в восстановленном предприятии.',
              'company_restored')
        return redirect(url_for('directory_company_profile', company_id=existing.id))

    # Создаём компанию в БД
    company = Company(name=org_name, created_by=current_user.id)
    db.session.add(company)
    db.session.flush()  # получаем company.id
    db.session.add(CompanyMembership(
        user_id=current_user.id, company_id=company.id,
        participation_type='creator', is_primary=False))

    # Перемещаем файлы в постоянное хранилище
    final_dir = os.path.join(UPLOADS, str(company.id), ts)
    os.makedirs(os.path.dirname(final_dir), exist_ok=True)
    shutil.move(tmp_dir, final_dir)
    final_paths = [os.path.join(final_dir, os.path.basename(p)) for p in tmp_paths]

    db.session.commit()  # commit company before starting thread
    _durable_backup('создание предприятия с исходными ОСВ')
    log_event('company_new_started', {'name': org_name, 'company_id': company.id})

    # Запускаем обработку в фоне
    back_url = url_for('company_detail', company_id=company.id)
    job_id = create_job(
        company.id,
        [os.path.basename(path) for path in final_paths],
    )
    threading.Thread(
        target=_bg_company_new,
        args=(job_id, company.id, final_dir, final_paths, ts,
              back_url, org_name, current_user.id),
        daemon=True,
    ).start()

    return redirect(url_for('company_detail', company_id=company.id, job=job_id))


# ---------------------------------------------------------------------------
# Страница предприятия
# ---------------------------------------------------------------------------
def _upload_history_rows(uploads, company_id):
    """Build useful, provenance-aware labels for the upload history table."""
    from key_indicators import _period_bounds

    def _short_filename(value, head=20, tail=14):
        name = os.path.basename(str(value or '').strip()) or 'Файл'
        if len(name) <= head + tail + 1:
            return name
        return f'{name[:head]}…{name[-tail:]}'

    def _short_text(value, limit=28):
        text = str(value or '').strip() or '—'
        return text if len(text) <= limit else f'{text[:limit - 1]}…'

    def _bounds(item, fallback_label=''):
        start = end = None
        try:
            if item.get('report_start'):
                start = datetime.strptime(
                    item['report_start'], '%Y-%m-%d').date()
            if item.get('report_end'):
                end = datetime.strptime(
                    item['report_end'], '%Y-%m-%d').date()
        except (TypeError, ValueError):
            start = end = None
        if start is None or end is None:
            try:
                start, end = _period_bounds(
                    item.get('period') or fallback_label or '')
            except (TypeError, ValueError):
                start, end = None, None
        return start, end

    def _period_type(start, end, granularity=None):
        if granularity == 'day':
            return 'День'
        if granularity == 'year':
            return 'Год'
        if granularity == 'month':
            return 'Месяц'
        if start is None or end is None or end < start:
            return '—'
        if start == end:
            return 'День'
        if start.year == end.year and (
                start.month, start.day, end.month, end.day) == (1, 1, 12, 31):
            return 'Год'
        return 'Месяц'

    def _coverage(start, end, accounts):
        if start is None or end is None or end < start:
            return '—'
        total_days = (end - start).days + 1
        query = db.session.query(
            func.count(func.distinct(DailyRow.date))
        ).filter(
            DailyRow.company_id == company_id,
            DailyRow.date >= start,
            DailyRow.date <= end,
        )
        if accounts:
            query = query.filter(DailyRow.acc.in_(sorted(accounts)))
        return f'{int(query.scalar() or 0)} из {total_days}'

    rows = []
    for upload in uploads:
        manifest = upload.source_manifest
        manifest_items = manifest if isinstance(manifest, list) else []
        file_rows = []
        for item in manifest_items:
            account = str(item.get('account') or '').strip()
            item_accounts = {account} if account else set()
            start, end = _bounds(item)
            relative_path = str(item.get('relative_path') or '').strip()
            filename = os.path.basename(relative_path) or 'Файл'
            observed_dates = {
                str(value) for value in (item.get('observed_dates') or [])
                if value
            }
            file_rows.append({
                'filename': filename,
                'display_filename': _short_filename(filename),
                'account_label': account or '—',
                'period_label': item.get('period') or '—',
                'period_type': _period_type(
                    start, end, item.get('granularity')),
                'coverage_label': (
                    f'{len(observed_dates)} из {(end - start).days + 1}'
                    if observed_dates and start is not None and end is not None
                    else '—' if item.get('granularity') == 'month'
                    else _coverage(start, end, item_accounts)
                ),
            })

        accounts = {
            str(item.get('account')).strip()
            for item in manifest_items
            if item.get('account') not in (None, '')
        }
        if manifest is None:
            accounts.update(
                str(account).strip()
                for account in upload.accounts_list
                if str(account).strip()
            )

        valid_starts_ends = [
            _bounds(item) for item in manifest_items
        ]
        starts = [start for start, _ in valid_starts_ends if start is not None]
        ends = [end for _, end in valid_starts_ends if end is not None]
        if starts and ends:
            start, end = min(starts), max(ends)
        else:
            start, end = _bounds({}, upload.period_label or '')
        granularities = {
            item.get('granularity') for item in manifest_items
            if item.get('granularity')
        }
        aggregate_granularity = (
            next(iter(granularities)) if len(granularities) == 1 else None
        )
        observed_dates = {
            str(value)
            for item in manifest_items
            for value in (item.get('observed_dates') or [])
            if value
        }

        rows.append({
            'upload': upload,
            'file_rows': file_rows,
            'file_label': (
                file_rows[0]['display_filename'] if len(file_rows) == 1
                else f'Пакет ({upload.files_count or len(file_rows)} файлов)'
            ) if file_rows else (
                f'Пакет ({upload.files_count} файлов)'
                if (upload.files_count or 0) > 1 else 'Файл'
            ),
            'full_file_label': (
                file_rows[0]['filename'] if len(file_rows) == 1 else ''
            ),
            'full_account_label': (
                ', '.join(sorted(accounts)) if accounts else '—'
            ),
            'account_label': _short_text(
                ', '.join(sorted(accounts)) if accounts else '—', 18),
            'period_display': _short_text(upload.period_label or '—', 32),
            'period_type': (
                'Смешанный' if len(granularities) > 1
                else _period_type(start, end, aggregate_granularity)
            ),
            'coverage_label': (
                f'{len(observed_dates)} из {(end - start).days + 1}'
                if observed_dates and start is not None and end is not None
                else '—' if aggregate_granularity == 'month'
                else _coverage(start, end, accounts)
            ),
        })
    return rows


@app.route('/company/<int:company_id>')
@login_required
def company_detail(company_id):
    company = db.session.get(Company, company_id) or abort(404)
    uploads = Upload.query.filter_by(company_id=company_id)\
                          .order_by(Upload.uploaded_at.desc()).all()
    completed_job_id = request.args.get('completed_job', '').strip()
    completed_job = get_job(completed_job_id) if completed_job_id else None
    if (not completed_job
            or completed_job.get('company_id') != company_id
            or completed_job.get('status') != 'done'):
        completed_job = None
    completed_result = completed_job.get('result') or {} if completed_job else {}
    year = request.args.get('year', None, type=int)
    ai_summary = AiSummary.query.filter_by(company_id=company_id).first()
    ai_provider_configured = AiProvider.query.filter_by(is_active=True).first() is not None
    from osv_ai_summary import (
        accepted_sources, serialized_file_summaries, source_set_revision,
    )
    ai_file_sources = accepted_sources(uploads)
    if ai_summary and ai_summary.source_revision != source_set_revision(
            ai_file_sources):
        ai_summary = None
    ai_file_summaries = serialized_file_summaries(company_id, ai_file_sources)

    upload_history_rows = _upload_history_rows(uploads, company_id)

    historical_imports = []
    historical_versions = []
    historical_preview = None
    historical_mappings_needed = []
    historical_section_options = []
    if app.config['HISTORICAL_IMPORT_ENABLED']:
        from historical_parser import SECTION_ALIASES
        historical_section_options = [
            (key, aliases[0]) for key, aliases in SECTION_ALIASES.items()]
        historical_imports = HistoricalImport.query.filter_by(
            company_id=company_id).order_by(
                HistoricalImport.created_at.desc()).all()
        historical_versions = HistoricalYearVersion.query.filter_by(
            company_id=company_id).order_by(
                HistoricalYearVersion.year.desc(),
                HistoricalYearVersion.version_no.desc()).all()
        preview_id = request.args.get('historical_preview', type=int)
        if preview_id:
            historical_preview = HistoricalImport.query.filter_by(
                id=preview_id, company_id=company_id).first()
            if historical_preview:
                parsed_preview = historical_preview.preview
                historical_mappings_needed.extend({
                    'kind': 'section',
                    'source': item['raw'],
                    'suggested': '',
                    'label': item['raw'],
                } for item in parsed_preview.get('unknown_sections', []))
                historical_mappings_needed.extend({
                    'kind': 'metric',
                    'source': item['raw'],
                    'suggested': item['suggested_key'],
                    'label': item['raw'],
                    'section': item['section'],
                } for item in parsed_preview.get('unknown_metrics', []))

    return render_template('company.html', company=company, uploads=uploads,
                           upload_history_rows=upload_history_rows,
                           audit=completed_result.get('audit'),
                           error=None,
                           discrepancy=completed_result.get('discrepancy', []),
                           annual_coverage=_annual_coverage(company_id, year),
                           ai_summary=ai_summary,
                           ai_file_summaries=ai_file_summaries,
                           ai_provider_configured=ai_provider_configured,
                           historical_enabled=app.config[
                               'HISTORICAL_IMPORT_ENABLED'],
                           historical_expert_enabled=(
                               app.config['HISTORICAL_IMPORT_ENABLED']
                               and app.config['HISTORICAL_EXPERT_ENABLED']),
                           historical_imports=historical_imports,
                           historical_versions=historical_versions,
                           historical_preview=historical_preview,
                           historical_mappings_needed=historical_mappings_needed,
                           historical_section_options=historical_section_options,
                           historical_error=request.args.get(
                               'historical_error', ''))


def _require_historical_import():
    if not app.config['HISTORICAL_IMPORT_ENABLED']:
        abort(404)


def _require_historical_expert():
    if (not app.config['HISTORICAL_IMPORT_ENABLED']
            or not app.config['HISTORICAL_EXPERT_ENABLED']):
        abort(404)


@app.route('/comparison')
@login_required
def company_comparison():
    """Изолированный режим сравнения только доступных пользователю РСО."""
    if (not app.config['COMPANY_COMPARISON_ENABLED']
            or not app.config['HISTORICAL_IMPORT_ENABLED']
            or not app.config['HISTORICAL_EXPERT_ENABLED']):
        abort(404)
    query = Company.query.filter(
        Company.status != 'archived', Company.is_rso.is_(True))
    if not current_user.is_admin:
        query = query.join(CompanyMembership).filter(
            CompanyMembership.user_id == current_user.id)
    visible = query.order_by(Company.name, Company.id).all()
    visible_by_id = {company.id: company for company in visible}
    selected_ids = list(dict.fromkeys(request.args.getlist('company', type=int)))
    if len(selected_ids) > 5:
        abort(400, description='Выберите не более пяти предприятий')
    if any(company_id not in visible_by_id for company_id in selected_ids):
        abort(403, description='Нет доступа к одному из предприятий')
    from comparison import MODES, build_comparison, build_metric_catalog
    selected_companies = [
        visible_by_id[company_id] for company_id in selected_ids]
    if len(selected_companies) >= 2:
        comparison = build_comparison(
            selected_companies,
            metric_refs=request.args.getlist('metric'),
            mode=request.args.get('mode', 'normalized'),
            profile_key=request.args.get('profile_key', 'methodical'),
            selected_years=request.args.getlist('year', type=int),
        )
    else:
        available_years = sorted({
            row[0] for row in db.session.query(HistoricalYearVersion.year).filter(
                HistoricalYearVersion.company_id.in_(visible_by_id),
                HistoricalYearVersion.is_active.is_(True),
            ).distinct().all()
        })[-4:]
        default_year = (
            2025 if 2025 in available_years
            else max(available_years) if available_years else None)
        catalog = build_metric_catalog(
            visible, available_years,
            request.args.get('profile_key', 'methodical'))
        comparison = {
            'companies': selected_companies,
            'years': [],
            'available_years': available_years,
            'metrics': catalog['metrics'],
            'groups': catalog['groups'],
            'selected_refs': [],
            'series': [],
            'table': [],
            'mode': 'normalized',
            'modes': MODES,
            'profile_key': request.args.get('profile_key', 'methodical'),
            'start_year': default_year,
            'end_year': default_year,
            'selected_years': [default_year] if default_year else [],
        }
    response = Response(render_template(
        'comparison.html', comparison=comparison,
        visible_companies=visible, selected_ids=selected_ids),
        mimetype='text/html')
    response.headers['Cache-Control'] = 'no-store, no-cache, must-revalidate'
    return response


@app.route('/company/<int:company_id>/historical/preview', methods=['POST'])
@login_required
def historical_preview_upload(company_id):
    """Сохранить XLSX и показать проверяемый предпросмотр без активации."""
    _require_historical_import()
    _require_company_access(company_id)
    source = request.files.get('historical_file')
    if source is None or not source.filename:
        return redirect(url_for(
            'company_detail', company_id=company_id,
            historical_error='Выберите файл XLSX.',
            _anchor='historical-analytics'))
    from historical_service import HistoricalImportError, stage_upload
    try:
        attempt = stage_upload(
            company_id, current_user.id, source, UPLOADS)
    except HistoricalImportError as exc:
        return redirect(url_for(
            'company_detail', company_id=company_id,
            historical_error=str(exc),
            _anchor='historical-analytics'))
    _durable_backup('предпросмотр годовой отчётности')
    log_event('historical_import_previewed', {
        'company_id': company_id,
        'import_id': attempt.id,
        'status': attempt.status,
        'filename': attempt.original_filename,
    })
    return redirect(url_for(
        'company_detail', company_id=company_id,
        historical_preview=attempt.id,
        _anchor='historical-analytics'))


@app.route(
    '/company/<int:company_id>/historical/<int:import_id>/confirm',
    methods=['POST'])
@login_required
def historical_confirm(company_id, import_id):
    """Подтвердить сопоставления и атомарно активировать новые версии."""
    _require_historical_import()
    _require_company_access(company_id)
    attempt = HistoricalImport.query.filter_by(
        id=import_id, company_id=company_id).first_or_404()
    from historical_parser import HistoricalParserError
    from historical_service import (
        HistoricalImportError, confirm_and_activate, parse_mapping_form)
    from sqlalchemy.exc import IntegrityError
    try:
        additions = parse_mapping_form(request.form)
        versions = confirm_and_activate(
            attempt,
            current_user.id,
            additions,
            confirm_partial=request.form.get('confirm_partial') == '1',
        )
    except (HistoricalImportError, HistoricalParserError, IntegrityError) as exc:
        db.session.rollback()
        message = (
            'Версия была одновременно изменена другим пользователем. '
            'Обновите страницу и повторите действие.'
            if isinstance(exc, IntegrityError) else str(exc))
        return redirect(url_for(
            'company_detail', company_id=company_id,
            historical_preview=import_id, historical_error=message,
            _anchor='historical-analytics'))
    _durable_backup('активация годовой отчётности')
    log_event('historical_import_activated', {
        'company_id': company_id,
        'import_id': import_id,
        'years': [version.year for version in versions],
        'versions': [version.version_no for version in versions],
        'partial': attempt.is_partial,
    })
    return redirect(url_for('company_detail', company_id=company_id))


@app.route(
    '/company/<int:company_id>/historical/<int:import_id>/cancel',
    methods=['POST'])
@login_required
def historical_cancel(company_id, import_id):
    """Отменить предпросмотр, не изменяя активные годовые срезы."""
    _require_historical_import()
    _require_company_access(company_id)
    attempt = HistoricalImport.query.filter_by(
        id=import_id, company_id=company_id).first_or_404()
    from historical_service import HistoricalImportError, cancel_attempt
    try:
        cancel_attempt(attempt)
    except HistoricalImportError as exc:
        return redirect(url_for(
            'company_detail', company_id=company_id,
            historical_preview=import_id, historical_error=str(exc),
            _anchor='historical-analytics'))
    _durable_backup('отмена годового предпросмотра', allow_removals=True)
    log_event('historical_import_cancelled', {
        'company_id': company_id, 'import_id': import_id})
    return redirect(url_for('company_detail', company_id=company_id))


@app.route(
    '/company/<int:company_id>/historical/<int:import_id>/source')
@login_required
def historical_source(company_id, import_id):
    """Скачать исходную книгу из истории загрузок."""
    _require_historical_import()
    _require_company_access(company_id)
    attempt = HistoricalImport.query.filter_by(
        id=import_id, company_id=company_id).first_or_404()
    if not os.path.isfile(attempt.stored_path):
        abort(404)
    return send_from_directory(
        os.path.dirname(attempt.stored_path),
        os.path.basename(attempt.stored_path),
        as_attachment=True,
        download_name=attempt.original_filename,
    )


@app.route('/company/<int:company_id>/historical/expert')
@login_required
def historical_expert(company_id):
    """Руководительский и экспертный режимы действующих годовых версий."""
    _require_historical_expert()
    company = _require_company_access(company_id)
    from historical_analytics import (
        build_expert_analytics, profile_payloads, validate_profile_key)
    saved_profile_key = Setting.get(
        f'historical_expert_profile:{company_id}', 'methodical')
    profile_key = validate_profile_key(
        request.args.get('profile_key', saved_profile_key))
    analytics = build_expert_analytics(company_id, profile_key)
    selectable_years = analytics['available_years'] or analytics['years']
    requested_years = request.args.getlist('years', type=int)
    has_year_selection = request.args.get('years_present') == '1'
    selected_years = (
        sorted(set(requested_years) & set(selectable_years))
        if has_year_selection
        else list(selectable_years)
    )
    selected_year = max(selected_years, default=None)
    executive_enabled = app.config['HISTORICAL_EXECUTIVE_ENABLED']
    view = request.args.get('view', 'executive' if executive_enabled else 'expert')
    if view == 'executive' and executive_enabled:
        from historical_executive import build_executive_summary
        summary = build_executive_summary(
            analytics, selected_years=selected_years)
        template_name = 'historical_executive.html'
    else:
        summary = None
        template_name = 'historical_expert.html'
    response = Response(render_template(
        template_name,
        company=company,
        analytics=analytics,
        summary=summary,
        profiles=profile_payloads(),
        selected_profile_key=profile_key,
        selected_year=selected_year,
        selected_years=selected_years,
        executive_enabled=executive_enabled,
        embedded=request.args.get('embedded') == '1',
    ), mimetype='text/html')
    response.headers['Cache-Control'] = 'no-store, no-cache, must-revalidate'
    return response


@app.route('/company/<int:company_id>/historical/expert/profile',
           methods=['POST'])
@login_required
def historical_expert_profile(company_id):
    """Сохранить именованный методический профиль предприятия."""
    _require_historical_expert()
    _require_company_access(company_id)
    from historical_analytics import PROFILES
    profile_key = request.form.get('profile_key', '').strip()
    if profile_key not in PROFILES:
        abort(400, description='Неизвестный методический профиль')
    Setting.set(f'historical_expert_profile:{company_id}', profile_key)
    db.session.commit()
    log_event('historical_expert_profile_changed', {
        'company_id': company_id,
        'profile_key': profile_key,
    })
    target = {'company_id': company_id}
    if request.form.get('embedded') == '1':
        target['embedded'] = '1'
    if request.form.get('view') in ('executive', 'expert'):
        target['view'] = request.form['view']
    selected_year = request.form.get('year', type=int)
    if selected_year:
        target['year'] = selected_year
    target['profile_key'] = profile_key
    return redirect(url_for('historical_expert', **target))


@app.route(
    '/company/<int:company_id>/historical/version/<int:version_id>/activate',
    methods=['POST'])
@login_required
def historical_version_activate(company_id, version_id):
    """Вернуть выбранную сохранённую версию календарного года."""
    _require_historical_import()
    _require_company_access(company_id)
    version = HistoricalYearVersion.query.filter_by(
        id=version_id, company_id=company_id).first_or_404()
    from historical_service import activate_version
    from sqlalchemy.exc import IntegrityError
    try:
        activate_version(version, current_user.id)
    except IntegrityError:
        db.session.rollback()
        return redirect(url_for(
            'company_detail', company_id=company_id,
            historical_error=(
                'Версия была одновременно изменена другим пользователем. '
                'Обновите страницу и повторите действие.')))
    _durable_backup('переключение годовой версии')
    log_event('historical_version_activated', {
        'company_id': company_id,
        'year': version.year,
        'version': version.version_no,
        'import_id': version.import_id,
    })
    return redirect(url_for('company_detail', company_id=company_id))


@app.route('/company/<int:company_id>/job/<job_id>')
@login_required
def job_status(company_id, job_id):
    """Опрос статуса фонового задания (polling endpoint)."""
    job = get_job(job_id)
    if not job or job.get('company_id') != company_id:
        return jsonify(status='not_found'), 404
    return jsonify(
        status=job['status'],
        error=job.get('error'),
        result=job.get('result'),
        files=job.get('files') or [],
    )


@app.route('/company/<int:company_id>/details', methods=['POST'])
@login_required
def company_details(company_id):
    """Сохраняет реквизиты предприятия (справочник РСО)."""
    company = db.session.get(Company, company_id) or abort(404)
    new_inn = clean_identifier(request.form.get('inn'), (10, 12))
    duplicates = duplicate_companies(new_inn, company.ogrn or '',
                                      company.kpp or '', exclude_id=company.id)
    if duplicates:
        abort(409, description='Предприятие с такими реквизитами уже существует')
    company.inn              = new_inn
    company.identity_key     = company_identity_key(
        company.inn, company.ogrn, company.kpp)
    company.full_name        = request.form.get('full_name',        '').strip()
    company.address          = request.form.get('address',          '').strip()
    company.director_name    = request.form.get('director_name',    '').strip()
    company.director_phone   = request.form.get('director_phone',   '').strip()
    company.accountant_name  = request.form.get('accountant_name',  '').strip()
    company.accountant_phone = request.form.get('accountant_phone', '').strip()
    company.note             = request.form.get('note',             '').strip()
    db.session.commit()
    log_event('company_details_updated', {'company_id': company_id, 'inn': company.inn})
    if _wants_json():
        return jsonify(ok=True)
    return redirect(url_for('company_detail', company_id=company_id))


@app.route('/company/registry/export')
@login_required
def company_registry_export():
    """Выгрузка реестра предприятий в CSV (логируется в audit_log)."""
    import io as _io
    import csv as _csv
    query = Company.query
    if not current_user.is_admin:
        query = query.join(CompanyMembership).filter(
            CompanyMembership.user_id == current_user.id)
    companies = query.order_by(Company.name).all()
    log_event('registry_export', {'count': len(companies)})
    out = _io.StringIO()
    w = _csv.writer(out, delimiter=';')
    w.writerow(['№', 'Сокр. наименование', 'Полное наименование', 'ИНН', 'Юр. адрес',
                'ФИО Руководителя', 'Тел. руководителя',
                'ФИО Гл. бухгалтера', 'Тел. гл. бухгалтера', 'Примечание'])
    for i, c in enumerate(companies, 1):
        w.writerow([i, c.name, c.full_name or '', c.inn or '', c.address or '',
                    c.director_name or '', c.director_phone or '',
                    c.accountant_name or '', c.accountant_phone or '', c.note or ''])
    csv_bytes = ('\ufeff' + out.getvalue()).encode('utf-8')
    return Response(
        csv_bytes,
        mimetype='text/csv; charset=utf-8',
        headers={'Content-Disposition': 'attachment; filename="Реестр_РСО.csv"'},
    )


@app.route('/company/<int:company_id>/ai-summary', methods=['POST'])
@login_required
def company_ai_summary(company_id):
    """Формирует ИИ-заключение по финансовым данным предприятия."""
    from ai_client import call_ai_provider as _call_ai
    company = db.session.get(Company, company_id) or abort(404)
    requested_company_revision = company.osv_source_revision or ''

    # Выбираем провайдера: явно помеченный default → иначе первый активный
    provider = (AiProvider.query.filter_by(is_default=True, is_active=True).first()
                or AiProvider.query.filter_by(is_active=True).first())
    if not provider or not provider.base_url or not provider.model:
        return jsonify(ok=False, error='ИИ-провайдер не настроен'), 400

    api_key = decrypt_value(provider.api_key) if provider.api_key else ''

    # ── Сбор данных ──────────────────────────────────────────────────────────
    uploads = Upload.query.filter_by(company_id=company_id)\
                          .order_by(Upload.uploaded_at.desc()).all()
    from osv_ai_summary import (
        accepted_sources, range_evidence, source_set_revision,
        verified_source_facts,
    )
    sources = accepted_sources(uploads)
    requested_revision = source_set_revision(sources)
    all_daily_rows = DailyRow.query.filter_by(company_id=company_id)\
        .order_by(DailyRow.date.asc(), DailyRow.id.asc()).all()
    def _format_ai_amount(value):
        return f'{float(value or 0):,.2f}'.replace(',', ' ')

    evidence = range_evidence(
        sources, extra_dates=[row.date for row in all_daily_rows])
    source_facts, unavailable_sources = verified_source_facts(
        company_id, sources, UPLOADS)

    # Факты по всему диапазону суточных наблюдений, без подмены пропусков.
    key_accs = {'51': 'Сч.51 Расчётный счёт',
                '62': 'Сч.62 Дебиторы (потребители)',
                '60': 'Сч.60 Кредиторы (поставщики)',
                '70': 'Сч.70 Расчёты с персоналом'}
    daily_facts = {}
    for acc in key_accs:
        rows = [row for row in all_daily_rows if row.acc == acc]
        if rows:
            daily_facts[acc] = {
                'start': rows[0].date, 'end': rows[-1].date,
                'opening': rows[0].snd - rows[0].snk,
                'closing': rows[-1].skd - rows[-1].skk,
                'obd': sum(row.obd or 0 for row in rows),
                'obk': sum(row.obk or 0 for row in rows),
                'observations': len({row.date for row in rows}),
            }

    # ── Формирование промпта ─────────────────────────────────────────────────
    lines = [f'Организация: {company.name}']
    if evidence['start'] and evidence['end']:
        lines.append(
            f'Фактический диапазон всех принятых ОСВ: '
            f'{evidence["start"]:%d.%m.%Y}–{evidence["end"]:%d.%m.%Y}')
        lines.append(
            f'Покрытые календарные месяцы: {len(evidence["covered_months"])}')
        lines.append(
            'Разрывы между периодами: ' + (
                ', '.join(
                    f'{start:%d.%m.%Y}'
                    if start == end else f'{start:%d.%m.%Y}–{end:%d.%m.%Y}'
                    for start, end in evidence['missing_intervals'][:20])
                + (f' и ещё {len(evidence["missing_intervals"]) - 20}'
                   if len(evidence['missing_intervals']) > 20 else '')
                if evidence['missing_intervals'] else 'не обнаружены'))
    else:
        lines.append('Фактический диапазон ОСВ не подтверждён метаданными источников.')
    lines.append(f'Принятых версий исходных файлов: {len(sources)}')
    all_accounts = sorted({item['account'] for item in sources if item['account']})
    if all_accounts:
        lines.append(f'Счета в принятых файлах: {", ".join(all_accounts)}')
    if source_facts:
        lines.append('\nПроверенные факты каждого принятого файла:')
        lines.extend(f'  {fact}' for fact in source_facts)
    if unavailable_sources:
        lines.append('\nНедоступные или изменённые источники (не использованы):')
        lines.extend(f'  {item}' for item in unavailable_sources)

    if daily_facts:
        lines.append('\nФакты по ключевым счетам из всех суточных наблюдений:')
        for acc, info in daily_facts.items():
            lines.append(
                f'  {key_accs[acc]}: {info["start"]:%d.%m.%Y}–'
                f'{info["end"]:%d.%m.%Y}, наблюдений {info["observations"]}; '
                f'сальдо {_format_ai_amount(info["opening"])} → '
                f'{_format_ai_amount(info["closing"])} руб.; '
                f'сумма наблюдаемых оборотов Дт '
                f'{_format_ai_amount(info["obd"])}, Кт '
                f'{_format_ai_amount(info["obk"])} руб.')

    lines.append(f'\nВсего пакетов загрузок: {len(uploads)}')

    system_msg = (
        'Ты финансовый аналитик, специализирующийся на анализе данных '
        'оборотно-сальдовых ведомостей (ОСВ) организаций жилищно-коммунального '
        'хозяйства России. Сформируй краткое профессиональное заключение строго '
        'на русском языке и в Markdown по структуре:\n'
        '### Общая оценка\n'
        'Один краткий абзац, не более 4 предложений.\n'
        '### Ключевые риски\n'
        'Маркированный список из 1–3 коротких пунктов.\n'
        '### Сильные стороны\n'
        'Маркированный список из 1–3 коротких пунктов; если данных недостаточно, '
        'прямо укажи это.\n'
        '### Рекомендации\n'
        'Маркированный список из 2–3 конкретных действий.\n'
        'Общий объём ответа — не более 1400 знаков с пробелами. '
        'Не используй таблицы и HTML. Не придумывай отсутствующие показатели, '
        'причины или суммы. Отделяй установленные факты от предположений. '
        'Если данных недостаточно для вывода, явно напиши об ограничении.'
    )
    user_msg = '\n'.join(lines)

    # ── Вызов ИИ ─────────────────────────────────────────────────────────────
    try:
        result = _call_ai(
            base_url=provider.base_url,
            api_key=api_key,
            model=provider.model,
            messages=[
                {'role': 'system', 'content': system_msg},
                {'role': 'user',   'content': user_msg},
            ],
            temperature=0.3,
            max_tokens=2200,
            enable_thinking=False,
            timeout=90,
        )
    except Exception as e:
        app.logger.error('AI summary error for company %s: %s', company_id, e)
        return jsonify(
            ok=False,
            error='ИИ-провайдер временно недоступен. Повторите запрос позже.',
        ), 502

    content = result['content']
    usage   = result['usage']

    # ── Сохранение (upsert) ───────────────────────────────────────────────────
    db.session.rollback()
    with db.session.begin():
        locked_company = Company.query.filter_by(id=company_id)\
            .with_for_update().first() or abort(404)
        current_sources = accepted_sources(
            Upload.query.filter_by(company_id=company_id).all())
        if (
            (locked_company.osv_source_revision or '')
                != requested_company_revision
            or source_set_revision(current_sources) != requested_revision
        ):
            return jsonify(
                ok=False,
                error=('Набор ОСВ изменился во время формирования. '
                       'Запустите заключение повторно.'),
            ), 409
        summary = AiSummary.query.filter_by(company_id=company_id)\
            .with_for_update().first()
        if summary is None:
            summary = AiSummary(company_id=company_id)
            db.session.add(summary)
        summary.content           = content
        summary.provider_name     = provider.name
        summary.model_name        = provider.model
        summary.prompt_tokens     = usage.get('prompt_tokens', 0)
        summary.completion_tokens = usage.get('completion_tokens', 0)
        summary.generated_at      = datetime.now(timezone.utc)
        summary.analysis_start    = evidence['start']
        summary.analysis_end      = evidence['end']
        summary.source_revision   = requested_revision
        db.session.flush()
        generated_at = summary.generated_at

    log_event('ai_summary_generated', {
        'company_id': company_id,
        'tokens': usage.get('prompt_tokens', 0) + usage.get('completion_tokens', 0),
    })

    return jsonify(
        ok=True,
        content=content,
        meta=dict(
            provider=provider.name,
            model=provider.model,
            tokens=usage.get('prompt_tokens', 0) + usage.get('completion_tokens', 0),
            generated_at=generated_at.strftime('%d.%m.%Y %H:%M'),
            range_label=(
                f'{evidence["start"]:%d.%m.%Y}–{evidence["end"]:%d.%m.%Y}'
                if evidence['start'] and evidence['end'] else 'не определён'),
        ),
    )


@app.route('/company/<int:company_id>/ai-file-summaries', methods=['GET'])
@login_required
def company_ai_file_summaries(company_id):
    db.session.get(Company, company_id) or abort(404)
    from osv_ai_summary import (accepted_sources, serialized_file_summaries)
    uploads = Upload.query.filter_by(company_id=company_id).all()
    sources = accepted_sources(uploads)
    return jsonify(ok=True, files=serialized_file_summaries(company_id, sources))


@app.route('/company/<int:company_id>/ai-file-summary/<source_id>',
           methods=['POST'])
@login_required
def company_ai_file_summary(company_id, source_id):
    db.session.get(Company, company_id) or abort(404)
    if not re.fullmatch(r'[0-9a-f]{64}', source_id or ''):
        abort(404)
    from osv_ai_summary import ensure_file_summaries, serialized_file_summaries
    sources = ensure_file_summaries(
        company_id, UPLOADS, source_ids=[source_id])
    if not any(item['source_id'] == source_id for item in sources):
        abort(404)
    item = next(item for item in serialized_file_summaries(
        company_id, sources) if item['source_id'] == source_id)
    return jsonify(ok=True, file=item), 202 if item['status'] in (
        'queued', 'running') else 200


# ---------------------------------------------------------------------------
# Единая загрузка файлов (.xls/.xlsx и/или .zip)
# ---------------------------------------------------------------------------
@app.route('/company/<int:company_id>/upload-all', methods=['POST'])
@login_required
def company_upload_all(company_id):
    company = db.session.get(Company, company_id) or abort(404)
    all_files = [f for f in request.files.getlist('files') if f.filename]

    def _err(msg):
        ups = Upload.query.filter_by(company_id=company_id)\
                          .order_by(Upload.uploaded_at.desc()).all()
        return render_template('company.html', company=company, uploads=ups,
                               upload_history_rows=_upload_history_rows(
                                   ups, company_id),
                               audit=None, error=msg, discrepancy=[],
                               annual_coverage=_annual_coverage(company_id))

    if not all_files:
        return _err('Файлы не выбраны.')

    ts = time.strftime('%Y%m%d_%H%M%S')

    # Сохраняем все файлы в staging-папку синхронно (быстро: только дисковый ввод-вывод)
    stage_dir = os.path.join(UPLOADS, str(company_id), f'stage_{ts}')
    os.makedirs(stage_dir, exist_ok=True)

    saved = 0
    staged_names = set()
    saved_names = []
    for f in all_files:
        name = _unique_attachment_name(f.filename, staged_names)
        ext = os.path.splitext(name)[1].lower()
        if ext not in ('.xls', '.xlsx', '.zip'):
            continue
        dest = os.path.join(stage_dir, name)
        f.save(dest)
        ok, err = _validate_upload(dest, name)
        if not ok:
            import shutil as _sh_v; _sh_v.rmtree(stage_dir, ignore_errors=True)
            return _err(err)
        saved += 1
        saved_names.append(name)

    if not saved:
        import shutil as _sh; _sh.rmtree(stage_dir, ignore_errors=True)
        return _err('Среди загруженных нет файлов .xls, .xlsx или .zip.')

    job_id = create_job(
        company_id,
        saved_names,
    )
    threading.Thread(
        target=_bg_upload_all,
        args=(job_id, company_id, company.name, stage_dir, ts, current_user.id),
        daemon=True,
    ).start()

    # Долговечные снимки делает сам worker после сохранения строк и итоговой
    # загрузки. Не блокируем быстрый ответ браузеру общим after_request backup.
    g.durable_backup_done = True
    return redirect(url_for('company_detail', company_id=company_id, job=job_id))


# ---------------------------------------------------------------------------
# Дашборд
# ---------------------------------------------------------------------------
@app.route('/company/<int:company_id>/key-indicators')
@login_required
def key_indicators_page(company_id):
    company = _require_company_access(company_id)
    return render_template(
        'key_indicators.html', company=company,
        data_url=url_for('key_indicators_data', company_id=company_id),
        ai_report_url=url_for('key_indicators_ai_report', company_id=company_id),
        mappings_url=url_for('key_indicator_revenue_mappings',
                             company_id=company_id),
        cost_mappings_url=url_for('key_indicator_cost_mappings',
                                  company_id=company_id),
        debtor_mappings_url=url_for('key_indicator_debtor_mappings',
                                    company_id=company_id),
        embedded=request.args.get('embedded') == '1',
    )


def _key_indicator_request_period():
    """Validate common operational KPI arguments."""
    period = request.args.get('period', 'month')
    endpoint = request.args.get('date') or None
    date_from = request.args.get('date_from') or None
    if period not in ('day', 'month', 'quarter', 'year'):
        raise ValueError('Неизвестный период')
    for value in (endpoint, date_from):
        if value is None:
            continue
        try:
            datetime.strptime(value, '%Y-%m-%d')
        except ValueError:
            raise ValueError('Дата должна иметь формат ГГГГ-ММ-ДД')
    if date_from is not None and period not in ('day', 'month', 'quarter'):
        raise ValueError('Начальная дата доступна для дня, месяца и квартала')
    return period, endpoint, date_from


def _key_indicator_mappings(company_id):
    """Read mappings used by both the JSON cards and factual report."""
    return (
        [{'account': row.account, 'analytic_key': row.analytic_key,
          'analytic_label': row.analytic_label, 'category': row.category}
         for row in RevenueAnalyticMapping.query.filter_by(company_id=company_id).all()],
        [{'account': row.account, 'analytic_key': row.analytic_key,
          'analytic_label': row.analytic_label, 'category': row.category,
          'target_account': row.target_account,
          'valid_from': row.valid_from.isoformat(),
          'valid_to': row.valid_to.isoformat()}
         for row in CostAnalyticMapping.query.filter_by(company_id=company_id).all()],
        [{'id': row.id, 'account': row.account, 'analytic_key': row.analytic_key,
          'analytic_label': row.analytic_label, 'category': row.category,
          'effective_from': row.effective_from.isoformat(),
          'effective_to': row.effective_to.isoformat()}
         for row in DebtorAnalyticMapping.query.filter_by(company_id=company_id).all()],
    )


def _operational_key_indicators_payload(company, period, endpoint, date_from):
    """Calculate the one authoritative period payload used by both views."""
    from key_indicators import build_key_indicators
    from key_indicators_service import load_company_source, build_period_passport

    source, updated_at = load_company_source(company.id, UPLOADS)
    mappings, cost_mappings, debtor_mappings = _key_indicator_mappings(company.id)
    collection_accruals = [{
        'date': row.accrual_date.isoformat(),
        'account': row.account,
        'analytic_key': row.analytic_key,
        'type': row.entry_type,
        'accrued': row.accrued,
        'accrual_document': row.accrual_document,
        'source': row.source_name,
        'version': row.source_version,
    } for row in ConsumerCollectionAccrual.query.filter_by(
        company_id=company.id).all()]
    collection_allocations = [{
        'date': row.payment_date.isoformat(),
        'payment_document': row.payment_document,
        'accrual_document': row.accrual_document,
        'type': row.entry_type,
        'allocated_amount': row.allocated_amount,
        'source': row.source_name,
        'version': row.source_version,
    } for row in ConsumerPaymentAllocation.query.filter_by(
        company_id=company.id).all()]
    collection_norms = [{
        'category': row.category,
        'plan_percent': row.plan_percent,
        'approved': row.approved,
        'approved_at': row.approved_at.isoformat(),
        'valid_from': row.valid_from.isoformat(),
        'valid_to': row.valid_to.isoformat(),
        'source': row.source_name,
        'version': row.source_version,
    } for row in ConsumerCollectionNorm.query.filter_by(
        company_id=company.id).all()]
    payload = build_key_indicators(
        source, period=period, date=endpoint, date_from=date_from,
        revenue_mappings=mappings, cost_mappings=cost_mappings,
        debtor_mappings=debtor_mappings,
        collection_accruals=collection_accruals,
        collection_allocations=collection_allocations,
        collection_norms=collection_norms)
    payload['company'] = {'id': company.id, 'name': company.name}
    payload['updated_at'] = updated_at
    payload['source_passport'] = build_period_passport(source, payload)
    return payload


@app.route('/company/<int:company_id>/key-indicators/data')
@login_required
def key_indicators_data(company_id):
    from key_indicators_service import (
        SourceBusyError, SourceReadError,
    )
    company = _require_company_access(company_id)
    try:
        period, endpoint, date_from = _key_indicator_request_period()
        payload = _operational_key_indicators_payload(
            company, period, endpoint, date_from)
    except SourceBusyError as exc:
        response = jsonify(error=str(exc))
        response.status_code = 503
        response.headers['Retry-After'] = '3'
        return response
    except SourceReadError as exc:
        return jsonify(error=str(exc)), 422
    except ValueError as exc:
        return jsonify(error=str(exc)), 400
    except Exception:
        app.logger.exception('Key indicators source failed for company %s', company_id)
        return jsonify(error='Не удалось рассчитать показатели. Повторите запрос позже.'), 500
    response = jsonify(payload)
    response.headers['Cache-Control'] = 'private, no-store'
    if ('gzip' in request.headers.get('Accept-Encoding', '').lower()
            and len(response.get_data()) >= 1024 * 256):
        response.set_data(gzip.compress(response.get_data(), compresslevel=5))
        response.headers['Content-Encoding'] = 'gzip'
        response.headers['Vary'] = 'Accept-Encoding'
    return response


def _ai_report_period_label(payload, annual=False):
    """Present the exact normalized period, rather than a source-file label."""
    if annual:
        period_label = payload.get('period_label')
        if period_label:
            return str(period_label)
        year_from, year_to = payload.get('year_from'), payload.get('year_to')
        if year_from is not None and year_to is not None:
            return (str(year_from) if year_from == year_to
                    else f'{year_from}–{year_to}')
        return str(payload.get('year') or 'год не указан')
    period = payload.get('period') or {}

    def displayed(value):
        try:
            return datetime.strptime(value, '%Y-%m-%d').strftime('%d.%m.%Y')
        except (TypeError, ValueError):
            return None

    start, end = displayed(period.get('start')), displayed(period.get('end'))
    if start and end:
        return start if start == end else f'{start} — {end}'
    return period.get('label') or 'период не указан'


def _annual_key_indicator_request_range():
    """Validate the annual endpoint's backward-compatible year arguments."""
    raw_year = request.args.get('year')
    raw_from = request.args.get('year_from')
    raw_to = request.args.get('year_to')
    if (raw_from is None) != (raw_to is None):
        raise ValueError(
            'Параметры year_from и year_to должны передаваться вместе')

    def parse(raw, label):
        if raw is None or not re.fullmatch(r'\d{4}', raw):
            raise ValueError(f'{label} должен иметь формат ГГГГ')
        value = int(raw)
        if not 1900 <= value <= 2100:
            raise ValueError(f'{label} должен иметь формат ГГГГ')
        return value

    year = parse(raw_year, 'Год') if raw_year is not None else None
    year_from = parse(raw_from, 'Начальный год') if raw_from is not None else None
    year_to = parse(raw_to, 'Конечный год') if raw_to is not None else None
    if year_from is not None:
        if year_from > year_to:
            raise ValueError('Начальный год не может быть больше конечного')
        if year is not None and year != year_to:
            raise ValueError(
                'Параметр year должен совпадать с концом выбранного диапазона')
    return year, year_from, year_to


def _ai_report_value(value, unit):
    """Formatting only; calculations remain in the KPI service."""
    if unit == 'RUB':
        return f'{value / 1000000:,.2f}'.replace(',', ' ') + ' млн ₽'
    if unit == 'percent':
        return f'{value:,.2f}'.replace(',', ' ') + ' %'
    return f'{value:,.2f}'.replace(',', ' ')


def _ai_report_narrative(payload):
    """Conservative statements that do not infer facts absent from ОСВ."""
    cards = {item.get('id'): item for item in payload.get('cards') or []}
    result = []
    for card_id in ('revenue', 'profit', 'cash', 'ar'):
        card = cards.get(card_id) or {}
        value = card.get('value')
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            # The KPI profit card can explicitly fall back to gross profit
            # when selling/administrative expenses are unavailable.  Do not
            # call that fallback a sales result.
            title = card.get('title') or card_id
            if card_id == 'profit' and card.get('profit_from_sales') is None:
                title = 'Валовая прибыль'
            result.append(f'{title}: {_ai_report_value(value, card.get("unit"))}.')
    return result[:4]


def _ai_report_cards(payload):
    """Add presentation-only coverage markers without changing KPI facts."""
    result = []
    for item in payload.get('cards') or []:
        card = dict(item)
        partial = False
        missing = []
        if card.get('id') == 'cash':
            partial = card.get('coverage') == 'partial'
            missing = list(card.get('missing_accounts') or [])
        elif card.get('id') == 'credits':
            coverage = card.get('coverage') or {}
            partial = (card.get('value') is not None
                       and coverage.get('complete') is False)
            missing = list(coverage.get('missing') or [])
        card['report_partial_balance'] = partial
        card['report_missing_accounts'] = missing
        result.append(card)
    return result


def _ai_brief_payload(payload, period_label):
    """Prepare optional AI facts solely from the selected server calculation."""
    cards = {item.get('id'): item for item in payload.get('cards') or []}
    revenue, cost, profit = (cards.get(key) or {} for key in
                             ('revenue', 'cogs', 'profit'))
    data = {
        'company_name': payload.get('company', {}).get('name', ''),
        'period': period_label,
    }
    gross = next((item.get('value') for item in revenue.get('details') or []
                  if item.get('label') == 'Выручка брутто'), None)
    vat = next((item.get('value') for item in revenue.get('details') or []
                if item.get('label') == 'НДС'), None)
    selling, administrative = profit.get('selling_expense'), profit.get('admin_expense')
    values = (gross, vat, cost.get('value'), selling, administrative)
    if all(isinstance(value, (int, float)) and not isinstance(value, bool)
           for value in values):
        data['pl'] = {
            'rev_total': gross, 'nds_total': vat, 'cost_total': cost['value'],
            'adm_total': selling + administrative,
        }
    cash = cards.get('cash') or {}
    # /ai/brief has no representation for an incomplete cash composition.
    # Omit it rather than inviting the provider to call a partial sum total.
    if cash.get('coverage') == 'complete' and isinstance(cash.get('value'), (int, float)):
        data['cash'] = {'end': cash['value']}
    # Do not offer an AI request that contains only a company name and period:
    # that would invite an invented conclusion when no compatible facts exist.
    return data if ('pl' in data or 'cash' in data) else None


@app.route('/company/<int:company_id>/key-indicators/ai-report')
@login_required
def key_indicators_ai_report(company_id):
    """Render a factual report from the same calculation as the KPI cards.

    A saved dashboard is deliberately not embedded: its balances belong to the
    saved ОСВ boundary and cannot safely describe an arbitrary selected date.
    """
    from key_indicators_service import SourceBusyError, SourceReadError

    company = _require_company_access(company_id)
    mode = request.args.get('mode', 'operational')
    if mode not in ('operational', 'annual'):
        return render_template('key_indicators_ai_report.html', company=company,
                               error='Неизвестный режим отчёта'), 400
    try:
        if mode == 'annual':
            from annual_key_indicators import build_annual_key_indicators
            _require_historical_import()
            year, year_from, year_to = _annual_key_indicator_request_range()
            payload = build_annual_key_indicators(
                company, requested_year=year,
                year_from=year_from, year_to=year_to)
        else:
            period, endpoint, date_from = _key_indicator_request_period()
            payload = _operational_key_indicators_payload(
                company, period, endpoint, date_from)
    except SourceBusyError as exc:
        response = Response(render_template(
            'key_indicators_ai_report.html', company=company, error=str(exc)),
            status=503)
        response.headers['Retry-After'] = '3'
        return response
    except SourceReadError as exc:
        return render_template('key_indicators_ai_report.html', company=company,
                               error=str(exc)), 422
    except ValueError as exc:
        return render_template('key_indicators_ai_report.html', company=company,
                               error=str(exc)), 400
    except Exception:
        app.logger.exception('Key indicator AI report failed for company %s', company_id)
        return render_template(
            'key_indicators_ai_report.html', company=company,
            error='Не удалось сформировать отчёт. Повторите запрос позже.'), 500

    annual = mode == 'annual'
    period_label = _ai_report_period_label(payload, annual)
    response = render_template(
        'key_indicators_ai_report.html', company=company, payload=payload,
        annual=annual, period_label=period_label,
        narrative=_ai_report_narrative(payload),
        report_cards=_ai_report_cards(payload),
        ai_brief_payload=(None if annual else _ai_brief_payload(payload, period_label)))
    return Response(response, headers={'Cache-Control': 'private, no-store'})


def _metric_layout_response(payload, status=200):
    """JSON response for a layout request, never reusable by a cache."""
    response = jsonify(payload)
    response.headers['Cache-Control'] = 'no-store'
    return response, status


@app.route(
    '/company/<int:company_id>/key-indicators/layout/<mode>/<metric_id>',
    methods=['GET', 'PUT'])
@login_required
def metric_layout(company_id, mode, metric_id):
    """Read or update one user's key-indicator layout preference."""
    _require_company_access(company_id)
    try:
        validate_mode_metric(mode, metric_id)
    except LayoutValidationError as exc:
        return _metric_layout_response({'error': str(exc)}, 400)

    # Both public spellings identify one reserved preference record; aliases
    # must not create two independent dashboard layouts.
    stored_metric_id = ('__dashboard__' if metric_id in RESERVED_DASHBOARD_IDS
                        else metric_id)
    query = MetricLayout.query.filter_by(
        user_id=current_user.id,
        company_id=company_id,
        mode=mode,
        metric_id=stored_metric_id,
    )
    if request.method == 'GET':
        row = query.one_or_none()
        if row is None:
            return _metric_layout_response({
                'layout': {'version': 1, 'items': {}},
                'revision': 0,
            })
        try:
            layout = json.loads(row.layout_json)
        except (TypeError, ValueError):
            # Do not silently turn a corrupt persisted preference into a
            # different preference.
            return _metric_layout_response(
                {'error': 'Сохранённая раскладка повреждена'}, 500)
        return _metric_layout_response({
            'layout': layout,
            'revision': row.revision,
        })

    if not request.is_json:
        return _metric_layout_response(
            {'error': 'Тело запроса должно быть JSON-объектом'}, 400)
    payload = request.get_json(silent=True)
    try:
        normalized = validate_layout_payload(
            payload, mode=mode, metric_id=metric_id)
    except LayoutValidationError as exc:
        return _metric_layout_response({'error': str(exc)}, 400)

    # PostgreSQL honours FOR UPDATE. The conditional UPDATE below also makes
    # the revision check atomic on SQLite, where FOR UPDATE is ignored.
    row = query.with_for_update().one_or_none()
    expected_revision = normalized['revision']
    new_revision = 1
    if row is None:
        if expected_revision != 0:
            return _metric_layout_response(
                {'error': 'Конфликт revision: раскладка уже изменена'}, 409)
        row = MetricLayout(
            user_id=current_user.id,
            company_id=company_id,
            mode=mode,
            metric_id=stored_metric_id,
            layout_json=json.dumps(
                normalized['layout'], ensure_ascii=False, separators=(',', ':')),
            revision=1,
        )
        db.session.add(row)
    else:
        if row.revision != expected_revision:
            return _metric_layout_response(
                {'error': 'Конфликт revision: раскладка уже изменена'}, 409)
        new_revision = expected_revision + 1
        changed = db.session.execute(
            MetricLayout.__table__.update().where(
                MetricLayout.id == row.id,
                MetricLayout.revision == expected_revision,
            ).values(
                layout_json=json.dumps(
                    normalized['layout'], ensure_ascii=False,
                    separators=(',', ':')),
                revision=new_revision,
                updated_at=utcnow(),
            )
        ).rowcount
        if changed != 1:
            db.session.rollback()
            return _metric_layout_response(
                {'error': 'Конфликт revision: раскладка уже изменена'}, 409)

    try:
        db.session.commit()
    except IntegrityError:
        # A concurrent first insert can win the unique key between the
        # lookup and commit. Never report success or overwrite that row.
        db.session.rollback()
        current = query.one_or_none()
        if current is not None:
            return _metric_layout_response(
                {'error': 'Конфликт revision: раскладка уже изменена'}, 409)
        raise

    return _metric_layout_response({
        'layout': normalized['layout'],
        'revision': new_revision,
    })


@app.route('/company/<int:company_id>/key-indicators/revenue-mappings',
           methods=['POST'])
@admin_required
def key_indicator_revenue_mappings(company_id):
    company = db.session.get(Company, company_id) or abort(404)
    if company.is_archived:
        abort(409, description='Архивное предприятие доступно только для чтения')
    body = request.get_json(silent=True) or {}
    items = body.get('mappings')
    if not isinstance(items, list):
        return jsonify(error='Поле mappings должно быть списком'), 400
    allowed_categories = {
        'population', 'uk', 'budget', 'legal', 'subsidy'}
    normalized = {}
    for item in items:
        if not isinstance(item, dict):
            return jsonify(error='Некорректная запись сопоставления'), 400
        account = str(item.get('account') or '').strip()
        analytic_key = str(item.get('analytic_key') or '').strip()
        label = str(item.get('analytic_label') or '').strip()
        category = str(item.get('category') or '').strip()
        if (account not in ('90.01', '90.03')
                or not re.fullmatch(r'[0-9a-f]{64}', analytic_key)
                or not label or len(label) > 1000):
            return jsonify(error='Некорректная аналитика 90.01/90.03'), 400
        if category and category not in allowed_categories:
            return jsonify(error='Неизвестная категория выручки'), 400
        normalized[(account, analytic_key)] = (label, category or None)
    existing = {
        (row.account, row.analytic_key): row
        for row in RevenueAnalyticMapping.query.filter_by(
            company_id=company_id).all()
    }
    for key, (label, category) in normalized.items():
        row = existing.get(key)
        if category is None:
            if row is not None:
                db.session.delete(row)
            continue
        if row is None:
            db.session.add(RevenueAnalyticMapping(
                company_id=company_id, account=key[0], analytic_key=key[1],
                analytic_label=label, category=category,
                updated_by=current_user.id))
        else:
            row.analytic_label = label
            row.category = category
            row.updated_by = current_user.id
    db.session.commit()
    _durable_backup(
        'сопоставление аналитик выручки',
        allowed_removed_tables={'revenue_analytic_mappings'})
    saved = RevenueAnalyticMapping.query.filter_by(
        company_id=company_id).count()
    return jsonify(saved=saved)

@app.route('/company/<int:company_id>/key-indicators/debtor-mappings',
           methods=['POST'])
@admin_required
def key_indicator_debtor_mappings(company_id):
    company = db.session.get(Company, company_id) or abort(404)
    if company.is_archived:
        abort(409, description='Архивное предприятие доступно только для чтения')
    body = request.get_json(silent=True) or {}
    items = body.get('mappings')
    if not isinstance(items, list):
        return jsonify(error='Поле mappings должно быть списком'), 400
    allowed_categories = {'population', 'uk', 'budget', 'legal', 'subsidy'}
    normalized = {}
    for item in items:
        if not isinstance(item, dict):
            return jsonify(error='Некорректная запись сопоставления'), 400
        account = str(item.get('account') or '').strip()
        analytic_key = str(item.get('analytic_key') or '').strip()
        label = str(item.get('analytic_label') or '').strip()
        category = str(item.get('category') or '').strip()
        try:
            effective_from = datetime.strptime(
                str(item.get('effective_from') or ''), '%Y-%m-%d').date()
            effective_to = datetime.strptime(
                str(item.get('effective_to') or ''), '%Y-%m-%d').date()
        except ValueError:
            return jsonify(error='Период действия должен иметь формат ГГГГ-ММ-ДД'), 400
        if (account != '62.01' and not re.fullmatch(r'76(?:\.\d+)+', account)
                or not re.fullmatch(r'[0-9a-f]{64}', analytic_key)
                or not label or len(label) > 1000
                or effective_from > effective_to):
            return jsonify(error='Некорректная аналитика или период ДЗ'), 400
        mapping_id = item.get('id')
        if mapping_id is not None and (
                isinstance(mapping_id, bool)
                or not isinstance(mapping_id, int) or mapping_id < 1):
            return jsonify(error='Некорректный идентификатор сопоставления'), 400
        if category and category not in allowed_categories:
            return jsonify(error='Неизвестная категория должников'), 400
        key = (account, analytic_key, effective_from, effective_to)
        if key in normalized and normalized[key][1] != (category or None):
            return jsonify(error='Конфликт категорий для одной аналитики и периода'), 400
        normalized[key] = (label, category or None, mapping_id)
    existing = {
        (row.account, row.analytic_key, row.effective_from, row.effective_to): row
        for row in DebtorAnalyticMapping.query.filter_by(
            company_id=company_id).all()
    }
    for key, (label, category, mapping_id) in normalized.items():
        row = (db.session.get(DebtorAnalyticMapping, mapping_id)
               if mapping_id is not None else existing.get(key))
        if mapping_id is not None and row is None:
            db.session.rollback()
            return jsonify(error='Сопоставление не найдено'), 404
        if row is not None and (row.company_id != company_id
                                or row.account != key[0]
                                or row.analytic_key != key[1]):
            return jsonify(error='Сопоставление не принадлежит аналитике предприятия'), 409
        if category is None:
            if row is not None:
                db.session.delete(row)
            continue
        if row is None:
            overlap = DebtorAnalyticMapping.query.filter(
                DebtorAnalyticMapping.company_id == company_id,
                DebtorAnalyticMapping.account == key[0],
                DebtorAnalyticMapping.analytic_key == key[1],
                DebtorAnalyticMapping.effective_from <= key[3],
                DebtorAnalyticMapping.effective_to >= key[2],
            ).first()
            if overlap is not None:
                db.session.rollback()
                return jsonify(
                    error='Для аналитики уже действует сопоставление; измените существующую запись'), 409
            db.session.add(DebtorAnalyticMapping(
                company_id=company_id, account=key[0], analytic_key=key[1],
                effective_from=key[2], effective_to=key[3],
                analytic_label=label, category=category,
                updated_by=current_user.id))
        else:
            row.analytic_label = label
            row.category = category
            row.updated_by = current_user.id
    db.session.commit()
    _durable_backup(
        'сопоставление аналитик дебиторской задолженности',
        allowed_removed_tables={'debtor_analytic_mappings'})
    saved = DebtorAnalyticMapping.query.filter_by(company_id=company_id).count()
    return jsonify(saved=saved)


@app.route('/company/<int:company_id>/key-indicators/collection-source',
           methods=['POST'])
@admin_required
def key_indicator_collection_source(company_id):
    """Atomically replace linked collection facts and approved norms."""
    company = db.session.get(Company, company_id) or abort(404)
    if company.is_archived:
        abort(409, description='Архивное предприятие доступно только для чтения')
    body = request.get_json(silent=True) or {}
    accruals = body.get('accruals')
    allocations = body.get('allocations')
    norms = body.get('norms')
    if (not isinstance(accruals, list)
            or not isinstance(allocations, list)
            or not isinstance(norms, list)):
        return jsonify(
            error='Поля accruals, allocations и norms должны быть списками'), 400
    allowed_categories = {'population', 'uk', 'budget', 'legal', 'subsidy'}
    normalized_accruals = []
    normalized_allocations = []
    normalized_norms = []
    try:
        for item in accruals:
            if not isinstance(item, dict):
                raise ValueError('Некорректная строка регистра')
            accrual_date = datetime.strptime(
                str(item.get('date') or ''), '%Y-%m-%d').date()
            account = str(item.get('account') or '').strip()
            analytic_key = str(item.get('analytic_key') or '').strip()
            document = str(item.get('accrual_document') or '').strip()
            source_name = str(item.get('source') or '').strip()
            source_version = str(item.get('version') or '').strip()
            accrued = item.get('accrued')
            entry_type = str(item.get('type') or '').strip()
            if (account != '62.01' and not re.fullmatch(r'76(?:\.\d+)+', account)
                    or not re.fullmatch(r'[0-9a-f]{64}', analytic_key)
                    or not document or len(document) > 500
                    or not source_name or len(source_name) > 200
                    or not source_version or len(source_version) > 120
                    or entry_type not in {'accrual', 'correction'}
                    or isinstance(accrued, bool) or not isinstance(accrued, (int, float))
                    or accrued < 0):
                raise ValueError('Некорректная строка начисления')
            normalized_accruals.append(ConsumerCollectionAccrual(
                company_id=company_id, accrual_date=accrual_date,
                account=account, analytic_key=analytic_key,
                entry_type=entry_type, accrued=float(accrued),
                accrual_document=document,
                source_name=source_name,
                source_version=source_version))
        for item in allocations:
            if not isinstance(item, dict):
                raise ValueError('Некорректная строка распределения оплаты')
            payment_date = datetime.strptime(
                str(item.get('date') or ''), '%Y-%m-%d').date()
            payment_document = str(item.get('payment_document') or '').strip()
            accrual_document = str(item.get('accrual_document') or '').strip()
            source_name = str(item.get('source') or '').strip()
            source_version = str(item.get('version') or '').strip()
            amount = item.get('allocated_amount')
            entry_type = str(item.get('type') or '').strip()
            if (not payment_document or not accrual_document
                    or len(payment_document) > 500 or len(accrual_document) > 500
                    or not source_name or len(source_name) > 200
                    or not source_version or len(source_version) > 120
                    or entry_type not in {'payment', 'refund'}
                    or isinstance(amount, bool) or not isinstance(amount, (int, float))
                    or amount < 0):
                raise ValueError('Некорректная строка распределения оплаты')
            normalized_allocations.append(ConsumerPaymentAllocation(
                company_id=company_id, payment_date=payment_date,
                payment_document=payment_document,
                accrual_document=accrual_document,
                entry_type=entry_type, allocated_amount=float(amount),
                source_name=source_name,
                source_version=source_version))
        for item in norms:
            if not isinstance(item, dict):
                raise ValueError('Некорректная строка норматива')
            category = str(item.get('category') or '').strip()
            source_name = str(item.get('source') or '').strip()
            source_version = str(item.get('version') or '').strip()
            valid_from = datetime.strptime(
                str(item.get('valid_from') or ''), '%Y-%m-%d').date()
            valid_to = datetime.strptime(
                str(item.get('valid_to') or ''), '%Y-%m-%d').date()
            approved_at = datetime.fromisoformat(
                str(item.get('approved_at') or '').replace('Z', '+00:00'))
            plan = item.get('plan_percent')
            if (category not in allowed_categories or valid_from > valid_to
                    or item.get('approved') is not True
                    or isinstance(plan, bool) or not isinstance(plan, (int, float))
                    or not source_name or len(source_name) > 200
                    or not source_version or len(source_version) > 120):
                raise ValueError('Некорректный утверждённый норматив')
            normalized_norms.append(ConsumerCollectionNorm(
                company_id=company_id, category=category,
                plan_percent=float(plan), approved=True,
                approved_at=approved_at, valid_from=valid_from,
                valid_to=valid_to, source_name=source_name,
                source_version=source_version))
    except (TypeError, ValueError) as exc:
        return jsonify(error=str(exc)), 400
    accrual_documents = {
        row.accrual_document for row in normalized_accruals}
    if any(row.accrual_document not in accrual_documents
           for row in normalized_allocations):
        return jsonify(error='Распределение ссылается на неизвестное начисление'), 400
    ConsumerCollectionAccrual.query.filter_by(company_id=company_id).delete()
    ConsumerPaymentAllocation.query.filter_by(company_id=company_id).delete()
    ConsumerCollectionNorm.query.filter_by(company_id=company_id).delete()
    db.session.add_all(
        normalized_accruals + normalized_allocations + normalized_norms)
    try:
        db.session.commit()
    except IntegrityError:
        db.session.rollback()
        return jsonify(error='Источник содержит дублирующиеся строки'), 409
    _durable_backup(
        'источник собираемости по категориям',
        allowed_removed_tables={
            'consumer_collection_accruals', 'consumer_payment_allocations',
            'consumer_collection_norms'})
    return jsonify(accruals=len(normalized_accruals),
                   allocations=len(normalized_allocations),
                   norms=len(normalized_norms))
@app.route('/company/<int:company_id>/key-indicators/cost-mappings',
           methods=['POST'])
@admin_required
def key_indicator_cost_mappings(company_id):
    company = db.session.get(Company, company_id) or abort(404)
    if company.is_archived:
        abort(409, description='Архивное предприятие доступно только для чтения')
    items = (request.get_json(silent=True) or {}).get('mappings')
    if not isinstance(items, list):
        return jsonify(error='Поле mappings должно быть списком'), 400
    allowed_categories = {
        'fuel', 'electricity', 'payroll', 'depreciation', 'repair',
        'reagents', 'other'}
    normalized = {}
    for item in items:
        if not isinstance(item, dict):
            return jsonify(error='Некорректная запись сопоставления'), 400
        account = str(item.get('account') or '').strip()
        analytic_key = str(item.get('analytic_key') or '').strip()
        label = str(item.get('analytic_label') or '').strip()
        category = str(item.get('category') or '').strip()
        target = str(item.get('target_account') or '').strip()
        try:
            valid_from = datetime.strptime(
                str(item.get('valid_from') or ''), '%Y-%m-%d').date()
            valid_to = datetime.strptime(
                str(item.get('valid_to') or ''), '%Y-%m-%d').date()
        except ValueError:
            return jsonify(error='Период действия должен иметь формат ГГГГ-ММ-ДД'), 400
        if (account not in ('20', '25', '26', '44')
                or not re.fullmatch(r'[0-9a-f]{64}', analytic_key)
                or not label or len(label) > 1000 or valid_from > valid_to):
            return jsonify(error='Некорректная аналитика затрат'), 400
        if category and category not in allowed_categories:
            return jsonify(error='Неизвестная статья затрат'), 400
        if category and target not in ('90.02', '90.07', '90.08'):
            return jsonify(error='Неизвестный признанный итог'), 400
        key = (account, analytic_key, valid_from, valid_to)
        normalized[key] = (label, category or None, target)
    normalized_items = list(normalized.items())
    for index, (left_key, left_value) in enumerate(normalized_items):
        for right_key, right_value in normalized_items[index + 1:]:
            if left_key[:2] != right_key[:2]:
                continue
            overlaps = left_key[2] <= right_key[3] and left_key[3] >= right_key[2]
            left_choice = left_value[1:]
            right_choice = right_value[1:]
            if overlaps and left_choice != right_choice:
                return jsonify(
                    error='Периоды противоречивых правил не должны пересекаться'), 400
    existing = {
        (row.account, row.analytic_key, row.valid_from, row.valid_to): row
        for row in CostAnalyticMapping.query.filter_by(company_id=company_id).all()
    }
    for key, (label, category, target) in normalized.items():
        row = existing.get(key)
        for old_key, old_row in list(existing.items()):
            same_analytic = old_key[:2] == key[:2]
            overlaps = old_key[2] <= key[3] and old_key[3] >= key[2]
            if same_analytic and overlaps and old_key != key:
                db.session.delete(old_row)
        if category is None:
            if row is not None:
                db.session.delete(row)
            continue
        if row is None:
            db.session.add(CostAnalyticMapping(
                company_id=company_id, account=key[0], analytic_key=key[1],
                analytic_label=label, valid_from=key[2], valid_to=key[3],
                category=category, target_account=target,
                updated_by=current_user.id))
        else:
            row.analytic_label = label
            row.category = category
            row.target_account = target
            row.updated_by = current_user.id
    db.session.commit()
    _durable_backup(
        'сопоставление аналитик затрат',
        allowed_removed_tables={'cost_analytic_mappings'})
    return jsonify(saved=CostAnalyticMapping.query.filter_by(
        company_id=company_id).count())
@app.route('/company/<int:company_id>/key-indicators/annual-data')
@login_required
def annual_key_indicators_data(company_id):
    """Accepted annual-history data for the key-indicator screen."""
    from annual_key_indicators import build_annual_key_indicators

    _require_historical_import()
    company = _require_company_access(company_id)
    try:
        year, year_from, year_to = _annual_key_indicator_request_range()
        payload = build_annual_key_indicators(
            company, requested_year=year,
            year_from=year_from, year_to=year_to)
    except ValueError as exc:
        return jsonify(error=str(exc)), 400
    response = jsonify(payload)
    response.headers['Cache-Control'] = 'private, no-store'
    return response


@app.route('/dashboard/<int:company_id>/<name>')
@login_required
def dashboard(company_id, name):
    name = os.path.basename(name)
    if not name.endswith('.html'):
        abort(404)
    fpath = os.path.join(OUTPUT, str(company_id), name)
    if not os.path.isfile(fpath):
        abort(404)
    with open(fpath, encoding='utf-8') as f:
        saved_html = f.read()
    with open(os.path.join(BASE, 'templates', 'dashboard_template.html'),
              encoding='utf-8') as template_file:
        current_template = template_file.read()
    upgraded = upgrade_saved_dashboard_template(saved_html, current_template)
    html = (upgraded if upgraded is not None else saved_html)
    html = upgrade_legacy_salary_daily_labels(html)
    # Убираем устаревшую кнопку и в старых снимках без полного обновления
    # шаблона. Меняем только ответ, не сохранённые финансовые отчёты.
    html = re.sub(
        r"""<a\b[^>]*\bclass=["']hd-back["'][^>]*>\s*☰\s*Все предприятия\s*</a>\s*""",
        '', html,
    )
    # Вставляем CSRF-токен, company_id и флаг AI прямо в уже собранный HTML
    ai_ok = AiProvider.query.filter_by(is_active=True).first() is not None
    html = html.replace("'__CSRF__'",   f"'{generate_csrf()}'",    1)
    html = html.replace('__CID__',      str(company_id),            1)
    html = html.replace('__AI_OK__',    'true' if ai_ok else 'false', 1)
    html = html.replace('__AI_DISP__',  'inline-flex' if ai_ok else 'none', 1)
    resp = Response(html, mimetype='text/html')
    resp.headers['Cache-Control'] = 'no-store, no-cache, must-revalidate'
    return resp


# ---------------------------------------------------------------------------
# Загрузка суточного ZIP-архива
# ---------------------------------------------------------------------------
def _find_monthly_xls(company_id):
    """Возвращает пути ко всем XLS-файлам из месячных загрузок компании."""
    company_dir = os.path.join(UPLOADS, str(company_id))
    if not os.path.isdir(company_dir):
        return []
    paths = []
    for ts_dir in os.listdir(company_dir):
        if ts_dir.startswith(
                ('zip_', 'daily_', 'tmp_', 'stage_')):  # не базовые папки
            continue
        full = os.path.join(company_dir, ts_dir)
        if not os.path.isdir(full):
            continue
        for fname in os.listdir(full):
            if fname.lower().endswith(('.xls', '.xlsx')):
                paths.append(os.path.join(full, fname))
    return paths


def _rebuild_dashboard_with_daily(company_id, company_name=''):
    """
    Перестраивает дашборд, объединяя:
      - строки аналитики (rows) из месячных XLS-файлов
      - временно́й ряд (daily) из месячных файлов + DailyRow из БД

    Возвращает (dataset, audit_list) или (None, audit_list) при ошибке.
    """
    def _dr_entry(dr):
        """Словарь daily-записи из DailyRow — включает sn/sk для суточных графиков."""
        return dict(
            date=dr.date.isoformat(),
            month=dr.date.strftime('%Y-%m'),
            path1='', path2='', path3='',
            snd=dr.snd, snk=dr.snk,
            obd=dr.obd, obk=dr.obk,
            skd=dr.skd, skk=dr.skk,
            source_type=dr.source_type or 'unknown',
        )

    monthly_paths = _find_monthly_xls(company_id)
    if not monthly_paths:
        # Нет файлов на диске — строим только из суточных записей БД (если есть)
        daily_db = DailyRow.query.filter_by(company_id=company_id)\
                                 .order_by(DailyRow.date).all()
        if not daily_db:
            return None, [{'file': '—', 'acc': '—',
                           'status': 'Нет данных: загрузите ОСВ-файлы (.xls/.xlsx) '
                                     'или суточный ZIP-архив', 'ok': False}]
        data = {}
        for dr in daily_db:
            acc = dr.acc
            if acc not in data:
                data[acc] = dict(acc=acc, period='', dims=[], rows=[], daily=[],
                                 warnings=[], check_ok=None,
                                 top=dict(snd=0, snk=0, obd=0, obk=0, skd=0, skk=0))
            data[acc]['daily'].append(_dr_entry(dr))
        try:
            dataset = build_dataset(data, company_name=company_name)
        except ValueError as e:
            return None, [{'file': '—', 'acc': '—',
                           'status': f'Ошибка сборки дашборда: {e}', 'ok': False}]
        return dataset, []

    data, audit = parse_many(monthly_paths)
    if not data:
        return None, audit

    # Какие месяцы уже покрыты месячными файлами — чтобы не задваивать
    monthly_months: dict[str, set[str]] = {}   # acc -> {month, …}
    for acc, d in data.items():
        monthly_months[acc] = {e['month'] for e in d['daily']}

    # Добавляем суточные записи из БД для периодов вне месячных файлов
    daily_db = DailyRow.query.filter_by(company_id=company_id)\
                             .order_by(DailyRow.date).all()
    for dr in daily_db:
        acc   = dr.acc
        month = dr.date.strftime('%Y-%m')
        if acc in monthly_months and month in monthly_months[acc]:
            continue
        if acc not in data:
            data[acc] = dict(acc=acc, period='', dims=[], rows=[], daily=[],
                             warnings=[], check_ok=None,
                             top=dict(snd=0, snk=0, obd=0, obk=0, skd=0, skk=0))
        data[acc]['daily'].append(_dr_entry(dr))

    try:
        dataset = build_dataset(data, company_name=company_name)
    except ValueError as e:
        return None, audit + [{'file': '—', 'acc': '—',
                                'status': f'Ошибка сборки дашборда: {e}', 'ok': False}]
    return dataset, audit


@app.route('/company/<int:company_id>/upload-zip', methods=['POST'])
@login_required
def company_upload_zip(company_id):
    company = db.session.get(Company, company_id) or abort(404)
    zfile = request.files.get('zipfile')

    def _back(error=None, audit=None):
        uploads = Upload.query.filter_by(company_id=company_id)\
                              .order_by(Upload.uploaded_at.desc()).all()
        return render_template('company.html', company=company, uploads=uploads,
                               upload_history_rows=_upload_history_rows(
                                   uploads, company_id),
                               audit=audit, error=error, discrepancy=[],
                               annual_coverage=_annual_coverage(company_id))

    if not zfile or not zfile.filename:
        return _back(error='Файл не выбран.')
    if not zfile.filename.lower().endswith('.zip'):
        return _back(error='Ожидается ZIP-архив (.zip).')

    # Сохраняем ZIP на диск
    ts = time.strftime('%Y%m%d_%H%M%S')
    zip_dir = os.path.join(UPLOADS, str(company_id), f'zip_{ts}')
    os.makedirs(zip_dir, exist_ok=True)
    zip_path = os.path.join(zip_dir, 'archive.zip')
    zfile.save(zip_path)

    # Парсим ZIP
    accepted_zip_dir = os.path.join(zip_dir, 'accepted')
    successes, errors = parse_zip(zip_path, accepted_dir=accepted_zip_dir)

    if not successes:
        msgs = '; '.join(e['error'] for e in errors[:3])
        return _back(error=f'Не удалось разобрать ни одного файла из архива. {msgs}')

    # Upsert в DailyRow
    upserted = 0
    for item in successes:
        dt    = item['date']
        acc   = item['parsed']['acc']
        top   = item['parsed']['top']
        row = DailyRow.query.filter_by(
            company_id=company_id, date=dt, acc=acc).first()
        if row is None:
            row = DailyRow(company_id=company_id, date=dt, acc=acc)
            db.session.add(row)
        row.snd = top['snd']; row.snk = top['snk']
        row.obd = top['obd']; row.obk = top['obk']
        row.skd = top['skd']; row.skk = top['skk']
        row.source_type = 'direct_daily_osv'
        row.uploaded_at = utcnow()
        upserted += 1
    db.session.commit()
    _durable_backup('суточные строки ZIP ОСВ')

    # Пересобираем дашборд (месячные данные + новые суточные)
    back_url = url_for('company_detail', company_id=company_id)
    dataset, audit = _rebuild_dashboard_with_daily(company_id,
                                                    company_name=company.name)

    audit_out = [{'file': e['file'], 'acc': '?',
                  'status': f'ОШИБКА: {e["error"]}', 'ok': False}
                 for e in errors]

    if dataset is None:
        db.session.rollback()
        return _back(error='Дашборд не построен: ' + (audit[-1]['status'] if audit else ''),
                     audit=audit_out)

    html = render_dashboard(dataset, back_url=back_url)
    out_dir = os.path.join(OUTPUT, str(company_id))
    os.makedirs(out_dir, exist_ok=True)
    out_name = 'dashboard_current.html'
    with open(os.path.join(out_dir, out_name), 'w', encoding='utf-8') as f:
        f.write(html)
    _save_kpi_cache(dataset, out_dir)

    # Диапазон дат из ZIP
    _MRU = {1:'января',2:'февраля',3:'марта',4:'апреля',5:'мая',6:'июня',
            7:'июля',8:'августа',9:'сентября',10:'октября',11:'ноября',12:'декабря'}
    dates = sorted({item['date'] for item in successes})
    d_from, d_to = dates[0], dates[-1]
    if d_from == d_to:
        period_lbl = f'{d_from.day} {_MRU[d_from.month]} {d_from.year}'
    else:
        period_lbl = f'{d_from.strftime("%d.%m")}–{d_to.strftime("%d.%m.%Y")}'

    upload = Upload(
        company_id=company_id,
        uploaded_by=current_user.id,
        period_label=period_lbl,
        accounts_loaded=json.dumps(sorted({item['parsed']['acc'] for item in successes})),
        closure_status=None,
        profit_sign='neutral',
        files_count=upserted,
        dashboard_file=out_name,
        source='manual',
        metrics_json=_dataset_metrics_json(dataset),
        source_manifest_json=dump_manifest(build_manifest(
            _accepted_daily_files([
                item['path'] for item in successes if item.get('path')
            ]), os.path.join(UPLOADS, str(company_id)))),
    )
    db.session.add(upload)
    db.session.commit()
    _schedule_osv_file_summaries(company_id)
    _durable_backup('суточный ZIP ОСВ')

    audit_out += [{'file': f'{item["date"]} сч.{item["parsed"]["acc"]}',
                   'acc': item['parsed']['acc'], 'status': 'добавлен', 'ok': True}
                  for item in successes]

    uploads = Upload.query.filter_by(company_id=company_id)\
                          .order_by(Upload.uploaded_at.desc()).all()
    return render_template('company.html', company=company, uploads=uploads,
                           upload_history_rows=_upload_history_rows(
                               uploads, company_id),
                           audit=audit_out, error=None, discrepancy=[],
                           annual_coverage=_annual_coverage(company_id))


def _annual_coverage(company_id, year=None):
    """Годовая карта покрытия: месячные ОСВ (Upload) + суточные (DailyRow)."""
    import calendar as _cal
    import re as _re
    from datetime import date as _d

    today = _d.today()
    if year is None:
        year = today.year

    # --- Все суточные строки ---
    all_rows = DailyRow.query.filter_by(company_id=company_id).all()
    all_dates = {r.date for r in all_rows}
    data_start = min(all_dates) if all_dates else None

    # --- Все загрузки ---
    all_uploads = Upload.query.filter_by(company_id=company_id)\
                              .order_by(Upload.uploaded_at.asc()).all()

    # Привязка Upload к (year, month) через period_label или uploaded_at
    RU_M = {'янв':1,'фев':2,'мар':3,'апр':4,'май':5,'июн':6,
             'июл':7,'авг':8,'сен':9,'окт':10,'ноя':11,'дек':12}

    def _upl_ym(u):
        """Возвращает список (year, month) — один элемент или диапазон.
        'Январь–Июнь 2026' → [(2026,1),(2026,2),(2026,3),(2026,4),(2026,5),(2026,6)]
        'Июнь 2026'        → [(2026,6)]
        """
        lbl = (u.period_label or '').lower()
        hits = _re.findall(r'(янв|фев|мар|апр|май|июн|июл|авг|сен|окт|ноя|дек)', lbl)
        yr_hit = _re.search(r'\b(20\d{2})\b', lbl)
        yr = int(yr_hit.group(1)) if yr_hit else u.uploaded_at.year
        if hits:
            m_start = RU_M[hits[0]]
            m_end   = RU_M[hits[-1]]
            if m_start <= m_end:
                return [(yr, m) for m in range(m_start, m_end + 1)]
            # Диапазон через год (редко) — берём только конечный месяц
            return [(yr, m_end)]
        # DD.MM–DD.MM → берём последний MM
        m2 = _re.search(r'\d+\.(\d+)[\s–\-]+\d+\.(\d+)', lbl)
        if m2:
            return [(u.uploaded_at.year, int(m2.group(2)))]
        return [(u.uploaded_at.year, u.uploaded_at.month)]

    # (year, month) → лучший Upload (с dashboard_file предпочтительнее)
    upl_map = {}
    month_accs = {}
    for u in all_uploads:
        for key in _upl_ym(u):
            if key not in upl_map or (u.dashboard_file and not upl_map[key].dashboard_file):
                upl_map[key] = u
            month_accs.setdefault(key, set()).update(str(acc) for acc in u.accounts_list if acc)

    # Старые Upload-записи могли сохраняться без accounts_loaded. В этом случае
    # используем уже существующий кэш разбора принятых ОСВ-файлов.
    source_ranges = {}
    source_date_accs = {}
    source_years = set()
    try:
        from key_indicators_service import (
            SourceBusyError, SourceReadError, load_company_source,
        )
        source, _ = load_company_source(company_id, UPLOADS)
        for raw_acc, account_data in source.get('accounts', {}).items():
            acc = str(raw_acc)
            report_start = account_data.get('report_start')
            report_end = account_data.get('report_end')
            try:
                start_date = _d.fromisoformat(report_start) if report_start else None
                end_date = _d.fromisoformat(report_end) if report_end else None
            except (TypeError, ValueError):
                start_date = end_date = None
            if start_date:
                source_years.add(start_date.year)
            if end_date:
                source_years.add(end_date.year)
            if start_date and end_date and start_date <= end_date:
                source_years.update(range(start_date.year, end_date.year + 1))
                source_ranges[acc] = (start_date, end_date)
            for row in account_data.get('daily', []):
                try:
                    row_date = _d.fromisoformat(str(row.get('date')))
                except (TypeError, ValueError):
                    continue
                source_years.add(row_date.year)
                source_date_accs.setdefault(row_date, set()).add(acc)
    except (SourceBusyError, SourceReadError, OSError, ValueError) as exc:
        # Метаданные Upload и DailyRow остаются полноценным источником покрытия,
        # даже если один из архивных XLS-файлов временно недоступен для разбора.
        app.logger.warning(
            'Не удалось восстановить счета покрытия компании %s: %s',
            company_id, exc,
        )
        source_ranges = {}
        source_date_accs = {}
        source_years = set()

    # Год данных определяется периодом отчёта, а не датой его загрузки.
    # Не добавляем текущий год, если данных за него нет.
    data_years = sorted(
        {d.year for d in all_dates} |
        {upload_year for upload_year, _ in upl_map} |
        source_years
    )
    if data_years and year not in data_years:
        year = data_years[-1]
    year_dates = {d for d in all_dates if d.year == year}
    year_rows = [r for r in all_rows if r.date.year == year]
    source_date_accs = {
        day: accounts for day, accounts in source_date_accs.items()
        if day.year == year
    }
    source_accs = {
        acc for acc, (start, end) in source_ranges.items()
        if start <= _d(year, 12, 31) and end >= _d(year, 1, 1)
    }
    for accounts in source_date_accs.values():
        source_accs.update(accounts)

    def _account_sort_key(acc):
        parts = _re.findall(r'\d+|\D+', str(acc))
        return tuple(
            (0, int(part)) if part.isdigit() else (1, part.casefold())
            for part in parts
        )

    year_accs = sorted(
        {r.acc for r in year_rows} |
        {acc for (acc_year, _), accounts in month_accs.items()
         if acc_year == year for acc in accounts} |
        source_accs,
        key=_account_sort_key,
    )
    acc_indexes = {acc: index for index, acc in enumerate(year_accs)}
    account_buttons = [
        dict(
            number=acc,
            name=ANALYTICAL_ACCOUNT_NAMES[acc],
            available=acc in acc_indexes,
            index=acc_indexes.get(acc),
        )
        for acc in sorted(ANALYTICAL_ACCOUNT_NAMES, key=_account_sort_key)
    ]
    date_acc_indexes = {}
    for row in year_rows:
        date_acc_indexes.setdefault(row.date, set()).add(acc_indexes[row.acc])
    for row_date, accounts in source_date_accs.items():
        date_acc_indexes.setdefault(row_date, set()).update(
            acc_indexes[acc] for acc in accounts if acc in acc_indexes
        )

    MON = ['Янв','Фев','Мар','Апр','Май','Июн','Июл','Авг','Сен','Окт','Ноя','Дек']

    months = []
    for m in range(1, 13):
        days_in_month = _cal.monthrange(year, m)[1]
        first_wd = _d(year, m, 1).weekday()  # 0=Пн
        upload = upl_map.get((year, m))

        # Статус месячной ОСВ
        month_start = _d(year, m, 1)
        if month_start > today:
            upl_status = 'future'
        elif upload is None:
            upl_status = 'no_upload'
        elif upload.closure_status is True:
            upl_status = 'closed_profit' if upload.profit_sign == 'profit' else (
                         'closed_loss'   if upload.profit_sign == 'loss'   else 'closed')
        elif upload.closure_status is False:
            upl_status = 'open'
        else:
            upl_status = 'uploaded'

        # Суточная сетка (None = пустая ячейка выравнивания)
        days = [None] * first_wd
        for dn in range(1, days_in_month + 1):
            day = _d(year, m, dn)
            if day > today:
                st = 'future'
            elif day in year_dates:
                st = 'ok'         # есть суточные данные
            elif upload is not None:
                st = 'monthly'    # месяц загружен, но суточной детализации нет
            elif data_start and day >= data_start:
                st = 'missing'    # пропуск в суточном ряду
            else:
                st = 'nodata'     # до начала каких-либо данных
            account_indexes = set(date_acc_indexes.get(day, ()))
            account_indexes.update(
                acc_indexes[acc] for acc in month_accs.get((year, m), ())
                if acc in acc_indexes
            )
            days.append(dict(
                day=dn,
                date=day.strftime('%d.%m'),
                status=st,
                is_today=(day == today),
                account_indexes=sorted(account_indexes),
            ))

        past  = [d for d in days if d and d['status'] != 'future']
        ok    = [d for d in days if d and d['status'] == 'ok']
        miss  = [d for d in days if d and d['status'] == 'missing']

        month_start_d = _d(year, m, 1)
        months_past_1 = 1 if month_start_d <= today else 0
        months_ok_1   = 1 if (upload is not None and month_start_d <= today) else 0

        months.append(dict(
            num=m, name=MON[m - 1],
            days=days,
            upload=upload,
            upl_status=upl_status,
            ok_days=len(ok),
            miss_days=len(miss),
            past_days=len(past),
            coverage_pct=int(len(ok) * 100 / len(past)) if past else 0,
            months_past=months_past_1,
            months_ok=months_ok_1,
        ))

    # Навигация по годам
    prev_year = max((y for y in data_years if y < year), default=None)
    next_year = min((y for y in data_years if y > year), default=None)

    total_ok          = sum(m['ok_days']    for m in months)
    total_past        = sum(m['past_days']  for m in months)
    total_months_ok   = sum(m['months_ok']  for m in months)
    total_months_past = sum(m['months_past'] for m in months)
    return dict(
        year=year,
        months=months,
        prev_year=prev_year,
        next_year=next_year,
        data_years=data_years,
        total_ok=total_ok,
        total_past=total_past,
        coverage_pct=int(total_ok * 100 / total_past) if total_past else 0,
        total_months_ok=total_months_ok,
        total_months_past=total_months_past,
        months_coverage_pct=int(total_months_ok * 100 / total_months_past)
                            if total_months_past else 0,
        has_daily=bool(all_dates),
        accs=year_accs,
        account_buttons=account_buttons,
        has_any=bool(all_dates or all_uploads or source_accs),
    )


@app.route('/company/<int:company_id>/coverage-report.xlsx', methods=['POST'])
@login_required
def coverage_report_excel(company_id):
    """Build a readable Excel matrix for a missing-data request."""
    company = _require_company_access(company_id)
    payload = request.get_json(silent=True) or {}
    raw_months = payload.get('months')
    raw_rows = payload.get('rows')
    if not isinstance(raw_months, list) or not isinstance(raw_rows, list):
        return jsonify(error='Некорректные данные отчёта'), 400

    months = []
    for value in raw_months[:12]:
        try:
            month = int(value)
        except (TypeError, ValueError):
            return jsonify(error='Некорректный месяц'), 400
        if month not in range(1, 13) or month in months:
            return jsonify(error='Некорректный список месяцев'), 400
        months.append(month)
    if not months:
        return jsonify(error='Не выбраны месяцы'), 400
    months.sort()

    normalized_rows = []
    seen_accounts = set()
    for item in raw_rows[:len(ANALYTICAL_ACCOUNT_NAMES)]:
        if not isinstance(item, dict):
            return jsonify(error='Некорректная строка отчёта'), 400
        account = str(item.get('account', ''))
        if account not in ANALYTICAL_ACCOUNT_NAMES or account in seen_accounts:
            return jsonify(error='Некорректный счёт'), 400
        seen_accounts.add(account)
        raw_days = item.get('daysByMonth', {})
        if not isinstance(raw_days, dict):
            return jsonify(error='Некорректные даты отчёта'), 400
        days_by_month = {}
        for month in months:
            values = raw_days.get(str(month), raw_days.get(month, []))
            if not isinstance(values, list):
                return jsonify(error='Некорректные дни отчёта'), 400
            days = []
            for value in values[:31]:
                try:
                    day = int(value)
                except (TypeError, ValueError):
                    return jsonify(error='Некорректный день'), 400
                if day not in range(1, 32):
                    return jsonify(error='Некорректный день'), 400
                if day not in days:
                    days.append(day)
            days_by_month[month] = sorted(days)
        normalized_rows.append((account, days_by_month))

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    month_names = [
        '', 'Январь', 'Февраль', 'Март', 'Апрель', 'Май', 'Июнь',
        'Июль', 'Август', 'Сентябрь', 'Октябрь', 'Ноябрь', 'Декабрь',
    ]
    year = payload.get('year')
    try:
        year = int(year)
    except (TypeError, ValueError):
        year = datetime.now().year

    def format_missing_days(days, month):
        if not days:
            return '—'
        last_day = calendar.monthrange(year, month)[1]
        if days == list(range(1, last_day + 1)):
            return f'с 1 по {last_day}'
        ranges = []
        start = previous = days[0]
        for day in days[1:]:
            if day == previous + 1:
                previous = day
                continue
            ranges.append(str(start) if start == previous else f'{start} - {previous}')
            start = previous = day
        ranges.append(str(start) if start == previous else f'{start} - {previous}')
        return ', '.join(ranges)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Запрос данных'
    last_column = 2 + len(months)
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
    sheet['A1'] = f'Запрос недостающих данных — {company.name}'
    sheet['A1'].font = Font(size=15, bold=True, color='FFFFFF')
    sheet['A1'].fill = PatternFill('solid', fgColor='1F7A9C')
    sheet['A1'].alignment = Alignment(horizontal='center', vertical='center')
    sheet.row_dimensions[1].height = 27
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_column)
    sheet['A2'] = f'Отчётный год: {year}. В ячейках указаны дни, за которые требуется предоставить данные.'
    sheet['A2'].alignment = Alignment(horizontal='left', vertical='center')

    headers = ['Счёт', 'Расшифровка'] + [month_names[month] for month in months]
    thin = Side(style='thin', color='C8D3DA')
    for column, header in enumerate(headers, 1):
        cell = sheet.cell(row=4, column=column, value=header)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='315468')
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for row_number, (account, days_by_month) in enumerate(normalized_rows, 5):
        values = [account, ANALYTICAL_ACCOUNT_NAMES[account]]
        values.extend(
            format_missing_days(days_by_month[month], month)
            for month in months
        )
        for column, value in enumerate(values, 1):
            cell = sheet.cell(row=row_number, column=column, value=value)
            cell.alignment = Alignment(
                horizontal='center' if column != 2 else 'left',
                vertical='top',
                wrap_text=True,
            )
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            if column == 1:
                cell.font = Font(bold=True)
            if row_number % 2 == 0:
                cell.fill = PatternFill('solid', fgColor='F3F7F9')

    sheet.column_dimensions['A'].width = 10
    sheet.column_dimensions['B'].width = 42
    for column in range(3, last_column + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 18
    sheet.freeze_panes = 'C5'
    sheet.auto_filter.ref = f'A4:{get_column_letter(last_column)}{max(4, 4 + len(normalized_rows))}'
    sheet.page_setup.orientation = 'landscape'
    sheet.sheet_view.showGridLines = False

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    safe_company_name = re.sub(r'[<>:"/\\|?*\x00-\x1f]+', '_', company.name)
    safe_company_name = re.sub(r'\s+', '_', safe_company_name).strip(' ._')[:100]
    safe_company_name = re.sub(r'_+', '_', safe_company_name)
    if not safe_company_name:
        safe_company_name = f'Предприятие_{company.id}'
    moscow_now = datetime.now(timezone(timedelta(hours=3)))
    download_filename = (
        f'Запрос_Данных_{safe_company_name}_{moscow_now:%d.%m.%Y}.xlsx'
    )
    return send_file(
        output,
        as_attachment=True,
        download_name=download_filename,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )


# ---------------------------------------------------------------------------
# Ручная пересборка дашборда
# ---------------------------------------------------------------------------
@app.route('/company/<int:company_id>/regen', methods=['POST'])
@login_required
def company_regen(company_id):
    company = db.session.get(Company, company_id) or abort(404)
    dataset, err = _rebuild_dashboard_with_daily(company_id, company.name)
    if dataset is None:
        app.logger.error(f'regen failed for company {company_id}: {err}')
    else:
        html = render_dashboard(dataset, back_url=url_for('company_detail', company_id=company_id))
        out_dir = os.path.join(OUTPUT, str(company_id))
        os.makedirs(out_dir, exist_ok=True)
        with open(os.path.join(out_dir, 'dashboard_current.html'), 'w', encoding='utf-8') as f:
            f.write(html)
        _save_kpi_cache(dataset, out_dir)
        log_event('dashboard_regen', {'company_id': company_id})
        _durable_backup('пересборка дашборда')
    return redirect(url_for('company_detail', company_id=company_id))


# Переименование предприятия
# ---------------------------------------------------------------------------
@app.route('/company/<int:company_id>/rename', methods=['POST'])
@login_required
def company_rename(company_id):
    company = db.session.get(Company, company_id) or abort(404)
    new_name = request.form.get('name', '').strip()
    if new_name:
        old_name = company.name
        company.name = new_name
        db.session.commit()
        log_event('company_renamed', {'id': company_id, 'old': old_name, 'new': new_name})
    return redirect(url_for('company_detail', company_id=company_id))


@app.route('/company/<int:company_id>/delete', methods=['POST'])
@admin_required
def company_delete(company_id):
    company = db.session.get(Company, company_id) or abort(404)
    company.status = 'archived'
    company.updated_at = datetime.now(timezone.utc)
    db.session.commit()
    log_event('company_archived', {'id': company_id, 'name': company.name})
    return redirect(url_for('directory_companies'))


# ---------------------------------------------------------------------------
# Универсальный справочник предприятий
# ---------------------------------------------------------------------------
def _visible_company_query():
    query = Company.query
    if not current_user.is_admin:
        query = query.join(CompanyMembership).filter(
            CompanyMembership.user_id == current_user.id)
    return query


def _dadata_token():
    encrypted = Setting.get('dadata_api_token', '').strip()
    return decrypt_value(encrypted) if encrypted else ''


def _safe_json(value, default=None):
    try:
        return json.loads(value or '')
    except (TypeError, ValueError):
        return {} if default is None else default


DEFAULT_ORGANIZATION_TYPES = ('РСО', 'УК', 'ТСЖ')
ORGANIZATION_TYPES_SETTING = 'directory_organization_types'


def _organization_type_options(extra=()):
    """Возвращает общий справочник типов предприятий без дублей."""
    stored = _safe_json(Setting.get(ORGANIZATION_TYPES_SETTING, '[]'), [])
    values = list(DEFAULT_ORGANIZATION_TYPES)
    if isinstance(stored, list):
        values.extend(stored)
    values.extend(extra or [])
    result = []
    seen = set()
    for value in values:
        value = str(value or '').strip()[:80]
        key = value.casefold()
        if value and key not in seen:
            result.append(value)
            seen.add(key)
    return result


def _snapshot_changes(company, normalized):
    fields = {
        'Наименование': ('full_name', 'full_name'),
        'ИНН': ('inn', 'inn'),
        'КПП': ('kpp', 'kpp'),
        'ОГРН': ('ogrn', 'ogrn'),
        'Юридический адрес': ('legal_address', 'address'),
        'Руководитель': ('director_name', 'director_name'),
        'Должность': ('director_position', 'director_position'),
        'Основной ОКВЭД': ('primary_okved', 'primary_okved'),
    }
    changes = []
    for label, (company_field, incoming_field) in fields.items():
        old = getattr(company, company_field, '') or ''
        new = normalized.get(incoming_field) or ''
        if new and old != new:
            changes.append({'label': label, 'old': old, 'new': new})
    incoming_contacts = {
        (str(item.get('type') or ''), str(item.get('value') or '').casefold())
        for item in normalized.get('contacts') or []
        if isinstance(item, dict) and item.get('value')
    }
    current_contacts = {
        (contact.contact_type or '', (contact.value or '').casefold())
        for contact in company.contacts
        if contact.source == 'DaData' and contact.value
    }
    if incoming_contacts and incoming_contacts != current_contacts:
        changes.append({
            'label': 'Контакты DaData',
            'old': ', '.join(sorted(value for _kind, value in current_contacts)),
            'new': ', '.join(sorted(value for _kind, value in incoming_contacts)),
        })
    return changes


@app.route('/directories/companies')
@login_required
def directory_companies():
    if current_user.is_admin:
        _retry_pending_company_storage_cleanup()
    q = request.args.get('q', '').strip()[:200]
    status_filter = request.args.get('status', 'active').strip()
    fill_filter = request.args.get('fill', '').strip()
    sort = request.args.get('sort', 'name').strip()
    requested_page = max(request.args.get('page', 1, type=int) or 1, 1)
    per_page = min(max(request.args.get('per_page', 25, type=int) or 25, 1), 100)

    if status_filter not in ('active', 'archived', 'all'):
        status_filter = 'active'
    if fill_filter not in ('full', 'partial'):
        fill_filter = ''
    if sort not in ('name', 'updated'):
        sort = 'name'

    visible_query = _visible_company_query()
    total_count = visible_query.order_by(None).count()
    active_count = visible_query.filter(Company.status != 'archived')\
        .order_by(None).count()
    archived_count = visible_query.filter(Company.status == 'archived')\
        .order_by(None).count()

    query = visible_query
    if q:
        escaped_q = q.replace('\\', '\\\\').replace('%', '\\%').replace('_', '\\_')
        pattern = f'%{escaped_q}%'
        query = query.filter(or_(
            Company.name.ilike(pattern, escape='\\'),
            func.coalesce(Company.full_name, '').ilike(pattern, escape='\\'),
            func.coalesce(Company.inn, '').ilike(pattern, escape='\\'),
            func.coalesce(Company.ogrn, '').ilike(pattern, escape='\\'),
            func.coalesce(Company.internal_code, '').ilike(pattern, escape='\\'),
        ))
    if status_filter != 'all':
        query = query.filter(Company.status == status_filter)

    has_name = func.length(func.trim(func.coalesce(Company.name, ''))) > 0
    has_inn = func.length(func.trim(func.coalesce(Company.inn, ''))) > 0
    has_full_name = (
        func.length(func.trim(func.coalesce(Company.full_name, ''))) > 0)
    address_value = func.coalesce(
        func.nullif(func.trim(Company.legal_address), ''),
        func.trim(func.coalesce(Company.address, '')),
        '')
    has_address = func.length(address_value) > 0
    is_full = and_(has_name, has_inn, has_full_name, has_address)
    if fill_filter == 'full':
        query = query.filter(is_full)
    elif fill_filter == 'partial':
        query = query.filter(~is_full)

    filtered_count = query.order_by(None).count()
    total_pages = max(1, (filtered_count + per_page - 1) // per_page)
    page = min(requested_page, total_pages)
    if sort == 'updated':
        query = query.order_by(
            func.coalesce(Company.updated_at, Company.created_at).desc(),
            Company.name.asc(), Company.id.asc())
    else:
        query = query.order_by(
            Company.status.asc(), Company.name.asc(), Company.id.asc())
    companies = query.offset((page - 1) * per_page).limit(per_page).all()

    company_ids = [company.id for company in companies]
    members_by_company = {}
    uploads_by_company = {}
    if company_ids:
        members_by_company = dict(
            db.session.query(
                CompanyMembership.company_id,
                func.count(CompanyMembership.id),
            ).filter(
                CompanyMembership.company_id.in_(company_ids)
            ).group_by(CompanyMembership.company_id).all()
        )
        uploads_by_company = dict(
            db.session.query(
                Upload.company_id,
                func.count(Upload.id),
            ).filter(
                Upload.company_id.in_(company_ids)
            ).group_by(Upload.company_id).all()
        )

    rows = []
    for company in companies:
        required = [company.name, company.inn, company.full_name,
                    company.legal_address or company.address]
        filled = sum(bool(value) for value in required)
        rows.append({
            'company': company,
            'filled': filled,
            'filled_label': f'{filled}/4',
            'members_count': members_by_company.get(company.id, 0),
            'uploads_count': uploads_by_company.get(company.id, 0),
        })
    page_start = (page - 1) * per_page + 1 if filtered_count else 0
    page_end = min(page * per_page, filtered_count)
    page_numbers = list(range(
        max(1, page - 2), min(total_pages, page + 2) + 1))
    return render_template(
        'directory_companies.html', rows=rows,
        active_count=active_count, archived_count=archived_count,
        total_count=total_count,
        filters={
            'q': q, 'status': status_filter, 'fill': fill_filter, 'sort': sort,
        },
        pagination={
            'page': page, 'per_page': per_page, 'total': filtered_count,
            'total_pages': total_pages, 'has_prev': page > 1,
            'has_next': page < total_pages, 'start': page_start,
            'end': page_end,
        },
        page_numbers=page_numbers,
        deleted=request.args.get('deleted') == '1',
        cleanup_pending=request.args.get('cleanup') == 'pending',
    )


@app.route('/directories/companies/new', methods=['GET', 'POST'])
@admin_required
@serialized_company_creation
def directory_company_new():
    preview = None
    error = None
    form = request.form
    if request.method == 'POST':
        action = form.get('action', 'create')
        inn = clean_identifier(form.get('inn'), (10, 12))
        ogrn = clean_identifier(form.get('ogrn'), (13, 15))
        if action == 'lookup':
            if Setting.get('dadata_enabled', '0') != '1':
                error = 'Интеграция DaData выключена в системных настройках.'
            else:
                try:
                    _party, preview = _find_dadata_party(inn=inn, ogrn=ogrn)
                    log_event('dadata_lookup', {'inn': inn, 'ogrn': ogrn, 'found': True})
                except Exception as exc:
                    error = str(exc)
                    log_event('dadata_lookup', {
                        'inn': inn, 'ogrn': ogrn, 'found': False,
                        'error': error[:300],
                    })
        elif action == 'create':
            name = form.get('name', '').strip()[:255]
            full_name = form.get('full_name', '').strip()
            kpp = clean_identifier(form.get('kpp'), (9,))
            is_rso = form.get('is_rso') == '1'
            if not name:
                error = 'Укажите внутреннее отображаемое имя предприятия.'
            elif any(
                    form.get(field, '').strip() and not value
                    for field, value in (('inn', inn), ('ogrn', ogrn), ('kpp', kpp))):
                error = ('Проверьте реквизиты: ИНН должен содержать 10 или 12 цифр, '
                         'ОГРН — 13 или 15, КПП — 9. Некорректные реквизиты '
                         'нельзя использовать для проверки дублей.')
            else:
                try:
                    existing = resolve_existing_company(name, inn, ogrn, kpp)
                    if existing:
                        return redirect(url_for(
                            'directory_company_profile', company_id=existing.id))
                except ValueError as exc:
                    db.session.rollback()
                    error = str(exc)
                if not error:
                    if form.get('confirm_dadata') == '1' and (
                            Setting.get('dadata_enabled', '0') != '1'):
                        error = 'Интеграция DaData выключена в системных настройках.'
                        return render_template(
                            'directory_company_new.html', preview=preview,
                            error=error, form=form, dadata_enabled=False)
                    company = Company(
                        name=name, full_name=full_name, inn=inn, ogrn=ogrn,
                        kpp=kpp, internal_code=form.get('internal_code', '').strip()[:80],
                        address=form.get('address', '').strip(),
                        legal_address=form.get('address', '').strip(),
                        is_rso=is_rso,
                        classifications_json=json.dumps({
                            'organization_types': ['РСО'] if is_rso else [],
                        }, ensure_ascii=False),
                        note=form.get('note', '').strip(),
                        identity_key=company_identity_key(inn, ogrn, kpp),
                        status='active', created_by=current_user.id,
                    )
                    if form.get('confirm_dadata') == '1' and (inn or ogrn):
                        try:
                            party, normalized = _find_dadata_party(
                                inn=inn, ogrn=ogrn)
                            existing = resolve_existing_company(
                                name, normalized.get('inn', ''),
                                normalized.get('ogrn', ''), normalized.get('kpp', ''))
                            if existing:
                                return redirect(url_for(
                                    'directory_company_profile', company_id=existing.id))
                            db.session.add(company)
                            db.session.flush()
                            snapshot = save_snapshot(
                                company, party, normalized, query_inn=inn,
                                query_ogrn=ogrn, requested_by=current_user.id)
                            db.session.flush()
                            accept_dadata(company, snapshot, current_user.id)
                            Setting.set('dadata_last_sync', utcnow().isoformat())
                        except Exception as exc:
                            db.session.rollback()
                            error = str(exc)
                            return render_template(
                                'directory_company_new.html', preview=preview,
                                error=error, form=form,
                                dadata_enabled=Setting.get('dadata_enabled', '0') == '1')
                    else:
                        db.session.add(company)
                    db.session.flush()
                    # ID предприятия сохраняется при восстановлении; новое
                    # предприятие получает код только после выдачи своего ID.
                    if not (company.internal_code or '').strip():
                        company.internal_code = str(company.id)
                    db.session.commit()
                    log_event('directory_company_created', {
                        'company_id': company.id, 'inn': company.inn,
                        'source': 'dadata' if company.dadata_source else 'manual',
                    })
                    return redirect(url_for(
                        'directory_company_profile', company_id=company.id))
    return render_template(
        'directory_company_new.html', preview=preview, error=error, form=form,
        dadata_enabled=Setting.get('dadata_enabled', '0') == '1')


@app.route('/directories/companies/<int:company_id>')
@login_required
def directory_company_profile(company_id):
    company = _require_company_access(company_id)
    classifications = company.classifications
    if not isinstance(classifications, dict):
        classifications = {}
    organization_types = classifications.get('organization_types', [])
    if not isinstance(organization_types, list):
        organization_types = []
    organization_types = [
        str(value).strip()[:80] for value in organization_types
        if str(value or '').strip()
    ]
    if company.is_rso and not any(
            value.casefold() == 'рсо' for value in organization_types):
        organization_types.insert(0, 'РСО')
    organization_type_options = _organization_type_options(organization_types)
    snapshots = CompanyDadataSnapshot.query.filter_by(company_id=company_id)\
        .order_by(CompanyDadataSnapshot.received_at.desc()).limit(20).all()
    latest = snapshots[0] if snapshots else None
    latest_normalized = _safe_json(latest.normalized_json) if latest else {}
    profile_normalized = _safe_json(
        company.dadata_profile.normalized_json) if company.dadata_profile else {}
    extended = _safe_json(
        company.dadata_profile.extended_json) if company.dadata_profile else {}
    extended_data = extended.get('data', {}) if isinstance(extended, dict) else {}
    organization_status = company.legal_status or profile_normalized.get('status')
    status_labels = {
        'ACTIVE': 'Действует',
        'LIQUIDATING': 'В процессе ликвидации',
        'LIQUIDATED': 'Ликвидировано',
        'BANKRUPT': 'Банкротство',
        'REORGANIZING': 'В процессе реорганизации',
    }
    branch_type_labels = {
        'MAIN': 'Основная организация',
        'BRANCH': 'Филиал',
        'REPRESENTATIVE_OFFICE': 'Представительство',
    }
    license_status_labels = {
        'ACTIVE': 'Действует',
        'SUSPENDED': 'Приостановлена',
        'REVOKED': 'Аннулирована',
        'EXPIRED': 'Срок действия истёк',
    }
    manager_type_labels = {
        'EMPLOYEE': 'Физическое лицо',
        'LEGAL': 'Юридическое лицо',
    }
    authority_labels = {
        'fts_registration': 'ФНС — регистрация',
        'fts_report': 'ФНС — отчётность',
        'pf': 'Пенсионный фонд',
        'sif': 'Социальный фонд',
    }
    license_rows = []
    for license_row in company.licenses:
        activities = _safe_json(license_row.activity, [])
        license_status = license_row.status or ''
        license_rows.append({
            'number': license_row.number or '—',
            'status': license_status_labels.get(
                license_status, license_status or '—'),
            'activities': activities if isinstance(activities, list) else [],
            'issue_authority': (
                _safe_json(license_row.details_json, {}).get('issue_authority')
                if license_row.details_json else None
            ),
        })
    return render_template(
        'directory_company_profile.html', company=company, snapshots=snapshots,
        organization_types=organization_types,
        organization_type_options=organization_type_options,
        latest=latest, latest_normalized=latest_normalized,
        profile_normalized=profile_normalized, extended=extended_data,
        license_rows=license_rows,
        organization_status_label=status_labels.get(
            organization_status, organization_status or '—'),
        branch_type_label=branch_type_labels.get(
            extended_data.get('branch_type'),
            extended_data.get('branch_type') or '—'),
        manager_type_labels=manager_type_labels,
        authority_labels=authority_labels,
        changes=_snapshot_changes(company, latest_normalized),
        users=User.query.order_by(User.username).all() if current_user.is_admin else [],
        linked={
            'uploads': Upload.query.filter_by(company_id=company_id).count(),
            'daily_rows': DailyRow.query.filter_by(company_id=company_id).count(),
            'ai_summaries': AiSummary.query.filter_by(company_id=company_id).count(),
            'mail_rules': EmailRule.query.filter_by(company_id=company_id).count(),
        },
        dadata_enabled=Setting.get('dadata_enabled', '0') == '1',
    )


@app.route('/directories/companies/<int:company_id>/edit', methods=['POST'])
@admin_required
def directory_company_edit(company_id):
    company = db.session.get(Company, company_id) or abort(404)
    old = {'name': company.name, 'internal_code': company.internal_code,
           'is_rso': company.is_rso}
    company.name = request.form.get('name', '').strip()[:255] or company.name
    company.internal_code = request.form.get('internal_code', '').strip()[:80]
    if not company.internal_code:
        company.internal_code = str(company.id)
    submitted_types = []
    seen_types = set()
    for value in request.form.getlist('organization_types'):
        value = value.strip()[:80]
        key = value.casefold()
        if value and key not in seen_types:
            submitted_types.append(value)
            seen_types.add(key)
    new_type = request.form.get('new_organization_type', '').strip()[:80]
    if new_type and new_type.casefold() not in seen_types:
        submitted_types.append(new_type)
        seen_types.add(new_type.casefold())
    if not submitted_types and request.form.get('is_rso') == '1':
        submitted_types = ['РСО']
    stored_types = _organization_type_options()
    stored_types_changed = False
    for value in submitted_types:
        if not any(existing.casefold() == value.casefold()
                   for existing in stored_types):
            stored_types.append(value)
            stored_types_changed = True
    if stored_types_changed:
        Setting.set(
            ORGANIZATION_TYPES_SETTING,
            json.dumps(stored_types, ensure_ascii=False),
        )
    company.is_rso = any(value.casefold() == 'рсо'
                         for value in submitted_types)
    company.note = request.form.get('note', '').strip()
    company.director_phone = request.form.get('director_phone', '').strip()[:100]
    company.accountant_name = request.form.get('accountant_name', '').strip()[:255]
    company.accountant_phone = request.form.get('accountant_phone', '').strip()[:100]
    company.classifications_json = json.dumps({
        'organization_types': submitted_types,
    }, ensure_ascii=False)
    company.profile_version = (company.profile_version or 0) + 1
    db.session.commit()
    log_event('directory_company_updated', {
        'company_id': company.id, 'before': old,
        'after': {'name': company.name, 'internal_code': company.internal_code,
                  'is_rso': company.is_rso,
                  'organization_types': submitted_types},
    })
    return redirect(url_for('directory_company_profile', company_id=company.id))


@app.route('/directories/companies/<int:company_id>/archive', methods=['POST'])
@admin_required
def directory_company_archive(company_id):
    company = Company.query.filter_by(id=company_id).with_for_update().first() or abort(404)
    company.status = 'archived' if request.form.get('archive') == '1' else 'active'
    db.session.commit()
    log_event('company_archived' if company.is_archived else 'company_restored',
              {'company_id': company.id})
    return redirect(url_for('directory_company_profile', company_id=company.id))


def _delete_company_records(company):
    """Удаляет предприятие и все ссылающиеся на него записи одной транзакцией."""
    company_id = company.id
    deleted_counts = {}
    # Bulk channel deletion bypasses ORM cascades. Explicit child cleanup also
    # keeps the existing SQLite regression environment consistent with PG FKs.
    channel_ids = db.session.query(OneCChannel.id).filter_by(company_id=company_id)
    for model in (OneCAttempt, OneCIdempotency):
        deleted_counts[model.__tablename__] = model.query.filter(
            model.channel_id.in_(channel_ids)).delete(synchronize_session=False)
    dependent_models = (
        OneCChannel,
        EmailRule,
        OsvFileAiSummary,
        AiSummary,
        DailyRow,
        RevenueAnalyticMapping,
        Upload,
        CompanyMembership,
        CompanyContact,
        CompanyDadataProfile,
        CompanyDadataSnapshot,
        CompanyOkved,
        CompanyLicense,
        CompanyFinance,
        MetricLayout,
    )
    for model in dependent_models:
        deleted_counts[model.__tablename__] = model.query.filter_by(
            company_id=company_id).delete(synchronize_session=False)
    db.session.delete(company)
    return deleted_counts


def _remove_company_storage(company_id):
    """Удаляет только каталоги конкретного предприятия в разрешённых корнях."""
    errors = []
    for root in (UPLOADS, OUTPUT):
        root_path = os.path.realpath(root)
        target = os.path.realpath(os.path.join(root_path, str(int(company_id))))
        if os.path.commonpath((root_path, target)) != root_path:
            errors.append(f'unsafe_path:{os.path.basename(root_path)}')
            continue
        try:
            shutil.rmtree(target)
        except FileNotFoundError:
            continue
        except OSError as exc:
            errors.append(f'{os.path.basename(root_path)}:{type(exc).__name__}')
    return errors


def _retry_pending_company_storage_cleanup(limit=10):
    """Повторяет незавершённую очистку файлов по долговечным записям аудита."""
    pending_actions = (
        'directory_company_delete_pending',
        'directory_company_delete_cleanup_pending',
    )
    try:
        rows = AuditLog.query.filter(
            AuditLog.action.in_(pending_actions)
        ).order_by(AuditLog.id.asc()).limit(limit).all()
        for row in rows:
            details = _safe_json(row.details)
            company_id = details.get('company_id')
            if not isinstance(company_id, int):
                continue
            errors = _remove_company_storage(company_id)
            details['storage_cleanup_errors'] = errors
            details['storage_cleanup_status'] = (
                'pending' if errors else 'complete')
            details['storage_cleanup_attempts'] = int(
                details.get('storage_cleanup_attempts') or 0) + 1
            row.action = (
                'directory_company_delete_cleanup_pending'
                if errors else 'directory_company_deleted')
            row.details = json.dumps(details, ensure_ascii=False, default=str)
        if rows:
            db.session.commit()
    except Exception:
        db.session.rollback()


@app.route('/directories/companies/<int:company_id>/delete', methods=['POST'])
@admin_required
def directory_company_delete(company_id):
    company = Company.query.filter_by(id=company_id).with_for_update().first() or abort(404)
    if not company.is_archived:
        abort(409, description='Перед полным удалением предприятие необходимо архивировать.')
    confirmation_name = request.form.get('confirmation_name', '').strip()
    if confirmation_name != company.name:
        abort(400, description='Для удаления необходимо точно ввести название предприятия.')

    company_name = company.name
    deletion_audit = AuditLog(
        user_id=current_user.id,
        action='directory_company_delete_pending',
        details=json.dumps({
            'company_id': company_id,
            'name': company_name,
            'storage_cleanup_status': 'pending',
            'storage_cleanup_attempts': 0,
        }, ensure_ascii=False),
        ip=request.remote_addr,
    )
    db.session.add(deletion_audit)
    try:
        deleted_counts = _delete_company_records(company)
        pending_details = _safe_json(deletion_audit.details)
        pending_details['deleted_records'] = deleted_counts
        deletion_audit.details = json.dumps(
            pending_details, ensure_ascii=False, default=str)
        db.session.commit()
    except Exception:
        db.session.rollback()
        raise

    storage_errors = _remove_company_storage(company_id)
    audit_details = _safe_json(deletion_audit.details)
    audit_details['storage_cleanup_errors'] = storage_errors
    audit_details['storage_cleanup_status'] = (
        'pending' if storage_errors else 'complete')
    audit_details['storage_cleanup_attempts'] = 1
    deletion_audit.action = (
        'directory_company_delete_cleanup_pending'
        if storage_errors else 'directory_company_deleted')
    deletion_audit.details = json.dumps(
        audit_details, ensure_ascii=False, default=str)
    try:
        db.session.commit()
    except Exception:
        db.session.rollback()
        storage_errors = storage_errors or ['audit_status_update_failed']
    _durable_backup('полное удаление предприятия', allow_removals=True)
    return redirect(url_for(
        'directory_companies', deleted='1',
        cleanup='pending' if storage_errors else None))


def _safe_float_setting(key, default, minimum, maximum):
    try:
        return min(maximum, max(minimum, float(Setting.get(key, str(default)))))
    except (TypeError, ValueError):
        return default


def _safe_int_setting(key, default, minimum, maximum):
    try:
        return min(maximum, max(minimum, int(Setting.get(key, str(default)))))
    except (TypeError, ValueError):
        return default


def _form_number(key, default, minimum, maximum, caster):
    try:
        return min(maximum, max(minimum, caster(request.form.get(key, default))))
    except (TypeError, ValueError):
        return default


def _record_dadata_diagnostic(event):
    """Best-effort append; only allowlisted operational fields are persisted."""
    try:
        table = DadataDiagnosticEvent.__table__
        values = {
            'created_at': utcnow(),
            'attempt_id': str(event.get('attempt_id', ''))[:32],
            'stage': str(event.get('stage', ''))[:32],
            'duration_ms': max(0, int(event.get('duration_ms') or 0)),
            'retry_no': max(0, int(event.get('retry_no') or 0)),
            'http_status': (int(event['http_status'])
                            if event.get('http_status') is not None else None),
            'category': str(event.get('category', ''))[:32],
            'outcome': str(event.get('outcome', ''))[:24],
        }
        # Отдельная транзакция: диагностика не фиксирует и не откатывает
        # незавершённые бизнес-изменения request-scoped сессии.
        with db.engine.begin() as connection:
            connection.execute(table.insert().values(**values))
            keep_ids = connection.execute(
                db.select(table.c.id).order_by(table.c.id.desc()).limit(200)
            ).scalars().all()
            if keep_ids:
                connection.execute(
                    table.delete().where(table.c.id < min(keep_ids)))
    except Exception:
        app.logger.warning('could not persist DaData diagnostic event',
                           exc_info=True)


def _dadata_transport_options():
    return {
        'timeout': _safe_float_setting('dadata_timeout', 12, 1, 120),
        'retries': _safe_int_setting('dadata_retries', 1, 0, 5),
        'backoff': _safe_float_setting('dadata_backoff', 0.35, 0, 30),
        'close_connection': Setting.get(
            'dadata_close_connection', '0') == '1',
        'diagnostic': _record_dadata_diagnostic,
    }


def _find_dadata_party(*, inn='', ogrn=''):
    """Единая граница всех рабочих обращений к DaData."""
    from dadata_client import find_party
    return find_party(
        _dadata_token(), inn=inn, ogrn=ogrn,
        **_dadata_transport_options())


def _dadata_diagnostic_items():
    stage_labels = {
        'sending': 'Отправка',
        'response': 'Ответ получен',
        'http_error': 'Ошибка HTTP',
        'network_error': 'Сетевая ошибка',
        'completed': 'Завершено',
    }
    category_labels = {
        'authorization': 'Авторизация',
        'rate_limit': 'Лимит запросов',
        'provider': 'Сбой DaData',
        'dns': 'DNS',
        'connection': 'Соединение',
        'timeout': 'Тайм-аут',
        'proxy': 'Proxy',
        'tls': 'TLS',
        'invalid_response': 'Некорректный ответ',
        'not_found': 'Не найдено',
        'http': 'HTTP',
    }
    outcome_labels = {
        'success': 'Успешно',
        'retry': 'Будет повтор',
        'error': 'Ошибка',
    }
    rows = DadataDiagnosticEvent.query.order_by(
        DadataDiagnosticEvent.id.desc()).limit(200).all()
    return [{
        'time': row.created_at.isoformat() if row.created_at else '',
        'attempt_id': row.attempt_id,
        'stage': stage_labels.get(row.stage, row.stage),
        'duration_ms': row.duration_ms,
        'retry_no': row.retry_no,
        'http_status': row.http_status,
        'category': category_labels.get(row.category, row.category or ''),
        'outcome': outcome_labels.get(row.outcome, row.outcome or ''),
    } for row in rows]


@app.route('/directories/companies/<int:company_id>/dadata/refresh', methods=['POST'])
@admin_required
def directory_company_dadata_refresh(company_id):
    company = db.session.get(Company, company_id) or abort(404)
    if Setting.get('dadata_enabled', '0') != '1':
        abort(409, description='Интеграция DaData выключена')
    inn = clean_identifier(request.form.get('inn') or company.inn, (10, 12))
    kpp = clean_identifier(request.form.get('kpp') or company.kpp, (9,))
    ogrn = clean_identifier(request.form.get('ogrn') or company.ogrn, (13, 15))
    try:
        party, normalized = _find_dadata_party(inn=inn, ogrn=ogrn)
        if kpp and not normalized.get('kpp'):
            normalized['kpp'] = kpp
        duplicates = duplicate_companies(
            normalized.get('inn', ''), normalized.get('ogrn', ''),
            normalized.get('kpp', ''), exclude_id=company.id)
        if duplicates:
            raise ValueError('Обновление конфликтует с записью ' + ', '.join(
                f'#{item.id} {item.name}' for item in duplicates))
        snapshot = save_snapshot(
            company, party, normalized, query_inn=inn, query_ogrn=ogrn,
            requested_by=current_user.id)
        company.dadata_sync_status = 'pending'
        company.dadata_last_error = ''
        Setting.set('dadata_last_sync', utcnow().isoformat())
        db.session.commit()
        log_event('dadata_refresh_received', {
            'company_id': company.id, 'snapshot_id': snapshot.id,
            'changes': len(_snapshot_changes(company, normalized)),
        })
    except Exception as exc:
        company.dadata_sync_status = 'error'
        company.dadata_last_error = str(exc)[:1000]
        save_snapshot(
            company, {}, {}, query_inn=inn, query_ogrn=ogrn,
            requested_by=current_user.id, status='error',
            error_text=str(exc))
        db.session.commit()
        log_event('dadata_refresh_failed', {
            'company_id': company.id, 'error': str(exc)[:300]})
    return redirect(url_for('directory_company_profile', company_id=company.id))


@app.route('/directories/companies/<int:company_id>/dadata/accept', methods=['POST'])
@admin_required
def directory_company_dadata_accept(company_id):
    company = db.session.get(Company, company_id) or abort(404)
    snapshot = db.session.get(
        CompanyDadataSnapshot, request.form.get('snapshot_id', type=int)) or abort(404)
    if snapshot.company_id != company.id or snapshot.status != 'success':
        abort(400)
    normalized = _safe_json(snapshot.normalized_json)
    duplicates = duplicate_companies(
        normalized.get('inn', ''), normalized.get('ogrn', ''),
        normalized.get('kpp', ''), exclude_id=company.id)
    if duplicates:
        abort(409, description='Нельзя принять реквизиты: найден дубликат')
    changes = _snapshot_changes(company, normalized)
    accept_dadata(company, snapshot, current_user.id)
    db.session.commit()
    log_event('dadata_changes_accepted', {
        'company_id': company.id, 'snapshot_id': snapshot.id,
        'fields': [item['label'] for item in changes],
    })
    return redirect(url_for('directory_company_profile', company_id=company.id))


@app.route('/directories/companies/<int:company_id>/contacts', methods=['POST'])
@admin_required
def directory_company_contact_add(company_id):
    company = db.session.get(Company, company_id) or abort(404)
    value = request.form.get('value', '').strip()
    if value:
        contact = CompanyContact(
            company_id=company.id,
            contact_type=request.form.get('contact_type', 'other')[:40],
            label=request.form.get('label', '').strip()[:120],
            name=request.form.get('contact_name', '').strip()[:255],
            value=value[:500], is_internal=True, source='manual')
        db.session.add(contact)
        db.session.commit()
        log_event('company_contact_added', {
            'company_id': company.id, 'contact_type': contact.contact_type})
    return redirect(url_for('directory_company_profile', company_id=company.id))


@app.route('/api/directories/companies')
@login_required
def directory_companies_api():
    companies = _visible_company_query().order_by(Company.name).all()
    return jsonify(items=[{
        'company_id': company.id, 'name': company.name,
        'full_name': company.full_name or '', 'inn': company.inn or '',
        'kpp': company.kpp or '', 'ogrn': company.ogrn or '',
        'status': company.status or 'active',
        'profile_url': url_for('directory_company_profile', company_id=company.id),
    } for company in companies])


@app.route('/api/companies/<int:company_id>/profile')
@login_required
def directory_company_profile_api(company_id):
    company = _require_company_access(company_id)
    classifications = company.classifications
    if not isinstance(classifications, dict):
        classifications = {}
    organization_types = classifications.get('organization_types', [])
    if not isinstance(organization_types, list):
        organization_types = []
    if company.is_rso and not any(
            str(value).casefold() == 'рсо' for value in organization_types):
        organization_types.insert(0, 'РСО')
    return jsonify({
        'company_id': company.id,
        'internal': {
            'name': company.name, 'internal_code': company.internal_code or '',
            'is_rso': bool(company.is_rso),
            'organization_types': organization_types,
            'note': company.note or '',
            'status': company.status or 'active',
        },
        'official': {
            'full_name': company.full_name or '', 'inn': company.inn or '',
            'kpp': company.kpp or '', 'ogrn': company.ogrn or '',
            'address': company.legal_address or company.address or '',
            'director_name': company.director_name or '',
            'director_position': company.director_position or '',
            'primary_okved': company.primary_okved or '',
            'source': company.dadata_source or 'manual',
            'received_at': company.dadata_received_at.isoformat()
                if company.dadata_received_at else None,
            'actuality_date': company.dadata_actuality_date.isoformat()
                if company.dadata_actuality_date else None,
        },
        'links': {
            'analytics': url_for('company_detail', company_id=company.id),
        },
    })


# ---------------------------------------------------------------------------
# Настройки приложения
# ---------------------------------------------------------------------------
IMAP_KEYS = [
    ('imap_host',     'Хост IMAP',              'imap.yandex.ru'),
    ('imap_port',     'Порт',                   '993'),
    ('imap_login',    'Логин (e-mail)',          ''),
    ('imap_password', 'Пароль / пароль приложения', ''),
    ('imap_folder',   'Папка',                  'INBOX'),
    ('imap_interval', 'Интервал опроса (мин)',   '30'),
]


@app.route('/settings/test-imap', methods=['POST'])
@admin_required
def settings_test_imap():
    import imaplib
    import ssl
    from flask import jsonify

    host     = request.form.get('imap_host', '').strip()
    port     = request.form.get('imap_port', '993').strip()
    login    = request.form.get('imap_login', '').strip()
    password = request.form.get('imap_password', '').strip()
    if not password:
        saved_password = Setting.get('imap_password', '')
        password = decrypt_value(saved_password) if saved_password else ''
    folder   = request.form.get('imap_folder', 'INBOX').strip() or 'INBOX'

    if not host or not login or not password:
        return jsonify(ok=False, msg='Заполните хост, логин и пароль')

    try:
        port_int = int(port)
    except ValueError:
        return jsonify(ok=False, msg=f'Неверный порт: {port}')

    try:
        ctx = ssl.create_default_context()
        with imaplib.IMAP4_SSL(host, port_int, ssl_context=ctx,
                               timeout=15) as M:
            M.login(login, password)
            status, data = M.select(folder, readonly=True)
            if status != 'OK':
                return jsonify(ok=False, msg=f'Папка «{folder}» не найдена')
            count = int(data[0]) if data and data[0] else 0
            M.logout()
        return jsonify(ok=True, msg=f'Подключение успешно. В папке «{folder}»: {count} писем')
    except imaplib.IMAP4.error as e:
        return jsonify(ok=False, msg=f'Ошибка аутентификации: {e}')
    except OSError as e:
        return jsonify(ok=False, msg=f'Сетевая ошибка: {e}')
    except Exception as e:
        return jsonify(ok=False, msg=f'Ошибка: {e}')


@app.route('/settings/sync-imap', methods=['POST'])
@admin_required
def settings_sync_imap():
    """Ручная синхронизация — тонкая обёртка над _do_imap_sync()."""
    result = _do_imap_sync(current_user.id)
    return jsonify(**result)


@app.route('/settings/test-dadata', methods=['POST'])
@admin_required
def settings_test_dadata():
    """Проверяет введённый или сохранённый токен, никогда не возвращая его."""
    from dadata_client import check_token
    token = request.form.get('dadata_api_token', '').strip() or _dadata_token()
    ok, msg = check_token(token, **_dadata_transport_options())
    now = utcnow().isoformat()
    Setting.set('dadata_last_check', now)
    Setting.set('dadata_last_error', '' if ok else msg[:1000])
    db.session.commit()
    log_event('dadata_connection_checked', {'ok': ok, 'error': '' if ok else msg[:300]})
    return jsonify(ok=ok, msg=msg, checked_at=now)


@app.route('/settings/dadata-diagnostics', methods=['GET', 'DELETE'])
@admin_required
def settings_dadata_diagnostics():
    if request.method == 'DELETE':
        DadataDiagnosticEvent.query.delete()
        db.session.commit()
        return jsonify(ok=True, events=[])
    return jsonify(events=_dadata_diagnostic_items())


@app.route('/ai/brief', methods=['POST'])
@login_required
def ai_brief():
    """Генерирует блок «Что происходит» через ИИ по данным из дашборда."""
    import json as _json
    import re as _re
    from ai_client import call_ai_provider as _call_ai

    req = request.get_json(force=True, silent=True) or {}

    provider = (AiProvider.query.filter_by(is_default=True, is_active=True).first()
                or AiProvider.query.filter_by(is_active=True).first())
    if not provider:
        return jsonify(ok=False, error='ИИ-провайдер не настроен'), 400

    api_key = decrypt_value(provider.api_key) if provider.api_key else ''

    def _mln(v):
        sign = '+' if v >= 0 else '−'
        return f'{sign}{abs(v)/1e6:.1f} млн ₽'.replace('.', ',')

    company_name = req.get('company_name', 'Предприятие')
    period       = req.get('period', '')
    months_list  = req.get('months', [])
    closing      = req.get('closing', [])
    pl           = req.get('pl', {})
    dz_tot       = req.get('dz_tot', {})
    kz_tot       = req.get('kz_tot', {})
    cash         = req.get('cash', {})
    kz_rows      = req.get('kz', [])
    hazard       = req.get('hazard', '')

    lines = [f'Организация: {company_name}']
    if period:
        lines.append(f'Отчётный период: {period}')
    if months_list:
        closed_m     = [m for m, s in zip(months_list, closing) if 'закрыт' in s.lower() and 'не закрыт' not in s.lower()]
        not_closed_m = [m for m, s in zip(months_list, closing) if 'не закрыт' in s.lower()]
        if closed_m:
            lines.append(f'Закрытые месяцы: {", ".join(closed_m)}')
        if not_closed_m:
            lines.append(f'⚠ Незакрытые (данные неполные): {", ".join(not_closed_m)}')

    if pl:
        rev   = pl.get('rev_total', 0)
        nds   = pl.get('nds_total', 0)
        cost  = pl.get('cost_total', 0)
        adm   = pl.get('adm_total', 0)
        oinc  = pl.get('oinc_total', 0)
        oexp  = pl.get('oexp_total', 0)
        revn  = rev - nds
        spend = cost + adm
        sales = revn - spend
        net   = sales + oinc - oexp
        lines.append('\nФинансовые результаты:')
        lines.append(f'  Выручка без НДС: {revn/1e6:.1f} млн ₽'.replace('.', ','))
        lines.append(f'  Расходы (производство + управление): {spend/1e6:.1f} млн ₽'.replace('.', ','))
        lines.append(f'  Результат от основной деятельности: {_mln(sales)}')
        if oinc or oexp:
            lines.append(f'  Прочие доходы: {oinc/1e6:.1f} млн ₽, прочие расходы: {oexp/1e6:.1f} млн ₽'.replace('.', ','))
            lines.append(f'  Итоговый результат: {_mln(net)}')
    else:
        sales = 0; net = 0

    if dz_tot:
        sn, sk = dz_tot.get('sn', 0), dz_tot.get('sk', 0)
        lines.append(f'\nДебиторская задолженность: начало {sn/1e6:.1f} млн ₽ → конец {sk/1e6:.1f} млн ₽ (изм. {_mln(sk-sn)})'.replace('.', ','))
    if kz_tot:
        sn, sk = kz_tot.get('sn', 0), kz_tot.get('sk', 0)
        lines.append(f'Кредиторская задолженность: начало {sn/1e6:.1f} млн ₽ → конец {sk/1e6:.1f} млн ₽ (изм. {_mln(sk-sn)})'.replace('.', ','))
        if kz_rows:
            top = kz_rows[0]
            lines.append(f'Крупнейший кредитор: {top[0]} — {top[2]/1e6:.1f} млн ₽'.replace('.', ','))
    if cash:
        cash_end  = cash.get('end', 0)
        cash_days = cash.get('days')
        cash_line = f'Деньги на счетах и в кассе: {cash_end/1e6:.1f} млн ₽'
        if isinstance(cash_days, (int, float)) and not isinstance(cash_days, bool):
            cash_line += f' (≈{max(cash_days,0):.0f} раб. дней платежей)'
        lines.append(cash_line.replace('.', ','))
    if hazard:
        h_clean = _re.sub(r'<[^>]+>', '', hazard)
        lines.append(f'\n⚠ Качество данных: {h_clean}')

    is_loss = (sales < 0) if pl else False

    system_msg = (
        'Ты финансовый аналитик — специалист по ОСВ организаций ЖКХ России. '
        'По приведённым финансовым данным сформируй блок «Что происходит». '
        'Верни ТОЛЬКО валидный JSON без пояснений:\n'
        '{\n'
        '  "lead": "Одно короткое предложение — главный вывод",\n'
        '  "points": [\n'
        '    "Тема. Краткий факт с цифрами",\n'
        '    "..."\n'
        '  ],\n'
        '  "foot": "Приоритеты: пункт 1; пункт 2.",\n'
        '  "is_loss": true\n'
        '}\n'
        'Правила: ровно 4–6 коротких пунктов; каждый начинается с короткого '
        'содержательного заголовка и точки, например «Ликвидность. Денежные '
        'средства...». Не пиши буквальное слово «Тема»; '
        'общий объём текста во всех полях не более 1800 знаков; не используй '
        'Markdown или HTML; is_loss=true, если основная деятельность убыточна; '
        'только русский язык. Обязательно закрой все строки, массив и объект JSON.'
    )

    try:
        result = _call_ai(
            base_url=provider.base_url,
            api_key=api_key,
            model=provider.model,
            messages=[
                {'role': 'system', 'content': system_msg},
                {'role': 'user',   'content': '\n'.join(lines)},
            ],
            temperature=0.25,
            max_tokens=2400,
            enable_thinking=False,
            timeout=60,
        )
    except Exception as e:
        app.logger.error('AI brief error: %s', e)
        return jsonify(ok=False, error=str(e)), 502

    content = result['content']

    # Извлекаем JSON — убираем markdown-фенсы, ищем первый {...}
    json_str = content
    m = _re.search(r'```(?:json)?\s*([\s\S]*?)```', content, _re.IGNORECASE)
    if m:
        json_str = m.group(1).strip()
    else:
        m2 = _re.search(r'\{[\s\S]*\}', content)
        if m2:
            json_str = m2.group(0)

    try:
        raw_brief = _json.loads(json_str)
        if not isinstance(raw_brief, dict):
            raise ValueError('ответ не является JSON-объектом')
        lead = raw_brief.get('lead', '')
        points = raw_brief.get('points', [])
        foot = raw_brief.get('foot', '')
        if not isinstance(lead, str) or not isinstance(foot, str):
            raise ValueError('lead и foot должны быть строками')
        if not isinstance(points, list) or not all(
                isinstance(point, str) for point in points):
            raise ValueError('points должен быть массивом строк')
        if not points:
            raise ValueError('массив points пуст')

        # Модель не управляет HTML в браузере: оставляем только чистый текст.
        def _plain(value, limit):
            value = _re.sub(r'<[^>]*>', '', value or '')
            return ' '.join(value.split())[:limit]

        brief_data = {
            'lead': _plain(lead, 500),
            'points': [_plain(point, 700) for point in points[:6] if point.strip()],
            'foot': _plain(foot, 500),
            'is_loss': bool(raw_brief.get('is_loss', is_loss)),
        }
        if not brief_data['points']:
            raise ValueError('после нормализации массив points пуст')
    except Exception as exc:
        app.logger.warning(
            'AI brief returned invalid JSON: %s; response prefix=%r',
            exc, content[:200])
        return jsonify(
            ok=False,
            error='ИИ вернул неполный ответ. Нажмите «Обновить» ещё раз.',
        ), 502

    return jsonify(ok=True, brief=brief_data)


@app.route('/settings/test-ai', methods=['POST'])
@admin_required
def settings_test_ai():
    """Проверка подключения к AI-провайдеру (AJAX)."""
    from ai_client import call_ai_provider as _call_ai
    base_url  = request.form.get('base_url', '').strip()
    api_key   = request.form.get('api_key', '').strip()
    model     = request.form.get('model', '').strip()

    if not base_url or not model:
        return jsonify(ok=False, msg='Укажите Base URL и модель')
    try:
        result = _call_ai(
            base_url=base_url,
            api_key=api_key,
            model=model,
            messages=[{'role': 'user', 'content': 'Ответь одним словом: OK'}],
            max_tokens=16,
            timeout=20,
        )
        answer = result['content'][:120] or '(пустой ответ)'
        return jsonify(ok=True, msg=f'Подключение успешно. Ответ модели: {answer}')
    except Exception as e:
        return jsonify(ok=False, msg=str(e))


@app.route('/settings', methods=['GET', 'POST'])
@admin_required
def settings():
    companies = Company.query.filter(Company.status != 'archived')\
        .order_by(Company.name).all()

    if request.method == 'POST':
        action = request.form.get('action', '')

        if action == 'save_imap':
            for key, _, _ in IMAP_KEYS:
                val = request.form.get(key, '').strip()
                if key == 'imap_password':
                    if not val:
                        continue
                    val = encrypt_value(val)
                Setting.set(key, val)
            db.session.commit()
            log_event('imap_settings_saved')

        elif action == 'add_rule':
            company_id      = request.form.get('company_id', type=int)
            sender_pattern  = request.form.get('sender_pattern', '').strip()
            subject_pattern = request.form.get('subject_pattern', '').strip()
            note            = request.form.get('note', '').strip()
            if company_id and (sender_pattern or subject_pattern):
                db.session.add(EmailRule(
                    company_id=company_id,
                    sender_pattern=sender_pattern,
                    subject_pattern=subject_pattern,
                    note=note,
                ))
                db.session.commit()

        elif action == 'delete_rule':
            rule_id = request.form.get('rule_id', type=int)
            rule = db.session.get(EmailRule, rule_id)
            if rule:
                db.session.delete(rule)
                db.session.commit()

        elif action == 'add_ai_provider':
            ai_name    = request.form.get('ai_name', '').strip()
            ai_url     = request.form.get('ai_base_url', '').strip()
            ai_key_raw = request.form.get('ai_api_key', '').strip()
            ai_model   = request.form.get('ai_model', '').strip()
            ai_active  = request.form.get('ai_is_active') == '1'
            ai_default = request.form.get('ai_is_default') == '1'
            ai_local   = request.form.get('ai_is_local') == '1'
            ai_vision  = request.form.get('ai_supports_vision') == '1'
            if ai_name and ai_url and ai_model:
                if ai_default:
                    AiProvider.query.update({'is_default': False})
                prov = AiProvider(
                    name=ai_name,
                    base_url=ai_url,
                    api_key=encrypt_value(ai_key_raw) if ai_key_raw else '',
                    model=ai_model,
                    is_active=ai_active,
                    is_default=ai_default,
                    is_local=ai_local,
                    supports_vision=ai_vision,
                )
                db.session.add(prov)
                db.session.commit()
                log_event('ai_provider_added', {'name': ai_name})

        elif action == 'update_ai_provider':
            provider_id  = request.form.get('provider_id', type=int)
            prov = db.session.get(AiProvider, provider_id)
            if prov:
                ai_name    = request.form.get('ai_name', '').strip()
                ai_url     = request.form.get('ai_base_url', '').strip()
                ai_key_raw = request.form.get('ai_api_key', '').strip()
                ai_model   = request.form.get('ai_model', '').strip()
                ai_active  = request.form.get('ai_is_active') == '1'
                ai_default = request.form.get('ai_is_default') == '1'
                if ai_name:
                    prov.name = ai_name
                if ai_url:
                    prov.base_url = ai_url
                if ai_model:
                    prov.model = ai_model
                if ai_key_raw:                          # пустое поле → ключ не перезаписывается
                    prov.api_key = encrypt_value(ai_key_raw)
                prov.is_active = ai_active
                prov.is_local = request.form.get('ai_is_local') == '1'
                prov.supports_vision = request.form.get('ai_supports_vision') == '1'
                if ai_default and not prov.is_default:
                    AiProvider.query.update({'is_default': False})
                prov.is_default = ai_default
                db.session.commit()
                log_event('ai_provider_updated', {'name': prov.name})

        elif action == 'delete_ai_provider':
            provider_id = request.form.get('provider_id', type=int)
            prov = db.session.get(AiProvider, provider_id)
            if prov:
                pname = prov.name
                db.session.delete(prov)
                db.session.commit()
                log_event('ai_provider_deleted', {'name': pname})

        elif action == 'set_default_ai_provider':
            provider_id = request.form.get('provider_id', type=int)
            AiProvider.query.update({'is_default': False})
            prov = db.session.get(AiProvider, provider_id)
            if prov:
                prov.is_default = True
                db.session.commit()
                log_event('ai_provider_set_default', {'name': prov.name})

        elif action == 'save_landing':
            label = request.form.get('landing_v09_label', '').strip()
            url_  = request.form.get('landing_v09_url', '').strip()
            Setting.set('landing_v09_label', label or 'Версия 0.9 alfa')
            Setting.set('landing_v09_url', url_)
            db.session.commit()
            log_event('landing_settings_saved')

        elif action == 'save_dadata':
            token = request.form.get('dadata_api_token', '').strip()
            Setting.set(
                'dadata_enabled',
                '1' if request.form.get('dadata_enabled') == '1' else '0')
            if token:
                Setting.set('dadata_api_token', encrypt_value(token))
            Setting.set(
                'dadata_close_connection',
                '1' if request.form.get('dadata_close_connection') == '1' else '0')
            Setting.set('dadata_timeout', str(_form_number(
                'dadata_timeout', 12, 1, 120, float)))
            Setting.set('dadata_retries', str(_form_number(
                'dadata_retries', 1, 0, 5, int)))
            Setting.set('dadata_backoff', str(_form_number(
                'dadata_backoff', 0.35, 0, 30, float)))
            db.session.commit()
            log_event('dadata_settings_saved', {
                'enabled': Setting.get('dadata_enabled', '0') == '1',
                'token_changed': bool(token),
            })

        action_locations = {
            'save_landing': ('general', 'landing'),
            'save_imap': ('mail', 'imap'),
            'add_rule': ('mail', 'rules'),
            'delete_rule': ('mail', 'rules'),
            'add_ai_provider': ('integrations', 'ai'),
            'update_ai_provider': ('integrations', 'ai'),
            'delete_ai_provider': ('integrations', 'ai'),
            'set_default_ai_provider': ('integrations', 'ai'),
            'save_dadata': ('integrations', 'dadata'),
        }
        tab, anchor = action_locations.get(action, ('general', 'landing'))
        target = url_for(
            'settings',
            tab=tab,
            embedded=request.args.get('embedded') or None,
        )
        return redirect(f'{target}#{anchor}')

    imap_cfg = {key: Setting.get(key, default) for key, _, default in IMAP_KEYS}
    imap_password_configured = bool(imap_cfg.get('imap_password'))
    # Секрет не возвращается в HTML; пустое поле означает «не менять».
    imap_cfg['imap_password'] = ''
    rules = EmailRule.query.order_by(EmailRule.id).all()
    ai_providers = AiProvider.query.order_by(AiProvider.created_at).all()
    landing_cfg = {
        'v09_label': Setting.get('landing_v09_label', 'Версия 0.9 alfa'),
        'v09_url':   Setting.get('landing_v09_url', ''),
    }
    dadata_cfg = {
        'enabled': Setting.get('dadata_enabled', '0') == '1',
        'token_configured': bool(Setting.get('dadata_api_token', '')),
        'last_check': Setting.get('dadata_last_check', ''),
        'last_sync': Setting.get('dadata_last_sync', ''),
        'last_error': Setting.get('dadata_last_error', ''),
        'close_connection': Setting.get('dadata_close_connection', '0') == '1',
        'timeout': Setting.get('dadata_timeout', '12'),
        'retries': Setting.get('dadata_retries', '1'),
        'backoff': Setting.get('dadata_backoff', '0.35'),
    }
    dadata_diagnostics = _dadata_diagnostic_items()
    watchdog_query = AuditLog.query.filter_by(
        action='http_worker_watchdog_recovery')
    watchdog_recent_events = watchdog_query.order_by(
        AuditLog.created_at.desc()).limit(10).all()
    watchdog_event_items = []
    for event in watchdog_recent_events:
        display_at = event.created_at
        try:
            event_details = json.loads(event.details or '{}')
            display_at = datetime.fromisoformat(event_details['occurred_at'])
        except (KeyError, TypeError, ValueError, json.JSONDecodeError):
            pass
        watchdog_event_items.append({
            'display_at': display_at,
            'action': event.action,
        })
    watchdog_recent_count = watchdog_query.filter(
        AuditLog.created_at >= datetime.now(timezone.utc) - timedelta(days=1),
    ).count()
    watchdog_total = watchdog_query.count()
    watchdog_status = {
        'total': watchdog_total,
        'recent_count': watchdog_recent_count,
        'is_frequent': watchdog_recent_count >= 3,
        'last_at': (watchdog_event_items[0]['display_at']
                    if watchdog_event_items else None),
        'events': watchdog_event_items,
    }
    return render_template('settings.html',
                           imap_keys=IMAP_KEYS,
                           imap_password_configured=imap_password_configured,
                           imap_cfg=imap_cfg,
                           rules=rules,
                           companies=companies,
                           ai_providers=ai_providers,
                           landing_cfg=landing_cfg,
                           dadata_cfg=dadata_cfg,
                           dadata_diagnostics=dadata_diagnostics,
                           watchdog_status=watchdog_status)


@app.route('/healthz/watchdog-recovery', methods=['POST'])
@csrf.exempt
def healthz_watchdog_recovery():
    """Accept a recovery event only from the local Gunicorn master."""
    if (request.remote_addr not in ('127.0.0.1', '::1')
            or request.headers.get('X-Watchdog-Event') != '1'):
        abort(404)
    payload = request.get_json(silent=True) or {}
    event_id = str(payload.get('event_id', ''))[:80]
    occurred_at = str(payload.get('occurred_at', ''))[:80]
    worker_pid = payload.get('worker_pid')
    if not event_id or not occurred_at or not isinstance(worker_pid, int):
        abort(400)
    duplicate = AuditLog.query.filter_by(
        action='http_worker_watchdog_recovery',
    ).filter(AuditLog.details.contains(f'"event_id": "{event_id}"')).first()
    if duplicate is None:
        persisted = log_event('http_worker_watchdog_recovery', {
            'event_id': event_id,
            'occurred_at': occurred_at,
            'worker_pid': worker_pid,
        }, ip='127.0.0.1')
        if not persisted:
            return '', 503
    return '', 204


# ---------------------------------------------------------------------------
# Управление пользователями (только для администраторов)
# ---------------------------------------------------------------------------
@app.route('/users')
@admin_required
def users_list():
    users = User.query.order_by(User.created_at.desc()).all()
    return render_template('users.html', users=users,
                           available_modules=registered_modules(),
                           companies=Company.query.order_by(Company.name).all(),
                           create_form=request.form,
                           created_username=request.args.get('created', '').strip())


@app.route('/users/create', methods=['POST'])
@admin_required
def users_create():
    username = request.form.get('username', '').strip()
    # Never strip passwords: spaces, including leading/trailing spaces, are
    # intentional password code points.
    password = request.form.get('password', '')
    password_repeat = request.form.get('password_repeat')
    if password_repeat is None:
        password_repeat = request.form.get(
            'password_confirmation',
            request.form.get('password_confirm', ''),
        )
    role = request.form.get('role', 'user')
    if role not in ('admin', 'user'):
        role = 'user'
    error = None
    if not username:
        error = 'Укажите логин.'
    elif password != password_repeat:
        error = 'Пароли не совпадают.'
    else:
        error = validate_password(password, username=username)
    if not error and User.query.filter_by(username=username).first():
        error = f'Пользователь «{username}» уже существует.'
    if error:
        users = User.query.order_by(User.created_at.desc()).all()
        return render_template('users.html', users=users, error=error,
                               available_modules=registered_modules(),
                               companies=Company.query.order_by(Company.name).all(),
                               create_form=request.form)
    new_user = User(
        username=username,
        password_hash=_hash_password(password),
        role=role,
        password_change_required=True,
        password_version=1,
    )
    db.session.add(new_user)
    db.session.flush()
    for key in valid_module_keys(request.form.getlist('modules')):
        db.session.add(UserModuleAccess(user_id=new_user.id, module_key=key))
    selected_company_ids = {
        int(value) for value in request.form.getlist('companies')
        if str(value).isdigit()
    }
    valid_company_ids = {
        item.id for item in Company.query.filter(
            Company.id.in_(selected_company_ids)
        ).all()
    } if selected_company_ids else set()
    for company_id in sorted(valid_company_ids):
        db.session.add(CompanyMembership(
            user_id=new_user.id,
            company_id=company_id,
            participation_type='employee',
            is_primary=False,
        ))
    _audit_in_transaction('user_created', {
        'username': username, 'role': role, 'password_change_required': True,
    })
    db.session.commit()
    return redirect(url_for('users_list', created=username), code=303)


@app.route('/users/<int:user_id>/delete', methods=['POST'])
@admin_required
def users_delete(user_id):
    user = db.session.get(User, user_id)
    if user and user.id != current_user.id:
        log_event('user_deleted', {'username': user.username})
        db.session.delete(user)
        db.session.commit()
    return redirect(url_for('users_list'))


@app.route('/users/<int:user_id>/set-password', methods=['POST'])
@admin_required
def users_set_password(user_id):
    user = db.session.get(User, user_id)
    password = request.form.get('password', '')
    password_repeat = request.form.get('password_repeat')
    if password_repeat is None:
        password_repeat = request.form.get(
            'password_confirmation',
            request.form.get('password_confirm', ''),
        )
    error = None
    if not user:
        error = 'Пользователь не найден.'
    elif password != password_repeat:
        error = 'Пароли не совпадают.'
    elif _check_password(password, user.password_hash):
        error = 'Новый пароль должен отличаться от текущего.'
    else:
        error = validate_password(password, username=user.username)
    if not error:
        _set_user_password(user, password, require_change=True)
        _audit_in_transaction(
            'user_password_changed', {'username': user.username},
            user_id=current_user.id)
        db.session.commit()
    else:
        flash(error)
    return redirect(url_for('users_list'))


@app.route('/users/<int:user_id>/set-role', methods=['POST'])
@admin_required
def users_set_role(user_id):
    user = db.session.get(User, user_id)
    role = request.form.get('role', 'user')
    if role not in ('admin', 'user'):
        role = 'user'
    # Нельзя понизить самого себя
    if user and user.id != current_user.id:
        old_role = user.role
        user.role = role
        db.session.commit()
        log_event('user_role_changed',
                  {'username': user.username, 'old': old_role, 'new': role})
    return redirect(url_for('users_list'))


@app.route('/users/<int:user_id>/set-modules', methods=['POST'])
@admin_required
def users_set_modules(user_id):
    """Назначения сохраняются независимо от роли и снова действуют при понижении."""
    user = db.session.get(User, user_id) or abort(404)
    selected = valid_module_keys(request.form.getlist('modules'))
    UserModuleAccess.query.filter_by(user_id=user.id).delete()
    for key in selected:
        db.session.add(UserModuleAccess(user_id=user.id, module_key=key))
    db.session.commit()
    log_event('user_modules_changed',
              {'username': user.username, 'modules': selected})
    return redirect(url_for('users_list'))


@app.route('/users/<int:user_id>/set-companies', methods=['POST'])
@admin_required
def users_set_companies(user_id):
    """Обновляет доступные предприятия без изменения внутренних типов участия."""
    user = db.session.get(User, user_id) or abort(404)
    selected = {
        int(value) for value in request.form.getlist('companies')
        if str(value).isdigit()
    }
    valid_ids = {
        item.id for item in Company.query.filter(Company.id.in_(selected)).all()
    } if selected else set()
    existing = {
        membership.company_id: membership
        for membership in CompanyMembership.query.filter_by(user_id=user.id).all()
    }
    for company_id, membership in existing.items():
        if company_id not in valid_ids:
            db.session.delete(membership)
    for company_id in sorted(valid_ids - existing.keys()):
        db.session.add(CompanyMembership(
            user_id=user.id, company_id=company_id,
            participation_type='employee',
            is_primary=False,
        ))
    db.session.commit()
    log_event('user_company_memberships_changed', {
        'user_id': user.id, 'company_ids': sorted(valid_ids),
    })
    return redirect(url_for('users_list'))


# ---------------------------------------------------------------------------
# Инициализация БД и первого пользователя
# ---------------------------------------------------------------------------
def _ensure_user_password_columns():
    """Install lifecycle columns before durable restore can query User rows."""
    from sqlalchemy import inspect as _inspect_password, text as _text_password
    inspector = _inspect_password(db.engine)
    existing = {column['name'] for column in inspector.get_columns('users')}
    columns = (
        ('password_changed_at', "TIMESTAMP WITH TIME ZONE"),
        ('password_change_required', "BOOLEAN NOT NULL DEFAULT FALSE"),
        ('password_version', "INTEGER NOT NULL DEFAULT 1"),
    )
    with db.engine.begin() as connection:
        for name, definition in columns:
            if name not in existing:
                connection.execute(_text_password(
                    f'ALTER TABLE users ADD COLUMN {name} {definition}'))
        connection.execute(_text_password(
            "UPDATE users SET password_version = 1 "
            "WHERE password_version IS NULL OR password_version < 1"))
        connection.execute(_text_password(
            "UPDATE users SET password_change_required = FALSE "
            "WHERE password_change_required IS NULL"))


def _ensure_initial_admin_after_restore():
    """Create the bootstrap account only after durable recovery has settled."""
    if User.query.first():
        return False
    initial_password = os.environ.get('INITIAL_ADMIN_PASSWORD', '')
    if not initial_password:
        raise RuntimeError(
            'No initial admin account was created. Set '
            'INITIAL_ADMIN_PASSWORD (a compliant secret) before '
            'starting a clean database.')
    if validate_password(initial_password, username='admin'):
        raise RuntimeError(
            'INITIAL_ADMIN_PASSWORD does not satisfy the password policy')
    admin = User(
        username='admin',
        password_hash=_hash_password(initial_password),
        role='admin',
        password_change_required=True,
        password_version=1,
    )
    db.session.add(admin)
    db.session.flush()
    _audit_in_transaction(
        'user_created',
        {'username': 'admin', 'bootstrap': True,
         'password_change_required': True},
        user_id=admin.id)
    db.session.commit()
    app.logger.info(
        'Created initial admin account from configured secret; '
        'password change is required')
    return True


def init_db(perform_durable=True):
    """Prepare the database, optionally including slow App Storage recovery.

    Gunicorn runs the schema/migrations synchronously, then performs the
    additive durable restore and startup snapshot in one background thread.
    CLI and deployment entrypoints keep the safer all-in-one default.
    """
    with app.app_context():
        db.create_all()
        # Limits tables are additive and may predate the current model set.
        # Complete this migration before durable restore reads restored rows.
        from limits_schema import migrate_limits_schema
        migrate_limits_schema()
        # User is part of the durable restore set; make its additive columns
        # available before restore queries the model on an old database.
        _ensure_user_password_columns()
        # Миграции must complete before durable restore so restored rows are
        # read with the current schema. Bootstrap is deliberately later still:
        # a restored original user always wins over an initial admin.
        # Миграция: новые колонки парсера документов и ИИ-провайдера.
        # Диалектобезопасно: колонки добавляются только если их нет
        # (инспекция схемы вместо IF NOT EXISTS — работает и на SQLite).
        from sqlalchemy import inspect as _inspect_docs, text as _text_docs
        _insp = _inspect_docs(db.engine)
        # Инвариант годовых версий: ровно одна активная версия на предприятие
        # и год. Для существующих аддитивных таблиц create_all индекс не создаст.
        try:
            with db.engine.begin() as _historical_conn:
                _historical_conn.execute(_text_docs(
                    "CREATE UNIQUE INDEX IF NOT EXISTS "
                    "uq_historical_active_year "
                    "ON historical_year_versions(company_id, year) "
                    "WHERE is_active"))
        except Exception as _historical_index_exc:
            app.logger.warning(
                'historical active-version index migration failed: %s',
                _historical_index_exc)

        def _add_missing(table, columns):
            existing = {c['name'] for c in _insp.get_columns(table)}
            with db.engine.connect() as _cd:
                for _col, _def in columns:
                    if _col in existing:
                        continue
                    _cd.execute(_text_docs(
                        f'ALTER TABLE {table} ADD COLUMN {_col} {_def}'))
                    app.logger.info('migration: %s.%s added', table, _col)
                _cd.commit()

        _add_missing('daily_rows', [
            ('source_type', "VARCHAR(40) NOT NULL DEFAULT 'unknown'"),
        ])

        # Task 117 is additive for deployed installations which already have
        # the analogue tables.  Columns carrying historical children remain
        # nullable during the upgrade; new imports always fill their version.
        _analog_columns = {
            'analog_passports': [
                ('extracted_json', "TEXT NOT NULL DEFAULT '{}'"),
                ('overrides_json', "TEXT NOT NULL DEFAULT '{}'"),
                ('provenance_json', "TEXT NOT NULL DEFAULT '{}'"),
                ('legacy_json', "TEXT NOT NULL DEFAULT '{}'"),
                ('cadastral_number', "VARCHAR(64) DEFAULT ''"),
                ('natural_conditions', "TEXT DEFAULT ''"),
                ('construction_object', "TEXT DEFAULT ''"),
                ('climate_source', "TEXT DEFAULT ''"),
                ('climate_reference_version', "VARCHAR(80) DEFAULT ''"),
                ('climate_reference_date', "DATE"),
            ],
            'analog_passport_versions': [
                ('cadastral_number', "VARCHAR(64) DEFAULT ''"),
                ('natural_conditions', "TEXT DEFAULT ''"),
                ('construction_object', "TEXT DEFAULT ''"),
                ('adjustment_reference_version', "VARCHAR(80) DEFAULT ''"),
                ('adjustment_reference_date', "DATE"),
                ('adjustment_source', "TEXT DEFAULT ''"),
                ('adjustment_price_level', "VARCHAR(120) DEFAULT ''"),
            ],
            'analog_sources': [
                ('raw_conclusion_path', "VARCHAR(1000) NOT NULL DEFAULT ''"),
                ('raw_ssr_path', "VARCHAR(1000) NOT NULL DEFAULT ''"),
            ],
            'analog_imports': [
                ('raw_conclusion_path', "VARCHAR(1000) DEFAULT ''"),
                ('raw_ssr_path', "VARCHAR(1000) DEFAULT ''"),
                ('target_passport_id', 'INTEGER'),
            ],
            'analog_network_attributes': [
                ('version_id', 'INTEGER'),
            ],
            'analog_organization_roles': [
                ('version_id', 'INTEGER'), ('extracted_json', "TEXT NOT NULL DEFAULT '{}'"),
                ('override_json', "TEXT NOT NULL DEFAULT '{}'"),
                ('effective_json', "TEXT NOT NULL DEFAULT '{}'"),
                ('provenance', "TEXT DEFAULT 'conclusionRaw'"),
            ],
            'analog_teps': [
                ('version_id', 'INTEGER'), ('extracted_json', "TEXT NOT NULL DEFAULT '{}'"),
                ('override_json', "TEXT NOT NULL DEFAULT '{}'"),
                ('effective_json', "TEXT NOT NULL DEFAULT '{}'"),
                ('normalized_value', 'NUMERIC'), ('normalized_unit', "VARCHAR(80) DEFAULT ''"),
            ],
            'analog_documents': [
                ('node_kind', "VARCHAR(16) NOT NULL DEFAULT 'document'"),
                ('parent_ordinal', "VARCHAR(40) DEFAULT ''"), ('version_id', 'INTEGER'),
                ('extracted_json', "TEXT NOT NULL DEFAULT '{}'"),
                ('override_json', "TEXT NOT NULL DEFAULT '{}'"),
                ('effective_json', "TEXT NOT NULL DEFAULT '{}'"),
                ('provenance', "TEXT DEFAULT 'conclusionRaw/Documents'"),
            ],
            'analog_cost_rows': [
                ('version_id', 'INTEGER'), ('extracted_json', "TEXT NOT NULL DEFAULT '{}'"),
                ('override_json', "TEXT NOT NULL DEFAULT '{}'"),
                ('effective_json', "TEXT NOT NULL DEFAULT '{}'"),
            ],
            'analog_ssr_chapters': [
                ('extracted_json', "TEXT NOT NULL DEFAULT '{}'"),
                ('override_json', "TEXT NOT NULL DEFAULT '{}'"),
                ('effective_json', "TEXT NOT NULL DEFAULT '{}'"),
            ],
            'analog_ssr_items': [
                ('extracted_json', "TEXT NOT NULL DEFAULT '{}'"),
                ('override_json', "TEXT NOT NULL DEFAULT '{}'"),
                ('effective_json', "TEXT NOT NULL DEFAULT '{}'"),
                ('provenance', "TEXT DEFAULT 'ССР 3.01'"),
            ],
            'analog_ssr_signatures': [],
        }
        for _analog_table, _analog_defs in _analog_columns.items():
            if _analog_table in _insp.get_table_names():
                _add_missing(_analog_table, _analog_defs)
        if all(_table in _insp.get_table_names() for _table in (
                'analog_network_attributes', 'analog_passports')):
            try:
                with db.engine.begin() as _analog_backfill:
                    for _child_table in (
                            'analog_network_attributes', 'analog_organization_roles',
                            'analog_teps', 'analog_documents', 'analog_cost_rows'):
                        if _child_table not in _insp.get_table_names():
                            continue
                        _analog_backfill.execute(_text_docs(
                            f'UPDATE {_child_table} SET version_id = '
                            f'(SELECT current_version_id FROM analog_passports '
                            f'WHERE analog_passports.id = {_child_table}.passport_id) '
                            'WHERE version_id IS NULL'))
            except Exception as _analog_backfill_exc:
                app.logger.warning('analog network version backfill failed: %s',
                                   _analog_backfill_exc)
        if 'analog_organization_roles' in _insp.get_table_names():
            try:
                with db.engine.begin() as _analog_role_conn:
                    if db.engine.dialect.name == 'postgresql':
                        _analog_role_conn.execute(_text_docs(
                            'ALTER TABLE analog_organization_roles '
                            'DROP CONSTRAINT IF EXISTS uq_analog_organization_role'))
                    _analog_role_conn.execute(_text_docs(
                        'CREATE UNIQUE INDEX IF NOT EXISTS '
                        'uq_analog_organization_role_version '
                        'ON analog_organization_roles(passport_id, version_id, role)'))
            except Exception as _analog_role_exc:
                app.logger.warning('analog role index migration failed: %s',
                                   _analog_role_exc)
        if db.engine.dialect.name == 'postgresql':
            try:
                with db.engine.begin() as _backfill_conn:
                    for _child_table in (
                            'analog_network_attributes', 'analog_organization_roles',
                            'analog_teps', 'analog_documents', 'analog_cost_rows'):
                        _backfill_conn.execute(_text_docs(
                            f'UPDATE {_child_table} a SET version_id = '
                            'p.current_version_id FROM analog_passports p '
                            'WHERE a.passport_id = p.id AND a.version_id IS NULL'))
                    for _child_table in (
                            'analog_network_attributes', 'analog_organization_roles',
                            'analog_teps', 'analog_documents', 'analog_cost_rows'):
                        _backfill_conn.execute(_text_docs(
                            f'ALTER TABLE {_child_table} ALTER COLUMN version_id SET NOT NULL'))
            except Exception as _backfill_exc:
                app.logger.warning('analog network version backfill deferred: %s',
                                   _backfill_exc)
            try:
                with db.engine.begin() as _network_unique:
                    _network_unique.execute(_text_docs(
                        'ALTER TABLE analog_network_attributes DROP CONSTRAINT IF EXISTS '
                        'uq_analog_network_attribute'))
                    _network_unique.execute(_text_docs(
                        'DROP INDEX IF EXISTS uq_analog_network_attribute'))
                    _network_unique.execute(_text_docs(
                        'CREATE UNIQUE INDEX IF NOT EXISTS '
                        'uq_analog_network_attribute_version '
                        'ON analog_network_attributes(passport_id, version_id, key)'))
            except Exception as _network_unique_exc:
                app.logger.warning('analog network uniqueness migration failed: %s',
                                   _network_unique_exc)
            for _table, _constraint, _column, _target in (
                ('analog_imports', 'fk_analog_import_target', 'target_passport_id',
                 'analog_passports(id)'),
                ('analog_network_attributes', 'fk_analog_network_version', 'version_id',
                 'analog_passport_versions(id)'),
                ('analog_organization_roles', 'fk_analog_role_version', 'version_id',
                 'analog_passport_versions(id)'),
                ('analog_teps', 'fk_analog_tep_version', 'version_id',
                 'analog_passport_versions(id)'),
                ('analog_cost_rows', 'fk_analog_cost_version', 'version_id',
                 'analog_passport_versions(id)'),
                ('analog_documents', 'fk_analog_document_version', 'version_id',
                 'analog_passport_versions(id)'),
                ('analog_ssr_chapters', 'fk_analog_chapter_version', 'version_id',
                 'analog_passport_versions(id)'),
                ('analog_ssr_items', 'fk_analog_item_version', 'version_id',
                 'analog_passport_versions(id)'),
                ('analog_ssr_signatures', 'fk_analog_signature_version', 'version_id',
                 'analog_passport_versions(id)'),
                ('analog_documents', 'fk_analog_document_parent', 'parent_id',
                 'analog_documents(id)'),
            ):
                try:
                    with db.engine.begin() as _fk_conn:
                        _fk_conn.execute(_text_docs(
                            f'ALTER TABLE {_table} ADD CONSTRAINT {_constraint} '
                            f'FOREIGN KEY ({_column}) REFERENCES {_target}'))
                except Exception:
                    # Existing deployments may already carry the equivalent
                    # constraint; the additive upgrade remains idempotent.
                    pass
            try:
                with db.engine.begin() as _analog_index_conn:
                    _analog_index_conn.execute(_text_docs(
                        'CREATE INDEX IF NOT EXISTS ix_analog_network_version_key '
                        'ON analog_network_attributes(version_id, key)'))
                    _analog_index_conn.execute(_text_docs(
                        'CREATE INDEX IF NOT EXISTS ix_analog_document_parent '
                        'ON analog_documents(parent_id)'))
            except Exception:
                pass

        _add_missing('users', [
            ('password_changed_at', "TIMESTAMP WITH TIME ZONE"),
            ('password_change_required', "BOOLEAN NOT NULL DEFAULT FALSE"),
            ('password_version', "INTEGER NOT NULL DEFAULT 1"),
        ])
        with db.engine.begin() as _user_password_conn:
            _user_password_conn.execute(_text_docs(
                "UPDATE users SET password_version = 1 "
                "WHERE password_version IS NULL OR password_version < 1"))
            _user_password_conn.execute(_text_docs(
                "UPDATE users SET password_change_required = FALSE "
                "WHERE password_change_required IS NULL"))

        _add_missing('ai_providers', [
            ('is_local', "BOOLEAN NOT NULL DEFAULT FALSE"),
            ('supports_vision', "BOOLEAN NOT NULL DEFAULT FALSE"),
        ])
        _add_missing('companies', [
            ('internal_code', "VARCHAR(80) DEFAULT ''"),
            ('identity_key', "VARCHAR(80)"),
            ('status', "VARCHAR(20) NOT NULL DEFAULT 'active'"),
            ('osv_source_revision', "VARCHAR(80) NOT NULL DEFAULT ''"),
            ('is_rso', "BOOLEAN NOT NULL DEFAULT TRUE"),
            ('curator_user_id', "INTEGER REFERENCES users(id)"),
            ('ogrn', "VARCHAR(20) DEFAULT ''"),
            ('kpp', "VARCHAR(20) DEFAULT ''"),
            ('opf_short', "VARCHAR(120) DEFAULT ''"),
            ('legal_status', "VARCHAR(40) DEFAULT ''"),
            ('registration_date', "DATE"),
            ('ogrn_date', "DATE"),
            ('actuality_date', "DATE"),
            ('liquidation_date', "DATE"),
            ('legal_address', "TEXT DEFAULT ''"),
            ('address_quality', "VARCHAR(40) DEFAULT ''"),
            ('address_ids_json', "TEXT DEFAULT '{}'"),
            ('director_position', "VARCHAR(255) DEFAULT ''"),
            ('primary_okved', "VARCHAR(32) DEFAULT ''"),
            ('classifications_json', "TEXT DEFAULT '{}'"),
            ('labels_json', "TEXT DEFAULT '[]'"),
            ('dadata_source', "VARCHAR(40) DEFAULT ''"),
            ('dadata_received_at', "TIMESTAMP WITH TIME ZONE"),
            ('dadata_actuality_date', "DATE"),
            ('dadata_sync_status', "VARCHAR(24) DEFAULT 'never'"),
            ('dadata_last_error', "TEXT DEFAULT ''"),
            ('profile_version', "INTEGER NOT NULL DEFAULT 1"),
            ('updated_at', "TIMESTAMP WITH TIME ZONE"),
        ])
        with db.engine.begin() as _company_conn:
            _company_conn.execute(_text_docs(
                "UPDATE companies SET legal_address = address "
                "WHERE COALESCE(legal_address, '') = '' AND COALESCE(address, '') != ''"))
            _company_conn.execute(_text_docs(
                "UPDATE companies SET status = 'active' "
                "WHERE status IS NULL OR status = ''"))
            _company_conn.execute(_text_docs(
                "UPDATE companies SET internal_code = CAST(id AS VARCHAR(80)) "
                "WHERE COALESCE(TRIM(internal_code), '') = ''"))
        for _company in Company.query.order_by(Company.id).all():
            _classifications = _safe_json(_company.classifications_json, {})
            if not isinstance(_classifications, dict):
                _classifications = {}
            _types = _classifications.get('organization_types')
            if not isinstance(_types, list):
                _types = []
            if not _types and _company.is_rso:
                _classifications['organization_types'] = ['РСО']
                _company.classifications_json = json.dumps(
                    _classifications, ensure_ascii=False)
        # Канонический ключ и совместимый scope для существующих установок.
        existing_keys = set()
        for _company in Company.query.order_by(Company.id).all():
            _key = company_identity_key(
                _company.inn or '', _company.ogrn or '', _company.kpp or '')
            if _key and _key not in existing_keys:
                _company.identity_key = _key
                existing_keys.add(_key)
            elif _key:
                _company.identity_key = None
        db.session.commit()
        try:
            with db.engine.begin() as _identity_conn:
                _identity_conn.execute(_text_docs(
                    "CREATE UNIQUE INDEX IF NOT EXISTS "
                    "uq_companies_identity_key ON companies(identity_key) "
                    "WHERE identity_key IS NOT NULL AND identity_key != ''"))
        except Exception as _identity_exc:
            app.logger.warning('company identity index migration failed: %s',
                               _identity_exc)
        if Setting.get('company_membership_backfill_v1', '') != 'done':
            _companies = Company.query.all()
            _grants = UserModuleAccess.query.filter_by(module_key='analytics').all()
            for _grant in _grants:
                for _company in _companies:
                    if not CompanyMembership.query.filter_by(
                            user_id=_grant.user_id,
                            company_id=_company.id).first():
                        db.session.add(CompanyMembership(
                            user_id=_grant.user_id, company_id=_company.id,
                            participation_type='legacy_access',
                            is_primary=False))
            for _company in _companies:
                if _company.created_by and not CompanyMembership.query.filter_by(
                        user_id=_company.created_by,
                        company_id=_company.id).first():
                    db.session.add(CompanyMembership(
                        user_id=_company.created_by, company_id=_company.id,
                        participation_type='creator', is_primary=False))
            Setting.set('company_membership_backfill_v1', 'done')
            db.session.commit()
        _add_missing('doc_files', [
            ('approved',    "BOOLEAN NOT NULL DEFAULT FALSE"),
            ('approved_at', "TIMESTAMP WITH TIME ZONE"),
            ('analysis_status', "VARCHAR(20) NOT NULL DEFAULT 'none'"),
        ])
        _add_missing('doc_packages', [
            ('analysis_status', "VARCHAR(20) NOT NULL DEFAULT 'none'"),
            ('analysis_error',  "TEXT DEFAULT ''"),
            ('analysis_total',  "INTEGER DEFAULT 0"),
            ('analysis_done',   "INTEGER DEFAULT 0"),
        ])
        _add_missing('doc_analysis_runs', [
            ('options_json', "TEXT DEFAULT '{}'"),
            ('result_json',  "TEXT DEFAULT '{}'"),
            ('catalog_run_id', "INTEGER REFERENCES doc_catalog_analysis_runs(id)"),
        ])
        _add_missing('doc_file_pages', [
            ('image_width', "INTEGER DEFAULT 0"),
            ('image_height', "INTEGER DEFAULT 0"),
            ('words_json', "TEXT DEFAULT '[]'"),
            ('coordinate_version', "VARCHAR(40) DEFAULT ''"),
        ])
        _add_missing('doc_findings', [
            ('source', "VARCHAR(20) DEFAULT ''"),
            ('boxes_json', "TEXT DEFAULT '[]'"),
            ('confidence', "DOUBLE PRECISION"),
            ('algorithm_version', "VARCHAR(40) DEFAULT ''"),
        ])
        _add_missing('doc_catalog_analysis_runs', [
            ('skipped_json', "TEXT DEFAULT '[]'"),
            ('catalog_name_snapshot', "VARCHAR(255) DEFAULT ''"),
            ('files_cancelled', "INTEGER NOT NULL DEFAULT 0"),
            ('cancel_requested_by', "INTEGER REFERENCES users(id)"),
            ('cancel_requested_at', "TIMESTAMP WITH TIME ZONE"),
            ('cancelled_at', "TIMESTAMP WITH TIME ZONE"),
        ])
        _add_missing('uploads', [
            ('metrics_json', "TEXT DEFAULT '{}'"),
            # Nullable by contract: existing Upload rows must not receive a
            # guessed file-to-upload association.
            ('source_manifest_json', "TEXT"),
        ])
        _add_missing('ai_summaries', [
            ('analysis_start', "DATE"),
            ('analysis_end', "DATE"),
            ('source_revision', "VARCHAR(64) DEFAULT ''"),
        ])
        try:
            stale_file_summaries = OsvFileAiSummary.query.filter(
                OsvFileAiSummary.status.in_(('queued', 'running'))).all()
            for _file_summary in stale_file_summaries:
                _file_summary.status = 'error'
                _file_summary.error_text = (
                    'Формирование прервано перезапуском приложения. '
                    'Запустите его повторно.')
            if stale_file_summaries:
                db.session.commit()
        except Exception as _file_summary_recovery_exc:
            db.session.rollback()
            app.logger.warning(
                'OSV file AI summary recovery failed: %s',
                _file_summary_recovery_exc)
        if db.engine.dialect.name == 'postgresql':
            try:
                from sqlalchemy import text as _text_catalog_run
                with db.engine.begin() as _catalog_conn:
                    _catalog_conn.execute(_text_catalog_run(
                        'ALTER TABLE doc_catalog_analysis_runs '
                        'ALTER COLUMN catalog_id DROP NOT NULL'))
                    _catalog_conn.execute(_text_catalog_run(
                        'UPDATE doc_catalog_analysis_runs r '
                        'SET catalog_name_snapshot = c.name '
                        'FROM doc_catalogs c '
                        'WHERE r.catalog_id = c.id '
                        "AND COALESCE(r.catalog_name_snapshot, '') = ''"))
            except Exception as _catalog_migration_exc:
                app.logger.warning(
                    'catalog run nullable migration failed: %s',
                    _catalog_migration_exc)
        # Daemon-поток не переживает рестарт. Не оставляем историю текущего
        # документа навсегда в состоянии queued/running.
        try:
            from database import (DocAnalysisRun, DocCatalogAnalysisRun,
                                  DocCheckResult, DocFreeAnalysisResult, utcnow)
            stale_catalog_runs = DocCatalogAnalysisRun.query.filter(
                DocCatalogAnalysisRun.status.in_(
                    ('queued', 'running', 'cancel_requested'))).all()
            for _catalog_run in stale_catalog_runs:
                if _catalog_run.status == 'cancel_requested':
                    _catalog_run.status = 'cancelled'
                    _catalog_run.error_text = ''
                    _catalog_run.cancelled_at = utcnow()
                    for _child in _catalog_run.file_runs:
                        if _child.status in (
                                'queued', 'running', 'cancel_requested'):
                            _child.status = 'cancelled'
                            _child.error_text = 'Анализ отменён пользователем'
                            _child.completed_at = utcnow()
                    _catalog_run.files_cancelled = sum(
                        1 for _child in _catalog_run.file_runs
                        if _child.status == 'cancelled')
                    _catalog_run.files_done = sum(
                        1 for _child in _catalog_run.file_runs
                        if _child.status in ('done', 'error'))
                else:
                    _catalog_run.status = 'error'
                    _catalog_run.error_text = (
                        'Анализ каталога прерван перезапуском приложения. '
                        'Запустите его повторно.')
                _catalog_run.completed_at = utcnow()
                if (_catalog_run.package
                        and _catalog_run.package.analysis_status == 'running'):
                    _catalog_run.package.analysis_status = (
                        'done' if _catalog_run.status == 'cancelled' else 'error')
                    _catalog_run.package.analysis_error = _catalog_run.error_text
            stale_runs = DocAnalysisRun.query.filter(
                DocAnalysisRun.status.in_(
                    ('queued', 'running', 'cancel_requested'))).all()
            for _run in stale_runs:
                was_cancelled = _run.status == 'cancel_requested'
                _run.status = 'cancelled' if was_cancelled else 'error'
                _run.error_text = (
                    'Анализ отменён пользователем' if was_cancelled else
                    'Анализ прерван перезапуском приложения. Запустите его повторно.')
                _run.completed_at = utcnow()
                _pkg = _run.package
                if _pkg and _pkg.analysis_status == 'running':
                    _pkg.analysis_status = 'error'
                    _pkg.analysis_error = _run.error_text
                _file = _run.file
                if _file and _file.analysis_status == 'processing':
                    has_previous = (
                        DocCheckResult.query.filter_by(file_id=_file.id).first()
                        or DocFreeAnalysisResult.query.filter_by(
                            file_id=_file.id).first())
                    _file.analysis_status = (
                        'done' if has_previous else
                        ('none' if was_cancelled else 'error'))
            if stale_runs or stale_catalog_runs:
                db.session.commit()
                app.logger.warning(
                    'recovered %s document and %s catalog analysis runs',
                    len(stale_runs), len(stale_catalog_runs))
        except Exception as _analysis_recovery_exc:
            db.session.rollback()
            app.logger.warning(
                'document analysis recovery failed: %s',
                _analysis_recovery_exc)
        # Универсальный парсер документов: зависшие после рестарта пакеты → error
        try:
            from doc_pipeline import recover_stale_packages
            recover_stale_packages()
        except Exception as _rexc:
            app.logger.warning(f'recover_stale_packages failed: {_rexc}')
        # Миграция: добавляем новые колонки если отсутствуют
        try:
            from sqlalchemy import text as _text
            with db.engine.connect() as _conn:
                _conn.execute(_text(
                    "ALTER TABLE uploads ADD COLUMN IF NOT EXISTS "
                    "source VARCHAR(20) DEFAULT 'manual'"
                ))
                _conn.commit()
        except Exception:
            pass
        # Миграция: добавляем колонку role если отсутствует
        try:
            from sqlalchemy import text as _text2
            with db.engine.connect() as _conn2:
                _conn2.execute(_text2(
                    "ALTER TABLE users ADD COLUMN IF NOT EXISTS "
                    "role VARCHAR(20) NOT NULL DEFAULT 'user'"
                ))
                # Первый созданный пользователь (наименьший id) — всегда admin
                _conn2.execute(_text2(
                    "UPDATE users SET role = 'admin' "
                    "WHERE id = (SELECT MIN(id) FROM users)"
                ))
                _conn2.commit()
        except Exception:
            pass
        # Миграция: реквизиты предприятий (справочник РСО)
        try:
            from sqlalchemy import text as _text3
            with db.engine.connect() as _conn3:
                for _col, _def in [
                    ("inn",              "VARCHAR(20)  DEFAULT ''"),
                    ("full_name",        "TEXT         DEFAULT ''"),
                    ("address",          "TEXT         DEFAULT ''"),
                    ("director_name",    "VARCHAR(255) DEFAULT ''"),
                    ("director_phone",   "VARCHAR(100) DEFAULT ''"),
                    ("accountant_name",  "VARCHAR(255) DEFAULT ''"),
                    ("accountant_phone", "VARCHAR(100) DEFAULT ''"),
                    ("note",             "TEXT         DEFAULT ''"),
                ]:
                    _conn3.execute(_text3(
                        f"ALTER TABLE companies ADD COLUMN IF NOT EXISTS {_col} {_def}"
                    ))
                _conn3.commit()
        except Exception:
            pass
        # A restored setting is not proof that this process has reconciled
        # the history window. Keep consumers stopped until restore/reconcile.
        _limits_history_ready.clear()
        Setting.set('limits_history_ready', 'restoring')
        db.session.commit()
        if not perform_durable:
            # Gunicorn completes recovery and bootstrap in its background
            # startup phase after this schema-only step.
            app.config['DURABLE_STARTUP_READY'] = False
            return
        if not _durable_restore():
            app.config['DURABLE_STARTUP_READY'] = False
            raise RuntimeError(
                'Durable restore failed; bootstrap was not attempted')
        _ensure_initial_admin_after_restore()
        _finish_durable_startup()
        app.config['DURABLE_STARTUP_READY'] = True
def _finish_durable_startup(restore=False):
    """Complete restore, bootstrap, and backup after schema readiness."""
    with app.app_context():
        if restore:
            if not _durable_restore():
                app.config['DURABLE_STARTUP_READY'] = False
                raise RuntimeError(
                    'Durable restore failed; bootstrap was not attempted')
            _ensure_initial_admin_after_restore()
        _import_bundled_mup_kgh_account55()
        _reconcile_limits_after_restore()
        if not app.config.get('TESTING'):
            from durable_data import has_snapshot
            limit_cleanup = Setting.get('limits_rotation_needs_backup', '') == '1'
            if Setting.get('limits_tombstones_external_pending') == '1':
                try:
                    from durable_data import backup_limit_tombstones
                    backup_limit_tombstones(db)
                    Setting.set('limits_tombstones_external_pending', '0')
                    Setting.set('limits_tombstone_backup_error', '')
                    db.session.commit()
                except Exception as exc:
                    db.session.rollback()
                    Setting.set('limits_tombstone_backup_error', str(exc))
                    db.session.commit()
            result = _durable_backup(
                'очистка лимитов после запуска' if limit_cleanup else
                'первичный снимок после запуска',
                required=not has_snapshot(),
                allowed_removed_tables=(
                    'limit_versions', 'limit_row_snapshots', 'limit_comparisons',
                    'limit_comparison_resolutions', 'limit_changes', 'limit_deliveries',
                    'limit_delivery_attempts', 'limit_delivery_parts') if limit_cleanup else ())
            if limit_cleanup and result and result.get('promoted'):
                Setting.set('limits_rotation_needs_backup', '0')
                db.session.commit()
        _ensure_limits_recovery()
        app.config['DURABLE_STARTUP_READY'] = True

def _reconcile_limits_after_restore():
    from limits_service import (cleanup_purged_files, cleanup_temporary_versions,
                                reconcile_history_window)
    try:
        cleanup_purged_files()
        cleanup_temporary_versions(int(os.environ.get('LIMITS_TEMP_TTL_HOURS', '72')))
        blocked = reconcile_history_window()
        Setting.set('limits_history_ready', 'blocked' if blocked else 'ready')
        Setting.set('limits_history_blocked', json.dumps(blocked))
        db.session.commit()
        if not blocked:
            _limits_history_ready.set()
        else:
            _limits_history_ready.clear()
        return blocked
    except Exception:
        db.session.rollback()
        _limits_history_ready.clear()
        Setting.set('limits_history_ready', 'blocked')
        db.session.commit()
        app.logger.exception('limits history reconciliation failed')
        raise
def _dataset_closure_status(dataset):
    """Закрытость по структурированным помесячным статусам."""
    labels = dataset.get('closing_label') or []
    if labels and all(label == 'закрыт' for label in labels):
        return True
    if any(label in ('НЕ ЗАКРЫТ', 'фин.рез. не списан') for label in labels):
        return False
    return None

def _rating_profiles_for_user():
    rows = RatingProfile.query.filter_by(owner_id=current_user.id)\
        .order_by(RatingProfile.name, RatingProfile.id).all()
    payload = [_rating_profile_payload()] + [_rating_profile_payload(row) for row in rows]
    if not any(item['is_default'] for item in payload):
        payload[0]['is_default'] = True
    return payload

@app.route('/rating-api/profiles', methods=['GET', 'POST'])
@login_required
def rating_profiles():
    from rating import ProfileValidationError, validate_profile
    if request.method == 'GET':
        return jsonify(profiles=_rating_profiles_for_user())
    try:
        config = validate_profile(request.get_json(silent=True) or {})
    except (ProfileValidationError, TypeError, ValueError) as exc:
        return jsonify(error=str(exc)), 400
    duplicate = RatingProfile.query.filter_by(
        owner_id=current_user.id, name=config['name']).first()
    if duplicate:
        return jsonify(error='Профиль с таким названием уже существует'), 409
    row = RatingProfile(owner_id=current_user.id, name=config['name'],
                        config_json=json.dumps(config, ensure_ascii=False))
    db.session.add(row)
    db.session.commit()
    return jsonify(profile=_rating_profile_payload(row)), 201

@app.route('/rating-api/catalog')
@login_required
def rating_catalog():
    from rating import METRICS
    return jsonify(metrics=METRICS, periods=[
        {'key': 'latest_available', 'name': 'Последний доступный период'},
        {'key': 'latest_closed', 'name': 'Последний закрытый период'},
        {'key': 'all_closed', 'name': 'Все доступные закрытые периоды'},
    ])

@app.route('/rating-api/profiles/<int:profile_id>', methods=['PUT', 'DELETE'])
@login_required
def rating_profile_item(profile_id):
    from rating import ProfileValidationError, validate_profile
    row = RatingProfile.query.filter_by(
        id=profile_id, owner_id=current_user.id).first_or_404()
    if request.method == 'DELETE':
        db.session.delete(row)
        db.session.commit()
        return ('', 204)
    body = request.get_json(silent=True) or {}
    try:
        config = validate_profile(body.get('config', body))
    except (ProfileValidationError, TypeError, ValueError) as exc:
        return jsonify(error=str(exc)), 400
    duplicate = RatingProfile.query.filter(
        RatingProfile.owner_id == current_user.id,
        RatingProfile.name == config['name'],
        RatingProfile.id != row.id).first()
    if duplicate:
        return jsonify(error='Профиль с таким названием уже существует'), 409
    row.name = config['name']
    row.config_json = json.dumps(config, ensure_ascii=False)
    if body.get('is_default'):
        RatingProfile.query.filter_by(owner_id=current_user.id).update(
            {'is_default': False})
        row.is_default = True
    db.session.commit()
    return jsonify(profile=_rating_profile_payload(row))

@app.route('/rating-api', methods=['GET'])
@login_required
def rating_results():
    from rating import DEFAULT_PROFILE, rank_companies, score_company, validate_profile
    profile_id = request.args.get('profile_id', 'base')
    if profile_id == 'base':
        profile = validate_profile(DEFAULT_PROFILE)
        profile_ref = 'base'
    else:
        try:
            numeric_id = int(profile_id)
        except ValueError:
            return jsonify(error='Некорректный профиль'), 400
        row = RatingProfile.query.filter_by(
            id=numeric_id, owner_id=current_user.id).first()
        if not row:
            return jsonify(error='Профиль не найден'), 404
        profile = validate_profile(row.config)
        profile_ref = row.id

    results = []
    companies_query = Company.query.filter(
        Company.status != 'archived', Company.is_rso.is_(True))
    if not current_user.is_admin:
        companies_query = companies_query.join(CompanyMembership).filter(
            CompanyMembership.user_id == current_user.id)
    for company in companies_query.order_by(Company.id).all():
        closed = [u for u in company.uploads if u.closure_status is True]
        if profile['period'] == 'latest_available':
            selected = company.uploads[:1]
        elif profile['period'] == 'latest_closed':
            selected = closed[:1]
        else:
            selected = closed
        freshness = max((u.uploaded_at for u in selected), default=None)
        snapshots = [upload.metrics for upload in selected if upload.metrics]
        # Безопасная совместимость со старыми установками: текущий файловый
        # кэш допустим только когда последняя вообще загрузка и выбранный
        # закрытый период — одна запись. Более новый открытый период исключает
        # такой fallback, чтобы не загрязнить рейтинг.
        if (not snapshots and profile['period'] != 'all_closed'
                and len(selected) == 1 and company.uploads
                and company.uploads[0].id == selected[0].id):
            kpi_path = os.path.join(OUTPUT, str(company.id), 'kpi_cache.json')
            try:
                with open(kpi_path, encoding='utf-8') as stream:
                    legacy = json.load(stream)
                if isinstance(legacy.get('metrics'), dict):
                    snapshots = [legacy['metrics']]
            except (OSError, ValueError, TypeError):
                pass
        values = {}
        # Для диапазона закрытых периодов используем среднее доступных
        # периодических значений каждой метрики; пропуски не становятся нулями.
        for key in profile['metrics']:
            observed = [snapshot.get(key) for snapshot in snapshots
                        if snapshot.get(key) is not None]
            values[key] = (sum(observed) / len(observed)) if observed else None
        no_period_reason = ('Нет доступного периода'
                            if profile['period'] == 'latest_available'
                            else 'Нет закрытого периода')
        missing = {key: (no_period_reason if not selected
                         else 'Недостаточно исходных счетов для формулы')
                   for key in profile['metrics']}
        result = score_company(values, profile, missing)
        result.update(company_id=company.id, company_name=company.name,
                      inn=company.inn or '',
                      freshness=freshness.isoformat() if freshness else None,
                      period=selected[0].period_label if len(selected) == 1 else
                      (f'{len(selected)} закрытых периодов' if selected else None),
                      detail_url=url_for('dashboard', company_id=company.id,
                                         name='dashboard_current.html') + '#rating')
        results.append(result)
    return jsonify(profile_id=profile_ref, profile=profile,
                   generated_at=datetime.now(timezone.utc).isoformat(),
                   results=rank_companies(results))

def _limits_adapter():
    from limits_max import MaxClient
    token = decrypt_value(Setting.get('limits_max_token', ''))
    return MaxClient(token=token or None)


# Самая внешняя WSGI-граница фиксирует реальный peer/scheme до ProxyFix и
# возможного DispatcherMiddleware. Не перемещать внутрь других middleware.
from one_c_transport import OneCTransportMiddleware
app.wsgi_app = OneCTransportMiddleware(app.wsgi_app)


if __name__ == '__main__':
    # PostgreSQL в Replit иногда поднимается позже Flask.
    # Повторяем попытку подключения до 10 раз с паузой 2 с.
    for _attempt in range(10):
        try:
            with app.app_context():
                init_db()
            break
        except Exception as _e:
            print(f'[startup] БД недоступна (попытка {_attempt + 1}/10): {_e}')
            if _attempt == 9:
                raise
            time.sleep(2)
    _ensure_imap_thread()
    app.run(host='0.0.0.0', port=int(os.environ.get('PORT', 8080)),
            debug=False, threaded=True)

def _restore_company_outputs(snapshot):
    """Restore outputs after a failed/cancelled filesystem transaction."""
    for path, content in snapshot.items():
        if content is None:
            try:
                os.remove(path)
            except FileNotFoundError:
                pass
            continue
        os.makedirs(os.path.dirname(path), exist_ok=True)
        with open(path, 'wb') as stream:
            stream.write(content)

def _reserve_upload_batch_dir(company_id, prefix):
    """Create a fresh batch directory without ever reusing an old one."""
    company_root = os.path.join(UPLOADS, str(company_id))
    os.makedirs(company_root, exist_ok=True)
    candidate = os.path.join(company_root, prefix)
    suffix = 0
    while True:
        try:
            os.makedirs(candidate)
            return candidate
        except FileExistsError:
            suffix += 1
            candidate = os.path.join(
                company_root, f'{prefix}_retry_{suffix}')

def _monthly_manifest_json(
        paths, parsed_data, audit, company_id, extra_accepted=()):
    """Serialize accepted monthly files, including an explicit empty array."""
    accepted = _accepted_monthly_files(paths, parsed_data, audit)
    accepted.extend(extra_accepted)
    return dump_manifest(build_manifest(
        accepted, os.path.join(UPLOADS, str(company_id))))

def _remove_pending_batch_dirs(batch_dirs):
    """Remove only directories reserved by the current worker invocation."""
    for path in reversed(batch_dirs):
        shutil.rmtree(path, ignore_errors=True)

def _unique_attachment_name(name, used_names):
    """Keep attachment/archive names while disambiguating duplicates."""
    original = os.path.basename(name) or 'attachment'
    stem, extension = os.path.splitext(original)
    candidate = original
    suffix = 1
    while candidate in used_names:
        suffix += 1
        candidate = f'{stem}_{suffix}{extension}'
    used_names.add(candidate)
    return candidate

@app.after_request
def _limits_private_cache(response):
    if request.path.startswith('/limit-notifications'):
        response.headers['Cache-Control'] = 'no-store, private'
        response.headers['X-Content-Type-Options'] = 'nosniff'
    return response

@app.route('/limit-notifications/datasets', methods=['POST'])
@admin_required
def limit_dataset_create():
    from limits_service import LimitsError, make_dataset
    try:
        dataset = make_dataset(request.form.get('name'), request.form.get('scope'), current_user.id)
        log_event('limits_dataset_created', {'dataset_id': dataset.id})
        return redirect(url_for('limit_notifications', dataset=dataset.id))
    except LimitsError as exc:
        flash(str(exc), 'error')
        return redirect(url_for('limit_notifications'))

@app.route('/limit-notifications', endpoint='limit_notifications')
@login_required
def limit_notifications():
    if Setting.get('limits_history_ready', 'ready') != 'ready':
        return render_template('limits_unavailable.html',
                               detail='Восстановление проверяет окно хранения версий и очередь уведомлений.'), 503
    datasets = LimitDataset.query.order_by(LimitDataset.updated_at.desc()).all()
    dataset_id = request.args.get('dataset', type=int)
    dataset = _limit_dataset_or_404(dataset_id) if dataset_id else (datasets[0] if datasets else None)
    version_query = LimitVersion.query.filter_by(dataset_id=dataset.id).filter(
        LimitVersion.status != 'purged') if dataset else None
    if version_query is not None and not current_user.is_admin:
        version_query = version_query.filter_by(status='accepted')
    versions = version_query.order_by(LimitVersion.uploaded_at.desc()).all() if version_query is not None else []
    comparison_ids = [row.id for row in LimitComparison.query.join(
        LimitVersion, LimitComparison.new_version_id == LimitVersion.id
    ).filter(LimitVersion.dataset_id == dataset.id).all()] if dataset else []
    changes = (LimitChange.query.filter(LimitChange.comparison_id.in_(comparison_ids)).count()
               if comparison_ids else 0)
    delivery_query = LimitDelivery.query.join(LimitVersion).filter(
        LimitVersion.dataset_id == dataset.id) if dataset else LimitDelivery.query.filter(False)
    queued = delivery_query.filter(LimitDelivery.status.in_(
        ('queued', 'retry', 'sending', 'unknown'))).count()
    current = db.session.get(LimitVersion, dataset.current_version_id) if dataset and dataset.current_version_id else None
    latest_comparison = LimitComparison.query.filter_by(new_version_id=current.id,
        old_version_id=current.previous_version_id).first() if current else None
    summary = json.loads(latest_comparison.summary_json or '{}') if latest_comparison else {}
    return render_template(
        'limits_dashboard.html', datasets=datasets, dataset=dataset, versions=versions,
        default_source_date=datetime.now(timezone(timedelta(hours=3))).date().isoformat(),
        changes=changes, queued=queued,
        max_connected=bool(Setting.get('limits_max_token', '')),
        max_errors=delivery_query.filter(LimitDelivery.status.in_(('error', 'unknown'))).count(),
        summary=summary, current_version=current,
        current_objects=LimitObject.query.filter_by(dataset_id=dataset.id, is_present=True).count() if dataset else 0,
        pending_count=sum(v.status in ('queued', 'parsing', 'review') for v in versions),
        max_enabled=Setting.get('limits_max_enabled') == '1',
        max_group_id=Setting.get('limits_max_group_chat_id', ''),
        bot_dataset_id=Setting.get('limits_bot_dataset_id', ''))

def _limit_parse_in_background(version_id):
    """Parser runs outside request; status/error is durable for restart diagnosis."""
    def work():
        with _limits_parse_slots, app.app_context():
            try:
                from limits_service import parse_version
                parse_version(version_id)
                _schedule_durable_safety_backup(f'limits_parse:{version_id}')
            except Exception:
                app.logger.exception('limits parser failed for version %s', version_id)
    thread = threading.Thread(target=work, name=f'limits-parse-{version_id}', daemon=True)
    thread.start()

@app.route('/limit-notifications/subscriptions/<int:subscription_id>/revoke', methods=['POST'])
@admin_required
def limit_subscription_revoke(subscription_id):
    from limits_service import LimitsError, revoke_subscription
    subscription = db.session.get(LimitSubscription, subscription_id) or abort(404)
    try:
        revoke_subscription(subscription.recipient_id, subscription.dataset_id, subscription.uin)
        log_event('limits_subscription_revoked', {'subscription_id': subscription.id})
        flash('Подписка отозвана. Очередь проверит это перед отправкой.', 'success')
    except LimitsError as exc:
        flash(str(exc), 'error')
    return redirect(url_for('limit_subscribers', dataset_id=subscription.dataset_id))

@app.route('/limit-notifications/recovery', methods=['GET', 'POST'])
@admin_required
def limit_history_recovery():
    """Minimal recovery controls: never expose pre-reconciliation financial data."""
    if Setting.get('limits_history_ready') == 'restoring':
        return jsonify(status='restoring', message='Дождитесь завершения восстановления.'), 503
    if request.method == 'POST':
        delivery_id = request.form.get('delivery_id', type=int)
        if delivery_id:
            from limits_bot import cancel_delivery
            cancel_delivery(delivery_id)
            log_event('limits_recovery_delivery_cancelled', {'delivery_id': delivery_id})
        _reconcile_limits_after_restore()
        if _limits_history_ready.is_set():
            _ensure_limits_recovery()
            return redirect(url_for('limit_notifications'))
    pending = LimitDelivery.query.filter(LimitDelivery.status.in_(
        ('queued', 'retry', 'sending', 'unknown', 'error'))).all()
    return render_template('limits_recovery.html', deliveries=pending,
                           blocked=json.loads(Setting.get('limits_history_blocked', '[]')))

@app.template_filter('limit_json')
def _limit_json_filter(value):
    from limits_views import decode
    return decode(value, {})

def _schedule_limits_rotation_backup():
    """Promote a removal-aware durable snapshot after 10-version rotation.

    A regular safety snapshot intentionally rejects table shrinkage.  Without
    this explicit path a later restore could reintroduce an intentionally
    purged version.  It is always background work, never an HTTP delay.
    """
    allowed = (
        'limit_versions', 'limit_row_snapshots', 'limit_comparisons',
        'limit_comparison_resolutions', 'limit_changes', 'limit_deliveries',
        'limit_delivery_attempts',
        'limit_delivery_parts',
    )
    def work():
        with app.app_context():
            try:
                from durable_data import backup_limit_tombstones
                backup_limit_tombstones(db)
                Setting.set('limits_tombstones_external_pending', '0')
                Setting.set('limits_tombstone_backup_error', '')
                db.session.commit()
            except Exception as exc:
                db.session.rollback()
                app.logger.exception('limits tombstone sidecar backup failed')
                Setting.set('limits_tombstone_backup_error', str(exc))
                db.session.commit()
            result = _durable_backup('очистка окна 10 версий лимитов',
                                     allowed_removed_tables=allowed)
            if result and result.get('promoted'):
                Setting.set('limits_rotation_needs_backup', '0')
                db.session.commit()
    threading.Thread(target=work, name='limits-rotation-backup', daemon=True).start()

@app.route('/limit-notifications/max-journal/<int:outbound_id>/<action>', methods=['POST'])
@admin_required
def limit_max_outbound_action(outbound_id, action):
    from limits_bot import LimitsBotError, retry_bot_outbound, cancel_bot_outbound
    try:
        if action == 'retry':
            retry_bot_outbound(outbound_id, confirm_unknown=request.form.get('confirm_unknown') == '1')
        elif action == 'cancel':
            cancel_bot_outbound(outbound_id, 'Отменено администратором веб-системы.')
        else:
            abort(404)
        log_event('limits_bot_outbox_action', {'outbound_id': outbound_id, 'action': action})
        flash('Состояние задания обновлено.', 'success')
    except LimitsBotError as exc:
        flash(str(exc), 'error')
    return redirect(url_for('limit_max_journal'))

@app.route('/limit-notifications/versions/<int:version_id>/source.xlsx')
@login_required
def limit_version_source(version_id):
    version = _limit_version_or_404(version_id)
    root = os.path.realpath(os.path.join(UPLOADS, 'limits')) + os.sep
    path = os.path.realpath(version.source_path)
    if not path.startswith(root) or not os.path.isfile(path):
        abort(404)
    return send_file(path, as_attachment=True, download_name=version.source_name,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@app.route('/limit-notifications/versions/<int:version_id>/accept', methods=['POST'])
@admin_required
def limit_version_accept(version_id):
    from limits_service import LimitsError, accept_version
    if request.form.get('confirm') != '1':
        return redirect(url_for('limit_version_review', version_id=version_id))
    version = _limit_version_or_404(version_id)
    try:
        accept_version(version.id, current_user.id)
        # The acceptance may have rotated an old version.  Do not let the
        # generic safety-net write a shrink-rejecting snapshot instead.
        g.durable_backup_done = True
        _schedule_limits_rotation_backup()
        log_event('limits_version_accepted', {'version_id': version.id})
        flash('Версия принята. Уведомления поставлены в долговечную очередь.', 'success')
    except LimitsError as exc:
        flash(str(exc), 'error')
    return redirect(url_for('limit_notifications', dataset=version.dataset_id))

@app.route('/limit-notifications/versions/<int:version_id>/reject', methods=['POST'])
@admin_required
def limit_version_reject(version_id):
    version = _limit_version_or_404(version_id)
    if version.status not in ('review', 'error'):
        abort(409, description='Отклонить можно только проверенную или ошибочную загрузку.')
    version.status = 'rejected'
    db.session.commit()
    log_event('limits_version_rejected', {'version_id': version_id,
                                        'reason': request.form.get('reason', '')})
    flash('Загрузка отклонена. Временные данные будут удалены по техническому TTL.', 'success')
    return redirect(url_for('limit_notifications', dataset=version.dataset_id))

@app.route('/limit-notifications/datasets/<int:dataset_id>/subscribers')
@admin_required
def limit_subscribers(dataset_id):
    dataset = _limit_dataset_or_404(dataset_id)
    search = request.args.get('q', '').strip()
    query = LimitSubscription.query.filter_by(dataset_id=dataset.id)
    if search:
        query = query.join(LimitRecipient).filter(or_(
            LimitSubscription.uin.ilike(f'%{search}%'),
            LimitRecipient.organization_inn.ilike(f'%{search}%'),
            LimitRecipient.surname_initials.ilike(f'%{search}%')))
    requests = LimitModerationRequest.query.filter_by(dataset_id=dataset.id).order_by(
        LimitModerationRequest.created_at.desc()).limit(200).all()
    return render_template('limits_subscribers.html', dataset=dataset,
                           subscriptions=query.order_by(LimitSubscription.created_at.desc()).limit(500).all(),
                           moderation_requests=requests, search=search)

@app.route('/limit-notifications/objects/<int:object_id>/message', methods=['GET', 'POST'])
@admin_required
def limit_manual_message(object_id):
    from limits_service import LimitsError, manual_delivery_preview, queue_manual_object_message
    signer = URLSafeTimedSerializer(app.secret_key, salt='limits-manual-preview')
    try:
        if request.method == 'POST' and request.form.get('confirm') == '1':
            try:
                preview = signer.loads(request.form.get('preview_token', ''), max_age=1800)
            except (BadSignature, SignatureExpired):
                raise LimitsError('Предпросмотр истёк или недействителен. Проверьте адресатов заново.')
            if preview.get('object_id') != object_id or preview.get('actor_id') != current_user.id:
                abort(403)
            content = request.form.get('content', '')
            if content != preview.get('content'):
                raise LimitsError('Текст изменён после предпросмотра. Проверьте сообщение заново.')
            count = queue_manual_object_message(
                object_id, content, expected_version_id=preview['version_id'],
                expected_recipient_ids=preview['recipient_ids'],
                campaign_key=f"{current_user.id}:{preview['campaign_key']}")
            log_event('limits_manual_message_queued', {
                'object_id': object_id, 'version_id': preview['version_id'],
                'campaign_key': f"manual:{current_user.id}:{preview['campaign_key']}",
                'recipient_count': count,
            })
            flash(f'В очередь добавлено адресатов: {count}.', 'success')
            return redirect(url_for('limit_object_card', object_id=object_id))
        obj, version, recipients, content = manual_delivery_preview(
            object_id, request.values.get('content', ''))
        preview_token = signer.dumps({
            'object_id': object_id, 'actor_id': current_user.id,
            'version_id': version.id,
            'recipient_ids': [r.id for r in recipients], 'content': content,
            'campaign_key': __import__('uuid').uuid4().hex,
        })
        return render_template('limits_message_preview.html', object=obj, version=version,
                               recipients=recipients, content=content, preview_token=preview_token)
    except LimitsError as exc:
        flash(str(exc), 'error')
        return redirect(url_for('limit_object_card', object_id=object_id))

@app.route('/limit-notifications/max/webhook', methods=['GET', 'POST'])
@csrf.exempt
def limit_max_webhook():
    """Fast webhook boundary: authenticate + durable dedup before acknowledging."""
    expected = decrypt_value(Setting.get('limits_max_webhook_secret', ''))
    received = request.headers.get('X-Max-Bot-Api-Secret', '')
    if not expected or not received or not __import__('hmac').compare_digest(expected, received):
        return jsonify(error='Недействительный секрет webhook'), 403
    if request.method == 'GET':
        return jsonify(ready=True), 200
    if request.content_length and request.content_length > 256 * 1024:
        return jsonify(error='Событие превышает допустимый размер'), 413
    payload = request.get_json(silent=True)
    if not isinstance(payload, dict):
        return jsonify(error='Ожидается JSON-событие MAX'), 400
    try:
        from limits_bot import ingest_event
        event, created = ingest_event(payload)
    except ValueError:
        db.session.rollback()
        return jsonify(error='Неверная структура события MAX'), 400
    except Exception:
        db.session.rollback()
        app.logger.error('limits MAX event persistence failed')
        return jsonify(error='Временная ошибка сохранения события'), 503
    # This endpoint deliberately does not send MAX messages or do group lookups.
    # A durable event is sufficient for provider retry semantics.
    g.skip_durable_backup = True
    return jsonify(accepted=True, duplicate=not created, event_id=event.id), 200

@app.route('/limit-notifications/max-settings', methods=['POST'])
@admin_required
def limit_max_settings():
    """Store connection material encrypted; never return token/secret to browser."""
    from limits_connection import configure_connection, DEFAULT_WEBHOOK_URL
    from limits_max import MaxAPIError
    token = request.form.get('token', '').strip()
    group_id = request.form.get('group_chat_id', '').strip()
    dataset_id = request.form.get('bot_dataset_id', type=int)
    enabled = request.form.get('enabled') == '1'
    if group_id and not re.fullmatch(r'-?\d+', group_id):
        abort(400, description='ID группы MAX должен быть целым числом.')
    if dataset_id:
        _limit_dataset_or_404(dataset_id)
    if enabled and (not (token or Setting.get('limits_max_token'))
                    or not group_id or not dataset_id):
        flash('Для включения нужны токен MAX, ID админ-группы и выбранный набор данных.', 'error')
        return redirect(url_for('limit_settings', dataset=request.form.get('context_dataset', type=int)))
    try:
        configure_connection(token=token, group_id=group_id, dataset_id=dataset_id,
            enabled=enabled, encrypt=encrypt_value, decrypt=decrypt_value,
            adapter_factory=_limits_adapter,
            webhook_url=app.config.get('LIMITS_MAX_WEBHOOK_URL',
                                      os.environ.get('LIMITS_MAX_WEBHOOK_URL', DEFAULT_WEBHOOK_URL)))
    except MaxAPIError:
        flash('Настройки сохранены, но MAX не подтвердил подключение. Отправки выключены.', 'error')
        return redirect(url_for('limit_settings', dataset=request.form.get('context_dataset', type=int)))
    log_event('limits_max_settings_changed', {'token_changed': bool(token),
                                              'group_configured': bool(group_id)})
    flash('Настройки сохранены. ' + (
        'MAX подтвердил регистрацию webhook. Бот включён.'
        if enabled else 'Отправки MAX выключены.'), 'success')
    return redirect(url_for('limit_settings', dataset=request.form.get('context_dataset', type=int)))

@app.route('/limit-notifications/moderation/<int:request_id>/decide', methods=['POST'])
@admin_required
def limit_moderation_decide(request_id):
    from limits_bot import decide_web_moderation
    decision = request.form.get('decision')
    if decision not in ('approved', 'rejected'):
        abort(400)
    try:
        decide_web_moderation(request_id, decision, current_user.id,
                              request.form.get('reason', ''))
        flash('Решение по заявке сохранено.', 'success')
    except ValueError as exc:
        db.session.rollback()
        flash(str(exc), 'error')
    return redirect(url_for('limit_notifications'))

@app.route('/limit-notifications/objects/<int:object_id>')
@login_required
def limit_object_card(object_id):
    obj = db.session.get(LimitObject, object_id) or abort(404)
    from limits_views import object_context
    selected_id = request.args.get('version_id', type=int) or request.args.get('version', type=int)
    if selected_id:
        selected = _limit_version_or_404(selected_id)
        if selected.dataset_id != obj.dataset_id or selected.status != 'accepted':
            abort(404)
    context = object_context(obj, selected_id)
    if not current_user.is_admin:
        # Do not fetch hidden personal details merely to hide them with CSS.
        context['messages'] = []
    return render_template('limits_object.html', **context)

def _limit_version_or_404(version_id):
    version = db.session.get(LimitVersion, version_id) or abort(404)
    if version.status == 'purged':
        abort(404)
    if not current_user.is_admin and version.status != 'accepted':
        abort(403)
    return version

@app.route('/limit-notifications/comparisons/<int:comparison_id>')
@login_required
def limit_comparison_view(comparison_id):
    comparison = db.session.get(LimitComparison, comparison_id) or abort(404)
    _limit_version_or_404(comparison.old_version_id)
    _limit_version_or_404(comparison.new_version_id)
    category = request.args.get('category', '').strip()
    query = LimitChange.query.filter_by(comparison_id=comparison.id)
    if category:
        query = query.filter_by(category=category)
    search = request.args.get('q', '').strip()
    if search:
        query = query.filter(LimitChange.uin.ilike(f'%{search}%'))
    total = query.count()
    pages = max(1, (total + 99)//100)
    page = max(1, min(request.args.get('page', 1, type=int), pages))
    return render_template('limits_comparison.html', comparison=comparison,
                           changes=query.order_by(LimitChange.uin, LimitChange.id).offset((page-1)*100).limit(100).all(),
                           issues=json.loads(comparison.issues_json or '[]'),
                           total=total, page=page, pages=pages, selected_category=category,
                           search=search,
                           categories=[x[0] for x in db.session.query(
                               LimitChange.category).filter_by(
                               comparison_id=comparison.id).distinct().all()])

@app.route('/limit-notifications/recipients/<int:recipient_id>/revoke-access', methods=['POST'])
@admin_required
def limit_recipient_revoke_access(recipient_id):
    recipient = db.session.get(LimitRecipient, recipient_id) or abort(404)
    from limits_bot import revoke_access
    revoke_access(recipient_id, current_user.id, request.form.get('reason', 'Отозвано администратором'))
    flash('Доступ получателя отозван. Очередь не отправит ему уже ожидающие сообщения.', 'success')
    return redirect(url_for('limit_notifications'))

@app.before_request
def _limits_history_gate():
    if not request.path.startswith('/limit-notifications'):
        return
    if request.endpoint in ('limit_max_webhook', 'limit_history_recovery'):
        return
    if not current_user.is_authenticated:
        return
    if not _limits_history_ready.is_set() or Setting.get('limits_history_ready') != 'ready':
        return render_template('limits_unavailable.html',
                               detail='Проверка восстановления и окна хранения ещё не завершена.'), 503

@app.route('/limit-notifications/comparisons/<int:comparison_id>/resolve', methods=['POST'])
@admin_required
def limit_comparison_resolve(comparison_id):
    from limits_service import LimitsError, resolve_comparison_issue
    try:
        resolve_comparison_issue(comparison_id, request.form.get('issue_key', ''),
                                 {'old_row': request.form.get('old_row', ''),
                                  'new_row': request.form.get('new_row', ''),
                                   'action': request.form.get('action', 'transfer'),
                                  'note': request.form.get('note', '')}, current_user.id)
        flash('Сопоставление сохранено и сравнение пересчитано.', 'success')
    except LimitsError as exc:
        flash(str(exc), 'error')
    return redirect(url_for('limit_comparison_view', comparison_id=comparison_id))

@app.template_filter('limit_value')
def _limit_value_filter(value):
    from limits_views import display_value
    return display_value(value)

@app.route('/limit-notifications/datasets/<int:dataset_id>/upload', methods=['POST'])
@admin_required
def limit_upload(dataset_id):
    from limits_service import LimitsError, stage_upload
    from datetime import date
    dataset = _limit_dataset_or_404(dataset_id)
    uploaded = request.files.get('file')
    try:
        source_date = date.fromisoformat(request.form.get('source_date', ''))
        if not uploaded or not uploaded.filename:
            raise LimitsError('Выберите исходный XLSX.')
        version, duplicate = stage_upload(
            dataset, uploaded, uploaded.filename, source_date, current_user.id, UPLOADS)
        if duplicate:
            flash('Этот файл уже есть в наборе; повторная версия и рассылка не созданы.', 'info')
        else:
            _limit_parse_in_background(version.id)
            flash('Файл принят в фоновую проверку.', 'success')
        log_event('limits_upload_staged', {'dataset_id': dataset.id, 'version_id': version.id,
                                           'duplicate': duplicate})
    except (LimitsError, ValueError) as exc:
        flash(str(exc), 'error')
    return redirect(url_for('limit_notifications', dataset=dataset.id))

@app.route('/limit-notifications/deliveries/<int:delivery_id>/cancel', methods=['POST'])
@admin_required
def limit_delivery_cancel(delivery_id):
    from limits_bot import LimitsBotError, cancel_delivery
    delivery = db.session.get(LimitDelivery, delivery_id) or abort(404)
    dataset_id = delivery.version.dataset_id
    try:
        cancel_delivery(delivery_id)
        log_event('limits_delivery_cancelled', {'delivery_id': delivery_id})
        flash('Задание отменено явно.', 'success')
    except LimitsBotError as exc:
        flash(str(exc), 'error')
    return redirect(url_for('limit_delivery_journal', dataset_id=dataset_id))

@app.route('/limit-notifications/comparisons/<int:comparison_id>/export.xlsx')
@login_required
def limit_comparison_export(comparison_id):
    from limits_service import export_comparison
    comparison = db.session.get(LimitComparison, comparison_id) or abort(404)
    _limit_version_or_404(comparison.old_version_id)
    _limit_version_or_404(comparison.new_version_id)
    return send_file(export_comparison(comparison_id), as_attachment=True,
                     download_name=f'limits-comparison-{comparison_id}.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@app.route('/limit-notifications/deliveries/<int:delivery_id>/retry', methods=['POST'])
@admin_required
def limit_delivery_retry(delivery_id):
    from limits_bot import LimitsBotError, retry_delivery
    delivery = db.session.get(LimitDelivery, delivery_id) or abort(404)
    dataset_id = delivery.version.dataset_id
    try:
        retry_delivery(delivery_id, confirm_unknown=request.form.get('confirm_unknown') == '1')
        log_event('limits_delivery_retry_requested', {'delivery_id': delivery_id})
        flash('Повтор разрешён и добавлен в очередь.', 'success')
    except LimitsBotError as exc:
        flash(str(exc), 'error')
    return redirect(url_for('limit_delivery_journal', dataset_id=dataset_id))

@app.route('/limit-notifications/max-events/<int:event_id>/retry', methods=['POST'])
@admin_required
def limit_max_event_retry(event_id):
    from limits_bot import LimitsBotError, retry_event
    try:
        retry_event(event_id, by_pk=True)
        log_event('limits_max_event_retried', {'event_id': event_id})
        flash('Событие поставлено на повторную обработку.', 'success')
    except LimitsBotError as exc:
        flash(str(exc), 'error')
    return redirect(url_for('limit_max_journal'))

@app.route('/limit-notifications/datasets/<int:dataset_id>/objects')
@login_required
def limit_objects(dataset_id):
    dataset = _limit_dataset_or_404(dataset_id)
    from limits_views import objects_context
    return render_template('limits_objects.html', **objects_context(dataset, request.args))

@app.route('/limit-notifications/max-journal')
@admin_required
def limit_max_journal():
    status = request.args.get('status', '').strip()
    search = request.args.get('q', '').strip()
    query = LimitBotOutbound.query
    events = LimitMaxEvent.query.filter(LimitMaxEvent.status.in_(
        ('received', 'processing', 'retry', 'error')))
    if status:
        query = query.filter_by(status=status)
    if search:
        query = query.filter(LimitBotOutbound.target_id.contains(search))
    page = max(1, request.args.get('page', 1, type=int) or 1)
    total = query.count()
    return render_template(
        'limits_max_journal.html',
        outbound=query.order_by(LimitBotOutbound.id.desc()).offset((page - 1) * 100).limit(100).all(),
        events=events.order_by(LimitMaxEvent.id.desc()).limit(100).all(),
        total=total, page=page, pages=max(1, (total + 99) // 100), status=status, search=search)

@app.route('/limit-notifications/versions/<int:version_id>/compare', methods=['POST'])
@admin_required
def limit_compare(version_id):
    from limits_service import LimitsError, create_comparison
    version_id = request.form.get('new_version_id', type=int) or version_id
    new = _limit_version_or_404(version_id)
    old_id = request.form.get('old_version_id', type=int) or new.previous_version_id
    try:
        comparison = create_comparison(old_id, new.id)
        log_event('limits_comparison_created', {'comparison_id': comparison.id})
        return redirect(url_for('limit_comparison_view', comparison_id=comparison.id))
    except LimitsError as exc:
        flash(str(exc), 'error')
        return redirect(url_for('limit_notifications', dataset=new.dataset_id))

@app.route('/limit-notifications/versions/<int:version_id>/review')
@admin_required
def limit_version_review(version_id):
    version = _limit_version_or_404(version_id)
    if version.status not in ('review', 'accepted'):
        abort(409)
    snapshot = json.loads(version.snapshot_json or '{}')
    rows = snapshot.get('rows', [])
    comparison = LimitComparison.query.filter_by(
        old_version_id=version.previous_version_id, new_version_id=version.id).first()
    return render_template('limits_version_review.html', version=version,
        dataset=version.dataset, row_count=len(rows), object_count=len({row['uin'] for row in rows}),
        comparison=comparison,
        summary=json.loads(comparison.summary_json or '{}') if comparison else {},
        warnings=json.loads(version.warnings_json or '[]'))

@app.route('/limit-notifications/settings')
@admin_required
def limit_settings():
    from limits_connection import DEFAULT_WEBHOOK_URL
    datasets = LimitDataset.query.order_by(LimitDataset.updated_at.desc()).all()
    dataset_id = request.args.get('dataset', type=int)
    dataset = _limit_dataset_or_404(dataset_id) if dataset_id else (datasets[0] if datasets else None)
    return render_template('limits_settings.html', datasets=datasets, dataset=dataset,
        max_connected=bool(Setting.get('limits_max_token', '')),
        max_connection_error=Setting.get('limits_max_connection_error', ''),
        max_webhook_url=app.config.get('LIMITS_MAX_WEBHOOK_URL',
                                      os.environ.get('LIMITS_MAX_WEBHOOK_URL', DEFAULT_WEBHOOK_URL)),
        max_enabled=Setting.get('limits_max_enabled') == '1',
        max_group_id=Setting.get('limits_max_group_chat_id', ''),
        bot_dataset_id=Setting.get('limits_bot_dataset_id', ''))

@app.route('/limit-notifications/versions/<int:version_id>/delete', methods=['GET', 'POST'])
@admin_required
def limit_version_delete(version_id):
    from limits_service import LimitsError, purge_version
    version = _limit_version_or_404(version_id)
    dataset = version.dataset
    dataset_id = dataset.id
    unfinished = LimitDelivery.query.filter(
        LimitDelivery.version_id == version.id,
        ~LimitDelivery.status.in_(('accepted', 'cancelled'))).count()
    blocked = ('Загрузка ещё обрабатывается. Дождитесь завершения проверки.'
               if version.status not in ('accepted', 'review', 'error', 'rejected') else
               'Есть незавершённые отправки. Завершите или явно отмените их в журнале уведомлений.'
               if unfinished else '')
    if request.method == 'POST':
        if request.form.get('confirm') != 'delete':
            flash('Подтвердите удаление загрузки на странице проверки.', 'error')
            return redirect(url_for('limit_version_delete', version_id=version_id))
        try:
            if blocked:
                raise LimitsError(blocked)
            purge_version(version_id)
            log_event('limits_version_deleted', {'version_id': version_id, 'dataset_id': dataset_id})
            _schedule_limits_rotation_backup()
            flash('Загрузка удалена. Исходный Excel и связанная история версии очищены; '
                  'если файл временно недоступен, очистка будет повторена автоматически.', 'success')
        except LimitsError as exc:
            db.session.rollback()
            flash(str(exc), 'error')
        return redirect(url_for('limit_notifications', dataset=dataset_id))
    comparison_count = LimitComparison.query.filter(
        (LimitComparison.old_version_id == version.id) |
        (LimitComparison.new_version_id == version.id)).count()
    fallback = None
    if dataset.current_version_id == version.id:
        fallback = LimitVersion.query.filter(
            LimitVersion.dataset_id == dataset.id, LimitVersion.status == 'accepted',
            LimitVersion.id != version.id).order_by(
                LimitVersion.accepted_at.desc(), LimitVersion.id.desc()).first()
    return render_template('limits_delete_version.html', dataset=dataset, version=version,
                           blocked=blocked, comparison_count=comparison_count, fallback=fallback)

@app.route('/limit-notifications/datasets/<int:dataset_id>/deliveries')
@admin_required
def limit_delivery_journal(dataset_id):
    dataset = _limit_dataset_or_404(dataset_id)
    status = request.args.get('status', '').strip()
    query = LimitDelivery.query.join(LimitVersion).filter(LimitVersion.dataset_id == dataset.id)
    if status:
        query = query.filter(LimitDelivery.status == status)
    search = request.args.get('q', '').strip()
    version_id = request.args.get('version_id', type=int) or request.args.get('version', type=int)
    if version_id:
        query = query.filter(LimitDelivery.version_id == version_id)
    if search:
        query = query.join(LimitObject).join(LimitRecipient).filter(or_(
            LimitObject.uin.ilike(f'%{search}%'),
            LimitRecipient.surname_initials.ilike(f'%{search}%'),
            LimitRecipient.max_user_id.ilike(f'%{search}%')))
    total = query.count()
    pages = max(1, (total + 99)//100)
    page = max(1, min(request.args.get('page', 1, type=int), pages))
    return render_template('limits_delivery_journal.html', dataset=dataset,
                           deliveries=query.order_by(LimitDelivery.created_at.desc()).offset((page-1)*100).limit(100).all(),
                           selected_status=status, search=search, selected_version=version_id,
                           total=total, page=page, pages=pages,
                           versions=LimitVersion.query.filter_by(dataset_id=dataset_id, status='accepted').all())
