# -*- coding: utf-8 -*-
import os
from pathlib import Path
import tempfile
import unittest
from decimal import Decimal

from openpyxl import Workbook, load_workbook

from limits_parser import (
    LimitsParserError, compare_snapshots, export_comparison, parse_xlsx,
)


ASSETS = Path(__file__).resolve().parents[1] / "attached_assets"
FORM_1 = ASSETS / "Форма_отчета_по_всем_утвержденным_объектам_-_Минэнерго_(1)_1789372306640.xlsx"
FORM_2 = ASSETS / "Форма_отчета_по_всем_утвержденным_объектам_-_Минэнерго_(2)_1789372306644.xlsx"


class LimitsParserRealWorkbookTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.old = parse_xlsx(FORM_1)
        cls.progress_events = []
        cls.new = parse_xlsx(FORM_2, progress=lambda *event: cls.progress_events.append(event))
        cls.comparison = compare_snapshots(cls.old, cls.new)

    def test_real_workbooks_have_semantic_schema_and_expected_uins(self):
        row_events = [event for event in self.progress_events if event[0] == 'rows']
        self.assertEqual(0, row_events[0][1])
        self.assertEqual(row_events[-1][2], row_events[-1][1])
        self.assertEqual(list(range(row_events[-1][2] + 1)), [event[1] for event in row_events])
        self.assertEqual(68, len(self.old["columns"]))
        self.assertEqual(2545, len(self.old["rows"]))
        self.assertEqual(2563, len(self.new["rows"]))
        self.assertEqual(2516, len({row["uin"] for row in self.old["rows"]}))
        self.assertEqual(2535, len({row["uin"] for row in self.new["rows"]}))
        self.assertEqual("annual_funding_2026_moscow_oblast_budget",
                         next(column["key"] for column in self.old["columns"]
                              if column["year"] == 2026 and
                              column["source"] == "moscow_oblast_budget" and
                              column["category"] == "annual_funding"))

    def test_real_difference_counts_and_examples_are_exact_decimals(self):
        summary = self.comparison["summary"]
        self.assertEqual(20, summary["new_objects"])
        self.assertEqual(1, summary["absent_objects"])
        self.assertEqual(3, summary["transfers"])
        self.assertEqual(893, summary["stable_changed_cells"])
        self.assertTrue(any(change["uin"] == "1000001491.1000001400"
                            and change["kind"] == "object_absent"
                            for change in self.comparison["changes"]))
        changed = [change for change in self.comparison["changes"]
                   if change["uin"] == "1000000938.1000000625"
                   and change.get("field") == "annual_funding_2026_moscow_oblast_budget"]
        self.assertTrue(any(change["old"]["value"] == "216209.81" and
                            change["new"]["value"] == "82746.07"
                            for change in changed))
        cost_changes = [change for change in self.comparison["changes"]
                        if change["uin"] == "1000000856.1000000539"]
        self.assertTrue(any(change.get("field") == "limit_cost" and
                            change["old"]["value"] == "25022.38" and
                            change["new"]["value"] == "75396.24"
                            for change in cost_changes))
        self.assertTrue(any(change.get("field") ==
                            "annual_funding_2026_municipal_budgets_moscow_oblast" and
                            change["old"]["value"] == "9408.42" and
                            change["new"]["value"] == "28348.99"
                            for change in cost_changes))
        dates = [change for change in self.comparison["changes"]
                 if change["uin"] == "1000000850.1000000533"
                 and change.get("field") == "work_completion_date"]
        self.assertTrue(any(change["old"]["value"] == "2025-11-29" and
                            change["new"]["value"] == "2026-11-29" and
                            change["category"] == "completion_year"
                            for change in dates))

    def test_export_has_exact_sheets_and_no_formula_cells(self):
        data = export_comparison(self.comparison)
        workbook = load_workbook(data, data_only=False)
        self.assertEqual([
            "Сводка", "Новые объекты", "Отсутствующие объекты",
            "Финансирование", "Сроки", "Реквизиты и коды", "Проблемы данных",
        ], workbook.sheetnames)
        self.assertTrue(all(cell.data_type != "f"
                            for sheet in workbook.worksheets
                            for row in sheet.iter_rows() for cell in row))


class LimitsComparisonSafetyTests(unittest.TestCase):
    def _snapshot(self, rows):
        columns = []
        for key, category in [
            ("uin", "identifier"), ("list_type", "programme"),
            ("ap_status", "status"), ("gp_code", "programme"),
            ("pp_code", "programme"), ("om_code", "programme"),
            ("dm_code", "programme"), ("limit_cost", "limit_cost"),
        ]:
            columns.append({"key": key, "label": key, "category": category,
                            "year": None, "source": None})
        return {"columns": columns, "rows": rows, "warnings": [], "metadata": {}}

    def _row(self, row, uin, om, money="0", name="name"):
        money = format(Decimal(money).normalize(), "f")
        if Decimal(money) == 0:
            money = "0"
        values = {
            "uin": (uin, "text"), "list_type": ("list", "text"),
            "ap_status": ("1", "text"), "gp_code": ("10", "text"),
            "pp_code": ("1", "text"), "om_code": (om, "text"),
            "dm_code": ("01", "text"), "limit_cost": (money, "money"),
        }
        return {
            "uin": uin, "row": row,
            "cells": {key: {"raw": value, "value": value, "kind": kind,
                            "cell": f"A{row}"}
                      for key, (value, kind) in values.items()},
        }

    def test_ambiguous_repeated_uin_requires_explicit_mapping(self):
        old = self._snapshot([self._row(5, "u", "01"), self._row(6, "u", "02")])
        new = self._snapshot([self._row(5, "u", "03"), self._row(6, "u", "04")])
        result = compare_snapshots(old, new)
        self.assertEqual(["u"], result["summary"]["blocked_uins"])
        self.assertEqual("ambiguous_row_match", result["issues"][0]["code"])
        resolved = compare_snapshots(old, new, {"u@5": "u@5", "u@6": "u@6"})
        self.assertFalse(resolved["issues"])
        self.assertEqual(2, resolved["summary"]["transfers"])

    def test_resolution_is_bijective_and_can_confirm_additions_removals(self):
        old = self._snapshot([self._row(5, "u", "01"), self._row(6, "u", "02")])
        new = self._snapshot([self._row(5, "u", "03"), self._row(6, "u", "04")])
        with self.assertRaisesRegex(LimitsParserError, "Multiple old rows"):
            compare_snapshots(old, new, {"u@5": "u@5", "u@6": "u@5"})
        with self.assertRaisesRegex(LimitsParserError, "unknown row"):
            compare_snapshots(old, new, {"u@99": "u@5"})
        independent = compare_snapshots(old, new, {"transfers": [
            {"old_row": "u@5", "action": "absent"},
            {"old_row": "u@6", "action": "absent"},
            {"new_row": "u@5", "action": "added"},
            {"new_row": "u@6", "action": "added"},
        ]})
        self.assertEqual(0, independent["summary"]["transfers"])
        self.assertEqual(2, independent["summary"]["new_rows"])
        self.assertEqual(2, independent["summary"]["absent_rows"])

    def test_blank_dash_zero_and_money_normalisation_are_not_conflated(self):
        old = self._snapshot([self._row(5, "u", "01", "0")])
        new = self._snapshot([self._row(5, "u", "01", "0.00")])
        self.assertEqual(0, compare_snapshots(old, new)["summary"]["changed_cells"])
        new["rows"][0]["cells"]["limit_cost"].update({"raw": "-", "value": None, "kind": "dash"})
        self.assertEqual(1, compare_snapshots(old, new)["summary"]["changed_cells"])

    def test_export_escapes_formula_like_raw_values(self):
        old = self._snapshot([])
        new = self._snapshot([self._row(5, "u", "01", "1")])
        result = compare_snapshots(old, new)
        result["changes"][0]["programme"]["gp_code"] = "=HYPERLINK(\"bad\")"
        data = export_comparison(result)
        book = load_workbook(data, data_only=False)
        self.assertFalse(any(cell.data_type == "f" for sheet in book.worksheets
                             for row in sheet.iter_rows() for cell in row))

    def test_export_includes_row_decisions_and_source_warnings(self):
        old = self._snapshot([self._row(5, "u", "01"), self._row(6, "u", "02")])
        new = self._snapshot([self._row(5, "u", "03"), self._row(6, "u", "04")])
        result = compare_snapshots(old, new, {"transfers": [
            {"old_row": "u@5", "action": "absent"},
            {"old_row": "u@6", "action": "absent"},
            {"new_row": "u@5", "action": "added"},
            {"new_row": "u@6", "action": "added"},
        ]})
        old["warnings"] = [{"code": "source_problem", "message": "проверить", "cell": "A1"}]
        book = load_workbook(export_comparison(result, old, new), data_only=False)
        actions = [row[2] for row in book["Реквизиты и коды"].iter_rows(
            min_row=2, values_only=True)]
        self.assertIn("row_added", actions)
        self.assertIn("row_absent", actions)
        problem_codes = [row[0] for row in book["Проблемы данных"].iter_rows(
            min_row=2, values_only=True)]
        self.assertIn("source_warning_source_problem", problem_codes)

    def test_formula_without_cache_is_rejected(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet["A1"] = "=1+1"
        handle = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        handle.close()
        workbook.save(handle.name)
        self.addCleanup(lambda: os.path.exists(handle.name) and os.unlink(handle.name))
        with self.assertRaisesRegex(LimitsParserError, "Formula without"):
            parse_xlsx(handle.name)


if __name__ == "__main__":
    unittest.main()