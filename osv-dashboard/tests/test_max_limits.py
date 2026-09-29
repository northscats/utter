"""MAX limits bot contracts; no real network requests."""
import json
import os
import tempfile
import unittest
from datetime import date, timedelta
from unittest.mock import patch

from flask import Flask
from sqlalchemy import inspect, text

from database import (
    LimitChange, LimitComparison, LimitDataset, LimitObject, LimitVersion,
    MaxLimitsApplication, MaxLimitsCursor, MaxLimitsOutbox, MaxLimitsProfile,
    MaxLimitsSubscription, MaxLimitsUin, Setting, User, db, utcnow,
)
from durable_data import COUNT_TABLES, RETIRED_LIMIT_TABLES
import max_limits
import max_registration
import max_digest
import max_bot_config
from max_dataset_scope import selected_dataset_id


class MaxLimitsTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='max-limits-')
        self.config_patch = patch.object(
            max_bot_config, 'CONFIG',
            __import__('pathlib').Path(self.temp.name) / 'private' / 'max-bot.json')
        self.config_patch.start()
        max_bot_config.save('A' * 40, '-345')
        max_bot_config.set_enabled(True)
        self.app = Flask(__name__)
        self.app.config.update(
            TESTING=True,
            SQLALCHEMY_DATABASE_URI='sqlite:///' + os.path.join(self.temp.name, 'bot.db'),
            SQLALCHEMY_TRACK_MODIFICATIONS=False,
        )
        db.init_app(self.app)
        self.context = self.app.app_context()
        self.context.push()
        db.create_all()
        user = User(username='max-test', password_hash='unused')
        db.session.add(user)
        db.session.flush()
        self.dataset = LimitDataset(name='Accepted', scope='', owner_id=user.id)
        db.session.add(self.dataset)
        db.session.flush()
        self.version = LimitVersion(
            dataset_id=self.dataset.id, uploaded_by=user.id,
            source_date=date(2025, 1, 1), status='accepted',
            source_name='source.xlsx', source_path='', sha256='a' * 64,
            snapshot_json=json.dumps({'rows': [{'uin': 'UIN-123', 'cells': {
                'object_name': {'value': 'Объект'}, 'address': {'value': 'Адрес'}}}]}),
        )
        db.session.add(self.version)
        db.session.flush()
        self.dataset.current_version_id = self.version.id
        self.obj = LimitObject(
            dataset_id=self.dataset.id, uin='UIN-123', title='Объект',
            address='Адрес', last_seen_version_id=self.version.id, is_present=True)
        db.session.add(self.obj)
        Setting.set('max_limits_version_id', str(self.version.id))
        db.session.commit()

    def tearDown(self):
        db.session.remove()
        db.drop_all()
        self.context.pop()
        self.config_patch.stop()
        self.temp.cleanup()

    @staticmethod
    def message(content, user_id=101, chat_type='dialog'):
        return {'update_type': 'message_created', 'message': {
            'sender': {'user_id': user_id, 'is_bot': False},
            'recipient': {'chat_type': chat_type},
            'body': {'text': content},
        }}

    def test_private_dialog_requires_registration_and_approval(self):
        self.assertIn('ФИО', max_limits.process_update(
            {'update_type': 'bot_started', 'user': {'user_id': 101}})[2]['text'])
        self.assertIn('номер', max_limits.process_update(
            self.message('Иванов Иван Иванович'))[2]['text'])
        self.assertTrue(max_limits.process_update(
            self.message('+79201234567'))[2]['file'])
        self.assertIn('Добавлен', max_limits.process_update(
            self.message('UIN-123'))[2]['text'])
        self.assertEqual(max_registration.approved_uins(101), [])
        self.assertNotIn('Объект', max_limits.process_update(
            self.message('/uin'))[2]['text'])
        self.assertIsNone(max_limits.process_update(
            self.message('UIN-123', chat_type='chat')))
        self.assertIsNone(max_limits.process_update({
            'update_type': 'message_created', 'message': {
                'sender': {'user_id': 101, 'is_bot': True},
                'recipient': {'chat_type': 'dialog'},
                'body': {'text': 'UIN-123'},
            }}))
        self.assertEqual(MaxLimitsUin.query.filter_by(user_id=101).one().status, 'draft')

    def test_information_shows_only_money_from_accepted_uin(self):
        self.version.snapshot_json = json.dumps({
            'columns': [
                {'key': 'funding', 'label': 'Финансирование, тыс. руб.',
                 'category': 'annual_funding', 'year': 2026},
                {'key': 'secret', 'label': 'Внутренняя заметка', 'category': 'requisite'},
            ],
            'rows': [
                {'uin': 'UIN-123', 'row': 8, 'cells': {
                    'funding': {'kind': 'money', 'value': '100'},
                    'secret': {'kind': 'text', 'value': 'не показывать'},
                }},
                {'uin': 'OTHER', 'row': 9, 'cells': {
                    'funding': {'kind': 'money', 'value': '999'},
                }},
            ],
        })
        db.session.commit()
        result = max_limits.current_information('UIN-123')
        self.assertIn('— 100', result)
        self.assertNotIn('999', result)
        self.assertNotIn('не показывать', result)

    def test_funding_year_filter_applies_to_digest_excel_browser_and_queued_messages(self):
        from io import BytesIO
        from openpyxl import load_workbook
        from max_object_browser import _financial_entries, _change_pages
        from max_funding_year import selection

        db.session.add_all([
            MaxLimitsUin(user_id=201, uin='UIN-123', status='approved'),
            MaxLimitsUin(user_id=201, uin='ONLY-OLD', status='approved'),
            LimitObject(dataset_id=self.dataset.id, uin='ONLY-OLD', title='Старый',
                        last_seen_version_id=self.version.id, is_present=True),
        ])
        new = LimitVersion(
            dataset_id=self.dataset.id, uploaded_by=self.version.uploaded_by,
            source_date=date(2026, 2, 1), status='accepted',
            source_name='new.xlsx', source_path='', sha256='f' * 64,
            previous_version_id=self.version.id)
        new.snapshot_json = json.dumps({
            'columns': [
                {'key': 'annual_funding_2024_budget', 'category': 'annual_funding',
                 'year': 2024, 'label': 'Финансирование 2024'},
                {'key': 'annual_funding_2026_budget', 'category': 'annual_funding',
                 'year': 2026, 'label': 'Финансирование 2026'},
            ],
            'rows': [{'uin': 'ONLY-OLD', 'cells': {}}, {'uin': 'UIN-123', 'cells': {
                'annual_funding_2024_budget': {'kind': 'money', 'value': '24'},
                'annual_funding_2026_budget': {'kind': 'money', 'value': '26'},
            }}],
        })
        db.session.add(new)
        db.session.flush()
        self.dataset.current_version_id = new.id
        Setting.set('max_limits_version_id', str(new.id))
        comparison = LimitComparison(
            old_version_id=self.version.id, new_version_id=new.id, status='ready')
        db.session.add(comparison)
        db.session.flush()
        for uin, year in [('UIN-123', 2024), ('UIN-123', 2025),
                          ('UIN-123', 2026), ('ONLY-OLD', 2024)]:
            db.session.add(LimitChange(
                comparison_id=comparison.id, uin=uin, category='annual_funding',
                field_key=f'annual_funding_{year}_budget',
                field_label=f'Финансирование {year}',
                source_json=json.dumps({'year': year}),
                old_json='10', new_json='20'))
        db.session.commit()
        self.assertEqual(selection(), '2026')
        self.assertEqual(max_digest.changed_uins_for_user(201), {'UIN-123'})
        self.assertEqual(max_digest.prepare_digest(), (1, 1))
        item = MaxLimitsOutbox.query.filter_by(user_id=201).one()
        self.assertNotIn('ONLY-OLD', max_digest.message_for_outbox(item))
        rows = list(load_workbook(BytesIO(max_digest.changes_excel_bytes(201)),
                                  read_only=True).active.values)
        self.assertEqual([row[3] for row in rows[1:]], ['Финансирование 2026'])
        self.assertEqual(len(_financial_entries(new, 'UIN-123')), 1)
        self.assertIn('2026', str(_change_pages(self.dataset.id, 'UIN-123')))
        self.assertNotIn('2024', str(_change_pages(self.dataset.id, 'UIN-123')))
        self.assertNotIn('2024', max_limits.current_information('UIN-123'))
        self.assertNotIn('2025', max_limits.accepted_changes('UIN-123'))
        # Delivery must re-check queued content if the administrator tightens access.
        db.session.add(MaxLimitsOutbox(
            user_id=201, uin='ONLY-OLD', version_id=new.id, text='Финансирование 2024'))
        db.session.commit()
        with patch.dict(os.environ, {'MAX_LIMITS_BOT_TOKEN': 'test-only'}):
            calls = []
            max_limits.drain_outbox(api_call=lambda *args, **kwargs:
                                    calls.append(args) or {'success': True})
        self.assertEqual(len(calls), 1)
        self.assertNotIn('2024', str(calls))
        Setting.set('max_limits_min_funding_year', 'all')
        db.session.commit()
        self.assertEqual(max_digest.changed_uins_for_user(201), {'UIN-123', 'ONLY-OLD'})
        rows = list(load_workbook(BytesIO(max_digest.changes_excel_bytes(201)),
                                  read_only=True).active.values)
        self.assertEqual(len(rows), 5)
        self.assertIn('2024', str(rows))

    def test_stop_cancels_pending_delivery(self):
        db.session.add(MaxLimitsUin(user_id=101, uin='UIN-123', status='approved'))
        db.session.commit()
        db.session.add(MaxLimitsOutbox(
            user_id=101, uin='UIN-123', version_id=self.version.id, text='queued'))
        db.session.commit()
        max_limits.process_update(self.message('/stop'))
        self.assertEqual(MaxLimitsUin.query.filter_by(user_id=101).one().status, 'inactive')
        self.assertIsNotNone(MaxLimitsOutbox.query.one().cancelled_at)
        max_limits.process_update(
            {'update_type': 'bot_stopped', 'user': {'user_id': 101}})

    def test_change_queue_matches_uin_and_retries(self):
        db.session.add_all([
            MaxLimitsUin(user_id=201, uin='OTHER', status='approved'),
            MaxLimitsUin(user_id=202, uin='UIN-123', status='approved'),
        ])
        new = LimitVersion(
            dataset_id=self.dataset.id, uploaded_by=self.version.uploaded_by,
            source_date=date(2025, 2, 1), status='accepted',
            source_name='new.xlsx', source_path='', sha256='b' * 64)
        db.session.add(new)
        db.session.flush()
        comparison = LimitComparison(
            old_version_id=self.version.id, new_version_id=new.id, status='ready')
        db.session.add(comparison)
        db.session.flush()
        db.session.add(LimitChange(
            comparison_id=comparison.id, uin='UIN-123', category='field',
            field_key='cost', field_label='Стоимость',
            old_json=json.dumps(10), new_json=json.dumps(12)))
        db.session.flush()
        new.previous_version_id = self.version.id
        new.snapshot_json = self.version.snapshot_json
        Setting.set('max_limits_version_id', str(new.id))
        self.assertEqual(max_limits.queue_accepted_changes(new, comparison), 1)
        db.session.commit()
        item = MaxLimitsOutbox.query.one()
        self.assertEqual(item.user_id, 202)
        calls = []
        with patch.dict(os.environ, {'MAX_LIMITS_BOT_TOKEN': 'test-only'}):
            self.assertEqual(max_limits.drain_outbox(
                api_call=lambda *args, **kwargs: calls.append((args, kwargs))), 1)
        self.assertEqual(calls[0][0][1:3], ('POST', '/messages'))
        self.assertEqual(calls[0][0][3], {'user_id': 202})
        self.assertIsNotNone(item.delivered_at)

    def test_manual_digest_queues_only_changed_approved_uins_and_rechecks_access(self):
        db.session.add_all([
            MaxLimitsUin(user_id=201, uin='UIN-123', status='approved'),
            MaxLimitsUin(user_id=201, uin='UNCHANGED', status='approved'),
            MaxLimitsUin(user_id=202, uin='UIN-123', status='pending'),
            MaxLimitsUin(user_id=203, uin='UIN-123', status='approved'),
        ])
        new = LimitVersion(
            dataset_id=self.dataset.id, uploaded_by=self.version.uploaded_by,
            source_date=date(2025, 2, 1), status='accepted',
            source_name='new.xlsx', source_path='', sha256='c' * 64,
            previous_version_id=self.version.id)
        db.session.add(new)
        db.session.flush()
        self.dataset.current_version_id = new.id
        new.snapshot_json = self.version.snapshot_json
        Setting.set('max_limits_version_id', str(new.id))
        comparison = LimitComparison(
            old_version_id=self.version.id, new_version_id=new.id, status='ready')
        db.session.add(comparison)
        db.session.flush()
        db.session.add_all([
            LimitChange(
                comparison_id=comparison.id, uin='UIN-123', category='field',
                field_key='year', field_label='Год реализации',
                old_json='2024', new_json='2025'),
            LimitChange(
                comparison_id=comparison.id, uin='UIN-123', category='field',
                field_key='funding', field_label='Объем финансирования',
                old_json='10', new_json='20'),
            LimitChange(
                comparison_id=comparison.id, uin='NOT_SUBSCRIBED', category='field',
                field_key='year', field_label='Год реализации',
                old_json='2024', new_json='2025'),
        ])
        db.session.commit()
        self.assertEqual(max_digest.prepare_digest('a' * 32), (2, 2))
        self.assertEqual(max_digest.prepare_digest('a' * 32), (0, 0))
        outbox = MaxLimitsOutbox.query.order_by(MaxLimitsOutbox.user_id).all()
        self.assertEqual([item.user_id for item in outbox], [201, 203])
        message = max_digest.message_for_outbox(outbox[0])
        self.assertIn('Всего ваших объектов — 1', message)
        self.assertIn('Объектов с изменившимися показателями — 1', message)
        self.assertIn('1. Объект\nУИН UIN-123', message)
        self.assertNotIn('финансирования', message)
        self.assertNotIn('UNCHANGED', message)
        self.assertNotIn('NOT_SUBSCRIBED', message)
        from openpyxl import load_workbook
        from io import BytesIO
        sheet = load_workbook(BytesIO(max_digest.changes_excel_bytes(201)),
                              read_only=True).active
        rows = list(sheet.values)
        self.assertEqual(len(rows), 3)
        self.assertEqual({row[0] for row in rows[1:]}, {'UIN-123'})
        self.assertEqual({row[3] for row in rows[1:]},
                         {'Год реализации', 'Финансирование'})
        self.assertEqual(max_registration.handle_callback(
            201, 'menu:changes-excel')['file'], 'changes')
        self.assertNotIn('file', max_registration.handle_callback(
            202, 'menu:changes-excel'))
        sent = []
        def fake_upload(_token, _api_call, data, filename):
            self.assertEqual(filename, 'limits_changes.xlsx')
            exported = list(load_workbook(BytesIO(data), read_only=True).active.values)
            self.assertEqual({row[0] for row in exported[1:]}, {'UIN-123'})
            return 'file-token'
        with patch.object(max_registration, 'upload_catalog', side_effect=fake_upload):
            max_limits.dispatch_update('A' * 40, {
                'update_type': 'message_callback',
                'callback': {'user': {'user_id': 201}, 'callback_id': 'cb-file',
                             'payload': 'menu:changes-excel'},
                'message': {'recipient': {'chat_type': 'dialog', 'user_id': 201}},
            }, api_call=lambda *args, **kwargs: sent.append((args, kwargs)) or {'success': True})
        self.assertEqual(sent[0][0][1:3], ('POST', '/answers'))
        self.assertEqual(sent[0][0][4]['message']['attachments'][0]['payload']['token'],
                         'file-token')
        MaxLimitsUin.query.filter_by(user_id=203).one().status = 'inactive'
        db.session.commit()
        calls = []
        with patch.dict(os.environ, {'MAX_LIMITS_BOT_TOKEN': 'test-only'}):
            self.assertEqual(max_limits.drain_outbox(
                api_call=lambda *args, **kwargs: calls.append(args)), 1)
        self.assertEqual(len(calls), 1)
        self.assertEqual(calls[0][3], {'user_id': 201})
        self.assertIsNotNone(outbox[1].cancelled_at)
        self.assertEqual(max_digest.prepare_digest('a' * 32), (0, 0))

    def test_digest_counts_all_approved_objects_and_lists_only_changed(self):
        for index in range(30):
            uin = f'OBJECT-{index:02d}'
            db.session.add(LimitObject(
                dataset_id=self.dataset.id, uin=uin,
                title=f'Объект {index}', last_seen_version_id=self.version.id,
                is_present=True))
            db.session.add(MaxLimitsUin(user_id=305, uin=uin, status='approved'))
        new = LimitVersion(
            dataset_id=self.dataset.id, uploaded_by=self.version.uploaded_by,
            source_date=date(2025, 2, 1), status='accepted',
            source_name='new.xlsx', source_path='', sha256='d' * 64,
            previous_version_id=self.version.id)
        db.session.add(new)
        db.session.flush()
        self.dataset.current_version_id = new.id
        new.snapshot_json = json.dumps({'rows': [
            {'uin': row.uin, 'cells': {'object_name': {'value': row.title}}}
            for row in LimitObject.query.filter_by(dataset_id=self.dataset.id).all()]})
        Setting.set('max_limits_version_id', str(new.id))
        comparison = LimitComparison(
            old_version_id=self.version.id, new_version_id=new.id, status='ready')
        db.session.add(comparison)
        db.session.flush()
        for index in range(20):
            db.session.add(LimitChange(
                comparison_id=comparison.id, uin=f'OBJECT-{index:02d}',
                category='field', field_key='year', field_label='Год',
                old_json='2024', new_json='2025'))
        db.session.commit()
        users, messages = max_digest.prepare_digest()
        self.assertEqual(users, 1)
        self.assertGreaterEqual(messages, 1)
        pages = [max_digest.message_for_outbox(item) for item in
                 MaxLimitsOutbox.query.order_by(MaxLimitsOutbox.id).all()]
        self.assertTrue(all(len(page) <= max_limits.MAX_TEXT for page in pages))
        combined = '\n'.join(pages)
        self.assertIn('Всего ваших объектов — 30', combined)
        self.assertIn('Объектов с изменившимися показателями — 20', combined)
        for index in range(20):
            self.assertIn(f'УИН OBJECT-{index:02d}', combined)
        for index in range(20, 30):
            self.assertNotIn(f'УИН OBJECT-{index:02d}', combined)

    def test_manual_digest_requires_current_accepted_comparison(self):
        with self.assertRaises(ValueError):
            max_digest.prepare_digest()

    def test_manual_digest_can_resend_for_new_filter_without_hiding_other_changes(self):
        db.session.add(MaxLimitsUin(user_id=201, uin='UIN-123', status='approved'))
        new = LimitVersion(
            dataset_id=self.dataset.id, uploaded_by=self.version.uploaded_by,
            source_date=date(2026, 2, 1), status='accepted',
            source_name='new.xlsx', source_path='', sha256='f' * 64,
            previous_version_id=self.version.id)
        db.session.add(new)
        db.session.flush()
        self.dataset.current_version_id = new.id
        new.snapshot_json = self.version.snapshot_json
        Setting.set('max_limits_version_id', str(new.id))
        comparison = LimitComparison(
            old_version_id=self.version.id, new_version_id=new.id, status='ready')
        db.session.add(comparison)
        db.session.flush()
        db.session.add_all([
            LimitChange(
                comparison_id=comparison.id, uin='UIN-123',
                category='annual_funding', field_key='annual_funding_2024_budget',
                field_label='Финансирование 2024', source_json='{"year": 2024}',
                old_json='10', new_json='20'),
            LimitChange(
                comparison_id=comparison.id, uin='UIN-123',
                category='requisites', field_key='object_name',
                field_label='Наименование объекта',
                old_json='"Старое"', new_json='"Новое"'),
        ])
        db.session.commit()
        self.assertEqual(max_digest.prepare_digest('a' * 32), (1, 1))
        first = MaxLimitsOutbox.query.one()
        self.assertEqual(first.uin, '__manual_digest__:2026:' + 'a' * 32 + ':0')
        self.assertIn('UIN-123', max_digest.message_for_outbox(first))
        self.assertEqual(max_digest.prepare_digest('a' * 32), (0, 0))
        first.cancelled_at = utcnow()
        Setting.set('max_limits_min_funding_year', 'all')
        db.session.commit()
        self.assertIsNone(max_digest.message_for_outbox(first))
        self.assertEqual(max_digest.prepare_digest('a' * 32), (1, 1))
        second = MaxLimitsOutbox.query.filter_by(uin='__manual_digest__:all:' + 'a' * 32 + ':0').one()
        self.assertIn('UIN-123', max_digest.message_for_outbox(second))
        self.assertEqual(max_digest.prepare_digest('a' * 32), (0, 0))
        second.cancelled_at = utcnow()
        Setting.set('max_limits_min_funding_year', '2026')
        db.session.commit()
        self.assertEqual(max_digest.prepare_digest('a' * 32), (1, 1))
        self.assertIsNone(first.cancelled_at)
        self.assertEqual(MaxLimitsOutbox.query.count(), 2)

    def test_excel_transfer_codes_have_real_before_and_after_columns(self):
        from io import BytesIO
        from openpyxl import load_workbook

        db.session.add(MaxLimitsUin(user_id=201, uin='UIN-123', status='approved'))
        new = LimitVersion(
            dataset_id=self.dataset.id, uploaded_by=self.version.uploaded_by,
            source_date=date(2025, 2, 1), status='accepted',
            source_name='new.xlsx', source_path='', sha256='e' * 64,
            previous_version_id=self.version.id)
        db.session.add(new)
        db.session.flush()
        self.dataset.current_version_id = new.id
        new.snapshot_json = self.version.snapshot_json
        Setting.set('max_limits_version_id', str(new.id))
        comparison = LimitComparison(
            old_version_id=self.version.id, new_version_id=new.id, status='ready')
        db.session.add(comparison)
        db.session.flush()
        db.session.add(LimitChange(
            comparison_id=comparison.id, uin='UIN-123',
            category='programme_transfer', field_key='row_transferred',
            old_json='null', new_json='null',
            source_json=json.dumps({
                'old_programme': {'gp_code': '01', 'om_code': '03', 'dm_code': '05'},
                'new_programme': {'gp_code': '01', 'om_code': '04', 'dm_code': '12'},
            })))
        db.session.commit()
        rows = list(load_workbook(
            BytesIO(max_digest.changes_excel_bytes(201)), read_only=True).active.values)
        self.assertEqual(rows[0][3:], ('Изменившийся показатель', 'Было', 'Стало'))
        self.assertEqual([row[3:] for row in rows[1:]],
                         [('Код ОМ', '03', '04'), ('Код ДМ', '05', '12')])
        self.assertEqual({row[0] for row in rows[1:]}, {'UIN-123'})

    def test_group_moderation_accepts_only_configured_chat_and_notifies(self):
        with patch.dict(os.environ, {'MAX_LIMITS_MODERATION_CHAT_ID': '-345'}):
            max_limits.process_update(self.message('/start'))
            max_limits.process_update(self.message('Иванов Иван Иванович'))
            max_limits.process_update(self.message('+79201234567'))
            max_limits.process_update(self.message('UIN-123'))
            pending = max_limits.process_update({
                'update_type': 'message_callback',
                'callback': {'user': {'user_id': 101}, 'callback_id': 'cb1',
                             'payload': 'apply:finish'},
                'message': {'recipient': {'chat_type': 'dialog', 'user_id': 101}},
            })
            self.assertIn('отправлена', pending[2]['text'])
            application = MaxLimitsApplication.query.one()
            sends = []
            max_registration.pending_moderation(
                lambda *args, **kwargs: sends.append((args, kwargs)), 'test-only')
            self.assertEqual(sends[0][0][3], {'chat_id': -345})
            self.assertIsNotNone(application.sent_at)
            event = lambda chat, action='approve': {
                'update_type': 'message_callback',
                'callback': {'user': {'user_id': 876}, 'callback_id': 'cb2',
                             'payload': f'mod:{application.id}:{action}'},
                'message': {'recipient': {'chat_type': 'chat', 'chat_id': chat}},
            }
            self.assertIsNone(max_limits.process_update(event(-999)))
            self.assertEqual(application.status, 'pending')
            denied = max_limits.process_update(
                event(-345), is_group_admin=lambda *_: False)
            self.assertIn('только администратор', denied[2]['text'])
            self.assertEqual(application.status, 'pending')
            self.assertIn('одобрена', max_limits.process_update(
                event(-345), is_group_admin=lambda *_: True)[2]['text'])
            self.assertEqual(application.status, 'approved')
            self.assertEqual(MaxLimitsUin.query.one().status, 'approved')
            self.assertTrue(max_limits.process_update(
                self.message('/uin'))[2].get('buttons'))
            self.assertIn('рассмотрена', max_limits.process_update(
                event(-345), is_group_admin=lambda *_: True)[2]['text'])
            self.assertIn('одобрена', db.session.get(MaxLimitsProfile, 101).notice)
            max_registration.pending_moderation(
                lambda *args, **kwargs: sends.append((args, kwargs)), 'test-only')
            self.assertEqual(sends[-1][0][3], {'user_id': 101})
            self.assertEqual(db.session.get(MaxLimitsProfile, 101).notice, '')

    def test_catalog_contains_accepted_objects_and_not_formulas(self):
        from openpyxl import load_workbook
        self.obj.title = '=SUM(1,2)'
        snapshot = json.loads(self.version.snapshot_json)
        snapshot['rows'][0]['cells']['object_name']['value'] = '=SUM(1,2)'
        self.version.snapshot_json = json.dumps(snapshot)
        db.session.commit()
        sheet = load_workbook(
            __import__('io').BytesIO(max_registration.catalog_bytes()),
            read_only=True).active
        rows = list(sheet.values)
        self.assertEqual(rows[0], ('УИН', 'Наименование объекта', 'Адрес'))
        self.assertEqual(rows[1][0], 'UIN-123')
        self.assertEqual(rows[1][1], "'=SUM(1,2)")

    def _other_dataset(self, uin='OTHER-UIN'):
        dataset = LimitDataset(
            name='Other dataset', scope='', owner_id=self.dataset.owner_id)
        db.session.add(dataset)
        db.session.flush()
        version = LimitVersion(
            dataset_id=dataset.id, uploaded_by=self.version.uploaded_by,
            source_date=date(2025, 1, 2), status='accepted',
            source_name='other.xlsx', source_path='', sha256='c' * 64,
            snapshot_json=json.dumps({'rows': [{'uin': uin, 'cells': {
                'object_name': {'value': 'Чужой объект'},
                'address': {'value': 'Чужой адрес'}}}]}))
        db.session.add(version)
        db.session.flush()
        dataset.current_version_id = version.id
        db.session.add(LimitObject(
            dataset_id=dataset.id, uin=uin, title='Чужой объект',
            address='Чужой адрес', last_seen_version_id=version.id, is_present=True))
        db.session.commit()
        return dataset, version

    def test_selected_dataset_is_explicit_and_existing(self):
        self.assertEqual(selected_dataset_id(), self.dataset.id)
        for value in ('', 'not-an-id', '-1', '²', '99999'):
            Setting.set('max_limits_version_id', value)
            db.session.commit()
            self.assertIsNone(selected_dataset_id())
            self.assertFalse(max_registration._valid_uin('UIN-123'))
            self.assertNotIn('UIN-123', self._catalog_uins())
            self.assertNotIn('Объект: Объект',
                             max_limits.current_information('UIN-123'))
        Setting.set('max_limits_version_id', str(self.version.id))
        db.session.commit()
        self.assertEqual(selected_dataset_id(), self.dataset.id)

    @staticmethod
    def _catalog_uins():
        from io import BytesIO
        from openpyxl import load_workbook
        book = load_workbook(BytesIO(max_registration.catalog_bytes()), read_only=True)
        try:
            return [row[0] for row in list(book.active.values)[1:]]
        finally:
            book.close()

    def test_catalog_and_registration_isolate_selected_dataset(self):
        other, _ = self._other_dataset()
        self.assertTrue(max_registration._valid_uin('UIN-123'))
        self.assertFalse(max_registration._valid_uin('OTHER-UIN'))
        self.assertEqual(self._catalog_uins(), ['UIN-123'])
        self.assertIn('Объект: Объект', max_limits.current_information('UIN-123'))
        self.assertNotIn('Чужой объект',
                         max_limits.current_information('OTHER-UIN'))
        Setting.set('max_limits_version_id', str(other.current_version_id))
        db.session.commit()
        self.assertFalse(max_registration._valid_uin('UIN-123'))
        self.assertTrue(max_registration._valid_uin('OTHER-UIN'))
        self.assertEqual(self._catalog_uins(), ['OTHER-UIN'])
        self.assertIn('Чужой объект',
                      max_limits.current_information('OTHER-UIN'))
        self.assertNotIn('Объект: Объект',
                         max_limits.current_information('UIN-123'))

    def test_notifications_and_legacy_changes_isolate_selected_dataset(self):
        other, old_other = self._other_dataset(uin='UIN-123')
        db.session.add(MaxLimitsUin(
            user_id=301, uin='UIN-123', status='approved'))
        db.session.flush()
        new_other = LimitVersion(
            dataset_id=other.id, uploaded_by=self.version.uploaded_by,
            source_date=date(2025, 2, 2), status='accepted',
            source_name='other-new.xlsx', source_path='', sha256='d' * 64)
        db.session.add(new_other)
        db.session.flush()
        comparison = LimitComparison(
            old_version_id=old_other.id, new_version_id=new_other.id, status='ready')
        db.session.add(comparison)
        db.session.flush()
        db.session.add(LimitChange(
            comparison_id=comparison.id, uin='UIN-123', category='field',
            field_key='other', field_label='Чужое изменение',
            old_json=json.dumps(1), new_json=json.dumps(2)))
        db.session.flush()
        self.assertEqual(max_limits.queue_accepted_changes(new_other, comparison), 0)
        self.assertEqual(MaxLimitsOutbox.query.count(), 0)
        self.assertNotIn('Чужое изменение', max_limits.accepted_changes('UIN-123'))

        # A previously queued notification must be cancelled even when it is
        # not yet due, and no API request may be made for the old dataset.
        old_item = MaxLimitsOutbox(
            user_id=301, uin='UIN-123', version_id=self.version.id, text='old',
            next_attempt_at=utcnow() + timedelta(days=1))
        db.session.add(old_item)
        new_other.previous_version_id = old_other.id
        new_other.snapshot_json = old_other.snapshot_json
        Setting.set('max_limits_version_id', str(new_other.id))
        db.session.commit()
        self.assertEqual(max_limits.queue_accepted_changes(new_other, comparison), 1)
        db.session.commit()
        self.assertIn('Чужое изменение', max_limits.accepted_changes('UIN-123'))
        calls = []
        with patch.dict(os.environ, {'MAX_LIMITS_BOT_TOKEN': 'test-only'}):
            self.assertEqual(max_limits.drain_outbox(
                api_call=lambda *args, **kwargs: calls.append((args, kwargs))), 1)
        self.assertIsNotNone(old_item.cancelled_at)
        self.assertIsNone(old_item.delivered_at)
        self.assertEqual(len(calls), 1)
        self.assertIn('Чужое изменение', calls[0][0][4]['text'])

    def test_all_navigation_and_subscription_actions_have_buttons(self):
        p = max_registration.profile_for(101)
        p.full_name = 'Иванов Иван Иванович'
        p.phone = '+79201234567'
        row = MaxLimitsUin(user_id=101, uin='UIN-123', status='approved')
        db.session.add(row)
        db.session.commit()
        menu = max_registration.start(101)
        payloads = {payload for line in menu['buttons'] for _, payload in line}
        self.assertEqual(payloads, {'menu:changes', 'menu:subscriptions',
                                    'menu:changes-excel'})
        subscriptions = max_registration.handle_callback(101, 'menu:subscriptions')
        sub_payloads = {payload for line in subscriptions['buttons'] for _, payload in line}
        self.assertEqual(sub_payloads, {'menu:uins', 'menu:add', 'menu:remove',
                                        'menu:stop', 'menu:home'})
        for action in ('menu:home', 'menu:subscriptions', 'menu:uins', 'menu:changes'):
            self.assertTrue(max_registration.handle_callback(101, action)['buttons'])
        self.assertTrue(max_registration.handle_callback(101, 'menu:add')['file'])
        max_registration.handle_callback(101, f'remove:{row.id}')
        self.assertEqual(row.status, 'approved')
        max_registration.handle_callback(101, f'remove-confirm:{row.id}')
        self.assertEqual(row.status, 'inactive')
        row.status = 'approved'
        db.session.commit()
        max_registration.handle_callback(101, 'menu:stop')
        self.assertEqual(row.status, 'approved')
        stopped = max_registration.handle_callback(101, 'stop:confirm')
        self.assertEqual(row.status, 'inactive')
        self.assertIn('menu:register', [p for line in stopped['buttons'] for _, p in line])
        restarted = max_registration.handle_callback(101, 'menu:register')
        self.assertIn('ФИО', restarted['text'])

    def test_denial_does_not_enable_access_and_remove_is_scoped(self):
        with patch.dict(os.environ, {'MAX_LIMITS_MODERATION_CHAT_ID': '-345'}):
            p = max_registration.profile_for(101)
            p.full_name = 'Иванов Иван Иванович'
            p.phone = '+79201234567'
            db.session.add(MaxLimitsUin(user_id=101, uin='UIN-123', status='draft'))
            db.session.commit()
            max_registration.handle_callback(101, 'apply:finish')
            app = MaxLimitsApplication.query.one()
            app.sent_at = __import__('datetime').datetime.utcnow()
            db.session.commit()
            max_registration.decide(-345, 876, f'mod:{app.id}:reject')
            self.assertEqual(max_registration.approved_uins(101), [])
            self.assertEqual(MaxLimitsUin.query.one().status, 'rejected')

    def test_group_admin_lookup_checks_current_roster_and_fails_closed(self):
        calls = []

        def api_call(token, method, path, params=None, timeout=None):
            calls.append((method, path, params))
            if params is None:
                return {'members': [{'user_id': 55, 'is_admin': False}],
                        'marker': 7}
            return {'members': [{'user_id': 55, 'is_admin': True}]}

        self.assertTrue(max_limits.group_admin_lookup('fake', -345, 55, api_call))
        self.assertEqual(calls, [
            ('GET', '/chats/-345/members/admins', None),
            ('GET', '/chats/-345/members/admins', {'marker': 7})])
        self.assertFalse(max_limits.group_admin_lookup(
            'fake', -345, 56, api_call))
        self.assertFalse(max_limits.group_admin_lookup(
            'fake', -345, 55, lambda *args, **kwargs: {'members': []}))
        self.assertFalse(max_limits.group_admin_lookup(
            'fake', -345, 55, lambda *args, **kwargs: {
                'members': [{'user_id': 55, 'is_bot': True, 'is_admin': True}]}))
        self.assertFalse(max_limits.group_admin_lookup(
            'fake', -345, 55, lambda *args, **kwargs: {}))

    def test_legacy_tables_removed_idempotently_without_financial_loss(self):
        with db.engine.begin() as connection:
            for table in ('telegram_limits_subscriptions', 'telegram_limits_outbox',
                          'telegram_limits_cursors'):
                connection.execute(text(f'CREATE TABLE "{table}" (id INTEGER PRIMARY KEY)'))
        Setting.set('telegram_limits_offset', '12')
        Setting.set('telegram_limits_username', 'retired')
        db.session.commit()
        max_limits.migrate_max_limits_schema()
        max_limits.migrate_max_limits_schema()
        self.assertIsNotNone(db.session.get(LimitVersion, self.version.id))
        self.assertIsNotNone(db.session.get(LimitObject, self.obj.id))
        self.assertIsNone(db.session.get(Setting, 'telegram_limits_username'))
        for table in ('telegram_limits_subscriptions', 'telegram_limits_outbox',
                      'telegram_limits_cursors'):
            self.assertNotIn(table, inspect(db.engine).get_table_names())
            self.assertIn(table, RETIRED_LIMIT_TABLES)
        self.assertIn('max_limits_outbox', COUNT_TABLES)

    def test_transport_uses_api2_and_header_not_query(self):
        class Response:
            def __enter__(self):
                return self

            def __exit__(self, *args):
                return False

            def read(self, limit):
                return b'{"updates":[],"marker":123}'

        with patch.object(max_limits, 'urlopen', return_value=Response()) as opener:
            result = max_limits.max_api('fake-token', 'GET', '/updates', {'marker': 12})
        request = opener.call_args.args[0]
        self.assertEqual(result['marker'], 123)
        self.assertEqual(request.get_header('Authorization'), 'fake-token')
        self.assertTrue(request.full_url.startswith('https://platform-api2.max.ru/updates?'))
        self.assertNotIn('fake-token', request.full_url)
        self.assertEqual(db.session.get(MaxLimitsCursor, 1), None)


if __name__ == '__main__':
    unittest.main()