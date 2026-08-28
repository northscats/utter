# -*- coding: utf-8 -*-
"""Интеграционный тест ИИ-проверки документов.

Изолированная среда: временная SQLite-БД и временное хранилище файлов —
тест НЕ трогает рабочую базу и uploads/.
"""
import io
import json
import os
import sys
import tempfile
import threading
import time
import zipfile
from http.server import BaseHTTPRequestHandler, HTTPServer

_TMP = tempfile.mkdtemp(prefix='docs_ai_test_')
os.environ['DATABASE_URL'] = f'sqlite:///{_TMP}/test.db'

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from app import app, init_db, encrypt_value
from database import AiProvider, DocAnalysisRun, DocFile, DocPackage, User, db

app.config['DOCS_STORAGE'] = os.path.join(_TMP, 'docs')
os.makedirs(app.config['DOCS_STORAGE'], exist_ok=True)
init_db()

# ── Фейковый OpenAI-совместимый сервер ─────────────────────────────────────
received_bodies = []
fake_response_mode = {'value': 'normal'}
fake_delay_seconds = {'value': 0}

class FakeAI(BaseHTTPRequestHandler):
    def do_POST(self):
        if fake_delay_seconds['value']:
            time.sleep(fake_delay_seconds['value'])
        body = json.loads(self.rfile.read(int(self.headers['Content-Length'])))
        received_bodies.append(body)
        system_msg = body['messages'][0]['content']
        user_msg = body['messages'][1]['content']
        if fake_response_mode['value'] == 'invalid':
            content = 'это не JSON'
        elif '{"answer"' in system_msg:
            content = json.dumps(
                {'answer': 'Срок действия найден. Существенных рисков не выявлено.'},
                ensure_ascii=False)
        else:
            import re
            rule_ids = [int(m) for m in re.findall(r'\[(\d+)\]', user_msg)]
            items = []
            for i, rid in enumerate(rule_ids):
                st = 'violation' if i % 2 else 'ok'
                items.append({'rule_id': rid, 'status': st,
                              'details': f'Обоснование для правила {rid}'})
            content = json.dumps(items, ensure_ascii=False)
        resp = {'choices': [{'message': {'content': content}}],
                'usage': {'prompt_tokens': 100, 'completion_tokens': 50}}
        data = json.dumps(resp).encode()
        self.send_response(200)
        self.send_header('Content-Type', 'application/json')
        self.send_header('Content-Length', str(len(data)))
        self.end_headers()
        self.wfile.write(data)
    def log_message(self, *a):
        pass

srv = HTTPServer(('127.0.0.1', 0), FakeAI)
threading.Thread(target=srv.serve_forever, daemon=True).start()
FAKE_URL = f'http://127.0.0.1:{srv.server_address[1]}/v1'

app.config['WTF_CSRF_ENABLED'] = False
client = app.test_client()

with app.app_context():
    admin_id = User.query.filter_by(role='admin').first().id
    # Внешний (нелокальный) провайдер — для негативного теста
    db.session.add(AiProvider(
        name='Внешний облачный', base_url='https://cloud.example.com/v1',
        api_key=encrypt_value('k'), model='gpt-x',
        is_active=True, is_default=True, is_local=False))
    db.session.commit()

with client.session_transaction() as s:
    s['_user_id'] = str(admin_id)
    s['_fresh'] = True

def check(name, cond, extra=''):
    print(('PASS' if cond else 'FAIL'), name, extra)
    if not cond:
        sys.exit(1)

# ── Каталог с правилами проверки ───────────────────────────────────────────
r = client.post('/docs/catalogs', json=dict(name='Договоры'))
cat_id = r.get_json()['id']
client.post(f'/docs/catalogs/{cat_id}/routing-rules',
            json=dict(rule_type='extension', pattern='txt'))
client.post(f'/docs/catalogs/{cat_id}/check-rules',
            json=dict(name='Печать', prompt_text='Есть ли отметка о печати?'))
client.post(f'/docs/catalogs/{cat_id}/check-rules',
            json=dict(name='Дата', prompt_text='Указана ли дата документа?'))

# ── Пакет (заголовок с XSS-инъекцией — регрессионный тест) ─────────────────
XSS_TITLE = "'-alert(1)-'<img src=x onerror=alert(2)>"
buf = io.BytesIO()
with zipfile.ZipFile(buf, 'w') as zf:
    zf.writestr('Договор 1.txt', 'Договор №1 от 01.08.2026. М.П. Подпись.')
    # Formula injection: имя файла начинается с «=»
    zf.writestr('=2+5+cmd.txt', 'Содержимое подозрительного файла')
buf.seek(0)
r = client.post('/docs/packages',
                data={'zipfile': (buf, 'dogovory.zip'), 'title': XSS_TITLE},
                content_type='multipart/form-data')
pkg_id = r.get_json()['id']
for _ in range(60):
    st = client.get(f'/docs/packages/{pkg_id}/status').get_json()
    if st['status'] in ('done', 'error'):
        break
    time.sleep(1)
check('пакет обработан', st['status'] == 'done', json.dumps(st, ensure_ascii=False))

# UI-страница не должна интерполировать пользовательские строки в JS-контекст
r = client.get('/docs/')
html = r.get_data(as_text=True)
check('UI открывается', r.status_code == 200 and 'Проверка документов' in html)
check('нет inline-onclick с интерполяцией', 'onclick="showResults' not in html
      and "onclick=\"analyze" not in html)

# ── 1. Внешний провайдер блокируется ───────────────────────────────────────
r = client.post(f'/docs/packages/{pkg_id}/analyze')
check('внешний провайдер отклонён', r.status_code == 400
      and 'локальн' in r.get_json()['error'].lower(), r.get_data(as_text=True)[:200])

# ── 2. Локальный провайдер настраивается ЧЕРЕЗ ФОРМУ НАСТРОЕК ─────────────
r = client.post('/settings', data=dict(
    action='add_ai_provider', ai_name='Локальный vLLM',
    ai_base_url=FAKE_URL, ai_api_key='local-key', ai_model='qwen',
    ai_is_active='1', ai_is_local='1'), follow_redirects=True)
check('провайдер добавлен через настройки', r.status_code == 200)
with app.app_context():
    loc = AiProvider.query.filter_by(name='Локальный vLLM').first()
    check('is_local сохранён формой', loc is not None and loc.is_local)
    loc_id = loc.id

r = client.post('/settings', data=dict(
    action='update_ai_provider', provider_id=str(loc_id),
    ai_name='Локальный vLLM', ai_base_url=FAKE_URL, ai_model='qwen',
    ai_is_active='1', ai_is_local='1'), follow_redirects=True)
with app.app_context():
    loc = db.session.get(AiProvider, loc_id)
    check('is_local сохраняется при редактировании', loc.is_local and loc.is_active)

r = client.post(f'/docs/packages/{pkg_id}/analyze')
check('анализ запущен', r.status_code == 202, r.get_data(as_text=True)[:200])

for _ in range(30):
    st = client.get(f'/docs/packages/{pkg_id}/status').get_json()
    if st['analysis_status'] in ('done', 'error'):
        break
    time.sleep(1)
check('анализ завершён', st['analysis_status'] == 'done',
      json.dumps(st, ensure_ascii=False))
check('прогресс анализа', st['analysis_done'] == 2 and st['analysis_total'] == 2,
      json.dumps(st, ensure_ascii=False))
check('ИИ-запросы получены', len(received_bodies) >= 1)
_user_prompt = received_bodies[0]['messages'][1]['content']
check('промпт содержит правила', 'Печать' in _user_prompt and 'Договор' in _user_prompt,
      repr(_user_prompt[:400]))

# ── 3. Результаты ──────────────────────────────────────────────────────────
res = client.get(f'/docs/packages/{pkg_id}/results').get_json()
check('title в результатах сохранён как данные', res['title'] == XSS_TITLE[:255])
_by_name = {f['orig_name']: f for f in res['files']}
fres = _by_name['Договор 1.txt']['results']
check('результаты по 2 правилам', len(fres) == 2, json.dumps(fres, ensure_ascii=False))
statuses = {r_['rule_id']: r_['status'] for r_ in fres}
check('статусы ok+violation', set(statuses.values()) == {'ok', 'violation'})
viol = [r_ for r_ in fres if r_['status'] == 'violation'][0]
check('детали нарушения', 'Обоснование' in json.dumps(viol['details'], ensure_ascii=False))

# ── 4. Анализ текущего документа пользовательским промптом ─────────────────
current_file_id = _by_name['Договор 1.txt']['file_id']
before_calls = len(received_bodies)
r = client.post(f'/docs/files/{current_file_id}/analyze', json={
    'mode': 'supplement',
    'user_prompt': 'Особенно проверь срок действия договора',
    'find_dates': True, 'find_signatures': False, 'find_seals': False,
})
check('дополнение правил запущено для текущего файла',
      r.status_code == 202 and r.get_json()['mode'] == 'supplement',
      r.get_data(as_text=True)[:300])
supplement_run_id = r.get_json()['id']
for _ in range(30):
    st = client.get(f'/docs/packages/{pkg_id}/status').get_json()
    if st['analysis_status'] in ('done', 'error'):
        break
    time.sleep(1)
check('дополнение правил завершено', st['analysis_status'] == 'done',
      json.dumps(st, ensure_ascii=False))
new_bodies = received_bodies[before_calls:]
check('пользовательский промпт передан ИИ',
      any('Особенно проверь срок действия договора' in
          b['messages'][0]['content'] for b in new_bodies))

r = client.post(f'/docs/files/{current_file_id}/analyze', json={
    'mode': 'free',
    'user_prompt': 'Найди срок действия и кратко оцени риски',
})
check('свободный анализ запущен',
      r.status_code == 202 and r.get_json()['mode'] == 'free',
      r.get_data(as_text=True)[:300])
free_run_id = r.get_json()['id']
for _ in range(30):
    st = client.get(f'/docs/packages/{pkg_id}/status').get_json()
    if st['analysis_status'] in ('done', 'error'):
        break
    time.sleep(1)
check('свободный анализ завершён', st['analysis_status'] == 'done',
      json.dumps(st, ensure_ascii=False))
fd = client.get(f'/docs/files/{current_file_id}').get_json()
runs = {x['id']: x for x in fd['analysis_runs']}
check('история содержит оба режима',
      supplement_run_id in runs and free_run_id in runs)
check('свободный ответ сохранён отдельно',
      runs[free_run_id]['answer'] and 'Срок действия' in runs[free_run_id]['answer'])
check('supplement хранит снимок правил и результата',
      runs[supplement_run_id]['options']['rules']
      and runs[supplement_run_id]['result']['rules'])
check('стандартные результаты не удалены свободным анализом',
      len(client.get(f'/docs/packages/{pkg_id}/results').get_json()['files'][0]['results']) >= 1)

# Новый невалидный ответ не должен уничтожить предыдущую успешную проверку.
before_invalid = client.get(f'/docs/packages/{pkg_id}/results').get_json()
before_rules = {f['file_id']: f['results'] for f in before_invalid['files']}[current_file_id]
fake_response_mode['value'] = 'invalid'
r = client.post(f'/docs/files/{current_file_id}/analyze', json={
    'mode': 'supplement', 'user_prompt': 'Повтори проверку',
})
check('запуск с будущим невалидным ответом принят', r.status_code == 202)
invalid_run_id = r.get_json()['id']
for _ in range(30):
    st = client.get(f'/docs/packages/{pkg_id}/status').get_json()
    if st['analysis_status'] in ('done', 'error'):
        break
    time.sleep(1)
fake_response_mode['value'] = 'normal'
check('невалидный ответ помечен ошибкой', st['analysis_status'] == 'error')
fd = client.get(f'/docs/files/{current_file_id}').get_json()
invalid_run = {x['id']: x for x in fd['analysis_runs']}[invalid_run_id]
check('ошибка сохранена в отдельном запуске',
      invalid_run['status'] == 'error' and invalid_run['error_text'])
after_invalid = client.get(f'/docs/packages/{pkg_id}/results').get_json()
after_rules = {f['file_id']: f['results'] for f in after_invalid['files']}[current_file_id]
check('предыдущий успешный результат сохранён при ошибке',
      after_rules == before_rules)

r = client.post(f'/docs/files/{current_file_id}/analyze',
                json={'mode': 'free', 'user_prompt': ''})
check('пустой промпт отклонён', r.status_code == 400)
r = client.post(f'/docs/files/{current_file_id}/analyze',
                json={'mode': 'other', 'user_prompt': 'тест'})
check('неизвестный режим отклонён', r.status_code == 400)

# Пакет резервируется до старта фонового потока: второй запрос получает 409.
fake_delay_seconds['value'] = 1
r = client.post(f'/docs/files/{current_file_id}/analyze', json={
    'mode': 'free', 'user_prompt': 'Медленный анализ',
})
check('медленный анализ запущен', r.status_code == 202)
r2 = client.post(f'/docs/files/{current_file_id}/analyze', json={
    'mode': 'free', 'user_prompt': 'Параллельный анализ',
})
check('параллельный запуск отклонён', r2.status_code == 409)
for _ in range(30):
    st = client.get(f'/docs/packages/{pkg_id}/status').get_json()
    if st['analysis_status'] in ('done', 'error'):
        break
    time.sleep(1)
fake_delay_seconds['value'] = 0
check('медленный анализ завершён', st['analysis_status'] == 'done')

# ── 5. Экспорт XLSX ────────────────────────────────────────────────────────
r = client.get(f'/docs/packages/{pkg_id}/export')
check('экспорт XLSX', r.status_code == 200 and r.data[:2] == b'PK')
from openpyxl import load_workbook
wb = load_workbook(io.BytesIO(r.data))
cells = [[c.value for c in row] for row in wb.active.iter_rows()]
check('XLSX содержит правило и статус',
      any('Печать' in str(row) for row in cells)
      and any('Нарушение' in str(row) or 'ОК' in str(row) for row in cells),
      str(cells)[:300])
# Formula injection: ни одна ячейка не должна начинаться с =, +, @
_bad = [(c.value, c.data_type) for row in wb.active.iter_rows() for c in row
        if isinstance(c.value, str) and c.value[:1] in ('=', '+', '@')]
_formula = [c for row in wb.active.iter_rows() for c in row if c.data_type == 'f']
check('XLSX без formula injection', not _bad and not _formula, str(_bad)[:200])
check('инъекционное имя экранировано',
      any(isinstance(c.value, str) and c.value.startswith("'=2+5") for row
          in wb.active.iter_rows() for c in row), str(cells)[:400])

# ── 6. Повторный запуск перезаписывает результаты ──────────────────────────
r = client.post(f'/docs/packages/{pkg_id}/analyze')
check('повторный анализ запущен', r.status_code == 202)
for _ in range(30):
    st = client.get(f'/docs/packages/{pkg_id}/status').get_json()
    if st['analysis_status'] in ('done', 'error'):
        break
    time.sleep(1)
res = client.get(f'/docs/packages/{pkg_id}/results').get_json()
_by_name2 = {f['orig_name']: f for f in res['files']}
check('результаты не дублируются', len(_by_name2['Договор 1.txt']['results']) == 2)

# ── 7. После рестарта зависший запуск получает конечную ошибку ──────────────
with app.app_context():
    stale = DocAnalysisRun(
        package_id=pkg_id, file_id=current_file_id, created_by=admin_id,
        mode='free', scope='current_file', user_prompt='Зависший запуск',
        status='queued')
    db.session.add(stale)
    pkg = db.session.get(DocPackage, pkg_id)
    df = db.session.get(DocFile, current_file_id)
    pkg.analysis_status = 'running'
    df.analysis_status = 'processing'
    db.session.commit()
    stale_id = stale.id
init_db()
with app.app_context():
    stale = db.session.get(DocAnalysisRun, stale_id)
    check('зависший запуск восстановлен как ошибка',
          stale.status == 'error' and stale.completed_at is not None
          and 'перезапуск' in stale.error_text.lower())

# ── 8. Доступ: обычный пользователь не видит чужой пакет/файл ──────────────
with app.app_context():
    from werkzeug.security import generate_password_hash
    other = User(username='test_user_docs', role='user',
                 password_hash=generate_password_hash('x'))
    db.session.add(other)
    db.session.commit()
    other_id = other.id
with client.session_transaction() as s:
    s['_user_id'] = str(other_id)
r1 = client.get(f'/docs/packages/{pkg_id}/results')
r2 = client.get(f'/docs/packages/{pkg_id}/export')
r3 = client.post(f'/docs/packages/{pkg_id}/analyze')
r4 = client.post(f'/docs/files/{current_file_id}/analyze',
                 json={'mode': 'free', 'user_prompt': 'Покажи документ'})
r5 = client.post('/docs/analysis-templates', json={
    'name': 'Пользовательский шаблон',
    'system_prompt': 'Проверь существенные условия',
})
check('чужой пакет недоступен', r1.status_code == 403 and r2.status_code == 403
      and r3.status_code == 403 and r4.status_code == 403,
      f'{r1.status_code}/{r2.status_code}/{r3.status_code}/{r4.status_code}')
check('обычный пользователь может создать шаблон', r5.status_code == 201,
      r5.get_data(as_text=True)[:200])

# ── Очистка временной среды ────────────────────────────────────────────────
import shutil
srv.shutdown()
shutil.rmtree(_TMP, ignore_errors=True)
print('\nВСЕ ТЕСТЫ ПРОЙДЕНЫ')
