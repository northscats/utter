# -*- coding: utf-8 -*-
"""
Парсер оборотно-сальдовых ведомостей (ОСВ) из 1С.

Поддерживает:
  - .xls  (легаси BIFF, типовая выгрузка 1С) — через xlrd (formatting_info=True,
           уровни группировки строк берутся из rowinfo_map)
  - .xlsx — через openpyxl (уровни из row_dimensions[].outlineLevel)
  - .zip  — архив с суточными XLS-файлами (см. parse_zip)

Ядро одно: получает список строк (значения + outline level) и восстанавливает
иерархию аналитик, отделяя строки "Обороты за ДД.ММ.ГГ" (динамика по датам).

Ожидаемый макет ОСВ (типовой 1С, БП 3.0):
  колонка A: Счет / аналитика / "Обороты за ..."
  колонки B..G: СН Дт | СН Кт | Обороты Дт | Обороты Кт | СК Дт | СК Кт
"""
import re
import os
import zipfile
import tempfile
import shutil
from datetime import date as _date

RE_DATE = re.compile(r'Обороты за (\d\d)\.(\d\d)\.(\d\d)')
RE_ACC  = re.compile(r'^\d{2}(\.\d{2})?$')
RE_ORG  = re.compile(r'(?:Организация|Организации|Предприятие)[:\s]+(.+)', re.IGNORECASE)
RE_ACC_HEADER = re.compile(
    r'(?:по\s+)?сч[её]т(?:у|а)?\s*(\d{2})(?=\D|$)',
    re.IGNORECASE,
)
RE_ACC_FILENAME = re.compile(r'(?:^|[\s_-])ОСВ[\s_-]+(\d{2})(?=\D|$)', re.IGNORECASE)

# Для извлечения даты из шапки суточных файлов: «за 5 июля 2026 г.»
_MONTHS_RU = {
    'января': 1, 'февраля': 2, 'марта': 3, 'апреля': 4,
    'мая': 5, 'июня': 6, 'июля': 7, 'августа': 8,
    'сентября': 9, 'октября': 10, 'ноября': 11, 'декабря': 12,
}
RE_DAILY_HDR = re.compile(
    r'за\s+(\d{1,2})\s+(' + '|'.join(_MONTHS_RU) + r')\s+(\d{4})',
    re.IGNORECASE,
)


def _num(v):
    if v is None or v == '':
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace('\xa0', '').replace(' ', '').replace(',', '.')
    try:
        return float(s)
    except ValueError:
        return 0.0


# ---------------------------------------------------------------- backends
def _read_xls(path):
    """-> list[(label, outline_level, [6 чисел])]"""
    import io
    import xlrd
    log = io.StringIO()  # глушим предупреждения xlrd о CODEPAGE
    book = xlrd.open_workbook(path, formatting_info=True, logfile=log)
    # 1С нередко выгружает .xls без записи CODEPAGE: xlrd берёт iso-8859-1,
    # и русский текст превращается в кракозябры. Такие файлы фактически cp1251.
    if book.codepage is None:
        book = xlrd.open_workbook(path, formatting_info=True,
                                  encoding_override='cp1251', logfile=log)
    sh = book.sheet_by_index(0)
    rows = []
    for r in range(sh.nrows):
        label = sh.cell_value(r, 0)
        ri = sh.rowinfo_map.get(r)
        lvl = ri.outline_level if ri is not None else 0
        vals = [_num(sh.cell_value(r, c)) if c < sh.ncols else 0.0
                for c in range(1, 7)]
        rows.append((label, lvl, vals))
    return rows


def _read_xlsx(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=False, data_only=True)
    ws = wb.active
    rows = []
    for r in range(1, ws.max_row + 1):
        label = ws.cell(r, 1).value
        dim = ws.row_dimensions.get(r)
        lvl = dim.outlineLevel if dim is not None else 0
        vals = [_num(ws.cell(r, c).value) for c in range(2, 8)]
        rows.append((label, lvl, vals))
    return rows


def read_rows(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == '.xls':
        return _read_xls(path)
    if ext == '.xlsx':
        return _read_xlsx(path)
    raise ValueError(f'Неподдерживаемый формат файла: {ext}')


# ---------------------------------------------------------------- название организации из шапки
def extract_org_name(path):
    """
    Пытается извлечь название организации из шапки ОСВ-файла 1С.
    Ищет в первых 15 строках паттерн «Организация: <название>».
    Возвращает строку или None (None — пользователь вводит вручную).
    """
    def _clean(s):
        # снимаем только парные внешние кавычки, чтобы не ломать
        # названия вида: Предприятие "Название"
        s = s.strip()
        for a, b in (('"', '"'), ("'", "'"), ('\u00ab', '\u00bb')):
            if len(s) > 1 and s.startswith(a) and s.endswith(b):
                s = s[1:-1].strip()
        return s

    try:
        rows = read_rows(path)
        for label, _, _ in rows[:15]:
            s = str(label or '').strip()
            m = RE_ORG.match(s)
            if m:
                name = _clean(m.group(1))
                if name:
                    return name
        # Фолбэк: в типовой шапке 1С название организации — строка
        # непосредственно перед «Оборотно-сальдовая ведомость…»
        prev = ''
        for label, _, _ in rows[:15]:
            s = str(label or '').strip()
            if s.lower().startswith('оборотно-сальдовая ведомость') and prev:
                return _clean(prev)
            if s:
                prev = s
    except Exception:
        pass
    return None


# ---------------------------------------------------------------- дата суточного файла из шапки
def extract_daily_date(path):
    """
    Извлекает дату из шапки суточного ОСВ-файла.
    Строка вида «Оборотно-сальдовая ведомость по счету 51 за 5 июля 2026 г.»
    Возвращает datetime.date или None.
    """
    try:
        rows = read_rows(path)
        for label, _, _ in rows[:5]:
            s = str(label or '').strip()
            m = RE_DAILY_HDR.search(s)
            if m:
                d, mon_name, y = int(m.group(1)), m.group(2).lower(), int(m.group(3))
                mon = _MONTHS_RU.get(mon_name)
                if mon:
                    return _date(y, mon, d)
    except Exception:
        pass
    return None


# ---------------------------------------------------------------- парсинг ZIP-архива суточных ОСВ
def parse_zip(zip_path, accepted_dir=None):
    """
    Распаковывает ZIP-архив с суточными ОСВ и парсит каждый XLS.

    Возвращает кортеж (successes, errors):
      successes — list of dict {date: datetime.date, parsed: результат parse_osv()}
      errors    — list of dict {file: имя файла, error: сообщение}
    """
    successes, errors = [], []

    with tempfile.TemporaryDirectory() as tmp:
        try:
            with zipfile.ZipFile(zip_path, 'r') as zf:
                zf.extractall(tmp)
        except zipfile.BadZipFile as e:
            return [], [{'file': os.path.basename(zip_path),
                         'error': f'Не удалось открыть ZIP: {e}'}]

        xls_files = []
        for root, _dirs, files in os.walk(tmp):
            for fname in sorted(files):
                if fname.lower().endswith(('.xls', '.xlsx')):
                    xls_files.append(os.path.join(root, fname))

        for fpath in xls_files:
            fname = os.path.basename(fpath)
            dt = extract_daily_date(fpath)
            if dt is None:
                errors.append({'file': fname,
                               'error': 'Не удалось определить дату из шапки'})
                continue
            try:
                parsed = parse_osv(fpath)
                item = {'date': dt, 'parsed': parsed, 'file': fname}
                if accepted_dir:
                    os.makedirs(accepted_dir, exist_ok=True)
                    stem, ext = os.path.splitext(fname)
                    destination = os.path.join(accepted_dir, fname)
                    suffix = 1
                    while os.path.exists(destination):
                        suffix += 1
                        destination = os.path.join(
                            accepted_dir, f'{stem}_{suffix}{ext}')
                    shutil.copy2(fpath, destination)
                    item['path'] = destination
                successes.append(item)
            except Exception as e:
                errors.append({'file': fname, 'error': str(e)})

    return successes, errors


# ---------------------------------------------------------------- ядро
def parse_osv(path):
    """
    Возвращает dict:
      acc        - номер счета ('20', '90', ...)
      period     - строка периода из шапки (если найдена)
      dims       - подписи измерений из шапки
      rows       - список словарей иерархии:
                   {lvl,label,p0..p3,snd,snk,obd,obk,skd,skk}
      daily      - список {date,month,path1,path2,path3,obd,obk}
      warnings   - список замечаний парсера
    """
    raw = read_rows(path)
    warnings = []

    # период из шапки (обычно 2-я строка: "за 1 полугодие 2026 г." / даты)
    period = ''
    for label, _, _ in raw[:6]:
        s = str(label or '')
        if 'полугодие' in s or 'квартал' in s or re.search(r'\d{4}', s):
            period = s.strip()
            break

    # строка заголовка таблицы: 'Счет'
    hdr = None
    for i, (label, _, _) in enumerate(raw[:15]):
        if str(label or '').strip() == 'Счет':
            hdr = i
            break
    if hdr is None:
        raise ValueError(f'{os.path.basename(path)}: не найден заголовок "Счет" — '
                         'это не типовая ОСВ 1С?')

    # номер счета = первая строка данных вида NN / NN.NN
    acc, data_start, dims = None, None, []
    for i in range(hdr + 1, min(hdr + 10, len(raw))):
        s = str(raw[i][0] or '').strip()
        if RE_ACC.match(s):
            acc = s.split('.')[0]
            data_start = i
            break
        if s:
            dims.append(s)
    empty_totals = None
    if acc is None:
        # A valid empty 1C report contains only its title and an "Итого" row,
        # without the usual standalone account row. Keep it as an explicit
        # zero-data report instead of blocking the whole company dashboard.
        title_text = ' '.join(
            str(label or '').strip() for label, _, _ in raw[:8] if label
        )
        match = RE_ACC_HEADER.search(title_text)
        if match is None:
            match = RE_ACC_FILENAME.search(os.path.basename(path))
        if match is None:
            raise ValueError(f'{os.path.basename(path)}: не удалось определить номер счета')
        acc = match.group(1)
        empty_totals = next(
            (vals for label, _, vals in raw[hdr + 1:]
             if str(label or '').strip().startswith('Итого')),
            None,
        )
        data_start = len(raw)
        dims = [dim for dim in dims if not dim.startswith('Итого')]
        warnings.append(
            f'сч.{acc}: пустая ОСВ — в отчёте нет строк с движениями и остатками')

    has_outline = any(lvl > 0 for _, lvl, _ in raw[data_start:])
    if data_start < len(raw) and not has_outline:
        warnings.append(f'сч.{acc}: в файле нет уровней группировки строк — '
                        'иерархия аналитик восстановлена частично')

    rows, daily = [], []
    path_by_lvl = {}
    for label, lvl, vals in raw[data_start:]:
        if label is None:
            continue
        label = str(label).strip()
        if not label or label.startswith('Итого'):
            continue
        snd, snk, obd, obk, skd, skk = vals
        m = RE_DATE.match(label)
        if m:
            d, mo, y = m.groups()
            account_codes = [
                str(value).strip() for value in path_by_lvl.values()
                if re.fullmatch(r'\d{2}(?:\.\d+)*', str(value).strip())
                and (str(value).strip() == acc
                     or str(value).strip().startswith(acc + '.'))
            ]
            daily.append(dict(
                date=f'20{y}-{mo}-{d}', month=f'20{y}-{mo}',
                account=max(account_codes, key=lambda value: value.count('.'))
                if account_codes else acc,
                path1=path_by_lvl.get(lvl - 1, ''),
                path2=path_by_lvl.get(lvl - 2, ''),
                path3=path_by_lvl.get(lvl - 3, ''),
                obd=obd, obk=obk))
        else:
            path_by_lvl[lvl] = label
            for k in [k for k in path_by_lvl if k > lvl]:
                del path_by_lvl[k]
            rows.append(dict(
                lvl=lvl, label=label,
                p0=path_by_lvl.get(0, ''), p1=path_by_lvl.get(1, ''),
                p2=path_by_lvl.get(2, ''), p3=path_by_lvl.get(3, ''),
                snd=snd, snk=snk, obd=obd, obk=obk, skd=skd, skk=skk))

    # контрольная проверка: СН + Об = СК по строке счета
    top = next((r for r in rows if r['label'] == acc), None)
    check_ok = None
    if top:
        diff = (top['snd'] - top['snk']) + (top['obd'] - top['obk']) \
               - (top['skd'] - top['skk'])
        check_ok = abs(diff) < 1.0
        if not check_ok:
            warnings.append(f'сч.{acc}: контрольное равенство СН+Об=СК '
                            f'не сходится на {diff:,.2f}')
    elif empty_totals is not None and any(value != 0 for value in empty_totals):
        check_ok = False
        warnings.append(
            f'сч.{acc}: противоречивая пустая ОСВ — строк с движениями и '
            'остатками нет, но строка «Итого» содержит ненулевые значения')
    elif period and empty_totals is not None:
        check_ok = True

    return dict(acc=acc, period=period, dims=dims, rows=rows, daily=daily,
                warnings=warnings, check_ok=check_ok,
                top=top or dict(snd=0, snk=0, obd=0, obk=0, skd=0, skk=0))


def parse_many(paths):
    """Парсит набор файлов -> (data по счетам, отчет аудита)."""
    data, audit = {}, []
    for p in sorted(paths):
        name = os.path.basename(p)
        if 'план' in name.lower() and 'счет' in name.lower():
            audit.append(dict(file=name, acc='—', status='пропущен (план счетов)',
                              ok=True))
            continue
        try:
            res = parse_osv(p)
            if res['acc'] in data:
                audit.append(dict(file=name, acc=res['acc'],
                                  status='ДУБЛЬ счета — файл пропущен', ok=False))
                continue
            data[res['acc']] = res
            st = 'разобран'
            if res['check_ok'] is False:
                st = 'разобран, но контрольное равенство НЕ сходится'
            if res['warnings']:
                st += '; ' + '; '.join(res['warnings'])
            audit.append(dict(file=name, acc=res['acc'], status=st,
                              ok=res['check_ok'] is not False))
        except Exception as e:  # noqa: BLE001
            audit.append(dict(file=name, acc='?', status=f'ОШИБКА: {e}', ok=False))
    return data, audit
