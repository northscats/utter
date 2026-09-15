# -*- coding: utf-8 -*-
"""Focused regressions for read-only annual report comparisons."""
import io
import os
import tempfile
import unittest
from datetime import datetime, timezone
from unittest.mock import patch

from openpyxl import Workbook


_TMP = tempfile.TemporaryDirectory(prefix="annual-report-comparison-")
os.environ["DATABASE_URL"] = "sqlite:///" + os.path.join(
    _TMP.name, "annual-report-comparison.db")

import app as app_module
from annual_key_indicators import build_annual_key_indicators
from database import (
    Company,
    HistoricalImport,
    HistoricalMetricValue,
    HistoricalRawValue,
    HistoricalYearVersion,
    User,
    db,
)
from historical_parser import canonical_metric_key
from historical_parser import _number_format_precision
from historical_parser import _source_precision
from historical_parser import source_precision_metadata
from historical_parser import parse_historical_xlsx
from historical_service import _persist_version_values


class AnnualReportComparisonTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = app_module.app
        cls.app.config.update(
            TESTING=True,
            WTF_CSRF_ENABLED=False,
            HISTORICAL_IMPORT_ENABLED=True,
        )
        cls.imap_patch = patch.object(app_module, "_ensure_imap_thread")
        cls.imap_patch.start()
        with cls.app.app_context():
            db.create_all()
            cls.admin = User(
                username="annual_report_comparison_admin",
                password_hash="x",
                role="admin",
            )
            db.session.add(cls.admin)
            db.session.flush()
            cls.company = Company(
                name="Сверка годовой книги",
                created_by=cls.admin.id,
            )
            cls.foreign = Company(
                name="Другое предприятие",
                created_by=cls.admin.id,
            )
            db.session.add_all([cls.company, cls.foreign])
            db.session.commit()
            cls.admin_id = cls.admin.id
            cls.company_id = cls.company.id
            cls.foreign_id = cls.foreign.id

    @classmethod
    def tearDownClass(cls):
        with cls.app.app_context():
            db.session.remove()
            db.drop_all()
        cls.imap_patch.stop()
        _TMP.cleanup()

    def setUp(self):
        self.ctx = self.app.app_context()
        self.ctx.push()
        db.session.rollback()
        HistoricalRawValue.query.delete()
        HistoricalMetricValue.query.delete()
        HistoricalYearVersion.query.delete()
        HistoricalImport.query.delete()
        db.session.commit()

    def tearDown(self):
        db.session.remove()
        self.ctx.pop()

    def _version(
        self, company_id=None, year=2024, version_no=1, active=True,
        status="active",
    ):
        company_id = self.company_id if company_id is None else company_id
        historical_import = HistoricalImport(
            company_id=company_id,
            uploaded_by=self.admin_id,
            original_filename=f"comparison-{company_id}-{year}.xlsx",
            stored_path="/must/not/be/read.xlsx",
            file_hash=f"comparison-{company_id}-{year}-{version_no}",
            status=status,
            created_at=datetime(year, 1, 2, tzinfo=timezone.utc),
        )
        version = HistoricalYearVersion(
            historical_import=historical_import,
            company_id=company_id,
            year=year,
            version_no=version_no,
            is_active=active,
            activated_at=datetime(year, 2, 3, tzinfo=timezone.utc),
            activated_by=self.admin_id,
        )
        db.session.add(version)
        db.session.flush()
        return version

    @staticmethod
    def _metric(
        version, section, label, value, unit,
        raw_text=None, period=None, upper=False, precision=None,
    ):
        metric = HistoricalMetricValue(
            year_version=version,
            section_key=section,
            section_label=section,
            metric_key=canonical_metric_key(label),
            metric_label=label,
            series_key="value",
            period_label=str(version.year) if period is None else period,
            value_numeric=value,
            unit=unit,
            value_text=(
                source_precision_metadata(precision, "text")
                if precision is not None else None
            ),
            is_upper_bound=upper,
        )
        metric.raw_value = HistoricalRawValue(
            raw_text=str(value) if raw_text is None else raw_text,
            sheet_name=section,
            cell_ref=f"D{version.year}",
        )
        db.session.add(metric)
        return metric

    def _autonomy_components(self, version, equity=404, assets=1000):
        self._metric(
            version, "liabilities", "Собственный капитал", equity,
            "thousand_rubles", raw_text=str(equity),
        )
        self._metric(
            version, "assets", "ВАЛЮТА БАЛАНСА", assets,
            "thousand_rubles", raw_text=str(assets),
        )

    @staticmethod
    def _cards(payload):
        return {card["id"]: card for card in payload["cards"]}

    def test_rounding_uses_raw_report_precision_and_keeps_both_values(self):
        version = self._version()
        self._autonomy_components(version)
        self._metric(
            version, "financial_stability", "Коэффициент автономии",
            0.4, "ratio", raw_text="0,4",
            precision=1,
        )
        db.session.commit()

        card = self._cards(
            build_annual_key_indicators(self.company, requested_year=2024)
        )["autonomy"]
        comparison = card["report_comparison"]
        self.assertEqual(comparison["status"], "rounding")
        self.assertEqual(comparison["year"], 2024)
        self.assertEqual(comparison["unit"], "ratio")
        self.assertAlmostEqual(comparison["calculated_value"], 0.404)
        self.assertEqual(comparison["reported_value"], 0.4)
        self.assertAlmostEqual(comparison["delta"], 0.004)
        self.assertEqual(comparison["reported_provenance"]["cell"], "D2024")
        self.assertEqual(
            comparison["calculated_provenance"][0]["component"], "assets")

    def test_missing_source_and_definition_mismatch_never_create_delta(self):
        version = self._version()
        self._autonomy_components(version)
        # A leverage coefficient is not evidence for autonomy even though it
        # is in the same published section.
        self._metric(
            version, "financial_stability", "Коэффициент финансового левериджа",
            1.2, "ratio",
        )
        db.session.commit()
        autonomy = self._cards(
            build_annual_key_indicators(self.company, requested_year=2024)
        )["autonomy"]["report_comparison"]
        self.assertEqual(autonomy["status"], "incompatible_definition")
        self.assertIsNone(autonomy["reported_value"])
        self.assertIsNone(autonomy["delta"])

        HistoricalRawValue.query.delete()
        HistoricalMetricValue.query.delete()
        db.session.commit()
        self._autonomy_components(version)
        db.session.commit()
        autonomy = self._cards(
            build_annual_key_indicators(self.company, requested_year=2024)
        )["autonomy"]["report_comparison"]
        self.assertEqual(autonomy["status"], "missing_source")
        self.assertIsNone(autonomy["reported_value"])
        self.assertIsNone(autonomy["delta"])

    def test_wrong_unit_and_period_suppress_published_value(self):
        version = self._version()
        self._autonomy_components(version)
        self._metric(
            version, "financial_stability", "Коэффициент автономии",
            0.4, "percent",
        )
        db.session.commit()
        comparison = self._cards(
            build_annual_key_indicators(self.company, requested_year=2024)
        )["autonomy"]["report_comparison"]
        self.assertEqual(comparison["status"], "incompatible_unit")
        self.assertIsNone(comparison["reported_value"])
        self.assertIsNone(comparison["delta"])
        self.assertEqual(comparison["reported_unit"], "percent")

        HistoricalRawValue.query.delete()
        HistoricalMetricValue.query.delete()
        db.session.commit()
        self._autonomy_components(version)
        self._metric(
            version, "financial_stability", "Коэффициент автономии",
            0.4, "ratio", period="2023 г.",
        )
        db.session.commit()
        comparison = self._cards(
            build_annual_key_indicators(self.company, requested_year=2024)
        )["autonomy"]["report_comparison"]
        self.assertEqual(comparison["status"], "incompatible_year")
        self.assertEqual(comparison["reported_year"], 2023)
        self.assertIsNone(comparison["reported_value"])
        self.assertIsNone(comparison["delta"])

        for interim_period in (
            "январь – июнь 2024", "2023–2024", "30.06.2024",
        ):
            HistoricalRawValue.query.delete()
            HistoricalMetricValue.query.delete()
            db.session.commit()
            self._autonomy_components(version)
            self._metric(
                version, "financial_stability", "Коэффициент автономии",
                0.4, "ratio", period=interim_period,
            )
            db.session.commit()
            comparison = self._cards(
                build_annual_key_indicators(self.company, requested_year=2024)
            )["autonomy"]["report_comparison"]
            self.assertEqual(comparison["status"], "incompatible_year")
            self.assertIsNone(comparison["reported_value"])
            self.assertIsNone(comparison["delta"])

    def test_point_endpoint_can_compare_but_icr_interval_cannot(self):
        first = self._version(year=2023)
        last = self._version(year=2024)
        for version, equity, assets, ebit, interest in (
            (first, 300, 1000, 100, 10),
            (last, 404, 1000, 20, 5),
        ):
            self._autonomy_components(version, equity=equity, assets=assets)
            self._metric(
                version, "financial_results",
                "EBIT (прибыль до % и налогов)", ebit,
                "thousand_rubles",
            )
            self._metric(
                version, "financial_results", "Проценты к уплате", interest,
                "thousand_rubles",
            )
        self._metric(
            last, "financial_stability", "Коэффициент автономии",
            0.4, "ratio", precision=1,
        )
        self._metric(
            last, "profitability", "Коэффициент покрытия процентов (ICR)",
            4, "ratio",
        )
        db.session.commit()

        cards = self._cards(build_annual_key_indicators(
            self.company, year_from=2023, year_to=2024))
        autonomy = cards["autonomy"]["report_comparison"]
        self.assertEqual(autonomy["status"], "rounding")
        self.assertEqual(autonomy["year"], 2024)
        self.assertEqual(autonomy["reported_value"], 0.4)
        icr = cards["interest_coverage"]["report_comparison"]
        self.assertEqual(icr["status"], "incompatible_period")
        self.assertIsNone(icr["reported_value"])
        self.assertIsNone(icr["delta"])
        self.assertEqual(icr["year"], 2024)

    def test_parser_persists_excel_number_format_precision_for_comparison(self):
        workbook = Workbook()
        assets = workbook.active
        assets.title = "1. Активы"
        assets.append(["Показатель", 2024])
        assets.append(["ВАЛЮТА БАЛАНСА", 1000])
        assets["B2"].number_format = "0"
        liabilities = workbook.create_sheet("2. Пассивы")
        liabilities.append(["Показатель", 2024])
        liabilities.append(["Собственный капитал", 404])
        liabilities["B2"].number_format = "0"
        stability = workbook.create_sheet("4. Фин. устойчивость")
        stability.append(["Показатель", 2024])
        stability.append(["Коэффициент автономии", 0.4])
        stability["B2"].number_format = "0.0"
        stability.append(["Коэффициент абсолютной ликвидности", "0,40"])
        stability.append(["Коэффициент финансового левериджа", 0.4])
        stability["B4"].number_format = "0.0%"
        stream = io.BytesIO()
        workbook.save(stream)

        with tempfile.NamedTemporaryFile(suffix=".xlsx") as source:
            source.write(stream.getvalue())
            source.flush()
            parsed = parse_historical_xlsx(source.name)
        row = next(
            row for section in parsed["sections"]
            if section["key"] == "financial_stability"
            for row in section["rows"]
            if row["label"] == "Коэффициент автономии"
        )
        parsed_value = row["values"][0]
        self.assertEqual(parsed_value["raw"], "0.4")
        self.assertEqual(parsed_value["raw_precision"], 1)
        self.assertEqual(
            parsed_value["raw_precision_source"], "excel_number_format")
        text_row = next(
            metric for section in parsed["sections"]
            if section["key"] == "financial_stability"
            for metric in section["rows"]
            if metric["label"] == "Коэффициент абсолютной ликвидности"
        )
        self.assertEqual(text_row["values"][0]["raw_precision"], 2)
        self.assertEqual(
            text_row["values"][0]["raw_precision_source"], "text")
        percent_format_row = next(
            metric for section in parsed["sections"]
            if section["key"] == "financial_stability"
            for metric in section["rows"]
            if metric["label"] == "Коэффициент финансового левериджа"
        )
        self.assertEqual(percent_format_row["values"][0]["value"], 0.4)
        self.assertEqual(percent_format_row["values"][0]["unit"], "ratio")
        self.assertIsNone(percent_format_row["values"][0]["raw_precision"])

        version = self._version()
        _persist_version_values(version, parsed, 2024)
        db.session.commit()
        persisted = HistoricalMetricValue.query.filter_by(
            year_version_id=version.id,
            section_key="financial_stability",
            metric_key=canonical_metric_key("Коэффициент автономии"),
        ).one()
        self.assertEqual(persisted.value_text, "__osv_source_precision__:excel_number_format:1")

        comparison = self._cards(
            build_annual_key_indicators(self.company, requested_year=2024)
        )["autonomy"]["report_comparison"]
        self.assertEqual(comparison["status"], "rounding")
        self.assertAlmostEqual(comparison["delta"], 0.004)

    def test_ambiguous_legacy_numeric_raw_does_not_claim_rounding(self):
        version = self._version()
        self._autonomy_components(version)
        # This is how an old import persisted a numeric Excel cell: the raw
        # string alone carries no evidence whether the display used one or
        # more decimal places.
        self._metric(
            version, "financial_stability", "Коэффициент автономии",
            0.4, "ratio", raw_text="0.4",
        )
        db.session.commit()
        comparison = self._cards(
            build_annual_key_indicators(self.company, requested_year=2024)
        )["autonomy"]["report_comparison"]
        self.assertEqual(comparison["status"], "comparable")
        self.assertAlmostEqual(comparison["delta"], 0.004)

    def test_precision_accepts_only_simple_numeric_formats(self):
        self.assertEqual(_number_format_precision("0"), 0)
        self.assertEqual(_number_format_precision("0.00"), 2)
        self.assertEqual(_number_format_precision("#,##0.00"), 2)
        self.assertEqual(_number_format_precision("0.0%"), 1)
        # Grouping commas are not decimal separators; fractions and
        # positive/negative multi-sections are not reliable evidence.
        self.assertIsNone(_number_format_precision("#,##0,00"))
        self.assertIsNone(_number_format_precision("# ?/?"))
        self.assertIsNone(_number_format_precision("0.00;[Red]-0.00"))
        self.assertEqual(_source_precision("0,40", "General"), (2, "text"))
        self.assertEqual(_source_precision(0.4, "0.0"), (
            1, "excel_number_format"))
        self.assertEqual(_source_precision(0.4, "0.0%"), (None, None))

    def test_comparison_uses_only_active_accepted_tenant_versions(self):
        foreign = self._version(self.foreign_id, year=2024)
        self._autonomy_components(foreign)
        self._metric(
            foreign, "financial_stability", "Коэффициент автономии",
            0.9, "ratio",
        )
        inactive = self._version(year=2024, version_no=1, active=False)
        self._autonomy_components(inactive)
        self._metric(
            inactive, "financial_stability", "Коэффициент автономии",
            0.8, "ratio",
        )
        active = self._version(year=2024, version_no=2, active=True)
        self._autonomy_components(active)
        self._metric(
            active, "financial_stability", "Коэффициент автономии",
            0.4, "ratio",
        )
        db.session.commit()

        payload = build_annual_key_indicators(self.company, requested_year=2024)
        comparison = self._cards(payload)["autonomy"]["report_comparison"]
        self.assertEqual(comparison["reported_value"], 0.4)
        self.assertEqual(
            comparison["reported_provenance"]["filename"],
            "comparison-1-2024.xlsx",
        )
        self.assertEqual(payload["sources"][0]["version"], "2")


if __name__ == "__main__":
    unittest.main()