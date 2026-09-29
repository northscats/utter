# -*- coding: utf-8 -*-
"""Acceptance tests for the transactional limits-version lifecycle.

This file intentionally owns its SQLite database.  It is useful to run by
itself while the service is being changed:

    PYTHONPATH=osv-dashboard python -m unittest test_limits_lifecycle

The parser is patched with small, semantic snapshots.  They retain the 68
column shape of the GIS REB export without requiring a workbook in every
test.
"""
import io
import json
import os
import sys
import tempfile
import unittest
from datetime import date, timedelta
from unittest.mock import patch

from openpyxl import load_workbook


_TMP = tempfile.mkdtemp(prefix="limits_lifecycle_test_")
os.environ["DATABASE_URL"] = f"sqlite:///{_TMP}/limits.db"
os.environ["INITIAL_ADMIN_PASSWORD"] = "Cedar!river7Orbit#lantern"
os.environ.setdefault(
    "ENCRYPTION_KEY",
    "MDAwMDAwMDAwMDAwMDAwMDAwMDAwMDAwMDAwMDAwMDA=",
)
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from app import app, init_db
from database import (
    LimitChange,
    LimitComparison,
    LimitDataset,
    LimitObject,
    LimitRowSnapshot,
    LimitVersion,
    User,
    db,
    utcnow,
)
from limits_service import (
    LimitsError,
    accept_version,
    create_comparison,
    export_comparison,
    parse_version,
    purge_version,
    resolve_comparison_issue,
    stage_upload,
    process_version,
    resolve_upload_dataset,
    resume_pending_versions,
)


def snapshot(value, uin="100.200", row_number=5):
    """Return one 68-column, parser-compatible synthetic export snapshot."""
    semantic = [
        ("uin", "identifier"),
        ("list_type", "programme"),
        ("ap_status", "status"),
        ("gp_code", "programme"),
        ("pp_code", "programme"),
        ("om_code", "programme"),
        ("dm_code", "programme"),
        ("limit_cost", "limit_cost"),
        ("annual_funding_2026_moscow_oblast_budget", "annual_funding"),
    ]
    columns = [
        {
            "key": key,
            "label": key,
            "category": category,
            "year": None,
            "source": None,
        }
        for key, category in semantic
    ]
    columns.extend(
        {
            "key": f"extra_{index:02d}",
            "label": f"extra_{index:02d}",
            "category": "requisite",
            "year": None,
            "source": None,
        }
        for index in range(59)
    )
    values = {
        "uin": (uin, "text"),
        "list_type": ("x", "text"),
        "ap_status": ("1", "text"),
        "gp_code": ("1", "text"),
        "pp_code": ("2", "text"),
        "om_code": ("3", "text"),
        "dm_code": ("4", "text"),
        "limit_cost": (str(value), "money"),
        "annual_funding_2026_moscow_oblast_budget": (str(value), "money"),
    }
    cells = {
        key: {
            "raw": raw,
            "value": raw,
            "kind": kind,
            "cell": f"{chr(65 + index % 26)}{row_number}",
        }
        for index, (key, (raw, kind)) in enumerate(values.items())
    }
    cells["annual_funding_2026_moscow_oblast_budget"]["cell"] = f"AP{row_number}"
    return {
        "columns": columns,
        "warnings": [],
        "metadata": {},
        "rows": [{"uin": uin, "row": row_number, "cells": cells}],
    }


def ambiguous_snapshot(rows):
    """Build two unmatched rows plus any stable row for export assertions."""
    result = snapshot("0")
    result["rows"] = []
    for row_number, uin, om_code, value in rows:
        row = snapshot(value, uin, row_number)["rows"][0]
        row["row"] = row_number
        row["cells"]["om_code"]["raw"] = om_code
        row["cells"]["om_code"]["value"] = om_code
        result["rows"].append(row)
    return result


class LimitsLifecycleAcceptanceTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        app.config.update(TESTING=True, WTF_CSRF_ENABLED=False)
        cls.fake_tombstone_sidecar = {}

        def fake_record(version_id, dataset_id, source_path='', source_sha256=''):
            cls.fake_tombstone_sidecar[int(version_id)] = {
                'version_id': int(version_id),
                'dataset_id': dataset_id,
                'source_path': source_path,
                'source_sha256': source_sha256,
            }
            return list(cls.fake_tombstone_sidecar.values())

        cls._tombstone_transport = patch(
            'durable_data.record_limit_tombstone',
            side_effect=fake_record)
        cls._tombstone_transport.start()
        cls.addClassCleanup(cls._tombstone_transport.stop)
        init_db()

    def setUp(self):
        self._app_context = app.app_context()
        self._app_context.push()
        db.session.rollback()
        db.session.remove()
        for table in reversed(db.metadata.sorted_tables):
            db.session.execute(table.delete())
        admin = User(username="limits-lifecycle-admin", password_hash="x", role="admin")
        db.session.add(admin)
        db.session.flush()
        dataset = LimitDataset(name="Lifecycle", owner_id=admin.id)
        db.session.add(dataset)
        db.session.commit()
        self.admin_id = admin.id
        self.dataset_id = dataset.id

    def tearDown(self):
        db.session.remove()
        self._app_context.pop()

    def _dataset(self):
        return db.session.get(LimitDataset, self.dataset_id)

    def _draft(self, value, previous=None, uin="100.200"):
        """Create and parse a queued version without touching a real XLSX."""
        path = os.path.join(_TMP, f"source-{self.dataset_id}-{value}-{uin.replace('.', '_')}.xlsx")
        with open(path, "wb") as stream:
            stream.write(f"source-{value}-{uin}".encode())
        version = LimitVersion(
            dataset_id=self.dataset_id,
            uploaded_by=self.admin_id,
            source_date=date(2026, 1, 1),
            source_name="synthetic.xlsx",
            source_path=path,
            sha256=(f"{value}-{uin}" * 64)[:64],
            status="queued",
            previous_version_id=previous,
        )
        db.session.add(version)
        db.session.commit()
        with patch("limits_parser.parse_xlsx", return_value=snapshot(value, uin)):
            parsed = parse_version(version.id)
        return db.session.get(LimitVersion, parsed.id)

    def _accept_first(self, value="10"):
        version = self._draft(value)
        accept_version(version.id, self.admin_id)
        return db.session.get(LimitVersion, version.id)

    def _direct_rotation_fixture(self):
        """Seed ten accepted snapshots and a review candidate for eviction tests."""
        dataset = self._dataset()
        obj = LimitObject(dataset_id=dataset.id, uin="100.200", title="Object")
        db.session.add(obj)
        db.session.flush()
        previous = None
        accepted = []
        for number in range(1, 11):
            path = os.path.join(_TMP, f"rotation-{self.dataset_id}-{number}.xlsx")
            with open(path, "wb") as stream:
                stream.write(f"rotation-{number}".encode())
            version = LimitVersion(
                dataset_id=self.dataset_id,
                uploaded_by=self.admin_id,
                source_date=date(2026, 1, number),
                source_name=f"{number}.xlsx",
                source_path=path,
                sha256=(f"rotation-{self.dataset_id}-{number}" * 8)[:64],
                status="accepted",
                previous_version_id=previous,
                accepted_by=self.admin_id,
                accepted_at=utcnow() + timedelta(seconds=number),
                snapshot_json=json.dumps(snapshot(str(number))),
            )
            db.session.add(version)
            db.session.flush()
            db.session.add(
                LimitRowSnapshot(
                    version_id=version.id,
                    object_id=obj.id,
                    source_row=5,
                    stable_key="100.200|x|1|1|2|3|4",
                    cells_json=json.dumps(snapshot(str(number))["rows"][0]["cells"]),
                )
            )
            accepted.append(version)
            previous = version.id
        obj.last_seen_version_id = accepted[-1].id
        obj.is_present = True
        dataset.current_version_id = accepted[-1].id
        candidate = LimitVersion(
            dataset_id=self.dataset_id,
            uploaded_by=self.admin_id,
            source_date=date(2026, 1, 12),
            source_name="11.xlsx",
            source_path=os.path.join(_TMP, f"rotation-{self.dataset_id}-11.xlsx"),
            sha256=(f"rotation-{self.dataset_id}-11" * 8)[:64],
            status="review",
            previous_version_id=accepted[-1].id,
            snapshot_json=json.dumps(snapshot("11")),
        )
        db.session.add(candidate)
        db.session.flush()
        db.session.add(
            LimitRowSnapshot(
                version_id=candidate.id,
                object_id=obj.id,
                source_row=5,
                stable_key="100.200|x|1|1|2|3|4",
                cells_json=json.dumps(snapshot("11")["rows"][0]["cells"]),
            )
        )
        comparison = LimitComparison(
            old_version_id=accepted[-1].id,
            new_version_id=candidate.id,
            status="ready",
            summary_json="{}",
            issues_json="[]",
        )
        db.session.add(comparison)
        db.session.flush()
        db.session.commit()
        return accepted, candidate, obj

    def test_draft_parse_is_atomic_and_has_no_limit_object_side_effect(self):
        version = self._draft("10")
        self.assertEqual("review", version.status)
        self.assertEqual(0, LimitObject.query.count())
        self.assertEqual(1, LimitRowSnapshot.query.filter_by(version_id=version.id).count())
        self.assertIsNone(LimitRowSnapshot.query.filter_by(version_id=version.id).one().object_id)
        self.assertIsNone(self._dataset().current_version_id)

    def test_first_baseline_is_accepted(self):
        first = self._accept_first()
        self.assertEqual("accepted", first.status)
        self.assertEqual(first.id, self._dataset().current_version_id)

    def test_automatic_pipeline_is_idempotent_and_audited(self):
        from database import AuditLog
        first = self._draft("10")
        with patch("max_limits.queue_accepted_changes") as queue:
            process_version(first.id)
            process_version(first.id)
            self.assertEqual(1, queue.call_count)
        self.assertEqual("accepted", first.status)
        self.assertEqual(self.admin_id, first.accepted_by)
        self.assertEqual(0, LimitComparison.query.count())
        self.assertEqual(1, AuditLog.query.filter_by(action="limits_auto_accepted").count())

    def test_stale_automatic_upload_recomputes_current_base(self):
        first = self._accept_first()
        second = self._draft("20", previous=first.id)
        third = self._draft("30", previous=first.id)
        process_version(second.id)
        process_version(third.id)
        self.assertEqual(second.id, third.previous_version_id)
        self.assertEqual(third.id, self._dataset().current_version_id)
        self.assertIsNotNone(LimitComparison.query.filter_by(
            old_version_id=second.id, new_version_id=third.id, status="ready").first())

    def test_restart_resumes_parsing_and_review_without_duplicate_effects(self):
        first = self._draft("10")
        first.status = "parsing"
        db.session.commit()
        with patch("limits_parser.parse_xlsx", return_value=snapshot("10")):
            self.assertEqual("accepted", resume_pending_versions()[0]["status"])
        self.assertEqual([], resume_pending_versions())
        self.assertEqual(1, LimitRowSnapshot.query.filter_by(version_id=first.id).count())

    def test_processing_error_keeps_current_data_and_records_reason(self):
        from limits_parser import LimitsParserError
        first = self._accept_first()
        second = self._draft("20", previous=first.id)
        second.status = "queued"
        db.session.commit()
        with patch("limits_parser.parse_xlsx", side_effect=LimitsParserError("Invalid workbook columns")):
            with self.assertRaisesRegex(ValueError, "Invalid workbook"):
                process_version(second.id)
        self.assertEqual("error", second.status)
        self.assertIn("Invalid workbook columns", second.error_text)
        self.assertEqual(first.id, self._dataset().current_version_id)

    def test_unexpected_processing_error_does_not_expose_server_details(self):
        first = self._accept_first()
        second = self._draft("20", previous=first.id)
        second.status = "queued"
        db.session.commit()
        with patch("limits_parser.parse_xlsx", side_effect=RuntimeError("private server detail")):
            with self.assertRaises(RuntimeError):
                process_version(second.id)
        self.assertEqual("error", second.status)
        self.assertIn("Обработка не завершена", second.error_text)
        self.assertNotIn("private server detail", second.error_text)
        self.assertEqual(first.id, self._dataset().current_version_id)

    def test_parallel_workers_publish_in_one_serial_sequence(self):
        import threading
        first = self._accept_first()
        ids = [self._draft(value, previous=first.id).id for value in ("20", "30")]
        errors = []
        def worker(version_id):
            with app.app_context():
                try:
                    process_version(version_id)
                except Exception as exc:
                    errors.append(str(exc))
                finally:
                    db.session.remove()
        workers = [threading.Thread(target=worker, args=(version_id,)) for version_id in ids]
        for worker_thread in workers:
            worker_thread.start()
        for worker_thread in workers:
            worker_thread.join(timeout=20)
        self.assertFalse(any(worker_thread.is_alive() for worker_thread in workers))
        self.assertEqual([], errors)
        db.session.expire_all()
        latest = db.session.get(LimitVersion, self._dataset().current_version_id)
        self.assertIn(latest.id, ids)
        self.assertEqual(set(ids) - {latest.id}, {latest.previous_version_id})

    def test_separate_processes_serialize_publication(self):
        import multiprocessing
        context = multiprocessing.get_context("fork")
        first = self._accept_first()
        ids = [self._draft(value, previous=first.id).id for value in ("20", "30")]
        db.session.remove()
        gate = context.Event()
        results = context.Queue()

        def worker(version_id):
            with app.app_context():
                db.session.remove()
                db.engine.dispose(close=False)
                gate.wait(timeout=10)
                try:
                    with patch("max_limits.queue_accepted_changes"):
                        process_version(version_id)
                    results.put(None)
                except Exception as exc:
                    results.put(str(exc))
                finally:
                    db.session.remove()

        workers = [context.Process(target=worker, args=(version_id,)) for version_id in ids]
        for worker_process in workers:
            worker_process.start()
        gate.set()
        for worker_process in workers:
            worker_process.join(timeout=20)
            if worker_process.is_alive():
                worker_process.terminate()
                worker_process.join()
                self.fail("Concurrent processing deadlocked")
            self.assertEqual(0, worker_process.exitcode)
        self.assertEqual([None, None], [results.get(timeout=5) for _ in workers])
        db.session.expire_all()
        latest = db.session.get(LimitVersion, self._dataset().current_version_id)
        self.assertIn(latest.id, ids)
        self.assertEqual(set(ids) - {latest.id}, {latest.previous_version_id})

    def test_unresolved_upload_continues_automatically_after_fix(self):
        old_data = ambiguous_snapshot([
            (5, "100.200", "01", "1"),
            (6, "100.200", "02", "2"),
        ])
        new_data = ambiguous_snapshot([(5, "100.200", "03", "3")])
        old = self._draft("10")
        old.snapshot_json = json.dumps(old_data)
        db.session.commit()
        process_version(old.id)
        new = self._draft("20", previous=old.id)
        new.snapshot_json = json.dumps(new_data)
        db.session.commit()
        process_version(new.id)
        self.assertEqual("review", new.status)
        self.assertEqual(old.id, self._dataset().current_version_id)
        comparison = LimitComparison.query.filter_by(
            old_version_id=old.id, new_version_id=new.id).one()
        issue = json.loads(comparison.issues_json)[0]
        resolve_comparison_issue(comparison.id, issue["key"],
                                 {"old_row": "100.200@5", "action": "absent"},
                                 self.admin_id)
        self.assertEqual("accepted", new.status)
        self.assertEqual(new.id, self._dataset().current_version_id)
        self.assertEqual([], json.loads(comparison.issues_json))

    def test_upload_source_is_persisted_and_never_merges_legacy_sources(self):
        from database import Setting
        extra = LimitDataset(name="Other", owner_id=self.admin_id)
        db.session.add(extra)
        db.session.commit()
        technical = resolve_upload_dataset(self.admin_id)
        self.assertNotIn(technical.id, (self.dataset_id, extra.id))
        Setting.set("max_limits_dataset_id", str(extra.id))
        db.session.commit()
        self.assertEqual(technical.id, resolve_upload_dataset(self.admin_id).id)
        self.assertEqual(3, LimitDataset.query.count())

    def test_upload_source_prefers_legacy_max_then_sole_dataset(self):
        from database import Setting
        self.assertEqual(self.dataset_id, resolve_upload_dataset(self.admin_id).id)
        db.session.delete(db.session.get(Setting, "limits_upload_dataset_id"))
        extra = LimitDataset(name="MAX", owner_id=self.admin_id)
        db.session.add(extra)
        db.session.flush()
        Setting.set("max_limits_dataset_id", str(extra.id))
        db.session.commit()
        self.assertEqual(extra.id, resolve_upload_dataset(self.admin_id).id)

    def test_empty_installation_creates_one_upload_source(self):
        db.session.delete(self._dataset())
        db.session.commit()
        first = resolve_upload_dataset(self.admin_id)
        self.assertEqual(first.id, resolve_upload_dataset(self.admin_id).id)
        self.assertEqual(1, LimitDataset.query.count())

    def test_selected_version_and_report_base_survive_rotation_and_deletion(self):
        from database import Setting
        accepted, candidate, _ = self._direct_rotation_fixture()
        base_id, selected_id = accepted[0].id, accepted[1].id
        Setting.set("max_limits_version_id", str(selected_id))
        db.session.commit()
        process_version(candidate.id)
        for version_id in (base_id, selected_id):
            self.assertIsNotNone(db.session.get(LimitVersion, version_id))
            with self.assertRaisesRegex(LimitsError, "MAX"):
                purge_version(version_id)
        self.assertEqual(str(selected_id), Setting.get("max_limits_version_id"))
        self.assertIsNone(db.session.get(LimitVersion, accepted[2].id))

    def _race_selection_with_retirement(self, rotation=False, selection_first=False):
        import threading
        from concurrent.futures import ThreadPoolExecutor
        import limits_service
        import max_dataset_scope
        from database import Setting
        accepted, candidate, _ = self._direct_rotation_fixture()
        victim_id, candidate_id = accepted[0].id, candidate.id
        entered = threading.Event()
        release = threading.Event()
        contender_started = threading.Event()
        db.session.remove()

        def select():
            with app.app_context():
                try:
                    # Deliberately cache the soon-to-be-deleted row before lock.
                    db.session.get(LimitVersion, victim_id)
                    contender_started.set()
                    max_dataset_scope.select_version(victim_id)
                    return "selected"
                except ValueError:
                    return "unavailable"
                finally:
                    db.session.remove()

        def retire():
            with app.app_context():
                try:
                    contender_started.set()
                    if rotation:
                        process_version(candidate_id)
                    else:
                        purge_version(victim_id)
                    return "retired"
                except LimitsError:
                    return "protected"
                finally:
                    db.session.remove()

        original = (max_dataset_scope.snapshot_objects if selection_first
                    else limits_service.protected_version_ids)
        def checkpoint(*args, **kwargs):
            result = original(*args, **kwargs)
            if not entered.is_set():
                entered.set()
                if not release.wait(timeout=10):
                    raise AssertionError("Retirement race checkpoint timed out")
            return result

        target = ("max_dataset_scope.snapshot_objects" if selection_first
                  else "limits_service.protected_version_ids")
        with patch(target, side_effect=checkpoint), ThreadPoolExecutor(max_workers=2) as pool:
            first = pool.submit(select if selection_first else retire)
            self.assertTrue(entered.wait(timeout=10))
            contender_started.clear()
            second = pool.submit(retire if selection_first else select)
            self.assertTrue(contender_started.wait(timeout=10))
            release.set()
            outcomes = [first.result(timeout=20), second.result(timeout=20)]
        db.session.expire_all()
        selected = Setting.get("max_limits_version_id", "")
        if selected:
            self.assertIsNotNone(db.session.get(LimitVersion, int(selected)))
        if selection_first:
            self.assertEqual(str(victim_id), selected)
            self.assertEqual("selected", outcomes[0])
            self.assertEqual("retired" if rotation else "protected", outcomes[1])
        else:
            self.assertEqual(["retired", "unavailable"], outcomes)
            self.assertIsNone(db.session.get(LimitVersion, victim_id))

    def test_selection_cannot_save_a_snapshot_during_purge(self):
        self._race_selection_with_retirement()

    def test_selection_cannot_save_a_snapshot_during_rotation(self):
        self._race_selection_with_retirement(rotation=True)

    def test_purge_waits_for_selection_commit_and_preserves_snapshot(self):
        self._race_selection_with_retirement(selection_first=True)

    def test_rotation_waits_for_selection_commit_and_preserves_snapshot(self):
        self._race_selection_with_retirement(rotation=True, selection_first=True)

    def test_second_accepted_update_preserves_financial_changes(self):
        first = self._accept_first()
        second = self._draft("15", previous=first.id)
        comparison = create_comparison(first.id, second.id)
        self.assertEqual("ready", comparison.status)
        accepted = accept_version(second.id, self.admin_id)
        self.assertEqual("accepted", accepted.status)
        self.assertEqual(second.id, self._dataset().current_version_id)
        self.assertTrue(LimitChange.query.filter_by(comparison_id=comparison.id).count())

    def test_duplicate_stage_returns_existing_version(self):
        from werkzeug.datastructures import FileStorage

        dataset = self._dataset()
        first, duplicate = stage_upload(
            dataset,
            FileStorage(stream=io.BytesIO(b"same workbook"), filename="limits.xlsx"),
            "limits.xlsx",
            date(2026, 1, 1),
            self.admin_id,
            _TMP,
        )
        second, duplicate_again = stage_upload(
            dataset,
            FileStorage(stream=io.BytesIO(b"same workbook"), filename="limits.xlsx"),
            "limits.xlsx",
            date(2026, 1, 1),
            self.admin_id,
            _TMP,
        )
        self.assertFalse(duplicate)
        self.assertTrue(duplicate_again)
        self.assertEqual(first.id, second.id)
        self.assertEqual(1, LimitVersion.query.count())

    def test_resolved_added_absent_rows_export_with_metadata_and_sources(self):
        old_data = ambiguous_snapshot([
            (5, "100.200", "01", "1"),
            (6, "100.200", "02", "2"),
            (7, "300.400", "01", "3"),
            (9, "200.300", "01", "10"),
        ])
        new_data = ambiguous_snapshot([
            (5, "100.200", "03", "3"),
            (7, "300.400", "03", "4"),
            (8, "300.400", "04", "5"),
            (9, "200.300", "01", "20"),
        ])
        old = LimitVersion(
            dataset_id=self.dataset_id,
            uploaded_by=self.admin_id,
            source_date=date(2026, 1, 1),
            source_name="old.xlsx",
            source_path=os.path.join(_TMP, "old-export.xlsx"),
            sha256="e" * 64,
            status="accepted",
            snapshot_json=json.dumps(old_data),
        )
        new = LimitVersion(
            dataset_id=self.dataset_id,
            uploaded_by=self.admin_id,
            source_date=date(2026, 2, 1),
            source_name="new.xlsx",
            source_path=os.path.join(_TMP, "new-export.xlsx"),
            sha256="f" * 64,
            status="review",
            snapshot_json=json.dumps(new_data),
        )
        db.session.add_all([old, new])
        db.session.commit()
        db.session.get(LimitDataset, self.dataset_id).current_version_id = old.id
        new.previous_version_id = old.id
        db.session.commit()
        comparison = create_comparison(old.id, new.id)
        issue = json.loads(comparison.issues_json)[0]
        self.assertEqual("ambiguous_row_match", issue["code"])
        from app import _limits_history_ready
        from database import Setting
        _limits_history_ready.set()
        Setting.set('limits_history_ready', 'ready')
        db.session.commit()
        app.config['WTF_CSRF_ENABLED'] = False
        client = app.test_client()
        with client.session_transaction() as session:
            session['_user_id'] = str(self.admin_id)
            session['_fresh'] = True
            session['password_version'] = 1
        decisions = [
            {"old_row": "100.200@5", "action": "absent"},
            {"new_row": "300.400@8", "action": "added"},
        ]
        for decision in decisions:
            comparison = db.session.get(LimitComparison, comparison.id)
            issue = next(
                item for item in json.loads(comparison.issues_json)
                if item["uin"] == ("100.200" if decision.get("old_row") else "300.400")
            )
            page = client.get(f'/limit-notifications/comparisons/{comparison.id}')
            self.assertEqual(200, page.status_code)
            selected = decision.get('old_row') or decision['new_row']
            self.assertIn(f'value="{selected}"', page.get_data(as_text=True))
            response = client.post(f'/limit-notifications/comparisons/{comparison.id}/resolve',
                                   data=dict(decision, issue_key=issue['key']))
            self.assertEqual(302, response.status_code)
        comparison = db.session.get(LimitComparison, comparison.id)
        self.assertEqual("ready", comparison.status)
        self.assertEqual([], json.loads(comparison.issues_json))
        response = client.post(f'/limit-notifications/versions/{new.id}/accept', data={'confirm':'1'})
        self.assertEqual(302, response.status_code)
        self.assertEqual('accepted', db.session.get(LimitVersion, new.id).status)

        workbook = load_workbook(export_comparison(comparison.id), data_only=True)
        self.assertEqual([
            "Сводка", "Новые объекты", "Отсутствующие объекты",
            "Финансирование", "Сроки", "Реквизиты и коды", "Проблемы данных",
        ], workbook.sheetnames)
        summary = {
            row[0]: row[1]
            for row in workbook["Сводка"].iter_rows(min_row=2, values_only=True)
            if row[0]
        }
        self.assertEqual(f"{old.id} / {old.source_date}", summary["Версия до"])
        self.assertEqual(f"{new.id} / {new.source_date}", summary["Версия после"])

        requisite_rows = list(
            workbook["Реквизиты и коды"].iter_rows(min_row=2, values_only=True)
        )
        self.assertEqual(
            ["row_absent", "row_added"],
            sorted(row[2] for row in requisite_rows if row[2] in {"row_absent", "row_added"}),
        )
        self.assertEqual(
            {
                ("100.200@5", ""),
                ("", "300.400@8"),
            },
            {(row[11] or "", row[12] or "") for row in requisite_rows
             if row[2] in {"row_absent", "row_added"}},
        )

        funding_rows = list(
            workbook["Финансирование"].iter_rows(min_row=2, values_only=True)
        )
        changed_funding = next(row for row in funding_rows if row[0] == "200.300")
        self.assertEqual("10", changed_funding[7])
        self.assertEqual("20", changed_funding[8])
        self.assertEqual("AP9", changed_funding[11])
        self.assertEqual("AP9", changed_funding[12])
        self.assertEqual(["УИН", "Категория", "Поле", "Наименование", "Год",
                          "Источник", "Единица", "Было", "Стало", "Разница",
                          "Изменение, %", "Ячейка до", "Ячейка после"],
                         [cell.value for cell in workbook["Финансирование"][1]])
        self.assertEqual(1, workbook["Проблемы данных"].max_row)

    def test_rotation_deletes_old_version_and_comparisons_both_sides(self):
        accepted, candidate, obj = self._direct_rotation_fixture()
        old = accepted[0]
        other = accepted[1]
        old_path = old.source_path
        reverse = LimitComparison(
            old_version_id=other.id,
            new_version_id=old.id,
            status="ready",
            summary_json="{}",
            issues_json="[]",
        )
        db.session.add(reverse)
        db.session.flush()
        db.session.add(
            LimitChange(
                comparison_id=reverse.id,
                object_id=obj.id,
                uin=obj.uin,
                category="money",
                field_key="limit_cost",
                field_label="Cost",
                old_json='"1"',
                new_json='"2"',
            )
        )
        db.session.commit()
        with open(old_path, "ab") as stream:
            stream.write(b"source exists")
        accept_version(candidate.id, self.admin_id)
        self.assertIsNone(db.session.get(LimitVersion, old.id))
        self.assertFalse(os.path.exists(old_path))
        self.assertIsNotNone(db.session.get(LimitObject, obj.id))
        self.assertFalse(
            LimitComparison.query.filter(
                (LimitComparison.old_version_id == old.id)
                | (LimitComparison.new_version_id == old.id)
            ).first()
        )
        self.assertIsNone(db.session.get(LimitVersion, other.id).previous_version_id)

    def test_stale_comparison_is_rejected_after_concurrent_accept(self):
        first, _, obj = self._direct_rotation_fixture()
        # The fixture has ten accepted versions; use its current version as
        # the common base and prepare two competing review versions.
        current = first[-1]
        stale = LimitVersion(
            dataset_id=self.dataset_id,
            uploaded_by=self.admin_id,
            source_date=date(2026, 2, 1),
            source_name="stale.xlsx",
            source_path=os.path.join(_TMP, "stale.xlsx"),
            sha256="s" * 64,
            status="review",
            previous_version_id=current.id,
            snapshot_json=json.dumps(snapshot("20")),
        )
        concurrent = LimitVersion(
            dataset_id=self.dataset_id,
            uploaded_by=self.admin_id,
            source_date=date(2026, 2, 2),
            source_name="concurrent.xlsx",
            source_path=os.path.join(_TMP, "concurrent.xlsx"),
            sha256="c" * 64,
            status="review",
            previous_version_id=current.id,
            snapshot_json=json.dumps(snapshot("21")),
        )
        db.session.add_all([stale, concurrent])
        db.session.flush()
        for version in (stale, concurrent):
            db.session.add(
                LimitRowSnapshot(
                    version_id=version.id,
                    object_id=obj.id,
                    source_row=5,
                    stable_key="100.200|x|1|1|2|3|4",
                    cells_json=json.dumps(snapshot("20")["rows"][0]["cells"]),
                )
            )
            db.session.add(
                LimitComparison(
                    old_version_id=current.id,
                    new_version_id=version.id,
                    status="ready",
                    summary_json="{}",
                    issues_json="[]",
                )
            )
        db.session.commit()
        accept_version(concurrent.id, self.admin_id)
        with self.assertRaises(LimitsError):
            accept_version(stale.id, self.admin_id)
        self.assertEqual(concurrent.id, self._dataset().current_version_id)
        self.assertEqual("review", db.session.get(LimitVersion, stale.id).status)


if __name__ == "__main__":
    unittest.main()