# -*- coding: utf-8 -*-
import os
import tempfile
import unittest
import zipfile

from openpyxl import Workbook

from historical_parser import (
    HistoricalParserError,
    canonical_metric_key,
    parse_historical_xlsx,
)


SAMPLE = os.path.join(
    os.path.dirname(os.path.dirname(__file__)),
    "attached_assets",
    "Balashikha_Kommunalnye_Sistemy_2022_2025_1788419445681.xlsx",
)


class HistoricalParserSampleTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.result = parse_historical_xlsx(SAMPLE)

    def test_all_thematic_sheets_and_years_are_discovered(self):
        self.assertEqual(15, self.result["preview"]["section_count"])
        self.assertEqual([2022, 2023, 2024, 2025], self.result["preview"]["years"])
        self.assertFalse(self.result["unknown_sections"])
        self.assertIn("assets", {s["key"] for s in self.result["sections"]})
        self.assertIn("bankruptcy_forecast", {s["key"] for s in self.result["sections"]})

    def test_summary_is_reconciliation_only(self):
        self.assertEqual(0, self.result["reconciliation"]["emitted_values"])
        self.assertGreater(self.result["reconciliation"]["compared_values"], 0)
        self.assertNotIn(
            "Сводная таблица",
            {section["sheet_name"] for section in self.result["sections"]},
        )

    def test_raw_missing_upper_bound_and_money(self):
        values = self.result["values"]
        upper = next(value for value in values if value["raw"] == "<0,01")
        self.assertEqual(0.01, upper["value"])
        self.assertTrue(upper["is_upper_bound"])
        missing = next(value for value in values if value["raw"] == "–")
        self.assertIsNone(missing["value"])
        self.assertTrue(missing["is_missing"])
        revenue = next(
            value for value in values
            if value["section_key"] == "financial_results"
            and value["metric_key"] == canonical_metric_key("Выручка")
            and value["year"] == 2022
        )
        self.assertEqual("thousand_rubles", revenue["unit"])
        self.assertEqual("6 669 900", revenue["raw"])

    def test_required_groups_validate(self):
        self.assertTrue(self.result["validation"]["ok"])
        checks = {c["group"]: c["ok"] for c in self.result["validation"]["checks"]}
        self.assertEqual(
            {"balance": True, "financial_results": True, "liquidity": True},
            checks,
        )


class HistoricalParserSyntheticTests(unittest.TestCase):
    def _save(self, workbook):
        handle = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        handle.close()
        workbook.save(handle.name)
        self.addCleanup(lambda: os.path.exists(handle.name) and os.unlink(handle.name))
        return handle.name

    def test_bare_years_unknowns_and_confirmed_aliases(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Новая секция"
        ws.append(["Показатель", 2023, 2024])
        ws.append(["Новый показатель", "—", "<0,01"])
        path = self._save(wb)
        unknown = parse_historical_xlsx(path)
        self.assertEqual(1, unknown["preview"]["unknown_section_count"])
        self.assertEqual([2023, 2024], unknown["preview"]["years"])

        confirmed = parse_historical_xlsx(path, {
            "sections": {"Новая секция": "liquidity"},
            "metrics": {
                "liquidity": {"Новый показатель": "custom_confirmed_metric"}
            },
        })
        row = confirmed["sections"][0]["rows"][0]
        self.assertTrue(row["is_known"])
        self.assertEqual("custom_confirmed_metric", row["key"])
        self.assertEqual(2, len(confirmed["aliases_applied"]))

    def test_money_in_label_is_normalized_to_thousand_rubles(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "15. Динамика"
        ws.append(["Показатель", "2022 → 2023"])
        ws.append(["Изменение выручки, млн руб.", "1,5"])
        parsed = parse_historical_xlsx(self._save(wb))
        value = parsed["sections"][0]["rows"][0]["values"][0]
        self.assertEqual(1500.0, value["value"])
        self.assertEqual("thousand_rubles", value["unit"])
        self.assertEqual("1,5", value["raw"])

    def test_corrupt_xlsx_has_clear_error(self):
        handle = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        handle.write(b"not an xlsx")
        handle.close()
        self.addCleanup(lambda: os.path.exists(handle.name) and os.unlink(handle.name))
        with self.assertRaisesRegex(HistoricalParserError, "XLSX"):
            parse_historical_xlsx(handle.name)

    def test_required_rows_with_only_missing_values_do_not_validate(self):
        wb = Workbook()
        assets = wb.active
        assets.title = "1. Активы"
        assets.append(["Показатель", "2025"])
        assets.append(["ВАЛЮТА БАЛАНСА", "–"])
        liabilities = wb.create_sheet("2. Пассивы")
        liabilities.append(["Показатель", "2025"])
        liabilities.append(["ВАЛЮТА БАЛАНСА", "–"])
        financial = wb.create_sheet("8. Фин. результаты")
        financial.append(["Показатель", "2025"])
        financial.append(["Выручка", "–"])
        financial.append(["Чистая прибыль (убыток)", "–"])
        liquidity = wb.create_sheet("6. Ликвидность")
        liquidity.append(["Показатель", "2025"])
        liquidity.append(["Коэффициент текущей (общей) ликвидности", "–"])
        liquidity.append(["Коэффициент быстрой (промежуточной) ликвидности", "–"])
        liquidity.append(["Коэффициент абсолютной ликвидности", "–"])
        parsed = parse_historical_xlsx(self._save(wb))
        self.assertFalse(parsed["validation"]["ok"])
        self.assertTrue(all(
            not item["ok"] for item in parsed["validation"]["checks"]))

    def test_archive_member_limit_is_checked_before_openpyxl(self):
        handle = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        handle.close()
        self.addCleanup(lambda: os.path.exists(handle.name) and os.unlink(handle.name))
        with zipfile.ZipFile(handle.name, "w") as archive:
            for index in range(2001):
                archive.writestr(f"padding/{index}.txt", "")
        with self.assertRaisesRegex(HistoricalParserError, "too many archive members"):
            parse_historical_xlsx(handle.name)


if __name__ == "__main__":
    unittest.main()