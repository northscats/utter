# -*- coding: utf-8 -*-
"""Safe, non-destructive parsing of supplier XLSX forms for the Ministry cabinet."""
import hashlib
import math
import re
import zipfile
from datetime import date, datetime
from io import BytesIO

from openpyxl import load_workbook


MAX_WORKBOOK_BYTES = 25 * 1024 * 1024
MAX_EXPANDED_BYTES = 120 * 1024 * 1024
MAX_WORKBOOK_ROWS = 120_000
MAX_WORKBOOK_COLUMNS = 80


class ReportImportError(ValueError):
    """A user-correctable supplier workbook validation failure."""


def _text(value):
    if value is None:
        return ''
    return str(value).strip()


def _cell_json(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        return str(value)
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def _header_key(value):
    return re.sub(r'\s+', ' ', _text(value).replace('\n', ' ')).strip().casefold()


def _column(headers, predicate):
    for index, header in enumerate(headers):
        if predicate(_header_key(header)):
            return index
    return None


def _number(value):
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, bool):
        return float(value)
    if isinstance(value, (int, float)):
        number = float(value)
        return number if math.isfinite(number) else None
    text = str(value).strip().replace('\u00a0', '').replace(' ', '').replace(',', '.')
    try:
        number = float(text)
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) else None


def _inn(value):
    if value is None:
        return ''
    digits = re.sub(r'\D', '', str(value))
    return digits if len(digits) in (10, 12) else ''


def _is_group_row(row):
    if row['sheet_kind'] == 'population':
        return True
    label = f"{row['debtor_name']} {row['consumer_detail']}".casefold().replace('ё', 'е')
    return (
        not row['debtor_name']
        or any(marker in label for marker in ('прочие', 'итого', 'всего'))
    )


def _readable_xlsx(raw):
    if len(raw) > MAX_WORKBOOK_BYTES:
        raise ReportImportError('Размер XLSX превышает допустимые 25 МБ.')
    if raw.startswith(bytes.fromhex('d0cf11e0a1b11ae1')):
        raise ReportImportError(
            'Книга не читается как XLSX и может быть зашифрована или защищена паролем. '
            'Попросите владельца прислать незашифрованную XLSX-копию; пароль здесь не обходится.')
    try:
        archive = zipfile.ZipFile(BytesIO(raw))
    except (zipfile.BadZipFile, OSError) as exc:
        raise ReportImportError(
            'Файл не является корректной XLSX-книгой. Если книга защищена паролем, '
            'попросите владельца прислать незашифрованную копию.') from exc
    with archive:
        entries = archive.infolist()
        if len(entries) > 500 or sum(item.file_size for item in entries) > MAX_EXPANDED_BYTES:
            raise ReportImportError('Книга превышает допустимый объём распакованных данных.')
        if any(item.flag_bits & 0x1 for item in entries):
            raise ReportImportError(
                'Зашифрованная книга не поддерживается. Попросите владельца прислать '
                'незашифрованную XLSX-копию; пароль не обходится.')
        names = set(archive.namelist())
        if '[Content_Types].xml' not in names or 'xl/workbook.xml' not in names:
            raise ReportImportError(
                'В архиве нет структуры XLSX-книги. Файл может быть зашифрован '
                'или иметь неверное расширение.')
    try:
        workbook = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:
        message = str(exc).lower()
        if any(word in message for word in ('encrypt', 'password', 'decrypt', 'unsupported')):
            raise ReportImportError(
                'Не удалось открыть книгу: она зашифрована или защищена. '
                'Попросите владельца предоставить обычную XLSX-копию.') from exc
        raise ReportImportError(
            'Не удалось прочитать XLSX-книгу. Проверьте её в Excel и сохраните повторно.') from exc
    return workbook


def parse_supplier_workbook(raw):
    """Parse recognized source sheets without collapsing duplicates or blank values."""
    if not raw:
        raise ReportImportError('Выберите XLSX-файл.')
    workbook = _readable_xlsx(raw)
    try:
        enterprise_sheet = None
        population_sheet = None
        for sheet in workbook.worksheets:
            title = _header_key(sheet.title)
            if 'предприяти' in title:
                enterprise_sheet = sheet
            elif 'населени' in title:
                population_sheet = sheet

        metadata_sheet = enterprise_sheet or workbook.worksheets[0]
        metadata = next(metadata_sheet.iter_rows(min_row=2, max_row=2, values_only=True), ())
        month_value = _text(metadata[6] if len(metadata) > 6 else '')
        year_value = _number(metadata[7] if len(metadata) > 7 else None)
        resource = _text(metadata[4] if len(metadata) > 4 else '')
        source_supplier = _text(metadata[1] if len(metadata) > 1 else '')
        source_supplier_inn = _inn(metadata[2] if len(metadata) > 2 else None)
        try:
            month = int(month_value)
        except (TypeError, ValueError):
            month = 0
        if month not in range(1, 13) or year_value is None:
            raise ReportImportError(
                'В строке заголовка не указан корректный отчётный месяц и год.')
        year = int(year_value)
        if year < 2000 or year > 2100:
            raise ReportImportError('Год отчёта должен быть между 2000 и 2100.')
        if not resource:
            raise ReportImportError('В строке заголовка не указан вид ТЭР/услуги.')

        parsed_rows = []
        warnings = []
        recognized_sheets = []
        blank_source_rows = 0
        total_source_rows = 0

        sheets = []
        if enterprise_sheet is not None:
            sheets.append(('enterprise', enterprise_sheet))
        if population_sheet is not None:
            sheets.append(('population', population_sheet))
        if not sheets:
            for sheet in workbook.worksheets:
                sheets.append(('enterprise', sheet))

        for sheet_kind, sheet in sheets:
            if sheet.max_row and sheet.max_row > MAX_WORKBOOK_ROWS:
                raise ReportImportError(
                    f'Лист «{sheet.title}» превышает лимит в {MAX_WORKBOOK_ROWS:,} строк.')
            if sheet.max_column and sheet.max_column > MAX_WORKBOOK_COLUMNS:
                raise ReportImportError(
                    f'Лист «{sheet.title}» превышает лимит в {MAX_WORKBOOK_COLUMNS} столбцов.')
            header_row = None
            headers = []
            for row_number, values in enumerate(
                    sheet.iter_rows(min_row=1, max_row=min(sheet.max_row or 1, 15),
                                    values_only=True), start=1):
                first = _header_key(values[0] if values else '')
                if first in ('№ строки', 'номер строки'):
                    header_row = row_number
                    headers = list(values)
                    break
            if header_row is None:
                continue
            recognized_sheets.append(sheet.title)
            indices = {
                'row_no': _column(headers, lambda h: h in ('№ строки', 'номер строки')),
                'name': _column(headers, lambda h: h.startswith('наименование потребителя')),
                'inn': _column(headers, lambda h: h.startswith('инн потребителя')),
                'kpp': _column(headers, lambda h: h.startswith('кпп потребителя')),
                'status': _column(headers, lambda h: h.startswith('состояние')),
                'detail': _column(headers, lambda h: h.startswith('детализация потребителя')),
                'territory': _column(headers, lambda h: h.startswith('административная принадлежность')),
                'debt': _column(headers, lambda h: 'просроченная задолженность' in h
                                and 'всего' in h),
                'periods': _column(headers, lambda h: 'количество периодов просрочки' in h),
                'current_accrual': _column(headers, lambda h: h.startswith(
                    'начислено в отчетном периоде')),
                'ytd_accrual': _column(headers, lambda h: h.startswith(
                    'начислено с начала года')),
            }
            if indices['name'] is None and indices['detail'] is None:
                continue
            if indices['territory'] is None:
                warnings.append(f'На листе «{sheet.title}» нет поля административной принадлежности.')
            if indices['debt'] is None:
                warnings.append(f'На листе «{sheet.title}» не найден столбец просроченной задолженности.')

            for row_number, values in enumerate(
                    sheet.iter_rows(min_row=header_row + 2,
                                    max_col=min(sheet.max_column or len(headers),
                                                MAX_WORKBOOK_COLUMNS),
                                    values_only=True), start=header_row + 2):
                total_source_rows += 1
                original = [_cell_json(value) for value in values]
                if not any(value not in (None, '') for index, value in enumerate(original)
                           if index != indices['row_no']):
                    blank_source_rows += 1
                    continue
                def value_for(key):
                    index = indices.get(key)
                    return values[index] if index is not None and index < len(values) else None

                raw_inn = value_for('inn')
                raw_name = value_for('name')
                raw_detail = value_for('detail')
                name = _text(raw_name) if raw_name is not None else ''
                detail = _text(raw_detail)
                if sheet_kind == 'population':
                    name = ''
                debt_value = value_for('debt')
                parsed_row = {
                    'sheet_kind': sheet_kind,
                    'sheet_name': sheet.title,
                    'source_row_number': row_number,
                    'inn': _inn(raw_inn),
                    'kpp': re.sub(r'\D', '', str(value_for('kpp') or '')),
                    'debtor_name': name,
                    'territory': _text(value_for('territory')),
                    'consumer_detail': detail,
                    'status_text': _text(value_for('status')),
                    'debt_total': _number(debt_value),
                    'overdue_periods': _number(value_for('periods')),
                    'current_period_accrual': _number(value_for('current_accrual')),
                    'year_to_date_accrual': _number(value_for('ytd_accrual')),
                    'source_values': original,
                    'is_data_row': True,
                }
                parsed_row['is_group_row'] = _is_group_row(parsed_row)
                parsed_rows.append(parsed_row)

        if not recognized_sheets:
            raise ReportImportError(
                'Не найдены таблицы формы поставщика: ожидается столбец «№ строки» '
                'и наименование потребителя или детализация населения.')
        if not parsed_rows:
            warnings.append('В распознанных таблицах нет заполненных строк данных.')
        if blank_source_rows:
            warnings.append(
                f'Пустые строки шаблона пропущены: {blank_source_rows}. '
                'Незаполненные значения в сохранённых строках не превращаются в нули.')
        if not any(row['debt_total'] is not None for row in parsed_rows):
            warnings.append(
                'В строках нет распознанных значений просроченной задолженности; '
                'пустые поля не считаются нулём.')

        missing_inn = sum(
            not row['inn'] and not _is_group_row(row)
            for row in parsed_rows if row['sheet_kind'] == 'enterprise')
        if missing_inn:
            warnings.append(
                f'Строки без пригодного ИНН сохранены отдельно: {missing_inn}.')
        inn_counts = {}
        for row in parsed_rows:
            if row['inn']:
                inn_counts[row['inn']] = inn_counts.get(row['inn'], 0) + 1
        duplicate_inn = sum(count > 1 for count in inn_counts.values())
        if duplicate_inn:
            warnings.append(
                f'В файле повторяются ИНН ({duplicate_inn} значений); строки не объединялись.')

        return {
            'year': year,
            'month': month,
            'resource': resource[:120],
            'source_supplier_name': source_supplier[:255],
            'source_supplier_inn': source_supplier_inn,
            'rows': parsed_rows,
            'row_count': total_source_rows,
            'data_row_count': len(parsed_rows),
            'warnings': list(dict.fromkeys(warnings)),
            'file_sha256': hashlib.sha256(raw).hexdigest(),
        }
    finally:
        workbook.close()
