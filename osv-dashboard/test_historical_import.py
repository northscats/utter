# -*- coding: utf-8 -*-
"""End-to-end regressions for staged, versioned historical imports."""
import io
import json
import os
import sys
import tempfile
import threading
from pathlib import Path

from openpyxl import Workbook


_TMP = tempfile.mkdtemp(prefix="historical_import_test_")
os.environ["DATABASE_URL"] = f"sqlite:///{_TMP}/test.db"
os.environ["ENCRYPTION_KEY"] = (
    "MDAwMDAwMDAwMDAwMDAwMDAwMDAwMDAwMDAwMDAwMDA=")
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from app import app, init_db
from database import (
    Company,
    HistoricalImport,
    HistoricalMetricValue,
    HistoricalRawValue,
    HistoricalYearVersion,
    Upload,
    User,
    db,
)


app.config.update(TESTING=True, WTF_CSRF_ENABLED=False)
init_db()


def check(label, condition):
    print(("PASS" if condition else "FAIL"), label)
    if not condition:
        raise AssertionError(label)


def login_as(client, user_id):
    with client.session_transaction() as session:
        session["_user_id"] = str(user_id)
        session["_fresh"] = True


def post_file(client, company_id, data, name="history.xlsx"):
    return client.post(
        f"/company/{company_id}/historical/preview",
        data={"historical_file": (io.BytesIO(data), name)},
        content_type="multipart/form-data",
    )


def partial_book(unknown=False):
    wb = Workbook()
    assets = wb.active
    assets.title = "1. Активы"
    assets.append(["Показатель", "31.12.2025"])
    assets.append(["ВАЛЮТА БАЛАНСА", "1 000"])
    if unknown:
        assets.append(["Совершенно новый показатель", "<0,01"])
    liabilities = wb.create_sheet("2. Пассивы")
    liabilities.append(["Показатель", "31.12.2025"])
    liabilities.append(["ВАЛЮТА БАЛАНСА", "1 000"])
    finance = wb.create_sheet("8. Фин. результаты")
    finance.append(["Показатель", "2025"])
    finance.append(["Выручка", "500"])
    finance.append(["Чистая прибыль (убыток)", "25"])
    liquidity = wb.create_sheet("6. Ликвидность")
    liquidity.append(["Показатель", "2025"])
    liquidity.append(["Коэффициент текущей (общей) ликвидности", "1,2"])
    liquidity.append(["Коэффициент быстрой (промежуточной) ликвидности", "0,8"])
    liquidity.append(["Коэффициент абсолютной ликвидности", "0,1"])
    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()


with app.app_context():
    admin = User.query.filter_by(role="admin").first()
    company = Company(name="Историческое предприятие", created_by=admin.id)
    db.session.add(company)
    db.session.commit()
    admin_id, company_id = admin.id, company.id

client = app.test_client()
login_as(client, admin_id)
sample = Path(
    os.path.dirname(os.path.dirname(__file__)),
    "attached_assets",
    "Balashikha_Kommunalnye_Sistemy_2022_2025_1788419445681.xlsx",
).read_bytes()

# Пример: предпросмотр не затрагивает действующие данные.
r = post_file(client, company_id, sample, "Балашиха.xlsx")
check("пример принят в предпросмотр", r.status_code == 302)
with app.app_context():
    attempt = HistoricalImport.query.order_by(HistoricalImport.id.desc()).first()
    attempt_id = attempt.id
    check("в примере найдены все листы и годы",
          attempt.preview["preview"]["section_count"] == 15
          and attempt.preview["preview"]["years"] == [2022, 2023, 2024, 2025])
    check("сводная таблица не создаёт значений",
          attempt.preview["reconciliation"]["emitted_values"] == 0)
    check("до подтверждения нет активных версий",
          HistoricalYearVersion.query.count() == 0)
preview_page = client.get(
    f"/company/{company_id}?historical_preview={attempt_id}")
preview_body = preview_page.get_data(as_text=True)
check("карточка отображает полный предпросмотр",
      preview_page.status_code == 200
      and "Годовая историческая аналитика" in preview_body
      and "Разделов: 15" in preview_body
      and "Подтвердить и активировать версии" in preview_body)

r = client.post(
    f"/company/{company_id}/historical/{attempt_id}/confirm",
    data={"mapping_count": "0"},
)
check("пример подтверждён", r.status_code == 302)
with app.app_context():
    active = HistoricalYearVersion.query.filter_by(is_active=True).all()
    check("созданы четыре активных годовых среза",
          {row.year for row in active} == {2022, 2023, 2024, 2025})
    upper_metric = HistoricalMetricValue.query.filter_by(
        is_upper_bound=True).first()
    missing_metric = HistoricalMetricValue.query.filter_by(
        is_missing=True).first()
    check("верхняя граница сохранена отдельно от исходной строки",
          upper_metric is not None
          and upper_metric.value_numeric == 0.01
          and upper_metric.raw_value.raw_text == "<0,01")
    check("тире сохранено как исходный пропуск",
          missing_metric is not None
          and missing_metric.value_numeric is None
          and missing_metric.raw_value.raw_text == "–")
    check("исходная строка есть у каждого значения",
          HistoricalRawValue.query.count() == HistoricalMetricValue.query.count())
    db.session.add(Upload(
        company_id=company_id,
        uploaded_by=admin_id,
        period_label="январь – декабрь 2025",
        metrics_json=json.dumps({
            "liquidity": 0.93,
            "dz_days": 74.0,
            "adm_pct": 8.433378878613446,
        }),
    ))
    db.session.commit()
    from historical_analytics import build_expert_analytics
    from historical_executive import build_executive_summary
    methodical = build_expert_analytics(company_id, "methodical")
    osv_control = build_expert_analytics(company_id, "osv_control")
    executive = build_executive_summary(methodical, 2025)
    executive_2022 = build_executive_summary(methodical, 2022)
    executive_conflict = build_executive_summary(osv_control, 2025)
    check("чекпоинт охватывает все 15 тематических разделов",
          len(methodical["sections"]) == 15
          and {section["key"] for section in methodical["sections"]}
          == {
              "assets", "liabilities", "net_assets",
              "financial_stability", "working_capital", "liquidity",
              "asset_liability", "financial_results", "profitability",
              "turnover", "rating", "insolvent_structure",
              "creditworthiness", "bankruptcy_forecast", "dynamics",
          })
    liquidity = next(
        metric for metric in methodical["metrics"]
        if metric["section_key"] == "liquidity"
        and metric["label"].startswith("Коэффициент текущей"))
    check("методически разные показатели ОСВ не сравниваются",
          liquidity["values_by_year"][2025]["osv"] is None
          and liquidity["values_by_year"][2025]["discrepancy"] is None)
    control_metric = next(
        metric for metric in methodical["metrics"]
        if metric["key"] == "control_adm_pct")
    methodical_2025 = control_metric["values_by_year"][2025]
    controlled_metric = next(
        metric for metric in osv_control["metrics"]
        if metric["key"] == "control_adm_pct")
    controlled_2025 = controlled_metric["values_by_year"][2025]
    expected_historical_control = 779881.0 / 9708071.0 * 100.0
    check("профиль выбирает основной источник без смешивания значений",
          methodical_2025["primary_source"] == "historical"
          and methodical_2025["primary"]["value"]
          == expected_historical_control
          and methodical_2025["secondary"]["value"] == 8.433378878613446
          and controlled_2025["primary_source"] == "osv"
          and controlled_2025["primary"]["value"] == 8.433378878613446)
    check("расхождение использует точные числа, 1% и абсолютный порог профиля",
          methodical_2025["discrepancy"]["threshold"] == 0.5
          and not methodical_2025["discrepancy"]["significant"]
          and controlled_2025["discrepancy"]["threshold"] == 0.25
          and controlled_2025["discrepancy"]["significant"])
    db.session.add(Upload(
        company_id=company_id,
        uploaded_by=admin_id,
        period_label="январь – июнь 2025",
        metrics_json=json.dumps({"adm_pct": 99.0}),
    ))
    db.session.commit()
    after_partial = build_expert_analytics(company_id, "osv_control")
    partial_control = next(
        metric for metric in after_partial["metrics"]
        if metric["key"] == "control_adm_pct")
    check("неполная ОСВ не подменяет годовой сопоставимый источник",
          partial_control["values_by_year"][2025]["osv"]["value"]
          == 8.433378878613446)
    no_formula = next(
        metric for metric in methodical["metrics"]
        if not metric["formula_confirmed"])
    trace = liquidity["values_by_year"][2025]["historical"]["provenance"]
    check("каталог различает подтверждённую и неподтверждённую формулу",
          liquidity["formula_confirmed"] and no_formula["formula"] is None)
    check("происхождение содержит источник, версию, лист и ячейку",
          trace["source"] == "Загруженная отчётность"
          and trace["version"] == "v1"
          and trace["sheet"] == "6. Ликвидность"
          and isinstance(trace["row"], int)
          and bool(trace["cell"]))
    check("исторический рейтинг отделён от платформенного",
          methodical["historical_rating"]
          and not methodical["platform_rating_included"])
    check("модели банкротства представлены отдельными дорожками",
          [lane["key"] for lane in methodical["bankruptcy_models"]]
          == ["altman", "taffler", "saifullin"])
    check("руководительский чекпоинт строит восемь KPI за 2022–2025",
          executive["years"] == [2022, 2023, 2024, 2025]
          and executive["year"] == 2025
          and len(executive["cards"]) == 8
          and [card["title"] for card in executive["cards"]][:3]
          == ["Выручка", "Прибыль от продаж", "EBIT / покрытие процентов"])
    check("карточки содержат норму, медиану, источник и компактную динамику",
          all(
              {"trend", "norm_label", "median_delta", "source", "sparkline"}
              <= set(card)
              for card in executive["cards"])
          and any(card["sparkline"] for card in executive["cards"]))
    check("медиана раннего года не использует будущие значения",
          executive_2022["cards"][0]["median_delta"] == "+0,0%")
    check("сигналы ограничены пятью и конфликт источников прозрачен",
          len(executive_conflict["signals"]) <= 5
          and any(
              signal["label"] == "Конфликт источников"
              for signal in executive_conflict["signals"]))
executive_page = client.get(f"/company/{company_id}/historical/expert")
executive_body = executive_page.get_data(as_text=True)
check("руководительский режим является начальным и сохраняет выбранный год",
      executive_page.status_code == 200
      and "Резюме руководителя" in executive_body
      and "Главное за 2025 год" in executive_body
      and "view=expert" in executive_body
      and "year=2025" in executive_body)
embedded_page = client.get(
    f"/company/{company_id}/historical/expert?embedded=1&year=2024")
embedded_body = embedded_page.get_data(as_text=True)
check("встроенный режим сохраняется в переключателях и фильтрах",
      embedded_page.status_code == 200
      and "embedded=1" in embedded_body
      and 'name="embedded" value="1"' in embedded_body)
expert_page = client.get(
    f"/company/{company_id}/historical/expert?view=expert&year=2025")
expert_body = expert_page.get_data(as_text=True)
check("экспертный экран отображает профиль, годы и проверяемые значения",
      expert_page.status_code == 200
      and "Единый методический профиль" in expert_body
      and "Финансовое положение" in expert_body
      and "Исторический рейтинг" in expert_body
      and "Модели банкротства — без среднего балла" in expert_body
      and "Значение по загруженной отчетности; формула не подтверждена"
      in expert_body)
check("экспертный экран разделяет исходные значения и единицы",
      "Исходное значение" in expert_body
      and "Нормализованное" not in expert_body
      and "Единица изм." in expert_body
      and "26,1 %" in expert_body
      and "26,1% %" not in expert_body
      and "thousand_rubles" not in expert_body
      and "percent" not in expert_body
      and '"ratio"' not in expert_body)
profile_response = client.post(
    f"/company/{company_id}/historical/expert/profile",
    data={"profile_key": "osv_control"},
    follow_redirects=True,
)
check("смена профиля сохраняется для предприятия",
      profile_response.status_code == 200
      and 'value="osv_control" selected' in profile_response.get_data(
          as_text=True))
client.post(
    f"/company/{company_id}/historical/expert/profile",
    data={"profile_key": "methodical"},
)
with app.app_context():
    missing_version = HistoricalYearVersion.query.filter_by(
        company_id=company_id, year=2023, is_active=True).one()
    missing_version.is_active = False
    db.session.commit()
    with_gap = build_expert_analytics(company_id, "methodical")
    gap_metric = next(
        metric for metric in with_gap["metrics"]
        if metric["section_key"] == "liquidity"
        and metric["label"].startswith("Коэффициент текущей"))
    check("промежуточный год остаётся явным разрывом без интерполяции",
          2023 in with_gap["years"]
          and 2023 in with_gap["missing_years"]
          and gap_metric["values_by_year"][2023]["primary"] is None)
    gap_summary = build_executive_summary(with_gap, 2023)
    check("при пропущенном году вывод не формируется",
          gap_summary["year"] == 2023
          and not gap_summary["conclusion"]["allowed"]
          and gap_summary["quality"]["label"] == "Недостаточно")
    missing_version.is_active = True
    db.session.commit()
history_page = client.get(f"/company/{company_id}")
history_body = history_page.get_data(as_text=True)
check("карточка показывает действующие версии и исходный файл",
      history_page.status_code == 200
      and "Годовые версии" in history_body
      and "действующая" in history_body
      and "Балашиха.xlsx" in history_body)

# Повторная загрузка создаёт v2; возврат включает v1 без удаления v2.
r = post_file(client, company_id, sample, "Балашиха повтор.xlsx")
with app.app_context():
    repeat_id = HistoricalImport.query.order_by(
        HistoricalImport.id.desc()).first().id
client.post(
    f"/company/{company_id}/historical/{repeat_id}/confirm",
    data={"mapping_count": "0"},
)
with app.app_context():
    versions_2025 = HistoricalYearVersion.query.filter_by(
        company_id=company_id, year=2025).order_by(
            HistoricalYearVersion.version_no).all()
    v1_id = versions_2025[0].id
    check("повторная загрузка создаёт новую версию",
          [row.version_no for row in versions_2025] == [1, 2]
          and versions_2025[1].is_active)
client.post(
    f"/company/{company_id}/historical/version/{v1_id}/activate")
with app.app_context():
    versions_2025 = HistoricalYearVersion.query.filter_by(
        company_id=company_id, year=2025).order_by(
            HistoricalYearVersion.version_no).all()
    check("возврат не удаляет последующую версию",
          len(versions_2025) == 2
          and versions_2025[0].is_active
          and not versions_2025[1].is_active)
client.post(
    f"/company/{company_id}/historical/version/{v1_id}/activate")
with app.app_context():
    check("повторная активация действующей версии идемпотентна",
          HistoricalYearVersion.query.filter_by(
              company_id=company_id, year=2025, is_active=True).count() == 1
          and db.session.get(HistoricalYearVersion, v1_id).is_active)

# Частичная книга требует явного подтверждения, хотя обязательные блоки есть.
post_file(client, company_id, partial_book(), "частичная.xlsx")
with app.app_context():
    partial = HistoricalImport.query.order_by(HistoricalImport.id.desc()).first()
    partial_id = partial.id
    before = HistoricalYearVersion.query.filter_by(
        company_id=company_id, year=2025, is_active=True).one().id
r = client.post(
    f"/company/{company_id}/historical/{partial_id}/confirm",
    data={"mapping_count": "0"},
)
check("частичный импорт без подтверждения отклонён",
      "historical_error=" in r.headers.get("Location", ""))
with app.app_context():
    check("отклонённый импорт не переключил версию",
          HistoricalYearVersion.query.filter_by(
              company_id=company_id, year=2025, is_active=True).one().id == before)
client.post(
    f"/company/{company_id}/historical/{partial_id}/confirm",
    data={"mapping_count": "0", "confirm_partial": "1"},
)
with app.app_context():
    partial = db.session.get(HistoricalImport, partial_id)
    check("явно подтверждённый частичный импорт активирован",
          partial.status == "active" and partial.partial_confirmed)

# Неизвестная строка требует явного сопоставления.
post_file(client, company_id, partial_book(unknown=True), "неизвестная.xlsx")
with app.app_context():
    unknown = HistoricalImport.query.order_by(HistoricalImport.id.desc()).first()
    unknown_id = unknown.id
    check("неизвестная строка видима в предпросмотре",
          unknown.preview["unknown_metrics"][0]["raw"]
          == "Совершенно новый показатель")
r = client.post(
    f"/company/{company_id}/historical/{unknown_id}/confirm",
    data={"mapping_count": "0", "confirm_partial": "1"},
)
check("неизвестная строка без сопоставления отклонена",
      "historical_error=" in r.headers.get("Location", ""))
r = client.post(
    f"/company/{company_id}/historical/{unknown_id}/confirm",
    data={
        "mapping_count": "1",
        "mapping_source_0": "Совершенно новый показатель",
        "mapping_kind_0": "metric",
        "mapping_label_0": "Новый подтверждённый показатель",
        "mapping_canonical_0": "confirmed_custom_metric",
        "confirm_partial": "1",
    },
)
check("подтверждённое сопоставление импортировано", r.status_code == 302)

# Два параллельных подтверждения одной попытки: версии создаёт только одно.
with app.app_context():
    concurrent_company = Company(
        name="Параллельное подтверждение", created_by=admin_id)
    db.session.add(concurrent_company)
    db.session.commit()
    concurrent_company_id = concurrent_company.id
post_file(
    client, concurrent_company_id, partial_book(),
    "параллельная.xlsx")
with app.app_context():
    concurrent_attempt_id = HistoricalImport.query.filter_by(
        company_id=concurrent_company_id).one().id
barrier = threading.Barrier(2)
concurrent_results = []


def confirm_concurrently():
    from historical_service import HistoricalImportError, confirm_and_activate
    with app.app_context():
        item = db.session.get(HistoricalImport, concurrent_attempt_id)
        barrier.wait()
        try:
            confirm_and_activate(
                item, admin_id, [], confirm_partial=True)
            concurrent_results.append("success")
        except HistoricalImportError:
            db.session.rollback()
            concurrent_results.append("already_processed")


threads = [
    threading.Thread(target=confirm_concurrently),
    threading.Thread(target=confirm_concurrently),
]
for thread in threads:
    thread.start()
for thread in threads:
    thread.join(timeout=20)
check("параллельные подтверждения завершились",
      all(not thread.is_alive() for thread in threads))
with app.app_context():
    check("одну попытку подтверждает только один запрос",
          sorted(concurrent_results) == ["already_processed", "success"]
          and HistoricalYearVersion.query.filter_by(
              company_id=concurrent_company_id, year=2025).count() == 1
          and HistoricalYearVersion.query.filter_by(
              company_id=concurrent_company_id, year=2025,
              is_active=True).count() == 1)

# Повреждённая и отменённая загрузки не меняют активный срез.
with app.app_context():
    active_before = HistoricalYearVersion.query.filter_by(
        company_id=company_id, year=2025, is_active=True).one().id
post_file(client, company_id, b"not xlsx", "сломанный.xlsx")
with app.app_context():
    broken = HistoricalImport.query.order_by(HistoricalImport.id.desc()).first()
    check("повреждённый XLSX хранится как неуспешная попытка",
          broken.status == "failed" and bool(broken.error_text))
    check("ошибка не меняет активный срез",
          HistoricalYearVersion.query.filter_by(
              company_id=company_id, year=2025, is_active=True).one().id
          == active_before)
post_file(client, company_id, partial_book(), "отмена.xlsx")
with app.app_context():
    cancelled_id = HistoricalImport.query.order_by(
        HistoricalImport.id.desc()).first().id
client.post(
    f"/company/{company_id}/historical/{cancelled_id}/cancel")
with app.app_context():
    check("отмена не меняет активную версию",
          db.session.get(HistoricalImport, cancelled_id).status == "cancelled"
          and HistoricalYearVersion.query.filter_by(
              company_id=company_id, year=2025, is_active=True).one().id
          == active_before)

# Независимый переключатель возвращает прежний экспертный начальный экран.
with app.app_context():
    version_count = HistoricalYearVersion.query.filter_by(
        company_id=company_id).count()
app.config["HISTORICAL_EXECUTIVE_ENABLED"] = False
legacy_expert = client.get(f"/company/{company_id}/historical/expert")
check("руководительский режим можно отключить без потери экспертного экрана",
      legacy_expert.status_code == 200
      and "Историческая аналитика (Эксперт)" in legacy_expert.get_data(
          as_text=True)
      and "Резюме руководителя" not in legacy_expert.get_data(as_text=True))
with app.app_context():
    check("выключение руководительского режима не меняет версии",
          HistoricalYearVersion.query.filter_by(
              company_id=company_id).count() == version_count)
app.config["HISTORICAL_EXECUTIVE_ENABLED"] = True

# Общий переключатель исторической аналитики блокирует новый вход.
with app.app_context():
    osv_count = Upload.query.filter_by(company_id=company_id).count()
    version_count = HistoricalYearVersion.query.filter_by(
        company_id=company_id).count()
app.config["HISTORICAL_EXPERT_ENABLED"] = False
expert_hidden_body = client.get(
    f"/company/{company_id}").get_data(as_text=True)
expert_blocked = client.get(
    f"/company/{company_id}/historical/expert")
check("экспертный вход скрывается отдельно от импортированных версий",
      "Годовая историческая аналитика" in expert_hidden_body
      and ">Эксперт<" not in expert_hidden_body
      and expert_blocked.status_code == 404)
with app.app_context():
    check("скрытие экспертного режима сохраняет накопленные версии",
          HistoricalYearVersion.query.filter_by(
              company_id=company_id).count() == version_count)
app.config["HISTORICAL_EXPERT_ENABLED"] = True
app.config["HISTORICAL_IMPORT_ENABLED"] = False
body = client.get(f"/company/{company_id}").get_data(as_text=True)
blocked = post_file(client, company_id, partial_book(), "выключено.xlsx")
check("выключатель скрывает исторический контур",
      "Годовая историческая аналитика" not in body and blocked.status_code == 404)
with app.app_context():
    check("поток ОСВ не изменён", Upload.query.filter_by(
        company_id=company_id).count() == osv_count)

print("\nВСЕ ТЕСТЫ ИСТОРИЧЕСКОГО ИМПОРТА ПРОЙДЕНЫ")