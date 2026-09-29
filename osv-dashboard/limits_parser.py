# -*- coding: utf-8 -*-
"""Read, compare and export GIS REB limits workbooks without application state.

The public functions in this module return only JSON-compatible values.  In
particular, a monetary value is a canonical decimal *string*, never a float:
``"82746.07"``.  This keeps the module suitable for both background jobs and
audit exports without introducing a database dependency.
"""
from __future__ import annotations

from collections import Counter, defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
import re
import unicodedata
import zipfile
from zipfile import BadZipFile
from xml.etree import ElementTree

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font


class LimitsParserError(ValueError):
    """Raised when a workbook is unsafe or does not have the limits schema."""


MAX_XLSX_MEMBERS = 2_000
MAX_XLSX_UNCOMPRESSED = 100 * 1024 * 1024
MAX_XLSX_MEMBER_SIZE = 25 * 1024 * 1024
MAX_XLSX_COMPRESSION_RATIO = 500
MAX_XLSX_CELLS = 500_000
EXPECTED_COLUMNS = 68

_SPACE_RE = re.compile(r"[\s\u00a0]+")
_DASH_RE = re.compile(r"^[\s\-–—−]+$")
_YEAR_RE = re.compile(r"(?:на\s+01\s+01\s+|в\s+)(20\d{2})(?:\s+году)?")
_FORMULA_PREFIXES = ("=", "+", "-", "@")
# Real GIS REB UINs include values such as ``1000006610.Ш000000001``;
# validate their full two-part identifier form without inventing a
# digits-only restriction.
_UIN_RE = re.compile(r"^[^\s.]+\.[^\s.]+$")

# These keys are deliberately about the report semantics, not its Excel
# positions.  They are also the only programme identifiers used for matching.
PROGRAMME_KEYS = (
    "list_type", "ap_status", "gp_code", "pp_code", "om_code", "dm_code",
)
_BASIC_HEADERS = {
    "вид перечня": ("list_type", "programme"),
    "признак актуальности ап": ("ap_status", "status"),
    "код гп": ("gp_code", "programme"),
    "код пп": ("pp_code", "programme"),
    "код ом": ("om_code", "programme"),
    "код дм": ("dm_code", "programme"),
    "уин": ("uin", "identifier"),
    "наименование объекта из ап": ("object_name", "requisite"),
    "адрес": ("address", "requisite"),
    "наименование го": ("municipality_name", "requisite"),
    "вид собственности муниципальная собственность": (
        "municipal_ownership", "requisite"),
    "вид собственности собственность субъекта российской федерации": (
        "regional_ownership", "requisite"),
    "тип объекта нси": ("object_type", "requisite"),
    "отрасль объекта код": ("industry_code", "requisite"),
    "вид работ из справочника гис рэб": ("work_type", "requisite"),
    "цио грбс": ("cio_grbs", "requisite"),
}
_SOURCE_KEYS = {
    "средства федерального бюджета": "federal_budget",
    "средства бюджета московской области": "moscow_oblast_budget",
    "средства бюджета г москвы": "moscow_city_budget",
    "средства бюджетов муниципальных образований московской области":
        "municipal_budgets_moscow_oblast",
    "внебюджетные средства": "extrabudgetary",
    "средства фонда развития территорий": "territorial_development_fund",
}


def _normalise_text(value):
    """Normalize only harmless presentation differences for comparisons."""
    return _SPACE_RE.sub(" ", str(value or "").strip())


def _header_text(value):
    text = unicodedata.normalize("NFKC", _normalise_text(value)).lower()
    text = text.replace("ё", "е").replace("№", " ")
    text = re.sub(r"[()/,.:;—–−-]", " ", text)
    return _SPACE_RE.sub(" ", text).strip()


def _slug(value):
    return re.sub(r"_+", "_", re.sub(r"[^\w]+", "_", _header_text(value))).strip("_")


def _decimal_string(value):
    """Return a non-exponent canonical Decimal string or raise ValueError."""
    if isinstance(value, bool):
        raise ValueError("boolean is not a money value")
    if isinstance(value, Decimal):
        number = value
    elif isinstance(value, int):
        number = Decimal(value)
    elif isinstance(value, float):
        # str(), unlike Decimal(float), preserves the workbook's displayed
        # decimal representation rather than binary floating-point noise.
        number = Decimal(str(value))
    else:
        text = _normalise_text(value).replace("\u00a0", "").replace(" ", "")
        text = text.replace(",", ".")
        number = Decimal(text)
    if not number.is_finite():
        raise ValueError("non-finite money value")
    if number == 0:
        return "0"
    return format(number.normalize(), "f")


def _date_value(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = _normalise_text(value)
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    return text


def _cell_value(value, category):
    """Represent blank, dash and numeric zero as three distinct values."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return {"raw": "" if value is None else str(value), "value": None, "kind": "blank"}
    raw = str(value)
    if isinstance(value, str) and _DASH_RE.match(value):
        return {"raw": raw, "value": None, "kind": "dash"}
    if category in {"limit_cost", "funding_snapshot", "annual_funding"}:
        try:
            return {"raw": raw, "value": _decimal_string(value), "kind": "money"}
        except (InvalidOperation, ValueError):
            # A non-numeric value in a money column remains auditable and is
            # reported by the caller rather than quietly coerced to zero.
            return {"raw": raw, "value": _normalise_text(value), "kind": "text"}
    if category == "date":
        return {"raw": raw, "value": _date_value(value), "kind": "date"}
    return {"raw": raw, "value": _normalise_text(value), "kind": "text"}


def _archive_is_safe(path):
    try:
        with zipfile.ZipFile(path) as archive:
            members = archive.infolist()
            if len(members) > MAX_XLSX_MEMBERS:
                raise LimitsParserError("XLSX has too many archive members")
            total = 0
            for member in members:
                if member.file_size > MAX_XLSX_MEMBER_SIZE:
                    raise LimitsParserError("XLSX archive member is too large")
                total += member.file_size
                if member.compress_size and (
                    member.file_size / member.compress_size > MAX_XLSX_COMPRESSION_RATIO
                ):
                    raise LimitsParserError("XLSX archive compression ratio is unsafe")
            if total > MAX_XLSX_UNCOMPRESSED:
                raise LimitsParserError("XLSX archive is too large when unpacked")
    except (BadZipFile, OSError) as exc:
        raise LimitsParserError("Invalid XLSX archive") from exc


def _funding_source(header):
    normalized = _header_text(header)
    # Match known source phrases in the entire semantic header.  Splitting
    # straight after "тыс. руб." is incorrect because the period (for example
    # "в 2026 году") sits between that unit and the source.
    for label, key in _SOURCE_KEYS.items():
        if label in normalized:
            return key
    # Unknown sources remain deterministic and are made from the tail after
    # removing the report's unit and period wording.
    marker = "тыс руб"
    source = normalized.split(marker, 1)[-1].strip() if marker in normalized else ""
    source = re.sub(r"^(?:на\s+01\s+01\s+20\d{2}|в\s+20\d{2}\s+году)\s*", "", source)
    if not source:
        raise LimitsParserError(f"Funding header has no source: {header!r}")
    return _slug(source)


def _column_from_header(label):
    normalized = _header_text(label)
    if normalized.startswith("признак актуальности ап"):
        return "ap_status", "status", None, None
    if normalized in _BASIC_HEADERS:
        key, category = _BASIC_HEADERS[normalized]
        return key, category, None, None
    if normalized.startswith("предельная стоимость объекта") and "тыс руб" in normalized:
        return "limit_cost", "limit_cost", None, None
    if normalized.startswith("плановые даты дата начала работ"):
        return "work_start_date", "date", None, None
    if normalized.startswith("плановые даты дата завершения работ"):
        return "work_completion_date", "date", None, None
    if normalized.startswith("плановые даты дата открытия"):
        return "opening_date", "date", None, None
    if normalized.startswith("признак ввода открытия"):
        return "opening_marker", "requisite", None, None
    if not normalized.startswith("объем финансирования") or "тыс руб" not in normalized:
        raise LimitsParserError(f"Unrecognised limits header: {label!r}")
    year_match = _YEAR_RE.search(normalized)
    if not year_match:
        raise LimitsParserError(f"Funding header has no year: {label!r}")
    year = int(year_match.group(1))
    source = _funding_source(label)
    if "на 01 01" in normalized:
        return f"funding_as_of_{year}_{source}", "funding_snapshot", year, source
    if re.search(r"\bв\s+20\d{2}\s+году\b", normalized):
        return f"annual_funding_{year}_{source}", "annual_funding", year, source
    raise LimitsParserError(f"Funding header has no period type: {label!r}")


def _discover_columns(workbook):
    if len(workbook.worksheets) != 1:
        raise LimitsParserError("Limits workbook must contain exactly one worksheet")
    sheet = workbook.worksheets[0]
    for candidate_row in range(1, min(sheet.max_row, 60) + 1):
        cells = [cell for cell in sheet[candidate_row] if cell.value is not None]
        if not any(_header_text(cell.value) == "уин" for cell in cells):
            continue
        if len(cells) != EXPECTED_COLUMNS:
            continue
        try:
            columns = []
            seen = set()
            for cell in cells:
                key, category, year, source = _column_from_header(cell.value)
                if key in seen:
                    raise LimitsParserError(f"Duplicate semantic header: {key}")
                seen.add(key)
                columns.append({
                    "key": key, "label": str(cell.value), "category": category,
                    "year": year, "source": source, "_index": cell.column,
                })
        except LimitsParserError:
            continue
        required = {"uin", *PROGRAMME_KEYS, "limit_cost", "work_start_date",
                    "work_completion_date", "opening_date"}
        if required.issubset(seen) and len(columns) == EXPECTED_COLUMNS:
            return sheet, candidate_row, columns
    raise LimitsParserError(
        "Could not find all 68 semantic limits headers (including UIN and funding)"
    )


def _validate_formula_caches(path):
    """Reject formulas which have no cached value, before importing any data.

    Walking two openpyxl read-only workbooks was prohibitively expensive for
    the supplied 2,500-row forms.  XLSX stores formulas and their cache in the
    worksheet XML, so inspect that bounded archive directly instead.
    """
    try:
        with zipfile.ZipFile(path) as archive:
            worksheet_members = [
                member for member in archive.namelist()
                if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", member)
            ]
            for member in worksheet_members:
                with archive.open(member) as source:
                    for _, cell in ElementTree.iterparse(source, events=("end",)):
                        if cell.tag.rsplit("}", 1)[-1] != "c":
                            continue
                        children = {child.tag.rsplit("}", 1)[-1]: child for child in cell}
                        if ("f" in children and
                                ("v" not in children or children["v"].text is None)):
                            raise LimitsParserError(
                                f"Formula without a saved cached value at {member}:{cell.get('r')}"
                            )
                        cell.clear()
    except (BadZipFile, OSError, ElementTree.ParseError) as exc:
        raise LimitsParserError("XLSX formula cache could not be inspected") from exc


def parse_xlsx(path, progress=None):
    """Parse a GIS REB limits XLSX into an immutable, JSON-compatible snapshot."""
    if progress:
        progress('opening')
    _archive_is_safe(path)
    _validate_formula_caches(path)
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise LimitsParserError("XLSX could not be opened") from exc
    try:
        sheet, header_row, internal_columns = _discover_columns(workbook)
        if sheet.max_row * len(internal_columns) > MAX_XLSX_CELLS:
            raise LimitsParserError("XLSX exceeds the cell safety limit")
        columns = [{key: value for key, value in column.items() if key != "_index"}
                   for column in internal_columns]
        warnings, rows = [], []
        total_rows = max(0, sheet.max_row - header_row - 1)
        if progress:
            progress('rows', 0, total_rows)
        first_column = min(column["_index"] for column in internal_columns)
        last_column = max(column["_index"] for column in internal_columns)
        # Read-only worksheets cannot service repeated ``sheet.cell`` lookups
        # efficiently: each lookup can replay the XML stream.  A single row
        # iterator keeps parsing the supplied forms linear in their cells.
        for excel_row, row_cells in enumerate(sheet.iter_rows(
            min_row=header_row + 2, max_row=sheet.max_row,
            min_col=first_column, max_col=last_column,
        ), start=header_row + 2):
            cells_by_index = {cell.column: cell for cell in row_cells}
            values = {}
            nonblank = False
            for column in internal_columns:
                cell = cells_by_index[column["_index"]]
                parsed = _cell_value(cell.value, column["category"])
                parsed["cell"] = cell.coordinate
                values[column["key"]] = parsed
                nonblank = nonblank or parsed["kind"] != "blank"
                if column["category"] in {"limit_cost", "funding_snapshot", "annual_funding"} and parsed["kind"] == "text":
                    warnings.append({
                        "code": "non_numeric_money", "row": excel_row,
                        "cell": cell.coordinate, "message": "Money cell is not numeric",
                    })
            if progress:
                progress('rows', excel_row - header_row - 1, total_rows)
            if not nonblank:
                continue
            uin = values["uin"]["value"]
            if not uin or not _UIN_RE.fullmatch(str(uin)):
                warnings.append({
                    "code": "invalid_uin", "row": excel_row,
                    "cell": values["uin"]["cell"],
                    "message": "Non-empty row has missing or invalid UIN",
                })
            rows.append({"uin": uin, "row": excel_row, "cells": values})
        invalid = [warning for warning in warnings if warning["code"] == "invalid_uin"]
        if invalid:
            coordinates = ", ".join(warning["cell"] for warning in invalid[:5])
            raise LimitsParserError(
                f"Non-empty data rows have missing or invalid UIN: {coordinates}"
            )
        return {
            "columns": columns, "rows": rows, "warnings": warnings,
            "metadata": {
                "sheet": sheet.title, "header_row": header_row,
                "data_row": header_row + 2, "schema": "gis_reb_limits_v1",
            },
        }
    finally:
        workbook.close()


def _row_reference(row):
    return f"{row.get('uin') or ''}@{row['row']}"


def _stable_key(row):
    cells = row["cells"]
    return tuple(cells.get(key, {}).get("value") for key in PROGRAMME_KEYS)


def _programme(row):
    return {key: row["cells"].get(key, {}).get("value") for key in PROGRAMME_KEYS}


def _resolution_pairs(resolutions):
    """Parse explicit transfer, independent absent, and independent added decisions."""
    result = {"transfers": {}, "absent": set(), "added": set()}
    if not resolutions:
        return result
    if isinstance(resolutions, dict) and "transfers" in resolutions:
        transfer_items = resolutions["transfers"]
        if not isinstance(transfer_items, list):
            raise LimitsParserError("resolutions.transfers must be a list")
        for item in transfer_items:
            if not isinstance(item, dict):
                raise LimitsParserError("Each resolution must be an object")
            old_ref, new_ref, action = item.get("old_row"), item.get("new_row"), item.get("action")
            if action == "absent" and old_ref and not new_ref:
                old_ref = str(old_ref)
                if old_ref in result["absent"]:
                    raise LimitsParserError(f"Duplicate absent resolution for old row {old_ref}")
                result["absent"].add(old_ref)
            elif action == "added" and new_ref and not old_ref:
                new_ref = str(new_ref)
                if new_ref in result["added"]:
                    raise LimitsParserError(f"Duplicate added resolution for new row {new_ref}")
                result["added"].add(new_ref)
            elif old_ref and new_ref and action in (None, "transfer"):
                old_ref, new_ref = str(old_ref), str(new_ref)
                if old_ref in result["transfers"]:
                    raise LimitsParserError(f"Duplicate resolution for old row {old_ref}")
                result["transfers"][old_ref] = new_ref
            else:
                raise LimitsParserError(
                    "Resolution is a transfer (old_row + new_row), absent (old_row), or added (new_row)"
                )
    elif isinstance(resolutions, dict):
        for old_ref, new_ref in resolutions.items():
            result["transfers"][str(old_ref)] = str(new_ref)
    elif isinstance(resolutions, list):
        for item in resolutions:
            if not isinstance(item, dict):
                raise LimitsParserError("Each resolution must be an object")
            old_ref, new_ref, action = item.get("old_row"), item.get("new_row"), item.get("action")
            if action == "absent" and old_ref and not new_ref:
                old_ref = str(old_ref)
                if old_ref in result["absent"]:
                    raise LimitsParserError(f"Duplicate absent resolution for old row {old_ref}")
                result["absent"].add(old_ref)
            elif action == "added" and new_ref and not old_ref:
                new_ref = str(new_ref)
                if new_ref in result["added"]:
                    raise LimitsParserError(f"Duplicate added resolution for new row {new_ref}")
                result["added"].add(new_ref)
            elif old_ref and new_ref and action in (None, "transfer"):
                old_ref, new_ref = str(old_ref), str(new_ref)
                if old_ref in result["transfers"]:
                    raise LimitsParserError(f"Duplicate resolution for old row {old_ref}")
                result["transfers"][old_ref] = new_ref
            else:
                raise LimitsParserError(
                    "Resolution is a transfer (old_row + new_row), absent (old_row), or added (new_row)"
                )
    else:
        raise LimitsParserError("resolutions must be a mapping or transfer list")
    targets = list(result["transfers"].values())
    if len(targets) != len(set(targets)):
        raise LimitsParserError("Multiple old rows cannot resolve to one new row")
    if set(result["transfers"]) & result["absent"]:
        raise LimitsParserError("An old row cannot be both transferred and absent")
    if set(result["transfers"].values()) & result["added"]:
        raise LimitsParserError("A new row cannot be both transferred and added")
    return result


def _change(kind, category, uin, old_row=None, new_row=None, **extra):
    result = {
        "kind": kind, "category": category, "uin": uin,
        "old_row": old_row and _row_reference(old_row),
        "new_row": new_row and _row_reference(new_row),
        "programme": _programme(new_row or old_row) if (new_row or old_row) else {},
    }
    result.update(extra)
    return result


def _money_delta(old_cell, new_cell):
    if old_cell["kind"] != "money" or new_cell["kind"] != "money":
        return None, None
    old_value, new_value = Decimal(old_cell["value"]), Decimal(new_cell["value"])
    delta = _decimal_string(new_value - old_value)
    percent = None
    if old_value != 0:
        percent = _decimal_string((new_value - old_value) / old_value * Decimal("100"))
    return delta, percent


def _cell_category(column, old_cell, new_cell):
    category = column["category"]
    if category in {"annual_funding", "funding_snapshot"}:
        return category
    if (column["key"] == "work_completion_date" and
            old_cell.get("kind") == new_cell.get("kind") == "date" and
            old_cell.get("value", "")[:4] != new_cell.get("value", "")[:4]):
        return "completion_year"
    if category == "date":
        return "schedule"
    if category == "limit_cost":
        return "limit_cost"
    if category in {"programme", "status"}:
        return "programme"
    return "requisites"


def compare_snapshots(old, new, resolutions=None, progress=None):
    """Compare two parser snapshots using UIN, stable rows, then safe transfers.

    A transfer is auto-matched only when one unmatched old and one unmatched new
    row remain for a UIN.  Any many-to-many case is an issue and produces no
    inferred additions/removals until the caller supplies ``resolutions``.
    """
    old_rows = defaultdict(list)
    new_rows = defaultdict(list)
    for row in old.get("rows", []):
        if row.get("uin"):
            old_rows[row["uin"]].append(row)
    for row in new.get("rows", []):
        if row.get("uin"):
            new_rows[row["uin"]].append(row)
    parsed_resolutions = _resolution_pairs(resolutions)
    resolution_pairs = parsed_resolutions["transfers"]
    known_old_refs = {_row_reference(row) for rows in old_rows.values() for row in rows}
    known_new_refs = {_row_reference(row) for rows in new_rows.values() for row in rows}
    unknown_old = (set(resolution_pairs) | parsed_resolutions["absent"]) - known_old_refs
    unknown_new = (set(resolution_pairs.values()) | parsed_resolutions["added"]) - known_new_refs
    if unknown_old or unknown_new:
        raise LimitsParserError(
            "Resolution refers to unknown row(s): " +
            ", ".join(sorted(unknown_old | unknown_new))
        )
    changes, issues, matches = [], [], []
    all_uins = sorted(set(old_rows) | set(new_rows))
    blocked_uins = set()
    used_old_resolutions, used_new_resolutions = set(), set()

    if progress:
        progress('matching', 0, len(all_uins))
    for uin_index, uin in enumerate(all_uins):
        if progress:
            progress('matching', uin_index, len(all_uins))
        before, after = old_rows.get(uin, []), new_rows.get(uin, [])
        if not before:
            changes.append(_change("object_added", "object", uin, new_row=after[0]))
            continue
        if not after:
            changes.append(_change("object_absent", "object", uin, old_row=before[0]))
            continue

        unmatched_old, unmatched_new = set(range(len(before))), set(range(len(after)))
        old_by_key, new_by_key = defaultdict(list), defaultdict(list)
        for index, row in enumerate(before):
            old_by_key[_stable_key(row)].append(index)
        for index, row in enumerate(after):
            new_by_key[_stable_key(row)].append(index)
        for key in sorted(set(old_by_key) & set(new_by_key), key=str):
            if len(old_by_key[key]) == len(new_by_key[key]) == 1:
                left, right = old_by_key[key][0], new_by_key[key][0]
                matches.append((before[left], after[right], False))
                unmatched_old.discard(left)
                unmatched_new.discard(right)

        old_refs = {_row_reference(before[index]): index for index in unmatched_old}
        new_refs = {_row_reference(after[index]): index for index in unmatched_new}
        for old_ref, new_ref in resolution_pairs.items():
            if old_ref in old_refs or new_ref in new_refs:
                if old_ref not in old_refs or new_ref not in new_refs:
                    raise LimitsParserError(
                        "Resolution does not name two unmatched rows for one UIN"
                    )
                left, right = old_refs[old_ref], new_refs[new_ref]
                matches.append((before[left], after[right], True))
                unmatched_old.discard(left)
                unmatched_new.discard(right)
                used_old_resolutions.add(old_ref)
                used_new_resolutions.add(new_ref)

        for old_ref in parsed_resolutions["absent"]:
            if old_ref not in old_refs:
                continue
            left = old_refs[old_ref]
            if left not in unmatched_old:
                raise LimitsParserError("Absent resolution names a stable matched row")
            changes.append(_change("row_absent", "row", uin, old_row=before[left]))
            unmatched_old.discard(left)
            used_old_resolutions.add(old_ref)
        for new_ref in parsed_resolutions["added"]:
            if new_ref not in new_refs:
                continue
            right = new_refs[new_ref]
            if right not in unmatched_new:
                raise LimitsParserError("Added resolution names a stable matched row")
            changes.append(_change("row_added", "row", uin, new_row=after[right]))
            unmatched_new.discard(right)
            used_new_resolutions.add(new_ref)

        if len(unmatched_old) == len(unmatched_new) == 1:
            left, right = next(iter(unmatched_old)), next(iter(unmatched_new))
            matches.append((before[left], after[right], True))
            unmatched_old.clear()
            unmatched_new.clear()
        elif unmatched_old and not unmatched_new:
            for index in sorted(unmatched_old):
                changes.append(_change("row_absent", "row", uin, old_row=before[index]))
            unmatched_old.clear()
        elif unmatched_new and not unmatched_old:
            for index in sorted(unmatched_new):
                changes.append(_change("row_added", "row", uin, new_row=after[index]))
            unmatched_new.clear()
        elif unmatched_old or unmatched_new:
            blocked_uins.add(uin)
            issues.append({
                "code": "ambiguous_row_match", "uin": uin,
                "old_rows": sorted(_row_reference(before[index]) for index in unmatched_old),
                "new_rows": sorted(_row_reference(after[index]) for index in unmatched_new),
                "message": "Explicit old_row/new_row mapping is required before comparison",
            })
            continue

    expected_old_resolutions = set(resolution_pairs) | parsed_resolutions["absent"]
    expected_new_resolutions = set(resolution_pairs.values()) | parsed_resolutions["added"]
    if expected_old_resolutions != used_old_resolutions or expected_new_resolutions != used_new_resolutions:
        raise LimitsParserError("Resolution is extraneous or does not apply to an unmatched row")

    if progress:
        progress('matching', len(all_uins), len(all_uins))
        progress('comparing', 0, len(matches))
    new_columns = {column["key"]: column for column in new.get("columns", [])}
    for match_index, (before, after, transfer) in enumerate(matches):
        if progress:
            progress('comparing', match_index, len(matches))
        uin = after["uin"]
        if transfer:
            changes.append(_change(
                "row_transferred", "programme_transfer", uin, before, after,
                old_programme=_programme(before), new_programme=_programme(after),
            ))
        for key in sorted(set(before["cells"]) & set(after["cells"])):
            old_cell, new_cell = before["cells"][key], after["cells"][key]
            if (old_cell.get("kind"), old_cell.get("value")) == (
                new_cell.get("kind"), new_cell.get("value")
            ):
                continue
            column = new_columns.get(key, {"key": key, "label": key, "category": "requisite",
                                           "year": None, "source": None})
            delta, percent = _money_delta(old_cell, new_cell)
            changes.append(_change(
                "cell_changed", _cell_category(column, old_cell, new_cell), uin, before, after,
                field=key, label=column["label"], year=column.get("year"),
                source=column.get("source"), old=old_cell, new=new_cell,
                delta=delta, percent_change=percent,
                match_type="transfer" if transfer else "stable",
            ))

    changes.sort(key=lambda item: (
        item["uin"], item["kind"], item.get("field", ""), item["old_row"] or "",
        item["new_row"] or "",
    ))
    if progress:
        progress('comparing', len(matches), len(matches))
    counts = Counter(item["kind"] for item in changes)
    summary = {
        "new_objects": counts["object_added"],
        "absent_objects": counts["object_absent"],
        "new_rows": counts["row_added"],
        "absent_rows": counts["row_absent"],
        "transfers": counts["row_transferred"],
        "changed_cells": counts["cell_changed"],
        "stable_changed_cells": sum(
            1 for item in changes
            if item["kind"] == "cell_changed" and item.get("match_type") == "stable"
        ),
        "total_changes": len(changes),
        "issues": len(issues),
        "blocked_uins": sorted(blocked_uins),
    }
    return {
        "changes": changes, "issues": issues, "summary": summary,
        "metadata": {
            "old_version": old.get("metadata", {}).get("version"),
            "new_version": new.get("metadata", {}).get("version"),
            "old_schema": old.get("metadata", {}).get("schema"),
            "new_schema": new.get("metadata", {}).get("schema"),
        },
    }


def _excel_safe(value):
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        value = str(value)
    text = str(value)
    return "'" + text if text.startswith(_FORMULA_PREFIXES) else text


def _append_sheet(workbook, name, headings, rows):
    sheet = workbook.create_sheet(name)
    sheet.append(headings)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    for row in rows:
        sheet.append([_excel_safe(value) for value in row])
    for column in sheet.columns:
        letter = column[0].column_letter
        sheet.column_dimensions[letter].width = min(
            45, max(12, max(len(str(cell.value or "")) for cell in column) + 2)
        )
    return sheet


def export_comparison(comparison, old_snapshot=None, new_snapshot=None):
    """Export a comparison to the seven agreed Russian worksheets as BytesIO."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    summary = comparison.get("summary", {})
    metadata = comparison.get("metadata", {})
    _append_sheet(workbook, "Сводка", ["Показатель", "Значение"], [
        ("Версия до", metadata.get("old_version") or ""),
        ("Версия после", metadata.get("new_version") or ""),
        ("Схема до", metadata.get("old_schema") or ""),
        ("Схема после", metadata.get("new_schema") or ""),
        *[(key, value) for key, value in summary.items()],
    ])
    headings = ["УИН", "Строка", "ГП", "ПП", "ОМ", "ДМ"]
    def object_rows(kind, side):
        return [
            (item["uin"], item[f"{side}_row"], *[
                item["programme"].get(key, "") for key in ("gp_code", "pp_code", "om_code", "dm_code")
            ])
            for item in comparison.get("changes", []) if item["kind"] == kind
        ]
    _append_sheet(workbook, "Новые объекты", headings, object_rows("object_added", "new"))
    _append_sheet(workbook, "Отсутствующие объекты", headings, object_rows("object_absent", "old"))
    changed_headings = [
        "УИН", "Категория", "Поле", "Наименование", "Год", "Источник",
        "Единица", "Было", "Стало", "Разница", "Изменение, %", "Ячейка до", "Ячейка после",
    ]
    def detail_rows(categories):
        result = []
        for item in comparison.get("changes", []):
            if item["kind"] != "cell_changed" or item["category"] not in categories:
                continue
            before, after = item["old"], item["new"]
            result.append((
                item["uin"], item["category"], item["field"], item["label"],
                item.get("year") or "", item.get("source") or "",
                "тыс. руб." if item["category"] in {
                    "annual_funding", "funding_snapshot", "limit_cost"
                } else "", before["raw"],
                after["raw"], item["delta"] or "", item["percent_change"] or "",
                before["cell"], after["cell"],
            ))
        return result
    _append_sheet(workbook, "Финансирование", changed_headings,
                  detail_rows({"annual_funding", "funding_snapshot", "limit_cost"}))
    _append_sheet(workbook, "Сроки", changed_headings,
                  detail_rows({"schedule", "completion_year"}))
    requisites_rows = detail_rows({"requisites", "programme"})
    for item in comparison.get("changes", []):
        if item["kind"] not in {"row_transferred", "row_added", "row_absent"}:
            continue
        requisites_rows.append((
            item["uin"], item["category"], item["kind"], item["kind"], "", "", "",
            "", "", "", "", item.get("old_row") or "", item.get("new_row") or "",
        ))
    _append_sheet(workbook, "Реквизиты и коды", changed_headings, requisites_rows)
    problem_rows = [
        (issue.get("code", ""), issue.get("uin", ""), issue.get("message", ""),
         issue.get("old_rows", ""), issue.get("new_rows", ""))
        for issue in comparison.get("issues", [])
    ]
    for version, snapshot in (("до", old_snapshot), ("после", new_snapshot)):
        if snapshot:
            problem_rows.extend((
                f"source_warning_{warning.get('code', '')}", "",
                f"{version}: {warning.get('message', '')}",
                warning.get("cell", ""), warning.get("row", ""),
            ) for warning in snapshot.get("warnings", []))
    _append_sheet(workbook, "Проблемы данных",
                  ["Код", "УИН", "Описание", "Строки до", "Строки после"], problem_rows)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output