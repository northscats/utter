"""Isolated SQLite regressions; never touches the application's real database."""
import copy
import json
import os
import re
import tempfile
import unittest
from datetime import date
from io import BytesIO
from urllib.parse import unquote
from unittest.mock import patch

_TMP = tempfile.TemporaryDirectory(prefix='key-indicators-test-')
os.environ['DATABASE_URL'] = 'sqlite:///' + os.path.join(_TMP.name, 'test.db')

import app as app_module
from database import (Company, CompanyMembership, CostAnalyticMapping, DailyRow,
                      DebtorAnalyticMapping, RevenueAnalyticMapping, Upload,
                      User, UserModuleAccess, db)
import key_indicators_service as service
from key_indicators import build_indicator_source, revenue_analytic_key


def report(period, amount, dates):
    return {
        'acc': '51', 'period': period, 'rows': [], 'warnings': [],
        'top': {'snd': 0, 'snk': 0, 'obd': amount, 'obk': 0,
                'skd': amount, 'skk': 0},
        'daily': [{'date': day, 'month': day[:7], 'path1': '', 'path2': '',
                   'path3': '', 'obd': amount, 'obk': 0} for day in dates],
    }


class SourceAssemblyTests(unittest.TestCase):
    def test_empty_osv_uses_account_number_from_report_title(self):
        from openpyxl import Workbook
        from osvparser import parse_osv

        path = os.path.join(_TMP.name, 'ОСВ 67 за 01.07.2026 - 17.08.2026.xlsx')
        book = Workbook()
        sheet = book.active
        sheet.append(['АО «Тест»'])
        sheet.append([
            'Оборотно-сальдовая ведомость по счету 67 '
            'за 01.07.2026 - 17.08.2026'
        ])
        sheet.append([])
        sheet.append(['Счет', 'Сальдо на начало периода', None,
                      'Обороты за период', None, 'Сальдо на конец периода'])
        sheet.append(['Контрагенты', 'Дебет', 'Кредит',
                      'Дебет', 'Кредит', 'Дебет', 'Кредит'])
        sheet.append(['Период'])
        sheet.append(['Итого'])
        book.save(path)
        book.close()

        parsed = parse_osv(path)
        self.assertEqual(parsed['acc'], '67')
        self.assertEqual(parsed['rows'], [])
        self.assertEqual(parsed['daily'], [])
        self.assertEqual(parsed['top']['obd'], 0)
        self.assertEqual(len(parsed['warnings']), 1)
        self.assertIn('пустая ОСВ', parsed['warnings'][0])
        source = service.assemble_source([path], [])
        self.assertIn('67', source['accounts'])
        self.assertEqual(len(source['accounts']['67']['warnings']), 1)

    def test_annual_workbook_does_not_block_valid_osv(self):
        from openpyxl import Workbook
        annual = os.path.join(_TMP.name, 'annual-analysis.xlsx')
        book = Workbook()
        book.active.title = '1. Активы'
        book.active.append(['Показатель', '2024 г.', '2025 г.'])
        book.active.append(['Баланс', 100, 200])
        liabilities = book.create_sheet('2. Пассивы')
        liabilities.append(['Показатель', '2024 г.', '2025 г.'])
        liabilities.append(['Баланс', 100, 200])
        book.save(annual)
        book.close()
        osv = os.path.join(_TMP.name, 'valid-osv.xls')
        with open(osv, 'wb') as target:
            target.write(b'fixture')
        real_parser = service.parse_osv
        with patch.object(service, 'parse_osv', side_effect=lambda path:
                          report('Июнь 2026', 20, ['2026-06-20'])
                          if path == osv else real_parser(path)):
            source = service.assemble_source([annual, osv], [])
        self.assertEqual(list(source['accounts']), ['51'])
        self.assertEqual(source['accounts']['51']['daily'][0]['obd'], 20)
        self.assertTrue(any('годовую финансовую аналитику' in w for w in source['warnings']))

    def test_unknown_xlsx_is_not_silently_ignored(self):
        from openpyxl import Workbook
        path = os.path.join(_TMP.name, 'incomplete-osv.xlsx')
        book = Workbook()
        book.active.append(['Неполная выгрузка счёта 51'])
        book.save(path)
        book.close()
        with self.assertRaises(service.SourceReadError) as caught:
            service.assemble_source([path], [])
        self.assertIn('incomplete-osv.xlsx', str(caught.exception))

    def test_report_replacement_removes_old_zero_transaction_dates(self):
        first = os.path.join(_TMP.name, 'old.xls')
        latest = os.path.join(_TMP.name, 'new.xls')
        for path, mtime in ((first, 1), (latest, 2)):
            with open(path, 'wb') as target:
                target.write(b'test')
            os.utime(path, (mtime, mtime))
        reports = {
            first: report('Июнь 2026', 10, ['2026-06-10']),
            latest: report('Июнь 2026', 20, ['2026-06-20']),
        }
        daily = {
            'acc': '51', 'date': date(2026, 6, 21), 'snd': 0, 'snk': 0,
            'obd': 100, 'obk': 0, 'skd': 100, 'skk': 0,
        }
        with patch.object(service, 'parse_osv', side_effect=lambda p: copy.deepcopy(reports[p])):
            source = service.assemble_source([first, latest], [daily])
        rows = source['accounts']['51']['daily']
        self.assertEqual(sum(row['obd'] for row in rows), 20)
        self.assertEqual([row['date'] for row in rows], ['2026-06-20'])
        # The monthly turnover wins, but the independently observed daily
        # balance must remain available on June 21.
        balances = source['accounts']['51'].get('balance_observations', [])
        self.assertEqual(len(balances), 1)
        self.assertEqual(balances[0]['date'], '2026-06-21')

    def test_unreadable_source_is_explicit_error(self):
        path = os.path.join(_TMP.name, 'broken.xls')
        with open(path, 'wb') as target:
            target.write(b'invalid')
        with patch.object(service, 'parse_osv', side_effect=ValueError):
            with self.assertRaises(service.SourceReadError):
                service.assemble_source([path], [])

    def test_only_this_company_accepted_monthly_files(self):
        root = tempfile.mkdtemp(dir=_TMP.name)
        for folder in ('3/accepted', '3/historical', '3/zip_x', '3/stage_x',
                       '3/failed_x', '4/accepted'):
            os.makedirs(os.path.join(root, folder))
            with open(os.path.join(root, folder, 'report.xls'), 'wb') as f:
                f.write(b'test')
        paths = service.monthly_files(root, 3)
        self.assertEqual(len(paths), 1)
        self.assertIn('/3/accepted/', paths[0])

    def test_in_progress_batch_newer_than_committed_upload_is_not_read(self):
        from datetime import datetime
        root = tempfile.mkdtemp(dir=_TMP.name)
        for folder in ('20260601_120000', '20260602_120000'):
            target = os.path.join(root, '3', folder)
            os.makedirs(target)
            with open(os.path.join(target, 'report.xls'), 'wb') as f:
                f.write(b'test')
        paths = service.monthly_files(root, 3, datetime(2026, 6, 1, 12, 5))
        self.assertEqual(len(paths), 1)
        self.assertIn('20260601_120000', paths[0])


class KeyIndicatorRouteTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = app_module.app
        cls.app.config.update(TESTING=True, WTF_CSRF_ENABLED=False)
        cls.imap_patch = patch.object(app_module, '_ensure_imap_thread')
        cls.imap_patch.start()
        cls.addClassCleanup(cls.imap_patch.stop)
        with cls.app.app_context():
            db.create_all()
            admin = User(username='kpi_admin', password_hash='not-a-password', role='admin')
            allowed = User(username='kpi_allowed', password_hash='not-a-password', role='user')
            outsider = User(username='kpi_outsider', password_hash='not-a-password', role='user')
            no_module = User(username='kpi_no_module', password_hash='not-a-password', role='user')
            db.session.add_all([admin, allowed, outsider, no_module])
            db.session.flush()
            company = Company(name='Источник <script> не HTML', created_by=admin.id)
            other = Company(name='Недоступное предприятие', created_by=admin.id)
            db.session.add_all([company, other])
            db.session.flush()
            for user in (allowed, outsider):
                db.session.add(UserModuleAccess(user_id=user.id, module_key='analytics'))
            for user in (allowed, no_module):
                db.session.add(CompanyMembership(
                    user_id=user.id, company_id=company.id, participation_type='creator'))
            db.session.commit()
            cls.ids = {u.username: u.id for u in (admin, allowed, outsider, no_module)}
            cls.company = company.id
            cls.other = other.id

    def client(self, username=None):
        client = self.app.test_client()
        if username:
            with client.session_transaction() as session:
                session['_user_id'] = str(self.ids[username])
                session['_fresh'] = True
        return client

    def url(self, company=None, data=False):
        return f'/company/{company or self.company}/key-indicators' + ('/data' if data else '')

    def test_auth_and_company_boundaries_on_html_and_json(self):
        for data in (False, True):
            self.assertEqual(self.client().get(self.url(data=data)).status_code, 302)
            for user in ('kpi_outsider', 'kpi_no_module'):
                self.assertEqual(self.client(user).get(self.url(data=data)).status_code, 403)
            self.assertEqual(
                self.client('kpi_allowed').get(self.url(self.other, data)).status_code, 403)

    def test_page_and_safe_company_name(self):
        response = self.client('kpi_allowed').get(self.url())
        self.assertEqual(response.status_code, 200)
        self.assertIn('&lt;script&gt;', response.get_data(as_text=True))
        self.assertIn('key_indicators.js', response.get_data(as_text=True))
        self.assertIn('window.KI_CONFIG =', response.get_data(as_text=True))

    def test_ai_report_uses_selected_operational_period_not_saved_dashboard(self):
        source = build_indicator_source({
            '90': {
                'acc': '90', 'period': 'за январь 2026 г.', 'rows': [],
                'warnings': [], 'top': {
                    'snd': 0, 'snk': 0, 'obd': 0, 'obk': 0,
                    'skd': 0, 'skk': 0,
                },
                'daily': [
                    {'date': '2026-01-30', 'month': '2026-01',
                     'path1': '90.01', 'path2': '', 'path3': '',
                     'obd': 0, 'obk': 1000000},
                    {'date': '2026-01-31', 'month': '2026-01',
                     'path1': '90.01', 'path2': '', 'path3': '',
                     'obd': 0, 'obk': 2000000},
                ],
            },
            '51': {
                'acc': '51', 'period': 'за январь 2026 г.', 'rows': [],
                'warnings': [], 'top': {
                    'snd': 0, 'snk': 0, 'obd': 0, 'obk': 0,
                    'skd': 0, 'skk': 0,
                },
                'daily': [{
                    'date': '2026-01-31', 'month': '2026-01',
                    'path1': '', 'path2': '', 'path3': '',
                    'obd': 0, 'obk': 0, 'skd': 3000000, 'skk': 0,
                }],
            },
            '66': {
                'acc': '66', 'period': 'за январь 2026 г.', 'rows': [],
                'warnings': [], 'top': {
                    'snd': 0, 'snk': 0, 'obd': 0, 'obk': 0,
                    'skd': 0, 'skk': 0,
                },
                'daily': [{
                    'date': '2026-01-31', 'month': '2026-01',
                    'path1': '', 'path2': '', 'path3': '',
                    'obd': 0, 'obk': 0, 'skd': 0, 'skk': 4000000,
                }],
            },
        })
        report_url = f'/company/{self.company}/key-indicators/ai-report'
        with patch.object(service, 'load_company_source', return_value=(source, None)):
            first = self.client('kpi_allowed').get(
                report_url, query_string={'period': 'day', 'date': '2026-01-30'})
            second = self.client('kpi_allowed').get(
                report_url, query_string={'period': 'day', 'date': '2026-01-31'})
            missing = self.client('kpi_allowed').get(
                report_url, query_string={'period': 'day', 'date': '2026-02-01'})
            first_data = self.client('kpi_allowed').get(
                self.url(data=True),
                query_string={'period': 'day', 'date': '2026-01-30'})
            second_data = self.client('kpi_allowed').get(
                self.url(data=True),
                query_string={'period': 'day', 'date': '2026-01-31'})
            first_shell = self.client('kpi_allowed').get(
                self.url(), query_string={'period': 'day', 'date': '2026-01-30'})
            second_shell = self.client('kpi_allowed').get(
                self.url(), query_string={'period': 'day', 'date': '2026-01-31'})
        first_html = first.get_data(as_text=True)
        second_html = second.get_data(as_text=True)
        self.assertEqual(first.status_code, 200)
        self.assertEqual(second.status_code, 200)
        self.assertIn('Что происходит — данные за 30.01.2026', first_html)
        self.assertIn('Что происходит — данные за 31.01.2026', second_html)
        self.assertIn('1.00 млн ₽', first_html)
        self.assertIn('2.00 млн ₽', second_html)
        self.assertNotIn('dashboard_current.html', second_html)
        self.assertIn('Частичная известная сумма', second_html)
        self.assertIn('отсутствуют: 50, 52, 55', second_html)
        self.assertIn('отсутствуют: 67', second_html)
        self.assertEqual(missing.status_code, 400)
        self.assertIn('Отчёт недоступен', missing.get_data(as_text=True))
        fixture_path = '/tmp/ai-period-fixtures.json'
        def fixture_html(value):
            # Browser fixtures document rendered shape, never a request token.
            value = re.sub(r'csrfToken:\s*"[^"]*"', 'csrfToken: "[redacted]"', value)
            return re.sub(r"'X-CSRFToken':\s*\"[^\"]*\"",
                          "'X-CSRFToken': \"[redacted]\"", value)
        with open(fixture_path, 'w', encoding='utf-8') as fixture:
            json.dump({
                'first': {
                    'report_html': fixture_html(first_html),
                    'kpi_shell_html': fixture_html(first_shell.get_data(as_text=True)),
                    'kpi_payload': first_data.get_json(),
                },
                'second': {
                    'report_html': fixture_html(second_html),
                    'kpi_shell_html': fixture_html(second_shell.get_data(as_text=True)),
                    'kpi_payload': second_data.get_json(),
                },
            }, fixture, ensure_ascii=False, indent=2)
        self.assertTrue(os.path.isfile(fixture_path))

    def test_ai_report_names_gross_profit_fallback_without_sales_result_claim(self):
        source = build_indicator_source({
            '90': {
                'acc': '90', 'period': 'за январь 2026 г.', 'rows': [],
                'warnings': [], 'top': {
                    'snd': 0, 'snk': 0, 'obd': 0, 'obk': 0,
                    'skd': 0, 'skk': 0,
                },
                'daily': [
                    {'date': '2026-01-31', 'month': '2026-01',
                     'path1': '90.01', 'path2': '', 'path3': '',
                     'obd': 0, 'obk': 3000000},
                    {'date': '2026-01-31', 'month': '2026-01',
                     'path1': '90.02', 'path2': '', 'path3': '',
                     'obd': 1000000, 'obk': 0},
                    {'date': '2026-01-31', 'month': '2026-01',
                     'path1': '90.03', 'path2': '', 'path3': '',
                     'obd': 0, 'obk': 0},
                ],
            },
        })
        with patch.object(service, 'load_company_source', return_value=(source, None)):
            response = self.client('kpi_allowed').get(
                f'/company/{self.company}/key-indicators/ai-report',
                query_string={'period': 'day', 'date': '2026-01-31'})
        html = response.get_data(as_text=True)
        self.assertEqual(response.status_code, 200)
        self.assertIn('Валовая прибыль: 2.00 млн ₽', html)
        self.assertNotIn('Результат от продаж по расчёту показателей', html)

    def test_ai_report_keeps_manual_brief_action_for_complete_period_facts(self):
        source = build_indicator_source({
            '90': {
                'acc': '90', 'period': 'за январь 2026 г.', 'rows': [],
                'warnings': [], 'top': {
                    'snd': 0, 'snk': 0, 'obd': 0, 'obk': 0,
                    'skd': 0, 'skk': 0,
                },
                'daily': [
                    {'date': '2026-01-31', 'month': '2026-01',
                     'path1': account, 'path2': '', 'path3': '',
                     'obd': debit, 'obk': credit}
                    for account, debit, credit in (
                        ('90.01', 0, 3000000), ('90.03', 0, 0),
                        ('90.02', 1000000, 0), ('90.07', 0, 0),
                        ('90.08', 0, 0),
                    )
                ],
            },
        })
        with patch.object(service, 'load_company_source', return_value=(source, None)):
            response = self.client('kpi_allowed').get(
                f'/company/{self.company}/key-indicators/ai-report',
                query_string={'period': 'day', 'date': '2026-01-31'})
        html = response.get_data(as_text=True)
        self.assertEqual(response.status_code, 200)
        self.assertIn('id="ai-brief-btn"', html)
        self.assertIn("fetch('/ai/brief'", html)
        self.assertIn('"period": "31.01.2026"', html)

    def test_ai_report_uses_annual_year_contract_separately(self):
        report_url = f'/company/{self.company}/key-indicators/ai-report'
        annual_payload = {
            'mode': 'annual', 'year': 2024, 'cards': [{
                'id': 'revenue', 'title': 'Выручка', 'value': 3000000,
                'unit': 'RUB', 'period_label': 'За год',
                'missing_reason': None,
            }], 'warnings': [],
        }
        with patch.object(app_module, '_require_historical_import'), patch(
                'annual_key_indicators.build_annual_key_indicators',
                return_value=annual_payload):
            response = self.client('kpi_admin').get(
                report_url, query_string={'mode': 'annual', 'year': '2024'})
        self.assertEqual(response.status_code, 200)
        html = response.get_data(as_text=True)
        self.assertIn('Что происходит — данные за 2024', html)
        self.assertIn('3.00 млн ₽', html)

    def test_ai_report_accepts_normalized_month_and_quarter_ranges(self):
        source = build_indicator_source({
            '90': {
                'acc': '90', 'period': 'за 1 полугодие 2026 г.', 'rows': [],
                'warnings': [], 'top': {
                    'snd': 0, 'snk': 0, 'obd': 0, 'obk': 0,
                    'skd': 0, 'skk': 0,
                },
                'daily': [
                    {'date': day, 'month': day[:7], 'path1': '90.01',
                     'path2': '', 'path3': '', 'obd': 0, 'obk': amount}
                    for day, amount in (
                        ('2026-03-31', 1000000), ('2026-04-30', 2000000),
                        ('2026-06-30', 4000000),
                    )
                ],
            },
        })
        report_url = f'/company/{self.company}/key-indicators/ai-report'
        with patch.object(service, 'load_company_source', return_value=(source, None)):
            month = self.client('kpi_allowed').get(report_url, query_string={
                'period': 'month', 'date_from': '2026-03-31',
                'date': '2026-04-30',
            })
            quarter = self.client('kpi_allowed').get(report_url, query_string={
                'period': 'quarter', 'date_from': '2026-03-31',
                'date': '2026-06-30',
            })
        self.assertEqual(month.status_code, 200)
        self.assertEqual(quarter.status_code, 200)
        self.assertIn('Что происходит — данные за 01.03.2026 — 30.04.2026',
                      month.get_data(as_text=True))
        self.assertIn('3.00 млн ₽', month.get_data(as_text=True))
        self.assertIn('Что происходит — данные за 01.01.2026 — 30.06.2026',
                      quarter.get_data(as_text=True))
        self.assertIn('7.00 млн ₽', quarter.get_data(as_text=True))

    def test_quarter_range_api_returns_daily_points(self):
        source = build_indicator_source({'51': report(
            'за 1 полугодие 2026 г.', 20,
            ['2026-03-31', '2026-04-30', '2026-06-30'])})
        with patch.object(service, 'load_company_source', return_value=(source, None)):
            response = self.client('kpi_allowed').get(
                self.url(data=True),
                query_string={'period': 'quarter', 'date_from': '2026-03-31',
                              'date': '2026-06-30'})
            self.assertEqual(response.status_code, 200)
            payload = response.get_json()
            self.assertEqual(payload['period']['start'], '2026-01-01')
            self.assertEqual(payload['period']['series_kind'], 'day')
            self.assertTrue(all(len(c['series']) == 181 for c in payload['cards']))
            bad = self.client('kpi_allowed').get(
                self.url(data=True),
                query_string={'period': 'quarter', 'date_from': '2026-06-30',
                              'date': '2026-03-31'})
            self.assertEqual(bad.status_code, 400)

    def test_coverage_days_include_loaded_account_indexes(self):
        with self.app.app_context():
            db.session.add_all([
                DailyRow(company_id=self.company, date=date(2026, 9, 1), acc='51'),
                DailyRow(company_id=self.company, date=date(2026, 9, 1), acc='62'),
                DailyRow(company_id=self.company, date=date(2026, 9, 2), acc='62'),
                Upload(
                    company_id=self.company,
                    uploaded_by=self.ids['kpi_admin'],
                    period_label='Август 2026',
                    accounts_loaded='["20"]',
                ),
            ])
            db.session.commit()
            coverage = app_module._annual_coverage(self.company, 2026)
            DailyRow.query.filter_by(company_id=self.company).delete()
            Upload.query.filter_by(company_id=self.company).delete()
            db.session.commit()

        self.assertEqual(coverage['accs'], ['20', '51', '62'])
        august = coverage['months'][7]
        august_days = {cell['day']: cell for cell in august['days'] if cell}
        self.assertEqual(august_days[1]['account_indexes'], [0])
        september = coverage['months'][8]
        days = {cell['day']: cell for cell in september['days'] if cell}
        self.assertEqual(days[1]['account_indexes'], [1, 2])
        self.assertEqual(days[2]['account_indexes'], [2])

    def test_coverage_lists_all_accounts_and_marks_missing(self):
        with app_module.app.app_context():
            coverage = app_module._annual_coverage(self.company, 2025)
        buttons = {item['number']: item for item in coverage['account_buttons']}

        self.assertEqual(
            list(buttons),
            ['01', '08', '10', '20', '26', '44', '50', '51', '52', '55',
             '60', '62', '66', '67', '68', '69', '70', '76', '90', '91'],
        )
        self.assertEqual(len(buttons), 20)
        self.assertFalse(buttons['01']['available'])

    def test_coverage_year_select_uses_data_periods_and_updates_coverage(self):
        from datetime import datetime
        source = {'accounts': {'51': {
            'report_start': '2021-01-01', 'report_end': '2023-12-31',
            'daily': [],
        }}}
        with self.app.app_context():
            db.session.add_all([
                DailyRow(company_id=self.company, date=date(2024, 5, 2), acc='62'),
                Upload(company_id=self.company, uploaded_by=self.ids['kpi_admin'],
                       uploaded_at=datetime(2030, 1, 1),
                       period_label='Август 2020', accounts_loaded='["20"]'),
            ])
            db.session.commit()
        try:
            with patch.object(service, 'load_company_source', return_value=(source, None)):
                with self.app.app_context():
                    default = app_module._annual_coverage(self.company)
                    old = app_module._annual_coverage(self.company, 2020)
                    middle = app_module._annual_coverage(self.company, 2022)
                self.assertEqual(default['data_years'], [2020, 2021, 2022, 2023, 2024])
                self.assertEqual(default['year'], 2024)
                self.assertEqual(default['accs'], ['62'])
                self.assertEqual(old['year'], 2020)
                self.assertEqual(old['accs'], ['20'])
                self.assertEqual(old['months'][7]['upl_status'], 'uploaded')
                self.assertEqual(middle['accs'], ['51'])
                response = self.client('kpi_admin').get(
                    f'/company/{self.company}?year=2020&embedded=1')
            self.assertEqual(response.status_code, 200)
            html = response.get_data(as_text=True)
            select = html.split('<select id="ac-year-select"', 1)[1].split('</select>', 1)[0]
            self.assertIn('onchange="setCoverageYear(this.value)"', select)
            self.assertIn('onclick="event.stopPropagation()"', select)
            self.assertIn('<option value="2020" selected>', select)
            for year in range(2021, 2025):
                self.assertIn(f'<option value="{year}"', select)
            self.assertNotIn('value="2030"', select)
            self.assertEqual(select.count('<option'), 5)
        finally:
            with self.app.app_context():
                DailyRow.query.filter_by(company_id=self.company).delete()
                Upload.query.filter_by(company_id=self.company).delete()
                db.session.commit()

    def test_coverage_year_select_empty_company(self):
        with patch.object(service, 'load_company_source', return_value=({'accounts': {}}, None)):
            with self.app.app_context():
                coverage = app_module._annual_coverage(self.company)
            response = self.client('kpi_admin').get(f'/company/{self.company}')
        self.assertEqual(coverage['data_years'], [])
        self.assertEqual(response.status_code, 200)
        select = response.get_data(as_text=True).split(
            '<select id="ac-year-select"', 1)[1].split('</select>', 1)[0]
        self.assertIn('disabled', select)
        self.assertIn('Нет данных', select)
        self.assertEqual(select.count('<option'), 1)

    def test_coverage_report_downloads_readable_excel_matrix(self):
        payload = {
            'year': 2026,
            'months': [1, 3],
            'rows': [{
                'account': '01',
                'accountName': 'Подменять серверное название нельзя',
                'daysByMonth': {
                    '1': list(range(1, 32)),
                    '3': [1, 2, 3, 4, 5, 8, 10, 11],
                },
            }],
        }
        response = self.client('kpi_allowed').post(
            f'/company/{self.company}/coverage-report.xlsx',
            json=payload,
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response.mimetype,
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        disposition = unquote(response.headers['Content-Disposition'])
        self.assertIn(
            'Запрос_Данных_Источник_script_не_HTML_',
            disposition,
        )
        self.assertTrue(disposition.endswith('.xlsx'))

        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(response.data), read_only=True)
        sheet = workbook['Запрос данных']
        self.assertEqual(sheet['A4'].value, 'Счёт')
        self.assertEqual(sheet['C4'].value, 'Январь')
        self.assertEqual(sheet['D4'].value, 'Март')
        self.assertEqual(sheet['A5'].value, '01')
        self.assertEqual(sheet['B5'].value, 'Основные средства')
        self.assertEqual(sheet['C5'].value, 'с 1 по 31')
        self.assertEqual(sheet['D5'].value, '1 - 5, 8, 10 - 11')
        workbook.close()

    def test_coverage_recovers_accounts_from_accepted_osv_source(self):
        source = {
            'accounts': {
                '51': {
                    'report_start': '2026-07-01',
                    'report_end': '2026-08-17',
                    'daily': [{'date': '2026-07-02'}],
                },
                '62': {
                    'report_start': '2026-07-01',
                    'report_end': '2026-08-17',
                    'daily': [{'date': '2026-07-03'}],
                },
            },
        }
        with self.app.app_context(), patch.object(
                service, 'load_company_source', return_value=(source, None)):
            coverage = app_module._annual_coverage(self.company, 2026)

        self.assertEqual(coverage['accs'], ['51', '62'])
        july = coverage['months'][6]
        days = {cell['day']: cell for cell in july['days'] if cell}
        self.assertEqual(days[2]['account_indexes'], [0])
        self.assertEqual(days[3]['account_indexes'], [1])

    def test_empty_company_returns_real_unavailable_cards(self):
        with patch.object(service, 'load_company_source',
                          return_value=(build_indicator_source({}), None)):
            response = self.client('kpi_admin').get(self.url(data=True))
        self.assertEqual(response.status_code, 200)
        payload = response.get_json()
        self.assertEqual(len(payload['cards']), 10)
        self.assertTrue(all(card['value'] is None for card in payload['cards']))
        self.assertIsNone(payload['updated_at'])

    def test_invalid_period_rejected_before_source_load(self):
        with patch.object(service, 'load_company_source') as loader:
            for query in ('?period=nope', '?date=garbage'):
                response = self.client('kpi_allowed').get(self.url(data=True) + query)
                self.assertEqual(response.status_code, 400)
            loader.assert_not_called()

    def test_busy_and_read_errors_visible(self):
        for error, expected in (
            (service.SourceBusyError('Занято'), 503),
            (service.SourceReadError('Ошибка источника'), 422),
        ):
            with patch.object(service, 'load_company_source', side_effect=error):
                response = self.client('kpi_admin').get(self.url(data=True))
            self.assertEqual(response.status_code, expected)
            self.assertIn('error', response.get_json())

    def test_only_admin_can_save_company_revenue_mappings(self):
        url = f'/company/{self.company}/key-indicators/revenue-mappings'
        payload = {'mappings': [{
            'account': '90.01',
            'analytic_key': revenue_analytic_key(['Особое название']),
            'analytic_label': '«Особое название»',
            'category': 'population',
        }]}
        self.assertEqual(
            self.client('kpi_allowed').post(url, json=payload).status_code, 403)
        response = self.client('kpi_admin').post(url, json=payload)
        self.assertEqual(response.status_code, 200)
        with self.app.app_context():
            row = RevenueAnalyticMapping.query.filter_by(
                company_id=self.company,
                analytic_key=revenue_analytic_key(['Особое название'])).one()
            self.assertEqual(row.category, 'population')

    def test_mapping_save_preserves_analytics_from_other_periods(self):
        url = f'/company/{self.company}/key-indicators/revenue-mappings'
        client = self.client('kpi_admin')
        first = {'account': '90.01',
                 'analytic_key': revenue_analytic_key(['Январь']),
                 'analytic_label': '«Январь»',
                 'category': 'population'}
        second = {'account': '90.01',
                  'analytic_key': revenue_analytic_key(['Февраль']),
                  'analytic_label': '«Февраль»',
                  'category': 'budget'}
        self.assertEqual(client.post(
            url, json={'mappings': [first]}).status_code, 200)
        self.assertEqual(client.post(
            url, json={'mappings': [second]}).status_code, 200)
        with self.app.app_context():
            labels = {
                row.analytic_label for row in RevenueAnalyticMapping.query.filter_by(
                    company_id=self.company).all()
            }
        self.assertTrue({'«Январь»', '«Февраль»'}.issubset(labels))

    def test_only_admin_can_save_periodic_cost_mapping(self):
        url = f'/company/{self.company}/key-indicators/cost-mappings'
        payload = {'mappings': [{
            'account': '20',
            'analytic_key': revenue_analytic_key(['Газ']),
            'analytic_label': '«Газ»',
            'category': 'fuel',
            'target_account': '90.02',
            'valid_from': '2026-01-01',
            'valid_to': '2026-12-31',
        }]}
        self.assertEqual(
            self.client('kpi_allowed').post(url, json=payload).status_code, 403)
        response = self.client('kpi_admin').post(url, json=payload)
        self.assertEqual(response.status_code, 200)
        with self.app.app_context():
            row = CostAnalyticMapping.query.filter_by(
                company_id=self.company).one()
            self.assertEqual(row.category, 'fuel')
            self.assertEqual(row.target_account, '90.02')
            self.assertEqual(row.valid_from, date(2026, 1, 1))

    def test_cost_mapping_rejects_invalid_period(self):
        url = f'/company/{self.company}/key-indicators/cost-mappings'
        response = self.client('kpi_admin').post(url, json={'mappings': [{
            'account': '20', 'analytic_key': revenue_analytic_key(['Газ']),
            'analytic_label': '«Газ»', 'category': 'fuel',
            'target_account': '90.02', 'valid_from': '2026-02-01',
            'valid_to': '2026-01-01',
        }]})
        self.assertEqual(response.status_code, 400)

    def test_cost_mapping_rejects_conflicting_overlap(self):
        url = f'/company/{self.company}/key-indicators/cost-mappings'
        common = {
            'account': '20', 'analytic_key': revenue_analytic_key(['Газ']),
            'analytic_label': '«Газ»', 'target_account': '90.02',
            'valid_from': '2026-01-01', 'valid_to': '2026-12-31',
        }
        response = self.client('kpi_admin').post(url, json={'mappings': [
            dict(common, category='fuel'),
            dict(common, category='other', valid_from='2026-06-01'),
        ]})
        self.assertEqual(response.status_code, 400)

    def test_only_admin_saves_periodic_debtor_mapping(self):
        url = f'/company/{self.company}/key-indicators/debtor-mappings'
        payload = {'mappings': [{
            'account': '76.05',
            'analytic_key': revenue_analytic_key(['Субсидия']),
            'analytic_label': '«Субсидия»',
            'category': 'subsidy',
            'effective_from': '2026-01-01',
            'effective_to': '2026-01-31',
        }]}
        self.assertEqual(
            self.client('kpi_allowed').post(url, json=payload).status_code, 403)
        self.assertEqual(
            self.client('kpi_admin').post(url, json=payload).status_code, 200)
        with self.app.app_context():
            row = DebtorAnalyticMapping.query.filter_by(
                company_id=self.company,
                analytic_key=revenue_analytic_key(['Субсидия'])).one()
            self.assertEqual(row.account, '76.05')
            self.assertEqual(row.effective_from.isoformat(), '2026-01-01')

    def test_day_edit_and_revoke_preserve_persisted_month_interval(self):
        url = f'/company/{self.company}/key-indicators/debtor-mappings'
        key = revenue_analytic_key(['Интервальный должник'])
        client = self.client('kpi_admin')
        original = {
            'account': '76.09', 'analytic_key': key,
            'analytic_label': '«Интервальный должник»',
            'category': 'budget',
            'effective_from': '2026-02-01',
            'effective_to': '2026-02-28',
        }
        self.assertEqual(client.post(
            url, json={'mappings': [original]}).status_code, 200)
        with self.app.app_context():
            row = DebtorAnalyticMapping.query.filter_by(
                company_id=self.company, account='76.09',
                analytic_key=key).one()
            mapping_id = row.id
        accidental = dict(original, category='legal',
                          effective_from='2026-02-14',
                          effective_to='2026-02-14')
        conflict = client.post(url, json={'mappings': [accidental]})
        self.assertEqual(conflict.status_code, 409)
        edited = dict(original, id=mapping_id, category='legal',
                      effective_from='2026-02-14',
                      effective_to='2026-02-14')
        self.assertEqual(client.post(
            url, json={'mappings': [edited]}).status_code, 200)
        with self.app.app_context():
            row = db.session.get(DebtorAnalyticMapping, mapping_id)
            self.assertEqual(row.category, 'legal')
            self.assertEqual(row.effective_from.isoformat(), '2026-02-01')
            self.assertEqual(row.effective_to.isoformat(), '2026-02-28')
        revoked = dict(edited, category='')
        self.assertEqual(client.post(
            url, json={'mappings': [revoked]}).status_code, 200)
        with self.app.app_context():
            self.assertIsNone(db.session.get(DebtorAnalyticMapping, mapping_id))


if __name__ == '__main__':
    unittest.main()
